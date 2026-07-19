import json
from datetime import datetime, timezone
from pathlib import Path

from tests.agent_blueprint_fakes import *  # noqa: F403
from tests.source_contract_helpers import read_agent_blueprints_frontend_source


def test_agent_blueprint_api_guards_version_blueprint_mismatch():
    api_source = Path("src/api/agent_blueprints_api.py").read_text(encoding="utf-8")
    workspace_source = Path("src/services/agent_blueprint_workspace.py").read_text(encoding="utf-8")
    legacy_migration_source = Path("src/services/agent_legacy_migration.py").read_text(encoding="utf-8")
    webhooks_source = Path("src/ai_agent_webhooks.py").read_text(encoding="utf-8")
    ai_agent_source = Path("src/ai_agent.py").read_text(encoding="utf-8")
    chats_api_source = Path("src/chats_api.py").read_text(encoding="utf-8")
    document_llm_source = Path("src/services/agent_document_llm.py").read_text(encoding="utf-8")
    email_llm_source = Path("src/services/agent_email_llm.py").read_text(encoding="utf-8")
    review_analysis_source = Path("src/services/agent_review_reply_analysis.py").read_text(encoding="utf-8")
    table_analysis_source = Path("src/services/agent_table_analysis.py").read_text(encoding="utf-8")
    capability_handlers_source = Path("src/services/agent_capability_handlers.py").read_text(encoding="utf-8")
    action_policy_source = Path("src/core/action_policy.py").read_text(encoding="utf-8")
    trigger_runtime_source = Path("src/services/agent_trigger_runtime.py").read_text(encoding="utf-8")
    worker_source = Path("src/worker.py").read_text(encoding="utf-8")
    compose_source = Path("docker-compose.yml").read_text(encoding="utf-8")
    telegram_webhook_source = Path("src/ai_agent_webhooks.py").read_text(encoding="utf-8")
    builder_api_source = Path("src/api/agent_builder_api.py").read_text(encoding="utf-8")
    agents_page_source = read_agent_blueprints_frontend_source()
    admin_page_source = Path("frontend/src/pages/dashboard/AdminPage.tsx").read_text(encoding="utf-8")

    assert "VERSION_BLUEPRINT_MISMATCH" in api_source
    assert "_load_blueprint_version_for_blueprint" in api_source
    assert "build_agent_blueprint_orchestrator" in api_source
    assert "run_status" in api_source
    assert "approval_queue" in api_source
    assert "last_run_status" in api_source
    assert "pending_approvals_count" in api_source
    assert "sources_count" in api_source
    assert "versions_count" in api_source
    assert "/api/agent-blueprints/draft" in api_source
    assert "build_agent_builder_state" in api_source
    assert "direct_draft_envelope" in api_source
    assert "_direct_selected_connection_bindings" in api_source
    assert "_direct_missing_required_connection_choices" in api_source
    assert "_apply_direct_selected_connection_bindings" in api_source
    assert "_selected_provider_routes" in api_source
    assert "_missing_required_provider_routes" in api_source
    assert "_required_provider_route_bindings" in api_source
    assert "_apply_selected_provider_routes" in api_source
    assert "_apply_answer_connection_bindings" in builder_api_source
    assert "builder_answer_connection_bindings" in builder_api_source
    assert "connection_answer_bindings" in Path("src/services/agent_builder_session.py").read_text(encoding="utf-8")
    assert "answer_config" in api_source
    assert "connectionResourceFacts" in agents_page_source
    assert "Ресурс из диалога" in agents_page_source
    assert "Поняли ресурс" in agents_page_source
    assert "BuilderExecutionBoundaryPanel" in agents_page_source
    assert "Execution boundary" in agents_page_source
    assert "OpenClaw action refs" in agents_page_source
    assert "openclaw_action_plan" in api_source
    assert "_preview_openclaw_action_plan" in api_source
    runner_source = Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "_preview_summary_openclaw_action_plan" in runner_source
    assert "_preview_summary_policy_envelope" in runner_source
    assert "\"approval_gate\": approval_gate" in runner_source
    assert "OpenClaw actions в safe preview" in agents_page_source
    assert "OpenClawPreviewActionPlanPanel" in agents_page_source
    assert "side effects выключены" in agents_page_source
    assert "BuilderRequiredConnectionsPanel" in agents_page_source
    assert "Доступы перед созданием агента" in agents_page_source
    assert "Ресурс из диалога" in agents_page_source
    assert "способ выбран" in agents_page_source
    assert "builderConnectionCardStatus" in agents_page_source
    assert "AGENT_CONNECTION_CHOICE_REQUIRED" in api_source
    assert "AGENT_PROVIDER_ROUTE_REQUIRED" in api_source
    assert "AGENT_PROVIDER_ROUTES_CONFIRMATION_REQUIRED" in api_source
    assert "\"post_create_handoff\": post_create_handoff" in api_source
    assert "metadata[\"agent_builder_preview\"] = preview" in api_source
    assert "metadata[\"openclaw_planner_context\"] = planner_context" in api_source
    assert "\"connection_summary\": preview.get(\"connection_summary\")" in api_source
    assert "_load_direct_builder_connection_inventory" in api_source
    assert "/api/agent-blueprints/legacy-migration/apply" in api_source
    assert "apply_legacy_ai_agent_migration" in api_source
    assert "learning_events" in api_source
    assert "version_events" in api_source
    assert "legacy_migration" in api_source
    assert "build_agent_blueprint_draft" in api_source
    assert "_insert_version(cursor, blueprint_id, version_payload, user_data)" in api_source
    assert "/api/agent-blueprints/<blueprint_id>/setup" in api_source
    assert "/api/agent-blueprints/<blueprint_id>/sources" in api_source
    assert "/api/agent-blueprints/<blueprint_id>/sources/catalog" in api_source
    assert "/api/agent-blueprints/<blueprint_id>/sources/upload" in api_source
    assert "build_agent_datahub_catalog" in api_source
    assert "build_agent_source_from_upload" in api_source
    assert "/api/agent-blueprints/<blueprint_id>/preflight" in api_source
    assert "preflight_agent_blueprint_run" in api_source
    assert "localos_agent_preview_run_gate_v1" in api_source
    assert '"external_side_effects_allowed": False' in api_source
    assert '"next_binding_key": _connection_plan_next_binding_key(connection_plan)' in api_source
    assert "_build_agent_preview_run_input" in api_source
    assert "localos_agent_preview_input_v1" in api_source
    assert "AGENT_INTEGRATIONS_REQUIRED" in api_source
    assert 'return jsonify(result), 400' in api_source
    assert "/api/agent-runs/<run_id>/feedback" in api_source
    assert "trigger_type" in api_source
    assert "auto_activate" in api_source
    assert "auto_activation_gate = _build_activation_gate_summary" in api_source
    assert "if auto_activation_gate.get(\"can_activate\")" in api_source
    assert "auto_activation_blocked" in api_source
    assert "auto_activation_applied" in api_source
    assert "SET status = 'active'" in api_source
    assert "build_learning_loop_summary" in api_source
    assert "/api/agent-runs/<run_id>/support-export" in api_source
    assert "build_run_support_export" in api_source
    assert "/api/agent-blueprints/<blueprint_id>/versions/<version_id>/diff" in api_source
    assert "/api/agent-blueprints/<blueprint_id>/versions/<version_id>/activate" in api_source
    assert "localos_agent_activation_gate_v1" in api_source
    assert "localos_agent_preview_run_status_v1" in api_source
    assert "_build_activation_gate_summary" in api_source
    assert "_activation_preview_run_status" in api_source
    assert "primary_action_label" in api_source
    assert "human_blockers" in api_source
    assert "requires_preview_run" in api_source
    assert "AGENT_ACTIVATION_GATE_BLOCKED" in api_source
    assert "_activation_connection_plan_from_preflight" in api_source
    assert "/api/agent-blueprints/<blueprint_id>/versions/<version_id>/rollback" in api_source
    assert "_resolve_active_version" in api_source
    assert "_remember_active_version" in api_source
    assert "_version_was_active_before" in api_source
    assert "AGENT_ROLLBACK_GATE_BLOCKED" in api_source
    assert "candidate_requires_preview" in api_source
    assert "Запустите безопасный preview run перед активацией." in api_source
    assert "build_agent_version_diff" in workspace_source
    assert "agent_learning_loop_v1" in workspace_source
    assert "versioned_review" in workspace_source
    assert "_review_journal" in workspace_source
    assert "journal" in workspace_source
    assert "analyze_document_sources_with_llm" in workspace_source
    assert "draft_email_with_llm" in workspace_source
    assert "draft_review_replies_with_llm" in workspace_source

    feedback_endpoint_start = api_source.index("def create_agent_run_feedback")
    feedback_endpoint = api_source[feedback_endpoint_start:api_source.index("@agent_blueprints_bp.route", feedback_endpoint_start + 1)]
    gate_position = feedback_endpoint.index("auto_activation_gate = _build_activation_gate_summary")
    gate_check_position = feedback_endpoint.index("if auto_activation_gate.get(\"can_activate\")")
    remember_position = feedback_endpoint.index("_remember_active_version")
    assert gate_position < gate_check_position < remember_position
    assert "build_learning_loop_summary(feedback, version, new_version, diff, auto_activation_applied)" in feedback_endpoint

    create_version_start = api_source.index("def create_agent_blueprint_version")
    create_version_endpoint = api_source[create_version_start:api_source.index("@agent_blueprints_bp.route", create_version_start + 1)]
    assert "candidate_version" in create_version_endpoint
    assert "rebuild_from_description" in create_version_endpoint
    assert "build_agent_blueprint_draft" in create_version_endpoint
    assert "_save_blueprint_metadata" in create_version_endpoint
    assert "_remember_active_version" not in create_version_endpoint

    setup_start = api_source.index("def setup_agent_blueprint")
    setup_endpoint = api_source[setup_start:api_source.index("@agent_blueprints_bp.route", setup_start + 1)]
    assert "candidate_version" in setup_endpoint
    assert "_build_activation_gate_summary" in setup_endpoint
    assert "_remember_active_version" not in setup_endpoint

    custom_process_start = api_source.index("def save_agent_blueprint_custom_process")
    custom_process_endpoint = api_source[custom_process_start:api_source.index("@agent_blueprints_bp.route", custom_process_start + 1)]
    assert "candidate_version" in custom_process_endpoint
    assert "_build_activation_gate_summary" in custom_process_endpoint
    assert "_remember_active_version" not in custom_process_endpoint

    direct_draft_start = api_source.index("def create_agent_blueprint_draft")
    direct_draft_endpoint = api_source[direct_draft_start:api_source.index("@agent_blueprints_bp.route", direct_draft_start + 1)]
    assert "connection_inventory = _load_direct_builder_connection_inventory" in direct_draft_endpoint
    assert "selected_bindings = _direct_selected_connection_bindings" in direct_draft_endpoint
    assert "missing_connection_choices = _direct_missing_required_connection_choices" in direct_draft_endpoint
    assert "selected_provider_routes = _selected_provider_routes" in direct_draft_endpoint
    assert "missing_provider_routes = _missing_required_provider_routes" in direct_draft_endpoint
    assert "accepted_provider_routes" in direct_draft_endpoint
    assert "metadata = _apply_direct_selected_connection_bindings" in direct_draft_endpoint
    assert "metadata = _apply_selected_provider_routes" in direct_draft_endpoint
    assert "metadata[\"builder_provider_routes_accepted\"]" in direct_draft_endpoint
    assert "connection_preflight = build_agent_integration_preflight" in direct_draft_endpoint
    assert "post_create_handoff = _build_agent_post_connect_handoff" in direct_draft_endpoint
    assert "analyze_table_with_llm" in workspace_source
    assert "analyze_text_with_gigachat" in document_llm_source
    assert "agent_email_draft" in email_llm_source
    assert "agent_review_replies" in review_analysis_source
    assert "agent_table_analysis" in table_analysis_source
    assert "external_dispatch_performed" in document_llm_source
    assert "external_dispatch_performed" in email_llm_source
    assert "external_dispatch_performed" in review_analysis_source
    assert "publish_state" in review_analysis_source
    assert "external_dispatch_performed" in table_analysis_source
    assert "Улучшение версии" in agents_page_source
    assert "Candidate-версия" in agents_page_source
    assert "Зафиксировать улучшение" in agents_page_source
    assert "auto_activate: false" in agents_page_source
    assert 'title="Агенты"' in agents_page_source
    assert "Создать копию агента" in agents_page_source
    assert "Запустить ещё раз" in agents_page_source
    assert "Примерно" in agents_page_source
    assert "Запустить похожую" not in agents_page_source
    assert "Состояние миграции" in agents_page_source
    assert "Ручные решения" in agents_page_source
    assert "Обучение" in agents_page_source
    assert "активной" in agents_page_source
    assert "explainApproval" in agents_page_source
    assert "Применить миграцию" in agents_page_source
    assert "Открыть Мои агенты" in admin_page_source
    assert "AIAgentsManagement" not in admin_page_source
    assert "AIAgentSettings" not in agents_page_source
    assert "AIAgentsManagement" not in agents_page_source
    assert "business_has_product_agent_runtime" in legacy_migration_source
    assert "business_agent_enabled_for_channel" in legacy_migration_source
    assert "apply_legacy_ai_agent_migration" in legacy_migration_source
    assert "legacy_ai_agent_migration_v1" in legacy_migration_source
    assert "business_agent_enabled_for_channel" in webhooks_source
    assert "WHERE ai_agent_enabled = 1" not in webhooks_source
    assert "legacy_workflow_context" in ai_agent_source
    assert "deprecated_not_runtime_truth" in ai_agent_source
    assert "Legacy AIAgents.workflow no longer drives runtime state transitions" in ai_agent_source
    assert "Ответь на сообщение клиента, учитывая workflow" not in ai_agent_source
    assert "state.get('init_state')" not in ai_agent_source
    assert "legacy_workflow_context" in chats_api_source
    assert "state.get('init_state')" not in chats_api_source
    assert "provenance" in document_llm_source
    assert "provenance" in email_llm_source
    assert "provenance" in review_analysis_source
    assert "provenance" in table_analysis_source
    assert "/api/agent-builder/sessions" in builder_api_source
    assert "build_agent_builder_state" in builder_api_source
    assert "create_blueprint_from_agent_builder_session" in builder_api_source
    assert builder_api_source.index("billing = charge_agent_creation_credits") < builder_api_source.index("draft = compile_agent_blueprint")
    assert "planner_context = preview.get(\"openclaw_planner_context\")" in builder_api_source
    assert "planner_context=planner_context" in builder_api_source
    assert "metadata[\"openclaw_planner_context\"] = planner_context" in builder_api_source
    assert "metadata[\"openclaw_planner_loop\"] = planner_loop" in builder_api_source
    assert "metadata[\"builder_setup_flow\"] = setup_flow" in builder_api_source
    assert "AGENT_SETUP_INCOMPLETE" in builder_api_source
    assert "\"missing_questions\": missing_questions" in builder_api_source
    assert "\"connector_intelligence\": preview.get(\"connector_intelligence\")" in builder_api_source
    assert "selected_connection_bindings" in builder_api_source
    assert "_selected_connection_bindings" in builder_api_source
    assert "_apply_selected_connection_bindings" in builder_api_source
    assert "_missing_required_connection_choices" in builder_api_source
    assert "AGENT_CONNECTION_CHOICE_REQUIRED" in builder_api_source
    assert "metadata[\"builder_selected_connection_bindings\"]" in builder_api_source
    assert "connection_preflight" in builder_api_source
    assert "post_create_handoff" in builder_api_source
    assert "localos_agent_post_create_handoff_v1" in builder_api_source
    assert "next_binding_key" in builder_api_source
    assert "_build_handoff_connection_plan" in builder_api_source
    assert "\"connection_plan\": connection_plan" in builder_api_source
    assert "next_step" in builder_api_source
    assert "use_ai_compiler: true" in agents_page_source
    assert "connect_required_integrations" in agents_page_source
    assert "recentPostCreateHandoff" in agents_page_source
    assert "recentPostCreateHandoff?.connection_plan || agentConnectionPlan" in agents_page_source
    assert "binding_key" in agents_page_source
    assert "selectedBuilderConnectionBindings" in agents_page_source
    assert "selected_connection_bindings: selectedBuilderConnectionBindings" in agents_page_source
    assert "onSelectConnectionBinding" in agents_page_source
    assert "autoSelectBuilderConnectionBindings" in agents_page_source
    assert "Выбрано автоматически" in agents_page_source
    assert "missingConnectionChoices" in agents_page_source
    assert "Сначала выберите подключение" in agents_page_source
    assert "Почему агента пока нельзя создать" in agents_page_source
    assert "createBlockers" in agents_page_source
    assert "LocalOS должен собрать проверяемый workflow" in agents_page_source
    assert "provider_routes" in agents_page_source
    assert "route_summary" in agents_page_source
    assert "providerRouteLabel" in agents_page_source
    assert "providerRouteTone" in agents_page_source
    assert "selectedConnectionBindingKey" in agents_page_source
    assert "onConfigureBinding" in agents_page_source
    assert "Настроить этот доступ" in agents_page_source
    assert "Сейчас настраивается" in agents_page_source
    assert "connectionPlan={agentConnectionPlan}" in agents_page_source
    assert "selectedPlanItem" in agents_page_source
    assert "binding_key: selectedBinding?.key || ''" in agents_page_source
    assert "Использовать для шага" in agents_page_source
    assert "Выбрано" in agents_page_source
    assert "Использовать" in agents_page_source
    assert "Preflight и preview run" in agents_page_source
    assert "preview_mode: true" in agents_page_source
    assert "external_side_effects_allowed: false" in agents_page_source
    assert "Тест без отправки" in agents_page_source
    assert "preflightResponse.data?.next_binding_key" in agents_page_source
    assert "setWorkspaceMode('connections')" in agents_page_source
    assert "activationGate" in agents_page_source
    assert "gate.summary" in agents_page_source
    assert "gate.connection_plan" in agents_page_source
    assert "activationGate?.next_binding_key" in agents_page_source
    assert "openConnectionsFromActivationGate" in agents_page_source
    assert "onDeleteAgent={deleteSelectedAgent}" in agents_page_source
    assert "Архивировать агента" in agents_page_source
    assert "resultPrimaryItems" in agents_page_source
    assert "Готовый ответ" in agents_page_source
    assert "Открыть в контент-плане" in agents_page_source
    assert "Открыть отзывы" in agents_page_source
    assert "Сообщение отправлено в Telegram" in agents_page_source
    assert "output: hasStructuredResult" in agents_page_source
    assert "Intl.supportedValuesOf('timeZone')" in agents_page_source
    assert "Найти город или Europe/Paris" in agents_page_source
    assert agents_page_source.count("<TimezoneSelect") == 3
    assert "Архивировать" in agents_page_source
    assert "primary_action_label" in agents_page_source
    assert "Активировать версию" in agents_page_source
    assert "Запустить preview" in agents_page_source
    assert "preview_run_status" in agents_page_source
    assert "BuilderPlannerLoopPanel" in agents_page_source
    assert "OpenClaw planner" in agents_page_source
    assert "ConnectorIntelligencePanel" in agents_page_source
    assert "Доступность сервисов" in agents_page_source
    assert "connector_intelligence" in agents_page_source
    assert "BuilderTechnicalDiagnostics" in agents_page_source
    assert "Техническая диагностика LocalOS/OpenClaw" in agents_page_source
    assert "Обычный следующий шаг показан выше" in agents_page_source
    assert "BuilderConnectionSummaryPanel" in agents_page_source
    assert "Подключения для агента" in agents_page_source
    assert "connection_summary" in agents_page_source
    assert "AgentProductCockpit" in agents_page_source
    assert "getBlueprintBuilderPreview" in agents_page_source
    assert "detailsBlueprint={blueprintDetails?.blueprint}" in agents_page_source
    assert "BuilderConnectionReadinessPanel" in agents_page_source
    assert "Что нужно агенту для работы" in agents_page_source
    assert "BuilderConnectionResolverPanel" in agents_page_source
    assert "Как LocalOS подключит сервисы" in agents_page_source
    assert "connection_resolver" in agents_page_source
    assert "resolverStateTone" in agents_page_source
    assert "connection_readiness" in agents_page_source
    assert "setup_cta" in agents_page_source
    assert "Настроить подключение" in agents_page_source
    assert "preview?.connection_plan" in agents_page_source
    assert "compact" in agents_page_source
    assert "_build_preview_connection_plan" in Path("src/services/agent_builder_session.py").read_text(encoding="utf-8")
    assert "_build_connection_readiness" in Path("src/services/agent_builder_session.py").read_text(encoding="utf-8")
    assert "localos_agent_connection_readiness_v1" in Path("src/services/agent_builder_session.py").read_text(encoding="utf-8")
    assert "_build_connection_resolver" in Path("src/services/agent_builder_session.py").read_text(encoding="utf-8")
    assert "localos_agent_connection_resolver_v1" in Path("src/services/agent_builder_session.py").read_text(encoding="utf-8")
    assert "_build_connection_summary" in Path("src/services/agent_builder_session.py").read_text(encoding="utf-8")
    assert "localos_agent_connection_summary_v1" in Path("src/services/agent_builder_session.py").read_text(encoding="utf-8")
    assert "_build_connector_intelligence" in Path("src/services/agent_builder_session.py").read_text(encoding="utf-8")
    assert "localos_agent_connector_intelligence_v1" in Path("src/services/agent_builder_session.py").read_text(encoding="utf-8")
    assert "onPreviewRun" in agents_page_source
    assert "onAttachExistingIntegration" in agents_page_source
    assert "bindingActionHint" in agents_page_source
    assert "AgentConnectionPlanPanel" in agents_page_source
    assert "План подключений" in agents_page_source
    assert "applyPostConnectHandoff" in agents_page_source
    assert "response.data?.post_connect_handoff" in agents_page_source
    assert "handoff.next_binding_key" in agents_page_source
    assert "setSelectedConnectionBindingKey(handoff.next_binding_key)" in agents_page_source
    assert "handoff.workspace_mode === 'run'" in agents_page_source
    assert "setSelectedConnectionBindingKey('')" in agents_page_source
    assert "next_binding" in builder_api_source
    assert "next_route" in builder_api_source
    assert "_preferred_handoff_route" in builder_api_source
    assert "Следующий доступ" in agents_page_source
    assert "recentPostCreateHandoff.next_binding" in agents_page_source
    assert "recentPostCreateHandoff.next_route" in agents_page_source
    assert "connection_plan" in api_source
    assert "localos_agent_connection_plan_v1" in api_source
    assert "agent_binding_integrations" in api_source
    assert "binding_key" in api_source
    assert "_build_agent_post_connect_handoff" in api_source
    assert "_connection_plan_next_binding_key" in api_source
    assert "_connection_plan_route_summary" in api_source
    assert "connector_provider_routes" in api_source
    assert "post_connect_handoff" in api_source
    assert "ready_for_preview" in api_source
    assert "connect_mode" in Path("src/services/agent_provider_registry.py").read_text(encoding="utf-8")
    assert "primary_cta" in Path("src/services/agent_provider_registry.py").read_text(encoding="utf-8")
    assert "provider_action" in Path("src/services/agent_provider_registry.py").read_text(encoding="utf-8")
    assert "use_openclaw_boundary" in Path("src/services/agent_provider_registry.py").read_text(encoding="utf-8")
    assert "select_external_account_key" in Path("src/services/agent_provider_registry.py").read_text(encoding="utf-8")
    assert "planned_oauth_connector" in Path("src/services/agent_provider_registry.py").read_text(encoding="utf-8")
    assert "ProviderActionPill" in agents_page_source
    assert "providerActionLabel" in agents_page_source
    assert "providerActionDescription" in agents_page_source
    assert "['openclaw', 'maton', 'manual'].includes(route.provider || '')" in agents_page_source
    assert "agentPolicyFacts(item)" in agents_page_source
    assert "item.policy_summary" in agents_page_source
    assert "RecommendedProviderRouteNote" in agents_page_source
    assert "recommended_route_reason" in agents_page_source
    assert "selected_provider_routes: acceptedBuilderProviderRoutes ? selectedBuilderProviderRoutes : {}" in agents_page_source
    assert "accepted_provider_routes: acceptedBuilderProviderRoutes" in agents_page_source
    assert "acceptedBuilderProviderRoutes" in agents_page_source
    assert "builderRequiredProviderRouteKeys" in agents_page_source
    assert "Подтвердить подключения" in agents_page_source
    assert "Подключения подтверждены" in agents_page_source
    assert "Что нужно агенту для работы" in agents_page_source
    assert "Использовать этот способ" in agents_page_source
    assert "Способ выбран" in agents_page_source
    assert "autoSelectBuilderProviderRoutes" in agents_page_source
    assert "builderConnectionStatusCopy" in agents_page_source
    assert "builderConnectionNextStepCopy" in agents_page_source
    assert "AGENT_PROVIDER_ROUTE_REQUIRED" in builder_api_source
    assert "AGENT_PROVIDER_ROUTES_CONFIRMATION_REQUIRED" in builder_api_source
    assert "builder_provider_routes_accepted" in builder_api_source
    assert "chooseProviderRoute" in agents_page_source
    assert "/provider-routes" in agents_page_source
    assert "onChooseProviderRoute" in agents_page_source
    assert "Выберите сохранённый Maton.ai key для этого шага." in agents_page_source
    assert "route?.primary_cta" in agents_page_source
    assert "saveMatonIntegration" in agents_page_source
    assert "Maton.ai bridge" in agents_page_source
    assert "matonAuthRef" in agents_page_source
    assert "onSaveMatonIntegration" in agents_page_source
    assert "AgentConnectionDecisionBanner" in agents_page_source
    assert "buildAgentConnectionDecision" in agents_page_source
    assert "BuilderCreationDecisionBanner" in agents_page_source
    assert "buildBuilderCreationDecision" in agents_page_source
    assert "BuilderCompilerPolicyReviewPanel" in agents_page_source
    assert "compiler_policy_review" in agents_page_source
    assert "compiler_workflow_draft" in agents_page_source
    assert "compiler_approval_points" in agents_page_source
    assert "compiler_unsupported_requests" in agents_page_source
    assert "accepted_compiler_plan: acceptedBuilderCompilerPlan" in agents_page_source
    assert "builderCompilerPlanRequiresConfirmation" not in agents_page_source
    assert "const canCreateDraft = setupFlowAllowsDraft" in agents_page_source
    assert "acceptedBuilderCompilerPlan" in agents_page_source
    assert "Принять план" in agents_page_source
    assert "План принят" in agents_page_source
    assert "План агента" in agents_page_source
    assert "compiled workflow candidate" in agents_page_source
    assert "Что нужно изменить в логике" in agents_page_source
    assert "Ответьте на уточнение" in agents_page_source
    assert "connection_resolver" in agents_page_source
    assert "подключение" in agents_page_source
    assert "Создать агента и открыть preview" in agents_page_source
    assert "У бизнеса уже есть несколько подходящих коннектов" in agents_page_source
    assert "Подключения готовы" in agents_page_source
    assert "Настройте один следующий доступ" in agents_page_source
    assert "Сохранить и перейти к тесту" in agents_page_source
    assert "LocalOS покажет следующий шаг" in agents_page_source
    assert "GenericRunProgress" in agents_page_source
    assert 'title="Агенты"' in agents_page_source
    assert "getAgentListStatus" in agents_page_source
    assert "AgentSummaryPill" in agents_page_source
    assert "Последний run" in agents_page_source
    assert "решений" in agents_page_source
    assert "Данные агента" in agents_page_source
    assert "Что будет делать агент" in agents_page_source
    assert "Голос и стиль" in agents_page_source
    assert "AgentVoiceStylePanel" in agents_page_source
    assert "AIAgents показываются как голоса" in agents_page_source
    assert "Путь {humanizeCategory(category).toLowerCase()}-агента" in agents_page_source
    assert "Технический журнал" in agents_page_source
    assert "AgentRunObservabilityPanel" in agents_page_source
    assert "PreviewRunSummaryPanel" in agents_page_source
    assert "CompiledPreviewSimulationPanel" in agents_page_source
    assert "previewSimulationTone" in agents_page_source
    assert "Симуляция compiled workflow" in agents_page_source
    assert "внешних действий не было" in agents_page_source
    assert "ActivationGateDecisionCard" in agents_page_source
    assert "buildActivationGateDecision" in agents_page_source
    assert "AgentActivationPathStrip" in agents_page_source
    assert "buildActivationPathSteps" in agents_page_source
    assert "workflow проверен" in agents_page_source
    assert "нужно подключить" in agents_page_source
    assert "activationBlockerText" in agents_page_source
    assert "Ждёт решения человека" in agents_page_source
    assert "AgentFourAnswerStrip" in agents_page_source
    assert "Что делает" in agents_page_source
    assert "Готов ли" in agents_page_source
    assert "Чего не хватает" in agents_page_source
    assert "Последний run" in agents_page_source
    assert "Единый billing ledger" in agents_page_source
    assert "Оценка до запуска" in agents_page_source
    assert "Факт после запуска" in agents_page_source
    assert "unified_billing_ledger" in agents_page_source
    assert "без решения человека агент не продолжит внешний шаг" in agents_page_source
    assert "Почему ждём" in agents_page_source
    assert "Preview: {gate.preview_run_status?.ready ? 'пройден' : 'нужен'}" in agents_page_source
    assert "Preflight: {gate.preflight?.ready ? 'готов' : 'проверить'}" in agents_page_source
    assert "Compiled: {gate.compiled_validation?.ready ? 'валиден' : 'проверить'}" in agents_page_source
    assert "Policy: {gate.approval_policy_status?.ready ? 'готова' : 'проверить'}" in agents_page_source
    assert "approvals и limits готовы" in agents_page_source
    assert "нужен human gate" in agents_page_source
    assert "Activation gate" in agents_page_source
    assert "Что показал preview run" in agents_page_source
    assert "activation_hint" in agents_page_source
    assert "Следующий шаг" in agents_page_source
    assert "next_step_label" in agents_page_source
    assert "previewNextStepActionLabel" in agents_page_source
    assert "onNextStepAction" in agents_page_source
    assert "canActivateFromPreview" in agents_page_source
    assert "Activation gate готов" in agents_page_source
    assert "onActivateVersion(activationVersionId)" in agents_page_source
    assert "Открыть подключения" in agents_page_source
    assert "Проверить активацию" in agents_page_source
    assert "Action ledger" in agents_page_source
    assert "localos_agent_preview_summary_v1" in Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "_build_preview_summary" in Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "_preview_next_step" in Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "domain_requests" in Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "agent_sheet_operation_requests" in Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "agent_communication_requests" in Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "provider_handoff" in Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "provider_handoff" in agents_page_source
    assert "Approvals" in agents_page_source
    assert "why_waiting" in agents_page_source
    assert "agent_review_publish_requests" in Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "publish_requests" in agents_page_source
    assert "Support export" in agents_page_source
    assert "agent_run_observability_v1" in Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "billing_ledger" in Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "BillingActionItem" in agents_page_source
    assert "reserve ${item.billing_summary?.reserved_tokens" in agents_page_source
    assert "Использовано в последнем запуске" in agents_page_source
    assert "used_sources" in workspace_source
    assert "resultFieldLabels" in agents_page_source
    assert "Bookings" in capability_handlers_source
    assert "agent_communication_requests" in capability_handlers_source
    assert "reviewreplydrafts" in capability_handlers_source
    assert "agent_service_optimization_requests" in capability_handlers_source
    assert "agent_sheet_operation_requests" in capability_handlers_source
    assert "provider_write_performed=False" in capability_handlers_source
    assert "sheets.append_row_request" in action_policy_source
    assert "dispatch_telegram_message_to_agent_blueprints" in telegram_webhook_source
    assert "agent_trigger_events" in trigger_runtime_source
    assert "telegram.message.received" in trigger_runtime_source
    assert "dispatch_scheduled_agent_blueprints" in trigger_runtime_source
    assert "dispatch_due_scheduled_agent_blueprints" in trigger_runtime_source
    assert "schedule.daily" in trigger_runtime_source
    assert "scheduler" in trigger_runtime_source
    assert "AGENT_SCHEDULE_DISPATCH_ENABLED" in worker_source
    assert "_dispatch_agent_schedules_if_due" in worker_source
    assert "AGENT_SCHEDULE_DISPATCH_ENABLED: ${AGENT_SCHEDULE_DISPATCH_ENABLED:-false}" in compose_source
    assert "manual_publish_required=True" in capability_handlers_source
    assert "manual_apply_required=True" in capability_handlers_source
    assert "reserve_paid_action_credits" in capability_handlers_source
    assert "finalize_reserved_action_credits" in capability_handlers_source


def test_generic_document_runner_uses_sources_and_completes_internal_result():
    from services.agent_blueprint_draft_builder import build_agent_blueprint_draft
    from services.agent_blueprint_runner import AgentBlueprintRunner

    cursor = FakeCursor()
    draft = build_agent_blueprint_draft("Обработай договор и найди риски")
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Document agent",
        "category": "documents",
        "metadata_json": {
            **draft["metadata"],
            "agent_setup": {
                "workflow_description": "Проверить договор",
                "extraction_rules": "Найти сроки, оплату и ответственность",
                "processing_rules": "Не придумывать факты",
                "output_format": "Список рисков",
                "approval_boundaries": ["final_output", "external_delivery"],
            },
            "agent_sources": [
                {
                    "id": "src1",
                    "source_type": "text",
                    "name": "Договор",
                    "content_text": "Оплата 10000 рублей. Ответственность за просрочку: штраф 10%.",
                    "content_length": 68,
                    "extraction_state": "ready",
                }
            ],
        },
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "steps_json": draft["version_payload"]["steps"],
        "capability_allowlist_json": [],
    }

    result = AgentBlueprintRunner(cursor).start_run("ver1", {}, {"user_id": "user1"})

    assert result["success"] is True
    run = result["run"]
    assert run["status"] == "completed"
    assert [step["step_key"] for step in run["steps"]] == [
        "collect_inputs",
        "extract_context",
        "prepare_output",
        "approve_output",
        "save_result",
    ]
    output = [item for item in run["artifacts"] if item["artifact_type"] == "agent_output_draft"][0]
    assert output["payload_json"]["external_dispatch_performed"] is False
    assert output["payload_json"]["result"]["title"] == "Разбор документа"
    assert output["payload_json"]["result"]["facts"]
    assert output["payload_json"]["result"]["fields"]["Оплата"]
    assert output["payload_json"]["result"]["fields"]["Ответственность"]
    assert output["payload_json"]["dispatch_state"] == "not_dispatched"
    assert run["approvals"] == []


def test_generic_email_runner_prepares_draft_and_never_dispatches():
    from services.agent_blueprint_draft_builder import build_agent_blueprint_draft
    from services.agent_blueprint_runner import AgentBlueprintRunner

    cursor = FakeCursor()
    draft = build_agent_blueprint_draft("Подготовь письмо клиенту по контексту", "email")
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Email agent",
        "category": "email",
        "metadata_json": {
            **draft["metadata"],
            "agent_setup": {
                "workflow_description": "Подготовить письмо клиенту о новой услуге",
                "extraction_rules": "Взять услугу, выгоду и ограничение по тону",
                "processing_rules": "Писать дружелюбно, не обещать скидку без подтверждения",
                "output_format": "subject, body, checklist, missing_info",
                "approval_boundaries": ["final_output", "external_delivery"],
            },
            "agent_sources": [
                {
                    "id": "src1",
                    "source_type": "text",
                    "name": "Контекст письма",
                    "content_text": "Адресат: Анна. Услуга: уход для волос. Цель: пригласить на консультацию.",
                    "content_length": 77,
                    "extraction_state": "ready",
                }
            ],
        },
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "steps_json": draft["version_payload"]["steps"],
        "capability_allowlist_json": [],
    }

    result = AgentBlueprintRunner(cursor).start_run("ver1", {}, {"user_id": "user1"})

    assert result["success"] is True
    run = result["run"]
    assert run["status"] == "completed"
    assert not [step for step in run["steps"] if step["step_type"] == "capability"]
    output = [item for item in run["artifacts"] if item["artifact_type"] == "agent_output_draft"][0]
    payload = output["payload_json"]
    email_result = payload["result"]
    assert payload["category"] == "email"
    assert payload["external_dispatch_performed"] is False
    assert payload["dispatch_state"] == "not_dispatched"
    assert email_result["external_dispatch_performed"] is False
    assert email_result["delivery_state"] == "not_dispatched"
    assert email_result["subject"]
    assert email_result["body"]
    assert email_result["checklist"]
    assert email_result["provenance"] == ["Контекст письма"]
    assert run["approvals"] == []


def test_generic_table_runner_prepares_report_and_never_dispatches():
    from services.agent_blueprint_draft_builder import build_agent_blueprint_draft
    from services.agent_blueprint_runner import AgentBlueprintRunner

    cursor = FakeCursor()
    draft = build_agent_blueprint_draft("Проверь CSV и найди ошибки", "tables")
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Table agent",
        "category": "tables",
        "metadata_json": {
            **draft["metadata"],
            "agent_setup": {
                "workflow_description": "Проверить таблицу клиентов",
                "extraction_rules": "Найти пустые email, дубли и строки к проверке",
                "processing_rules": "Не изменять таблицу, только показать проблемы",
                "output_format": "summary, exceptions, rows_to_review, recommendations",
                "approval_boundaries": ["final_output", "external_delivery"],
            },
            "agent_sources": [
                {
                    "id": "src1",
                    "source_type": "text",
                    "name": "clients.csv",
                    "content_text": "name,email,phone\nАнна,anna@example.com,+1\nАнна,anna@example.com,+1\nБорис,,+2\n",
                    "content_length": 83,
                    "extraction_state": "ready",
                }
            ],
        },
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "steps_json": draft["version_payload"]["steps"],
        "capability_allowlist_json": [],
    }

    result = AgentBlueprintRunner(cursor).start_run("ver1", {}, {"user_id": "user1"})

    assert result["success"] is True
    run = result["run"]
    assert run["status"] == "completed"
    assert not [step for step in run["steps"] if step["step_type"] == "capability"]
    output = [item for item in run["artifacts"] if item["artifact_type"] == "agent_output_draft"][0]
    payload = output["payload_json"]
    table_result = payload["result"]
    assert payload["category"] == "tables"
    assert payload["external_dispatch_performed"] is False
    assert payload["dispatch_state"] == "not_dispatched"
    assert table_result["external_dispatch_performed"] is False
    assert table_result["delivery_state"] == "not_dispatched"
    assert table_result["summary"]
    assert table_result["exceptions"]
    assert table_result["rows_to_review"]
    assert table_result["recommendations"]
    assert table_result["provenance"] == ["clients.csv"]
    assert run["approvals"] == []


def test_generic_reviews_runner_prepares_reply_drafts_and_never_publishes():
    from services.agent_blueprint_draft_builder import build_agent_blueprint_draft
    from services.agent_blueprint_runner import AgentBlueprintRunner

    cursor = FakeCursor()
    draft = build_agent_blueprint_draft("Подготовь ответы на отзывы", "reviews")
    reviews_text = (
        "author_name,rating,text\n"
        "Анна,5,Очень понравился сервис\n"
        "Иван,2,Долго ждал и администратор был груб\n"
    )
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Reviews agent",
        "category": "reviews",
        "metadata_json": {
            **draft["metadata"],
            "agent_setup": {
                "workflow_description": "Подготовить ответы на отзывы",
                "extraction_rules": "Определить тон, проблему клиента и безопасный ответ",
                "processing_rules": "Не обещать скидку, компенсацию или публикацию без подтверждения",
                "output_format": "reply_drafts, manual_review_reasons, checklist",
                "approval_boundaries": ["final_output", "external_delivery"],
            },
            "agent_sources": [
                {
                    "id": "src1",
                    "source_type": "text",
                    "name": "Отзывы",
                    "content_text": reviews_text,
                    "content_length": len(reviews_text),
                    "extraction_state": "ready",
                }
            ],
        },
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "steps_json": draft["version_payload"]["steps"],
        "capability_allowlist_json": [],
    }

    result = AgentBlueprintRunner(cursor).start_run("ver1", {}, {"user_id": "user1"})

    assert result["success"] is True
    run = result["run"]
    assert run["status"] == "completed"
    assert not [step for step in run["steps"] if step["step_type"] == "capability"]
    output = [item for item in run["artifacts"] if item["artifact_type"] == "agent_output_draft"][0]
    payload = output["payload_json"]
    review_result = payload["result"]
    assert payload["category"] == "reviews"
    assert payload["external_dispatch_performed"] is False
    assert payload["dispatch_state"] == "not_dispatched"
    assert review_result["external_dispatch_performed"] is False
    assert review_result["publish_state"] == "not_published"
    assert review_result["delivery_state"] == "not_dispatched"
    assert review_result["reply_drafts"]
    assert review_result["manual_review_reasons"]
    assert review_result["checklist"]
    assert review_result["provenance"] == ["Отзывы"]
    assert run["approvals"] == []


def test_message_result_needs_source_data_without_sheet_rows():
    from services.agent_blueprint_workspace import _render_output

    result = _render_output(
        "custom",
        {
            "workflow_description": "Открой таблицу поездок и подготовь сообщение владельцу",
            "processing_rules": "Не придумывать факты",
            "output_format": "Готовое сообщение для проверки",
        },
        [{"source_name": "business_profile", "summary": "ready; id: biz-1", "raw": {"id": "biz-1"}}],
        [],
        {},
    )

    assert result["status"] == "needs_source_data"
    assert "draft_text" not in result
    assert "не получил строку поездки" in result["summary"][0].lower()


def test_message_result_prompts_google_reconnect_when_sheet_auth_is_revoked():
    from services.agent_blueprint_workspace import _render_output

    result = _render_output(
        "custom",
        {
            "workflow_description": "Открой таблицу поездок и подготовь сообщение владельцу",
            "processing_rules": "Не придумывать факты",
            "output_format": "Готовое сообщение для проверки",
        },
        [
            {
                "source_name": "google_sheets_error",
                "summary": "Google token refresh failed with HTTP 400: invalid_grant",
                "raw": {
                    "provider_error": "GOOGLE_SHEETS_PROVIDER_NOT_READY",
                    "provider_error_message": "Google token refresh failed with HTTP 400: invalid_grant",
                    "next_action": "connect_or_repair_google_sheets_provider",
                },
            }
        ],
        [],
        {},
    )

    assert result["status"] == "needs_google_access"
    assert result["title"] == "Нужно переподключить Google-доступ"
    assert "Таблица выбрана" in result["summary"][0]
    assert "Переподключите Google-доступ" in result["next_questions"][0]
    assert "invalid_grant" in result["technical_reason"]


def test_message_result_for_disabled_google_sheets_api_guides_project_setup():
    from services.agent_blueprint_workspace import _render_output

    result = _render_output(
        "custom",
        {
            "workflow_description": "Открой таблицу поездок и подготовь сообщение владельцу",
            "processing_rules": "Не придумывать факты",
            "output_format": "Готовое сообщение для проверки",
        },
        [
            {
                "source_name": "google_sheets_error",
                "summary": "Google Sheets read failed with HTTP 403",
                "raw": {
                    "provider_error": "GOOGLE_SHEETS_PROVIDER_NOT_READY",
                    "provider_error_message": (
                        "Google Sheets read failed with HTTP 403: "
                        "Google Sheets API has not been used in project 304042072643 before or it is disabled. "
                        "Enable it by visiting https://console.developers.google.com/apis/api/sheets.googleapis.com/overview?project=304042072643"
                    ),
                },
            }
        ],
        [],
        {},
    )

    assert result["status"] == "needs_google_api_enabled"
    assert result["title"] == "Нужно включить Google Sheets API"
    assert "Google-доступ подключён" in result["summary"][0]
    assert "304042072643" in result["next_questions"][0]
    assert "Переподключите Google-доступ" not in " ".join(result["next_questions"])


def test_message_result_for_google_sheets_invalid_range_guides_sheet_tab_setup():
    from services.agent_blueprint_workspace import _render_output

    result = _render_output(
        "custom",
        {
            "workflow_description": "Открой таблицу поездок и подготовь сообщение владельцу",
            "processing_rules": "Не придумывать факты",
            "output_format": "Готовое сообщение для проверки",
        },
        [
            {
                "source_name": "google_sheets_error",
                "summary": "Google Sheets read failed with HTTP 400",
                "raw": {
                    "provider_error": "GOOGLE_SHEETS_PROVIDER_NOT_READY",
                    "provider_error_message": (
                        "Google Sheets read failed with HTTP 400: "
                        "Unable to parse range: Sheet1!A1:Z"
                    ),
                },
            }
        ],
        [],
        {},
    )

    assert result["status"] == "needs_sheet_tab"
    assert result["title"] == "Нужно выбрать лист таблицы"
    assert "лист таблицы" in result["summary"][0]
    assert "Переподключите Google-доступ" not in " ".join(result["next_questions"])


def test_message_result_uses_google_sheets_rows_for_concrete_draft(monkeypatch):
    import services.agent_blueprint_workspace as workspace

    monkeypatch.setattr(workspace, "analyze_text_with_gigachat", lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("llm offline")))

    result = workspace._render_output(
        "custom",
        {
            "workflow_description": "Выбери поездку на 20 апреля и подготовь сообщение владельцу",
            "processing_rules": "Не придумывать факты",
            "output_format": "Готовое сообщение для проверки",
        },
        [
            {
                "source_name": "google_sheets",
                "summary": "date: 2026-04-20; route: Airport -> Center; passenger: Anna",
                "raw": {
                    "date": "2026-04-20",
                    "route": "Airport -> Center",
                    "passenger": "Anna",
                },
            }
        ],
        [],
        {},
    )

    assert result["title"] == "Черновик сообщения"
    assert "Airport -> Center" in result["draft_text"]
    assert result["analysis_source"] == "deterministic_fallback"


def test_runner_propagates_google_sheets_run_rows_into_output_artifact(monkeypatch):
    import services.agent_blueprint_workspace as workspace
    from services.agent_blueprint_runner import AgentBlueprintRunner

    monkeypatch.setattr(workspace, "analyze_text_with_gigachat", lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("llm offline")))

    cursor = FakeCursor()
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Trips agent",
        "category": "custom",
        "metadata_json": {
            "agent_setup": {
                "workflow_description": "Выбери поездку на 20 апреля и подготовь сообщение владельцу",
                "processing_rules": "Не придумывать факты",
                "output_format": "Готовое сообщение для проверки",
            },
            "agent_sources": [
                {
                    "id": "business-profile",
                    "source_type": "internal",
                    "internal_source": "business_profile",
                    "name": "Профиль бизнеса",
                }
            ],
        },
    }
    cursor.tables["agent_runs"]["run1"] = {
        "id": "run1",
        "blueprint_id": "bp1",
        "blueprint_version_id": "ver1",
        "business_id": "biz1",
        "status": "running",
        "input_json": {},
        "output_json": {},
        "created_by_user_id": "user1",
    }
    cursor.tables["agent_run_steps"]["read-step"] = {
        "id": "read-step",
        "run_id": "run1",
        "step_index": 0,
        "step_key": "read_google_sheets",
        "step_type": "capability",
        "status": "completed",
        "input_json": {},
        "output_json": {
            "capability": "google_sheets.read_rows",
            "orchestrator": {
                "success": True,
                "result": {
                    "source": "google_sheets",
                    "provider_read_performed": True,
                    "rows": [
                        {
                            "date": "2026-04-20",
                            "route": "Tallinn Airport -> Old Town",
                            "passenger": "Anna",
                        }
                    ],
                },
            },
        },
    }

    payload = AgentBlueprintRunner(cursor)._build_artifact_payload(
        cursor.tables["agent_runs"]["run1"],
        {
            "key": "prepare_output",
            "artifact_type": "agent_output_draft",
            "payload": {
                "category": "custom",
                "rows_from_step": "read_google_sheets",
                "format": "Готовое сообщение для проверки",
            },
        },
    )

    assert payload["items_used"] >= 1
    assert payload["result"]["title"] == "Черновик сообщения"
    assert payload["result"].get("status") != "needs_source_data"
    assert "Tallinn Airport -> Old Town" in payload["result"]["draft_text"]
    assert "business_profile" not in payload["result"]["draft_text"]
    assert payload["external_dispatch_performed"] is False


def test_agents_page_normal_result_panel_does_not_dump_raw_artifact_payload():
    source = read_agent_blueprints_frontend_source()

    assert "stringifyBusinessValue(artifact.payload_json)" not in source
    assert "Результат не был сохранён. Запустите тест ещё раз" in source
    assert "const isBlocked = result.state === 'blocker';" in source
    assert "const canApprove = Boolean(pendingApproval && !isBlocked);" in source
    assert "needsScenarioRebuildForSourceResult" in source
    assert "Пересобрать сценарий" in source
    assert "Этот агент создан старой версией сценария" in source
    assert "needsGoogleSheetsSourceSetup" in source
    assert "Указать Google-таблицу" in source
    assert "Укажите Google-таблицу и лист со списком поездок" in source
    assert "needsGoogleAccessReconnect" in source
    assert "focus: 'google_sheets'" in source
    assert "return_to: '/dashboard/agents'" in source
    assert "Переподключить Google-доступ" in source
    assert "google_auth" in source
    assert "googleAccessJustConnected" in source
    assert "Google-доступ подключён. Теперь запустите тест ещё раз" in source
    assert "Этот результат был получен до переподключения" in source
    assert "blocked_result" in source
    assert "hasFreshGoogleSheetsAccessAfterResult" in source
    assert "Google-доступ обновлён. Запустите тест ещё раз" in source
    assert "Почему нельзя подтвердить результат" in source


def test_settings_integrations_first_layer_separates_google_sheets_from_google_business():
    external = Path("frontend/src/components/ExternalIntegrations.tsx").read_text(encoding="utf-8")
    hub_state = Path("frontend/src/pages/dashboard/settings/settingsHubState.ts").read_text(encoding="utf-8")
    hub_copy = Path("frontend/src/pages/dashboard/settings/settingsHubCopy.ts").read_text(encoding="utf-8")

    assert 'data-testid="settings-integrations-scenario"' in external
    assert "Google Таблицы" in external
    assert "Этот доступ нужен агентам для чтения Google Таблиц. Он не публикует ничего наружу." in external
    assert "Google Документы: позже" in external
    assert "Google-доступ" in external
    assert "Карточка" in external
    assert "Таблицы" in external
    assert "google_sheets" in hub_state
    assert "/dashboard/settings/integrations?focus=google_sheets" in hub_state
    assert "Agent access to table rows." in hub_copy
    assert "Доступ агентов к строкам таблиц." in hub_copy


def test_agent_source_ingestion_extracts_text_pdf_docx_xlsx_and_rejects_unsafe_files():
    import io
    import zipfile

    from openpyxl import Workbook

    from services.agent_source_ingestion import build_agent_source_from_upload

    def build_test_pdf_bytes(text):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("utf-8")
        objects = [
            b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
            b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n",
            (
                b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >> endobj\n"
            ),
            b"4 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n",
            b"5 0 obj << /Length "
            + str(len(stream)).encode("ascii")
            + b" >> stream\n"
            + stream
            + b"\nendstream endobj\n",
        ]
        output = bytearray(b"%PDF-1.4\n")
        offsets = [0]
        for item in objects:
            offsets.append(len(output))
            output.extend(item)
        xref_offset = len(output)
        output.extend(b"xref\n0 6\n0000000000 65535 f \n")
        for offset in offsets[1:]:
            output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
        output.extend(b"trailer << /Root 1 0 R /Size 6 >>\n")
        output.extend(f"startxref\n{xref_offset}\n%%EOF\n".encode("ascii"))
        return bytes(output)

    text_source, text_error = build_agent_source_from_upload(
        FakeUpload("contract.txt", "text/plain", "Оплата 10000. Ответственность: штраф.".encode("utf-8")),
        "Договор",
    )
    assert text_error == {}
    assert text_source["content_text"].startswith("Оплата 10000")
    assert text_source["extraction_state"] == "ready"

    pdf_bytes = build_test_pdf_bytes("Payment 15000. Penalty 12 percent.")
    pdf_source, pdf_error = build_agent_source_from_upload(
        FakeUpload("contract.pdf", "application/pdf", pdf_bytes),
    )
    assert pdf_error == {}
    assert "Payment 15000" in pdf_source["content_text"]
    assert pdf_source["extraction_method"] == "pypdf"

    docx_buffer = io.BytesIO()
    archive = zipfile.ZipFile(docx_buffer, "w")
    try:
        archive.writestr(
            "word/document.xml",
            (
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                "<w:body><w:p><w:r><w:t>Документ содержит срок и оплату.</w:t></w:r></w:p></w:body>"
                "</w:document>"
            ),
        )
    finally:
        archive.close()
    docx_source, docx_error = build_agent_source_from_upload(
        FakeUpload(
            "contract.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            docx_buffer.getvalue(),
        )
    )
    assert docx_error == {}
    assert "срок и оплату" in docx_source["content_text"]
    assert docx_source["extraction_method"] == "docx_xml"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Риски"
    sheet.append(["Поле", "Значение"])
    sheet.append(["Штраф", "10%"])
    xlsx_buffer = io.BytesIO()
    workbook.save(xlsx_buffer)
    xlsx_source, xlsx_error = build_agent_source_from_upload(
        FakeUpload(
            "risks.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            xlsx_buffer.getvalue(),
        )
    )
    assert xlsx_error == {}
    assert "Штраф" in xlsx_source["content_text"]
    assert xlsx_source["extraction_method"] == "openpyxl"

    unsafe_source, unsafe_error = build_agent_source_from_upload(
        FakeUpload("payload.exe", "application/octet-stream", b"bad"),
    )
    assert unsafe_source == {}
    assert unsafe_error["code"] == "UNSUPPORTED_FILE_TYPE"
    assert "поддерживается" in unsafe_error["message"].lower()

    empty_source, empty_error = build_agent_source_from_upload(
        FakeUpload("empty.txt", "text/plain", b""),
    )
    assert empty_source == {}
    assert empty_error["code"] == "EMPTY_FILE"
    assert "пустой" in empty_error["message"].lower()


def test_agent_datahub_catalog_returns_available_internal_sources():
    from services.agent_datahub import build_agent_datahub_catalog

    cursor = FakeDatahubCursor()
    catalog = build_agent_datahub_catalog(
        cursor,
        "biz1",
        [{"source_type": "internal", "internal_source": "services"}],
    )

    by_key = {item["key"]: item for item in catalog}
    assert by_key["business_profile"]["available_count"] == 1
    assert by_key["services"]["available_count"] == 2
    assert by_key["services"]["connected"] is True
    assert by_key["reviews"]["preview"]
    assert by_key["prospectingleads"]["state"] == "empty"


def test_agent_document_llm_analysis_uses_generator_rules_and_provenance():
    from services.agent_document_llm import analyze_document_sources_with_llm

    captured = {}

    def fake_generator(prompt, *, business_id="", user_id=""):
        captured["prompt"] = prompt
        captured["business_id"] = business_id
        captured["user_id"] = user_id
        return json.dumps(
            {
                "title": "LLM contract analysis",
                "summary": ["Оплата 10000 рублей"],
                "risks": ["Штраф 10% за просрочку"],
                "facts": ["Срок 30 дней", "Оплата 10000 рублей"],
                "fields": {"Оплата": "10000 рублей", "Срок": "30 дней"},
                "next_questions": ["Кто подписывает договор?"],
                "rules_applied": ["Не придумывать факты"],
            },
            ensure_ascii=False,
        )

    result = analyze_document_sources_with_llm(
        {
            "workflow_description": "Проверить договор",
            "extraction_rules": "Суммы, сроки, штрафы",
            "processing_rules": "Не придумывать факты",
            "output_format": "Краткий отчёт",
        },
        [
            {
                "source_name": "contract.txt",
                "summary": "Оплата 10000 рублей. Срок 30 дней. Штраф 10%.",
                "raw": {"text": "Оплата 10000 рублей. Срок 30 дней. Штраф 10%."},
            }
        ],
        business_id="biz1",
        user_id="user1",
        generator=fake_generator,
    )

    assert result["llm_analysis_used"] is True
    assert result["analysis_source"] == "gigachat"
    assert result["external_dispatch_performed"] is False
    assert result["fields"]["Оплата"] == "10000 рублей"
    assert result["provenance"] == ["contract.txt"]
    assert captured["business_id"] == "biz1"
    assert captured["user_id"] == "user1"
    assert "Не придумывать факты" in captured["prompt"]
    assert "contract.txt" in captured["prompt"]


def test_agent_document_llm_analysis_falls_back_without_external_dispatch():
    from services.agent_document_llm import analyze_document_sources_with_llm

    def failing_generator(prompt, *, business_id="", user_id=""):
        raise RuntimeError("provider unavailable")

    result = analyze_document_sources_with_llm(
        {"processing_rules": "Показывать риски", "output_format": "Отчёт"},
        [{"source_name": "contract.txt", "summary": "Оплата 10000. Ответственность: штраф.", "raw": {}}],
        generator=failing_generator,
    )

    assert result["llm_analysis_used"] is False
    assert result["analysis_source"] == "deterministic_fallback"
    assert result["external_dispatch_performed"] is False
    assert result["provenance"] == ["contract.txt"]
    assert result["risks"]


def test_agent_email_llm_draft_uses_generator_rules_and_provenance():
    from services.agent_email_llm import draft_email_with_llm

    captured = {}

    def fake_generator(prompt, *, business_id="", user_id=""):
        captured["prompt"] = prompt
        captured["business_id"] = business_id
        captured["user_id"] = user_id
        return json.dumps(
            {
                "title": "Email draft",
                "subject": "Анна, приглашаем на консультацию",
                "body": "Здравствуйте, Анна! Приглашаем на консультацию по уходу для волос.",
                "checklist": ["Проверить имя", "Проверить оффер"],
                "assumptions": ["Контекст взят из источника"],
                "missing_info": ["Дата консультации"],
                "rules_applied": ["Не обещать скидку"],
            },
            ensure_ascii=False,
        )

    result = draft_email_with_llm(
        {
            "workflow_description": "Подготовить письмо клиенту",
            "extraction_rules": "Адресат, услуга, цель",
            "processing_rules": "Не обещать скидку",
            "output_format": "subject/body/checklist",
        },
        [
            {
                "source_name": "Контекст",
                "summary": "Адресат: Анна. Услуга: уход для волос.",
                "raw": {"text": "Адресат: Анна. Услуга: уход для волос. Цель: консультация."},
            }
        ],
        business_id="biz1",
        user_id="user1",
        generator=fake_generator,
    )

    assert result["llm_analysis_used"] is True
    assert result["analysis_source"] == "gigachat"
    assert result["external_dispatch_performed"] is False
    assert result["delivery_state"] == "not_dispatched"
    assert result["subject"] == "Анна, приглашаем на консультацию"
    assert result["checklist"]
    assert result["provenance"] == ["Контекст"]
    assert captured["business_id"] == "biz1"
    assert captured["user_id"] == "user1"
    assert "Не обещать скидку" in captured["prompt"]


def test_agent_email_llm_falls_back_without_external_dispatch():
    from services.agent_email_llm import draft_email_with_llm

    def failing_generator(prompt, *, business_id="", user_id=""):
        raise RuntimeError("provider unavailable")

    result = draft_email_with_llm(
        {"workflow_description": "Подготовить письмо", "processing_rules": "Дружелюбно"},
        [{"source_name": "Контекст", "summary": "Адресат: Анна. Услуга: консультация.", "raw": {}}],
        generator=failing_generator,
    )

    assert result["llm_analysis_used"] is False
    assert result["analysis_source"] == "deterministic_fallback"
    assert result["external_dispatch_performed"] is False
    assert result["delivery_state"] == "not_dispatched"
    assert result["subject"]
    assert result["body"]


def test_agent_table_analysis_uses_generator_rules_and_provenance():
    from services.agent_table_analysis import analyze_table_with_llm

    captured = {}

    def fake_generator(prompt, *, business_id="", user_id=""):
        captured["prompt"] = prompt
        captured["business_id"] = business_id
        captured["user_id"] = user_id
        return json.dumps(
            {
                "title": "Table report",
                "summary": ["Проверено 3 строки"],
                "exceptions": ["Строка 3: пустой email"],
                "rows_to_review": [
                    {"row": 3, "reason": "пустой email", "source_name": "clients.csv", "values": {"name": "Борис"}}
                ],
                "recommendations": ["Заполнить email"],
                "rules_applied": ["Не изменять таблицу"],
            },
            ensure_ascii=False,
        )

    result = analyze_table_with_llm(
        {
            "workflow_description": "Проверить клиентов",
            "extraction_rules": "Пустые email и дубли",
            "processing_rules": "Не изменять таблицу",
            "output_format": "exceptions report",
        },
        [
            {"source_name": "clients.csv", "summary": "name: Борис; email: ; phone: +2", "raw": {"name": "Борис", "email": "", "phone": "+2"}}
        ],
        business_id="biz1",
        user_id="user1",
        generator=fake_generator,
    )

    assert result["llm_analysis_used"] is True
    assert result["analysis_source"] == "gigachat"
    assert result["external_dispatch_performed"] is False
    assert result["delivery_state"] == "not_dispatched"
    assert result["exceptions"] == ["Строка 3: пустой email"]
    assert result["rows_to_review"][0]["row"] == 3
    assert result["provenance"] == ["clients.csv"]
    assert captured["business_id"] == "biz1"
    assert captured["user_id"] == "user1"
    assert "Не изменять таблицу" in captured["prompt"]


def test_agent_table_analysis_falls_back_without_external_dispatch():
    from services.agent_table_analysis import analyze_table_with_llm

    def failing_generator(prompt, *, business_id="", user_id=""):
        raise RuntimeError("provider unavailable")

    result = analyze_table_with_llm(
        {"workflow_description": "Проверить таблицу", "processing_rules": "Только отчёт"},
        [
            {"source_name": "clients.csv", "summary": "name: Анна; email: anna@example.com", "raw": {"name": "Анна", "email": "anna@example.com"}},
            {"source_name": "clients.csv", "summary": "name: Борис; email: ", "raw": {"name": "Борис", "email": ""}},
        ],
        generator=failing_generator,
    )

    assert result["llm_analysis_used"] is False
    assert result["analysis_source"] == "deterministic_fallback"
    assert result["external_dispatch_performed"] is False
    assert result["delivery_state"] == "not_dispatched"
    assert result["exceptions"]
    assert result["rows_to_review"]


def test_agent_review_reply_analysis_uses_generator_rules_and_provenance():
    from services.agent_review_reply_analysis import draft_review_replies_with_llm

    captured = {}

    def fake_generator(prompt, *, business_id="", user_id=""):
        captured["prompt"] = prompt
        captured["business_id"] = business_id
        captured["user_id"] = user_id
        return json.dumps(
            {
                "title": "Review replies",
                "summary": ["Подготовлено 2 черновика"],
                "reply_drafts": [
                    {
                        "review_id": "rev1",
                        "author_name": "Иван",
                        "rating": "2",
                        "sentiment": "negative",
                        "reply": "Иван, спасибо за обратную связь. Мы разберём ситуацию с ожиданием.",
                        "manual_review_reason": "Негативный отзыв требует проверки менеджером.",
                    }
                ],
                "manual_review_reasons": ["Негативный отзыв требует проверки менеджером."],
                "checklist": ["Проверить тон", "Не обещать компенсацию"],
                "rules_applied": ["Не обещать скидку"],
            },
            ensure_ascii=False,
        )

    result = draft_review_replies_with_llm(
        {
            "workflow_description": "Подготовить ответы на отзывы",
            "extraction_rules": "Тональность и причина недовольства",
            "processing_rules": "Не обещать скидку",
            "output_format": "reply_drafts/checklist",
        },
        [
            {
                "source_name": "Отзывы",
                "summary": "Иван поставил 2: долго ждал",
                "raw": {"id": "rev1", "author_name": "Иван", "rating": 2, "text": "Долго ждал"},
            }
        ],
        business_id="biz1",
        user_id="user1",
        generator=fake_generator,
    )

    assert result["llm_analysis_used"] is True
    assert result["analysis_source"] == "gigachat"
    assert result["external_dispatch_performed"] is False
    assert result["publish_state"] == "not_published"
    assert result["delivery_state"] == "not_dispatched"
    assert result["reply_drafts"][0]["reply"]
    assert result["manual_review_reasons"]
    assert result["provenance"] == ["Отзывы"]
    assert captured["business_id"] == "biz1"
    assert captured["user_id"] == "user1"
    assert "Не обещать скидку" in captured["prompt"]


def test_agent_review_reply_analysis_falls_back_without_publish():
    from services.agent_review_reply_analysis import draft_review_replies_with_llm

    def failing_generator(prompt, *, business_id="", user_id=""):
        raise RuntimeError("provider unavailable")

    result = draft_review_replies_with_llm(
        {"workflow_description": "Ответить на отзывы", "processing_rules": "Проверить негатив вручную"},
        [
            {
                "source_name": "Отзывы",
                "summary": "Ольга поставила 2: плохо и долго",
                "raw": {"id": "rev2", "author_name": "Ольга", "rating": 2, "text": "Плохо и долго"},
            }
        ],
        generator=failing_generator,
    )

    assert result["llm_analysis_used"] is False
    assert result["analysis_source"] == "deterministic_fallback"
    assert result["external_dispatch_performed"] is False
    assert result["publish_state"] == "not_published"
    assert result["delivery_state"] == "not_dispatched"
    assert result["reply_drafts"]
    assert result["manual_review_reasons"]
    assert result["provenance"] == ["Отзывы"]


def test_agent_date_parameter_commits_native_picker_value_on_blur():
    employee_source = Path("frontend/src/pages/dashboard/agents/employee.tsx").read_text(encoding="utf-8")

    assert 'type={field.format === \'date\' ? \'date\'' in employee_source
    assert 'onInput={(event) => onChange(key, event.currentTarget.value)}' in employee_source
    assert 'onBlur={(event) => onChange(key, event.currentTarget.value)}' in employee_source
