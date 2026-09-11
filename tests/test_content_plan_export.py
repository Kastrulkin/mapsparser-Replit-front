from datetime import datetime, timezone
from openpyxl import load_workbook
from services.content_plan_export import export_rows, render_export, fingerprint


def sample():
    return {'title': 'План Органики', 'scope_target_label': 'Органика', 'items': [
        {'scheduled_for': datetime(2026, 9, 11, tzinfo=timezone.utc), 'theme': '=HYPERLINK("https://example.com")', 'goal': 'Описание идеи', 'draft_text': 'Готовый текст', 'location_label': 'Центр', 'status': 'draft_generated'},
        {'theme': 'Идея без текста', 'goal': 'Это не готовая публикация', 'status': 'planned'},
    ]}


def test_saved_plan_xlsx_has_all_rows_and_no_formulas():
    sheet = load_workbook(render_export(sample(), 'xlsx')).active
    assert sheet.max_row == 6
    assert sheet['D5'].data_type == 's'
    assert sheet['D5'].value.startswith('=HYPERLINK')
    assert sheet['F5'].value == 'Готовый текст'
    assert sheet['F6'].value is None
    assert sheet['G6'].value == 'Идея'
    assert sheet.freeze_panes == 'A5'
    assert sheet.auto_filter.ref == 'A4:H6'


def test_snapshot_hash_and_draft_status():
    plan = sample()
    original = fingerprint(plan)
    plan['items'][0]['draft_text'] = 'Изменённый текст'
    assert fingerprint(plan) != original
    assert export_rows(plan)[1][6] == 'Идея'


def test_very_long_text_is_not_truncated_by_excel():
    plan = sample()
    plan['items'] = [plan['items'][0]]
    plan['items'][0]['draft_text'] = 'а' * 45000
    sheet = load_workbook(render_export(plan, 'xlsx')).active
    assert sheet['F5'].value + sheet['F6'].value == 'а' * 45000
