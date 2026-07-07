import json
from datetime import datetime, timezone
from pathlib import Path


def test_agent_blueprint_routes_are_owned_by_blueprint():
    import main

    expected = {
        "/api/agent-builder/sessions": {
            "POST": "agent_builder_api.create_agent_builder_session",
        },
        "/api/agent-builder/sessions/<session_id>/message": {
            "POST": "agent_builder_api.add_agent_builder_message",
        },
        "/api/agent-builder/sessions/<session_id>/create-blueprint": {
            "POST": "agent_builder_api.create_blueprint_from_agent_builder_session",
        },
        "/api/agent-blueprints": {
            "GET": "agent_blueprints_api.list_agent_blueprints",
            "POST": "agent_blueprints_api.create_agent_blueprint",
        },
        "/api/agent-blueprints/draft": {
            "POST": "agent_blueprints_api.create_agent_blueprint_draft",
        },
        "/api/agent-blueprints/legacy-migration-plan": {
            "GET": "agent_blueprints_api.get_agent_blueprint_legacy_migration_plan",
        },
        "/api/agent-blueprints/legacy-migration/apply": {
            "POST": "agent_blueprints_api.apply_agent_blueprint_legacy_migration",
        },
        "/api/agent-blueprints/<blueprint_id>": {
            "GET": "agent_blueprints_api.get_agent_blueprint",
            "DELETE": "agent_blueprints_api.archive_agent_blueprint",
        },
        "/api/agent-blueprints/<blueprint_id>/versions": {
            "POST": "agent_blueprints_api.create_agent_blueprint_version",
        },
        "/api/agent-blueprints/<blueprint_id>/versions/<version_id>/diff": {
            "GET": "agent_blueprints_api.get_agent_blueprint_version_diff",
        },
        "/api/agent-blueprints/<blueprint_id>/versions/<version_id>/activate": {
            "POST": "agent_blueprints_api.activate_agent_blueprint_version",
        },
        "/api/agent-blueprints/<blueprint_id>/versions/<version_id>/rollback": {
            "POST": "agent_blueprints_api.rollback_agent_blueprint_version",
        },
        "/api/agent-blueprints/<blueprint_id>/setup": {
            "POST": "agent_blueprints_api.setup_agent_blueprint",
        },
        "/api/agent-blueprints/<blueprint_id>/integrations": {
            "GET": "agent_blueprints_api.list_agent_blueprint_integrations",
            "POST": "agent_blueprints_api.save_agent_blueprint_integration",
        },
        "/api/agent-blueprints/<blueprint_id>/custom-process": {
            "POST": "agent_blueprints_api.save_agent_blueprint_custom_process",
        },
        "/api/agent-blueprints/<blueprint_id>/custom-process/preview": {
            "POST": "agent_blueprints_api.preview_agent_blueprint_custom_process",
        },
        "/api/agent-blueprints/<blueprint_id>/sources": {
            "POST": "agent_blueprints_api.add_agent_blueprint_source",
        },
        "/api/agent-blueprints/<blueprint_id>/sources/catalog": {
            "GET": "agent_blueprints_api.list_agent_blueprint_source_catalog",
        },
        "/api/agent-blueprints/<blueprint_id>/sources/upload": {
            "POST": "agent_blueprints_api.upload_agent_blueprint_source",
        },
        "/api/agent-blueprints/<blueprint_id>/review": {
            "GET": "agent_blueprints_api.review_agent_blueprint",
        },
        "/api/agent-blueprints/<blueprint_id>/preflight": {
            "POST": "agent_blueprints_api.preflight_agent_blueprint_run",
        },
        "/api/agent-blueprints/<blueprint_id>/runs": {
            "POST": "agent_blueprints_api.start_agent_blueprint_run",
        },
        "/api/agent-runs/<run_id>": {
            "GET": "agent_blueprints_api.get_agent_run",
        },
        "/api/agent-runs/<run_id>/support-export": {
            "GET": "agent_blueprints_api.get_agent_run_support_export",
        },
        "/api/agent-runs/<run_id>/finance-requests/apply": {
            "POST": "agent_blueprints_api.apply_agent_run_finance_requests",
        },
        "/api/agent-runs/<run_id>/feedback": {
            "POST": "agent_blueprints_api.create_agent_run_feedback",
        },
        "/api/agent-runs/<run_id>/approvals/<approval_id>/approve": {
            "POST": "agent_blueprints_api.approve_agent_run",
        },
        "/api/agent-runs/<run_id>/approvals/<approval_id>/reject": {
            "POST": "agent_blueprints_api.reject_agent_run",
        },
    }

    actual = {}
    for rule in main.app.url_map.iter_rules():
        methods = rule.methods - {"HEAD", "OPTIONS"}
        actual.setdefault(rule.rule, {})
        for method in methods:
            actual[rule.rule][method] = rule.endpoint

    for route, methods in expected.items():
        for method, endpoint in methods.items():
            assert actual.get(route, {}).get(method) == endpoint


def test_agent_blueprint_migration_creates_expected_tables():
    migration = Path("alembic_migrations/versions/20260523_add_agent_blueprint_layer.py").read_text(encoding="utf-8")
    for table_name in [
        "agent_blueprints",
        "agent_blueprint_versions",
        "agent_runs",
        "agent_run_steps",
        "agent_artifacts",
        "agent_approvals",
    ]:
        assert f"CREATE TABLE IF NOT EXISTS {table_name}" in migration
    assert "JSONB" in migration
    assert "20260523_001" in migration
    assert "20260521_001" in migration


def test_agent_builder_session_migration_creates_expected_table():
    migration = Path("alembic_migrations/versions/20260525_add_agent_builder_sessions.py").read_text(encoding="utf-8")

    assert "CREATE TABLE IF NOT EXISTS agent_builder_sessions" in migration
    assert "preview_json JSONB" in migration
    assert "missing_questions_json JSONB" in migration
    assert "20260525_001" in migration
    assert "20260523_001" in migration


def test_agent_domain_request_migration_creates_expected_tables():
    migration = Path("alembic_migrations/versions/20260609_add_agent_domain_request_tables.py").read_text(encoding="utf-8")

    assert "CREATE TABLE IF NOT EXISTS agent_communication_requests" in migration
    assert "CREATE TABLE IF NOT EXISTS agent_service_optimization_requests" in migration
    assert "recipients_json JSONB" in migration
    assert "suggestions_json JSONB" in migration
    assert "delivery_state TEXT NOT NULL DEFAULT 'not_dispatched'" in migration
    assert "apply_state TEXT NOT NULL DEFAULT 'not_applied'" in migration
    assert "20260609_001" in migration
    assert "20260525_001" in migration


def test_custom_agent_integration_migration_creates_expected_tables():
    migration = Path("alembic_migrations/versions/20260609_add_custom_agent_integration_tables.py").read_text(encoding="utf-8")

    assert "CREATE TABLE IF NOT EXISTS agent_integrations" in migration
    assert "CREATE TABLE IF NOT EXISTS agent_trigger_events" in migration
    assert "CREATE TABLE IF NOT EXISTS agent_sheet_operation_requests" in migration
    assert "provider_write_performed BOOLEAN NOT NULL DEFAULT FALSE" in migration
    assert "apply_state TEXT NOT NULL DEFAULT 'not_applied'" in migration
    assert "20260609_002" in migration
    assert "20260609_001" in migration


def test_agent_service_optimization_diff_migration_adds_visual_diff_column():
    migration = Path("alembic_migrations/versions/20260609_add_agent_service_optimization_diff.py").read_text(encoding="utf-8")

    assert "ADD COLUMN IF NOT EXISTS diff_json JSONB" in migration
    assert "20260609_003" in migration
    assert "20260609_002" in migration


def test_agent_communication_delivery_journal_migration_creates_handoff_table():
    migration = Path("alembic_migrations/versions/20260609_add_agent_communication_delivery_journal.py").read_text(encoding="utf-8")

    assert "CREATE TABLE IF NOT EXISTS agent_communication_delivery_journal" in migration
    assert "router_handoff_json JSONB" in migration
    assert "provider_write_performed BOOLEAN NOT NULL DEFAULT FALSE" in migration
    assert "20260609_004" in migration
    assert "20260609_003" in migration


def test_agent_review_publish_request_migration_creates_expected_table():
    migration = Path("alembic_migrations/versions/20260609_add_agent_review_publish_requests.py").read_text(encoding="utf-8")

    assert "CREATE TABLE IF NOT EXISTS agent_review_publish_requests" in migration
    assert "provider_request_json JSONB" in migration
    assert "audit_json JSONB" in migration
    assert "provider_write_performed BOOLEAN NOT NULL DEFAULT FALSE" in migration
    assert "publish_state TEXT NOT NULL" in migration
    assert "20260609_005" in migration
    assert "20260609_004" in migration


def test_default_supervised_outreach_template_has_approval_gates():
    from services.agent_blueprint_runner import default_supervised_outreach_version_payload

    payload = default_supervised_outreach_version_payload()
    steps = payload["steps"]

    assert payload["capability_allowlist"] == ["outreach.send_batch"]
    assert [step["key"] for step in steps] == [
        "source_leads",
        "shortlist",
        "approve_shortlist",
        "draft_messages",
        "approve_drafts",
        "send_limited_batch",
        "record_outcomes",
    ]
    assert steps[2]["type"] == "approval"
    assert steps[4]["type"] == "approval"
    assert steps[5]["type"] == "capability"
    assert steps[5]["requires_approval"] is True
    assert steps[5]["required_approval_type"] == "drafts"


def test_agent_blueprint_orchestrator_exposes_stage4_capability_map():
    from services.agent_blueprint_orchestrator import build_agent_blueprint_orchestrator
    from services.agent_capability_handlers import build_capability_catalog
    from services.agent_provider_registry import integration_execution_boundary

    orchestrator = build_agent_blueprint_orchestrator()
    expected = {
        "outreach.send_batch",
        "reviews.reply.draft",
        "reviews.reply.publish_request",
        "services.optimize",
        "news.generate",
        "appointments.read",
        "appointments.create_request",
        "communications.draft",
        "communications.send_reminder",
        "communications.send_offer",
        "support.export",
        "sheets.append_row_request",
        "google_sheets.read_rows",
        "finance.transaction.create",
        "partnership.audit_card",
        "partnership.match_services",
        "partnership.draft_offer",
        "billing.reserve",
        "billing.settle",
    }

    for capability in expected:
        assert capability in orchestrator.handlers

    assert "reviews.reply" in orchestrator.handlers
    assert "appointments.create" in orchestrator.handlers
    assert "communications.send" in orchestrator.handlers
    assert "google_sheets.append_row" in orchestrator.handlers
    assert "finance.manual_entry" in orchestrator.handlers
    assert "partners.match_services" in orchestrator.handlers
    assert "partners.draft_first_offer" in orchestrator.handlers
    catalog = build_capability_catalog()
    assert expected.issubset(set(catalog["capabilities"]))
    assert catalog["capabilities"]["reviews.reply"]["alias_for"] == "reviews.reply.draft"
    assert catalog["capabilities"]["google_sheets.append_row"]["alias_for"] == "sheets.append_row_request"
    assert catalog["capabilities"]["finance.manual_entry"]["alias_for"] == "finance.transaction.create"
    assert catalog["capabilities"]["partners.match_services"]["alias_for"] == "partnership.match_services"
    assert catalog["capabilities"]["partners.draft_first_offer"]["alias_for"] == "partnership.draft_offer"
    assert catalog["provider_registry"]["maton"]["status"] == "available"
    assert catalog["provider_registry"]["openclaw"]["status"] == "available"
    assert catalog["provider_registry"]["composio"]["status"] == "planned"
    reminder_providers = {
        item["provider"]
        for item in catalog["capabilities"]["communications.send_reminder"]["provider_candidates"]
    }
    assert {"maton", "openclaw", "native_localos", "composio", "manual"}.issubset(reminder_providers)
    sheet_read_providers = {
        item["provider"]
        for item in catalog["capabilities"]["sheets.append_row_request"]["provider_candidates"]
    }
    assert {"native_localos", "composio", "manual"}.issubset(sheet_read_providers)
    partnership_match_providers = {
        item["provider"]
        for item in catalog["capabilities"]["partnership.match_services"]["provider_candidates"]
    }
    assert {"native_localos", "openclaw", "manual"}.issubset(partnership_match_providers)
    assert integration_execution_boundary("google_sheets")["executor"] == "agent_sheet_provider_executor_v1"
    assert integration_execution_boundary("maton")["executor"] == "channel_router"
    assert integration_execution_boundary("localos_finance")["executor"] == "localos_finance_request_executor"
    assert integration_execution_boundary("composio")["external_write"] == "planned_provider_write"
    from services.agent_provider_registry import connector_provider_routes

    reminder_routes = connector_provider_routes("telegram", "communications.send_reminder")
    maton_route = next(item for item in reminder_routes if item["provider"] == "maton")
    openclaw_route = next(item for item in reminder_routes if item["provider"] == "openclaw")
    composio_route = next(item for item in reminder_routes if item["provider"] == "composio")
    assert maton_route["connect_mode"] == "external_account_key"
    assert maton_route["primary_cta"] == "Выбрать Maton key"
    assert maton_route["provider_action"]["kind"] == "select_external_account_key"
    assert maton_route["provider_action"]["ui_target"] == "external_business_accounts"
    assert openclaw_route["connect_mode"] == "openclaw_policy_boundary"
    assert openclaw_route["primary_cta"] == "Использовать OpenClaw boundary"
    assert openclaw_route["provider_action"]["kind"] == "use_openclaw_boundary"
    assert openclaw_route["provider_action"]["ui_target"] == "openclaw_policy_boundary"
    assert composio_route["connect_mode"] == "planned_oauth_connector"
    assert composio_route["primary_cta"] == "Будет доступно позже"
    assert composio_route["provider_action"]["kind"] == "planned_oauth_connector"
    assert composio_route["provider_action"]["available"] is False


def test_custom_process_preview_input_uses_bound_integrations_and_safe_telegram_payload():
    from api import agent_blueprints_api

    blueprint = {
        "metadata_json": {
            "custom_process": {
                "google_sheets": {
                    "integration_id": "sheet-integration-1",
                    "spreadsheet_id": "spreadsheet-1",
                    "sheet_name": "Leads",
                }
            }
        }
    }

    preview = agent_blueprints_api._build_custom_process_preview_input(
        blueprint,
        {
            "message_text": "Новая заявка: Анна",
            "telegram_username": "anna",
        },
    )

    assert preview["preview_mode"] is True
    assert preview["source_event"]["preview"] is True
    assert preview["source_event"]["source"] == "telegram_preview"
    assert preview["integration_id"] == "sheet-integration-1"
    assert preview["spreadsheet_id"] == "spreadsheet-1"
    assert preview["sheet_name"] == "Leads"
    assert preview["telegram"]["message_text"] == "Новая заявка: Анна"
    assert preview["telegram"]["username"] == "anna"


def test_agent_preview_run_input_is_safe_and_compiled_workflow_aware():
    from api import agent_blueprints_api
    from services.agent_blueprint_draft_builder import compile_agent_blueprint

    draft = compile_agent_blueprint(
        "Каждый день бери заказ из Google Sheets за вчера и готовь пост в Telegram. "
        "Результат нужен как черновик поста. Перед публикацией человек проверяет результат."
    )
    draft["metadata"]["custom_process"]["google_sheets"] = {
        "spreadsheet_id": "spreadsheet-1",
        "sheet_name": "Orders",
        "gid": "0",
    }
    draft["metadata"]["custom_process"]["telegram"] = {
        "telegram_target": "@riderra_updates",
        "target_type": "chat_or_channel",
    }
    draft["metadata"]["connector_action_handlers"] = {
        "google_sheets_read": {
            "schema": "localos_connector_action_handler_v1",
            "binding_key": "google_sheets_read",
            "route_provider": "openclaw",
            "handler": "openclaw_policy_boundary",
            "credential_source": "openclaw_m2m",
            "preflight_resolution": "provider_route_openclaw_boundary",
            "execution_boundary": "localos_policy_envelope",
            "approval_required": True,
            "audit_required": True,
            "external_side_effects_allowed_in_preview": False,
        }
    }
    version_payload = draft["version_payload"]
    blueprint = {
        "id": "bp1",
        "business_id": "biz1",
        "category": draft["category"],
        "description": "Google Sheets to Telegram",
        "metadata_json": draft["metadata"],
    }
    version = {
        "id": "version-1",
        "goal": version_payload["goal"],
        "trigger": version_payload["trigger"],
        "inputs_schema_json": version_payload["inputs_schema"],
        "steps_json": version_payload["steps"],
        "capability_allowlist_json": version_payload["capability_allowlist"],
        "approval_policy_json": version_payload["approval_policy"],
        "output_schema_json": version_payload["output_schema"],
    }

    preview_input = agent_blueprints_api._build_agent_preview_run_input(
        blueprint,
        version,
        {"input": {"preview_mode": True}},
    )

    assert preview_input["schema"] == "localos_agent_preview_input_v1"
    assert preview_input["preview_mode"] is True
    assert preview_input["external_side_effects_allowed"] is False
    assert preview_input["approval_required_for_external_actions"] is True
    assert preview_input["source_event"]["preview"] is True
    assert preview_input["google_sheets"]["read_only"] is True
    assert preview_input["google_sheets"]["spreadsheet_id"] == "spreadsheet-1"
    assert preview_input["google_sheets"]["sheet_name"] == "Orders"
    assert preview_input["google_sheets"]["gid"] == "0"
    assert preview_input["google_sheets"]["sample_rows"][0]["route"] == "Los Angeles airport → Santa Barbara"
    assert preview_input["telegram"]["draft_only"] is True
    assert preview_input["telegram"]["external_publish_performed"] is False
    assert preview_input["telegram"]["telegram_target"] == "@riderra_updates"
    assert [item["provider"] for item in preview_input["provider_bindings"]] == ["google_sheets", "telegram"]
    assert preview_input["policy_envelope"]["execution_boundary"] == "openclaw_action_orchestrator"
    assert preview_input["policy_envelope"]["external_side_effects_allowed_in_preview"] is False
    assert preview_input["connector_action_handlers"][0]["handler"] == "openclaw_policy_boundary"
    assert preview_input["openclaw_preview_routes"][0]["binding_key"] == "google_sheets_read"
    assert preview_input["openclaw_preview_routes"][0]["external_side_effects_allowed_in_preview"] is False
    assert preview_input["openclaw_action_plan"][0]["provider_action_ref"] == "openclaw.google_sheets.read_rows"
    assert preview_input["openclaw_action_plan"][0]["provider_policy"] == "localos_envelope"


def test_openclaw_and_capability_routes_are_registered():
    import main

    actual = {}
    for rule in main.app.url_map.iter_rules():
        methods = rule.methods - {"HEAD", "OPTIONS"}
        actual.setdefault(rule.rule, set()).update(methods)

    expected = {
        "/api/capabilities/execute": "POST",
        "/api/capabilities/catalog": "GET",
        "/api/capabilities/actions/<action_id>": "GET",
        "/api/capabilities/actions/<action_id>/decision": "POST",
        "/api/capabilities/actions/<action_id>/billing": "GET",
        "/api/capabilities/health": "GET",
        "/api/capabilities/support-export": "GET",
        "/api/openclaw/capabilities/execute": "POST",
        "/api/openclaw/capabilities/catalog": "GET",
        "/api/openclaw/capabilities/actions/<action_id>": "GET",
        "/api/openclaw/capabilities/actions/<action_id>/decision": "POST",
        "/api/openclaw/capabilities/health": "GET",
        "/api/openclaw/callbacks/outbox": "GET",
        "/api/openclaw/audit-timeline": "GET",
    }

    for route, method in expected.items():
        assert method in actual.get(route, set())


def test_legacy_ai_agent_migration_plan_marks_runtime_truth_and_deprecations():
    from services.agent_legacy_migration import (
        LEGACY_WORKFLOW_STATUS,
        build_business_ai_settings_deprecation_plan,
        build_legacy_run_preview_bridge,
    )

    settings_plan = build_business_ai_settings_deprecation_plan(
        {
            "ai_agent_enabled": True,
            "ai_agent_tone": "friendly",
            "ai_agent_restrictions": "no discounts",
            "ai_agents_config": "{}",
            "ai_agent_id": "voice-1",
        }
    )
    bridge = build_legacy_run_preview_bridge({"id": "voice-1"}, "biz-1")

    assert LEGACY_WORKFLOW_STATUS == "deprecated_not_runtime_truth"
    assert settings_plan["fields"]["ai_agent_enabled"]["status"] == "deprecated_migration_source"
    assert settings_plan["fields"]["ai_agent_id"]["target"] == "agent_blueprint_versions.persona_agent_id"
    assert bridge["status"] == "moved_to_shared_run_preview_contract"
    assert bridge["preview_contract"]["target_runtime"] == "agent_blueprints"
    assert bridge["preview_contract"]["external_dispatch_performed"] is False


def test_agent_blueprint_draft_builder_creates_safe_document_agent():
    from services.agent_blueprint_draft_builder import build_agent_blueprint_draft

    draft = build_agent_blueprint_draft("Обработай договор, найди риски и подготовь письмо клиенту")
    version_payload = draft["version_payload"]
    steps = version_payload["steps"]

    assert draft["category"] == "documents"
    assert draft["metadata"]["builder"] == "description_builder_v1"
    assert "uploaded_documents" in draft["summary"]["sources"]
    assert version_payload["capability_allowlist"] == []
    assert any(step["type"] == "approval" for step in steps)
    assert "external_delivery" in draft["summary"]["approval_boundaries"]


def test_agent_blueprint_draft_builder_respects_explicit_category():
    from services.agent_blueprint_draft_builder import build_agent_blueprint_draft

    draft = build_agent_blueprint_draft("Подготовь результат по этому контексту", "documents")

    assert draft["category"] == "documents"
    assert draft["name"] == "Подготовь результат по этому контексту"
    assert draft["version_payload"]["capability_allowlist"] == []
    assert draft["summary"]["external_dispatch_performed"] is False


def test_agent_compiler_creates_communications_reminder_blueprint():
    from services.agent_blueprint_draft_builder import compile_agent_blueprint

    draft = compile_agent_blueprint(
        "Сделай агента, который напоминает клиентам о записи и сообщает про пакетное предложение"
    )
    payload = draft["version_payload"]

    assert draft["category"] == "communications"
    assert draft["metadata"]["compiler"] == "agent_compiler_v1"
    assert payload["trigger"] == "appointment.reminder.before"
    assert payload["audience"] == "clients_with_upcoming_appointments"
    assert payload["audience_rules"]
    assert payload["consent_rules"]
    assert payload["message_template"]
    assert payload["persona"]
    assert payload["send_capability"] == "communications.send_reminder"
    assert payload["delivery_outcome_journal"]["journal_type"] == "communications_delivery_outcome"
    assert payload["mode"] == "approved_batch_only"
    assert payload["external_dispatch_performed"] is False
    assert payload["data_sources"] == ["appointments", "services", "packages", "business_profile"]
    assert [step["key"] for step in payload["steps"]] == [
        "collect_audience",
        "prepare_message",
        "validate_consent",
        "approve_message",
        "send_message",
        "record_outcome",
    ]
    assert payload["capability_allowlist"] == [
        "appointments.read",
        "communications.draft",
        "communications.send_reminder",
    ]
    assert payload["approval_policy"]["first_run"] == "manual_approval_required"
    assert payload["approval_policy"]["mass_send"] == "manual_approval_required"
    assert payload["approval_policy"]["mode"] == "approved_batch_only"
    assert payload["limits"]["daily_cap"] == 10
    assert payload["limits"]["autonomous_send_allowed"] is False
    assert "drafts" in payload["output_schema"]["properties"]
    assert "delivery_report" in payload["output_schema"]["properties"]
    assert "outcomes" in payload["output_schema"]["properties"]
    assert "delivery_outcome_journal" in payload["output_schema"]["properties"]


def test_agent_compiler_creates_custom_telegram_to_sheets_blueprint():
    from services.agent_blueprint_draft_builder import compile_agent_blueprint

    draft = compile_agent_blueprint("Когда пользователь пишет в Telegram бота, добавь строку в Google таблицу")
    payload = draft["version_payload"]

    assert draft["category"] == "custom"
    assert draft["metadata"]["custom_process"]["kind"] == "source_destination_workflow"
    assert draft["metadata"]["custom_process"]["trigger"] == "telegram.message.received"
    assert draft["metadata"]["custom_process"]["target"] == "sheets.append_row_request"
    assert payload["trigger"] == "telegram.message.received"
    assert payload["mode"] == "approved_capability_request"
    assert payload["capability_allowlist"] == ["sheets.append_row_request"]
    assert payload["approval_policy"]["sheet_update"] == "manual_approval_required"
    assert draft["metadata"]["compiled_process"]["schema"] == "compiled_source_destination_workflow_v1"
    assert draft["metadata"]["required_integration_bindings"][0]["key"] == "telegram_trigger"
    assert draft["metadata"]["required_integration_bindings"][1]["key"] == "google_sheets_append"
    assert draft["metadata"]["required_integration_bindings"][1]["required_config"] == ["spreadsheet_id", "sheet_name"]
    assert payload["required_integration_bindings"][1]["capability"] == "sheets.append_row_request"
    assert payload["limits"]["autonomous_external_write_allowed"] is False
    assert [step["key"] for step in payload["steps"]] == [
        "capture_telegram_trigger",
        "prepare_sheet_row",
        "approve_sheet_update",
        "request_google_sheets",
        "record_google_sheets_outcome",
    ]
    assert payload["steps"][3]["requires_approval"] is True
    assert payload["steps"][3]["required_approval_type"] == "sheet_update"
    assert payload["side_effects_performed"] is False


def test_agent_compiler_creates_source_destination_blueprint_for_sheets_to_finance():
    from services.agent_blueprint_draft_builder import compile_agent_blueprint

    draft = compile_agent_blueprint(
        "Каждый вечер проверяй Google Sheets, бери новые оплаты и создавай транзакции в финансах LocalOS"
    )
    payload = draft["version_payload"]

    assert draft["category"] == "custom"
    assert draft["metadata"]["custom_process"]["kind"] == "source_destination_workflow"
    assert draft["metadata"]["custom_process"]["archetype"] == "google_sheets_to_localos_finance"
    assert draft["metadata"]["custom_process"]["source"] == "google_sheets.read_rows"
    assert draft["metadata"]["custom_process"]["target"] == "finance.transaction.create"
    assert draft["metadata"]["custom_process"]["schedule"]["time"] == "19:00"
    assert payload["trigger"] == "schedule.daily"
    assert payload["capability_allowlist"] == ["google_sheets.read_rows", "finance.transaction.create"]
    assert payload["approval_policy"]["finance_transaction_import"] == "manual_approval_required"
    assert draft["metadata"]["compiled_process"]["schema"] == "compiled_source_destination_workflow_v1"
    assert payload["required_integration_bindings"][0]["key"] == "google_sheets_read"
    assert payload["required_integration_bindings"][1]["key"] == "localos_finance"
    assert [step["key"] for step in payload["steps"]] == [
        "read_google_sheets",
        "normalize_finance_rows",
        "approve_finance_transaction_import",
        "request_localos_finance",
        "record_localos_finance_outcome",
    ]
    assert payload["steps"][3]["capability"] == "finance.transaction.create"
    assert payload["steps"][3]["payload"]["rows_from_step"] == "read_google_sheets"
    assert payload["steps"][0]["provider"] == "openclaw"
    assert payload["steps"][0]["provider_action_ref"] == "openclaw.google_sheets.read_rows"
    assert payload["steps"][0]["provider_policy"] == "localos_envelope"
    assert payload["steps"][3]["payload"]["input_mappings"] == [
        {
            "target": "rows",
            "from_step": "read_google_sheets",
            "path": "orchestrator.result.rows",
            "required": True,
        }
    ]
    assert payload["limits"]["autonomous_localos_write_allowed"] is False
    assert draft["metadata"]["compiled_artifact_candidate"]["schema"] == "localos_compiled_artifact_candidate_v1"
    assert draft["metadata"]["compiled_artifact_candidate"]["status"] == "validation_passed"
    assert draft["metadata"]["compiled_validation"]["valid"] is True
    assert draft["metadata"]["compiled_artifact_candidate"]["dsl"]["schema"] == "localos_agent_workflow_dsl_v1"
    assert draft["metadata"]["compiled_artifact_candidate"]["activation_gate"]["requires_validation_passed"] is True


def test_agent_compiler_creates_source_destination_blueprint_for_sheets_to_telegram_post():
    from services.agent_blueprint_draft_builder import compile_agent_blueprint

    draft = compile_agent_blueprint(
        "Мне нужен агент, который из онлайн таблицы гугл со списком заказов берёт один из заказов "
        "на предыдущий день и создаёт пост в телеграм. В стиле - наши пассажиры насладились поездкой "
        "из аэропорта лос анджелеса в санта барбару."
    )
    payload = draft["version_payload"]

    assert draft["category"] == "custom"
    assert draft["metadata"]["custom_process"]["kind"] == "source_destination_workflow"
    assert draft["metadata"]["custom_process"]["archetype"] == "google_sheets_to_telegram"
    assert draft["metadata"]["custom_process"]["source"] == "google_sheets.read_rows"
    assert draft["metadata"]["custom_process"]["target"] == "communications.draft"
    assert draft["metadata"]["custom_process"]["schedule"]["time"] == "10:00"
    assert payload["trigger"] == "schedule.daily"
    assert payload["capability_allowlist"] == ["google_sheets.read_rows", "communications.draft"]
    assert payload["approval_policy"]["telegram_post_approval"] == "manual_approval_required"
    assert draft["metadata"]["compiled_process"]["schema"] == "compiled_source_destination_workflow_v1"
    assert payload["required_integration_bindings"][0]["key"] == "google_sheets_read"
    assert payload["required_integration_bindings"][1]["key"] == "telegram_delivery"
    assert payload["required_integration_bindings"][1]["provider"] == "telegram"
    assert payload["required_integration_bindings"][1]["required_config"] == ["bot_mode"]
    assert [step["key"] for step in payload["steps"]] == [
        "read_google_sheets",
        "prepare_telegram_post",
        "approve_telegram_post_approval",
        "request_telegram",
        "record_telegram_outcome",
    ]
    assert payload["steps"][1]["artifact_type"] == "telegram_post_draft"
    assert payload["steps"][3]["capability"] == "communications.draft"
    assert payload["steps"][0]["provider_action_ref"] == "openclaw.google_sheets.read_rows"
    assert payload["steps"][3]["provider_action_ref"] == "openclaw.telegram.create_draft"
    assert payload["steps"][3]["provider_policy"] == "localos_envelope"
    assert payload["steps"][3]["payload"]["message_type"] == "telegram_post_draft"
    assert payload["steps"][3]["payload"]["rows_from_step"] == "read_google_sheets"
    assert payload["steps"][3]["payload"]["input_mappings"] == [
        {
            "target": "rows",
            "from_step": "read_google_sheets",
            "path": "orchestrator.result.rows",
            "required": True,
        }
    ]
    assert payload["limits"]["autonomous_external_write_allowed"] is False
    assert draft["metadata"]["compiled_artifact_candidate"]["status"] == "validation_passed"
    assert draft["metadata"]["compiled_validation"]["valid"] is True


def test_agent_compiler_creates_source_only_blueprint_for_google_sheet_result_url():
    from services.agent_blueprint_draft_builder import compile_agent_blueprint

    draft = compile_agent_blueprint(
        "Агент открывает таблицу со списком поездок, выбирает одну из поездок 20 апреля 2022 года. "
        "https://docs.google.com/spreadsheets/d/1s79gWCm7A8X1drwN6yAscetf0adpRkamHCJyHCkyIqY/edit?gid=0#gid=0 "
        "пишет пост и сохраняет его в контент план на 27 июня 2026 года"
    )
    payload = draft["version_payload"]

    assert draft["category"] == "custom"
    assert draft["metadata"]["custom_process"]["kind"] == "source_to_result_workflow"
    assert draft["metadata"]["custom_process"]["source"] == "google_sheets.read_rows"
    assert draft["metadata"]["custom_process"]["target"] == "agent_output_draft"
    assert draft["metadata"]["compiled_process"]["schema"] == "compiled_source_to_result_workflow_v1"
    assert payload["mode"] == "source_to_reviewed_result"
    assert payload["capability_allowlist"] == ["google_sheets.read_rows"]
    assert payload["required_integration_bindings"][0]["key"] == "google_sheets_read"
    assert [step["key"] for step in payload["steps"]] == [
        "read_google_sheets",
        "prepare_output",
        "approve_output",
        "save_result",
    ]
    assert payload["steps"][0]["capability"] == "google_sheets.read_rows"
    assert payload["steps"][0]["provider_action_ref"] == "openclaw.google_sheets.read_rows"
    assert payload["steps"][1]["artifact_type"] == "agent_output_draft"
    assert payload["steps"][1]["payload"]["rows_from_step"] == "read_google_sheets"
    assert payload["approval_policy"]["final_output"] == "manual_approval_required"
    assert payload["limits"]["autonomous_external_write_allowed"] is False
    assert draft["metadata"]["compiled_artifact_candidate"]["status"] == "validation_passed"
    assert draft["metadata"]["compiled_validation"]["valid"] is True


def test_compiled_workflow_validation_rejects_write_without_approval():
    from services.agent_compiled_artifact import validate_compiled_artifact_candidate

    version_payload = {
        "goal": "Записать строку в таблицу",
        "trigger": "manual.run",
        "mode": "approved_capability_request",
        "inputs_schema": {"type": "object"},
        "steps": [
            {
                "key": "write_sheet",
                "type": "capability",
                "capability": "sheets.append_row_request",
                "requires_approval": False,
            }
        ],
        "capability_allowlist": ["sheets.append_row_request"],
        "approval_policy": {},
        "required_integration_bindings": [
            {
                "key": "google_sheets_append",
                "provider": "google_sheets",
                "capability": "sheets.append_row_request",
            }
        ],
        "limits": {"autonomous_external_write_allowed": False},
        "output_schema": {"type": "object"},
    }

    result = validate_compiled_artifact_candidate(version_payload, {"compiled_process": {"schema": "compiled_source_destination_workflow_v1"}})

    assert result["ready"] is False
    assert result["validation"]["status"] == "invalid"
    fields = [item["field"] for item in result["validation"]["errors"]]
    assert "steps[0].requires_approval" in fields
    assert "steps[0].required_approval_type" in fields


def test_compiled_workflow_validation_rejects_openclaw_action_capability_mismatch():
    from services.agent_compiled_artifact import validate_compiled_artifact_candidate

    version_payload = {
        "goal": "Прочитать таблицу",
        "trigger": "manual.run",
        "mode": "approved_capability_request",
        "inputs_schema": {"type": "object"},
        "steps": [
            {
                "key": "read_sheet",
                "type": "capability",
                "capability": "google_sheets.read_rows",
                "provider": "openclaw",
                "provider_action_ref": "openclaw.telegram.publish_message",
                "requires_approval": False,
            }
        ],
        "capability_allowlist": ["google_sheets.read_rows"],
        "approval_policy": {},
        "required_integration_bindings": [
            {
                "key": "google_sheets_read",
                "provider": "google_sheets",
                "capability": "google_sheets.read_rows",
            }
        ],
        "limits": {"autonomous_external_write_allowed": False},
        "output_schema": {"type": "object"},
    }

    result = validate_compiled_artifact_candidate(version_payload, {"compiled_process": {"schema": "compiled_source_destination_workflow_v1"}})

    assert result["ready"] is False
    assert result["validation"]["status"] == "invalid"
    assert "steps[0].provider_action_ref" in [item["field"] for item in result["validation"]["errors"]]


def test_compiled_workflow_validation_uses_metadata_snapshot_for_version_rows():
    from services.agent_blueprint_draft_builder import compile_agent_blueprint
    from services.agent_blueprint_workspace import build_version_payload_from_row
    from services.agent_compiled_artifact import validate_compiled_artifact_candidate

    draft = compile_agent_blueprint(
        "Каждый вечер проверяй Google Sheets, бери новые оплаты и создавай транзакции в финансах LocalOS"
    )
    version_row = {
        "goal": draft["version_payload"]["goal"],
        "inputs_schema_json": draft["version_payload"]["inputs_schema"],
        "steps_json": draft["version_payload"]["steps"],
        "capability_allowlist_json": draft["version_payload"]["capability_allowlist"],
        "approval_policy_json": draft["version_payload"]["approval_policy"],
        "output_schema_json": draft["version_payload"]["output_schema"],
    }

    result = validate_compiled_artifact_candidate(build_version_payload_from_row(version_row), draft["metadata"])

    assert result["ready"] is True
    assert result["candidate"]["dsl"]["trigger"] == "schedule.daily"
    assert result["candidate"]["dsl"]["limits"]["autonomous_localos_write_allowed"] is False
    assert result["candidate"]["dsl"]["required_integration_bindings"][0]["key"] == "google_sheets_read"


def test_agent_metrics_summary_reports_compiled_runtime_health():
    from services.agent_blueprint_draft_builder import compile_agent_blueprint
    from services.agent_metrics import build_agent_metrics_summary

    draft = compile_agent_blueprint(
        "Каждый вечер проверяй Google Sheets, бери новые оплаты и создавай транзакции в финансах LocalOS"
    )
    metadata = dict(draft["metadata"])
    metadata["billing"] = {
        "estimated_credits": 3,
        "actual_credits": 3,
        "idempotency_key": "charge1",
    }
    metrics = build_agent_metrics_summary(
        {"id": "bp1", "status": "draft"},
        [{"id": "ver1", "version_number": 1}],
        {"id": "ver1", "version_number": 1},
        [
            {
                "id": "run1",
                "status": "completed",
                "input_json": {"preview_mode": True},
                "output_json": {
                    "observability": {
                        "cost_tokens": {
                            "reserved_tokens": 10,
                            "settled_tokens": 7,
                            "released_tokens": 3,
                            "inflight_reserved_tokens": 0,
                            "total_cost": 0.14,
                        },
                        "billing_ledger": {
                            "actions": [
                                {
                                    "action_id": "action1",
                                    "capability": "localos.finance.write",
                                    "settled_tokens": 42,
                                    "total_cost": 0.012,
                                }
                            ],
                            "entries": [],
                        },
                    }
                },
            },
            {"id": "run2", "status": "failed", "output_json": {}},
        ],
        [{"id": "approval1", "approval_type": "finance_transaction_import"}],
        metadata,
    )

    assert metrics["schema"] == "agent_metrics_summary_v1"
    assert metrics["compiled"]["validation_valid"] is True
    assert metrics["compiled"]["runtime_llm_required"] is False
    assert metrics["versions"]["active_version_number"] == 1
    assert metrics["runs"]["by_status"] == {"completed": 1, "failed": 1}
    assert metrics["approvals"]["pending"] == 1
    assert metrics["cost_tokens"]["reserved_tokens"] == 10
    assert metrics["cost_tokens"]["total_cost"] == 0.14
    assert metrics["billing_breakdown"]["schema"] == "localos_agent_billing_breakdown_v1"
    assert metrics["billing_breakdown"]["items"][0]["key"] == "agent_creation"
    assert any(item["key"] == "external_action" for item in metrics["cost_tokens"]["breakdown"])
    assert metrics["unified_billing_ledger"]["schema"] == "localos_agent_unified_billing_ledger_v1"
    assert metrics["unified_billing_ledger"]["summary"]["estimated_credits"] == 4
    assert metrics["unified_billing_ledger"]["summary"]["actual_credits"] == 3
    assert metrics["unified_billing_ledger"]["summary"]["actual_tokens"] == 49
    assert metrics["unified_billing_ledger"]["items"][0]["key"] == "agent_creation"
    assert any(item["key"] == "preview_run" and item["actual_tokens"] == 7 for item in metrics["unified_billing_ledger"]["items"])
    assert any(item["key"] == "external_action" and item["actual_tokens"] == 42 for item in metrics["unified_billing_ledger"]["items"])


def test_agent_creation_cost_preview_exposes_unified_ledger_estimate_items():
    from services.agent_builder_billing import build_agent_creation_cost_preview

    preview = build_agent_creation_cost_preview()

    assert preview["schema"] == "localos_agent_billing_estimate_v1"
    assert preview["total_estimated_credits"] == 4
    assert preview["total_estimated_tokens"] == 8500
    assert [item["key"] for item in preview["items"]] == [
        "agent_creation",
        "preview_run",
        "production_run",
        "external_action",
        "operator_chat",
    ]
    assert preview["items"][0]["billing_mode"] == "reserve_then_charge"
    assert preview["items"][3]["billing_mode"] == "action_orchestrator_reserve_settle"


def test_existing_agent_templates_publish_compiled_artifact_candidate():
    from services.agent_blueprint_draft_builder import compile_agent_blueprint

    examples = [
        ("Найди клиентов для outreach", "outreach", "compiled_outreach_workflow_v1"),
        ("Проверяй договоры и ищи риски", "documents", "compiled_documents_workflow_v1"),
        ("Готовь ответы на отзывы", "reviews", "compiled_reviews_workflow_v1"),
        ("Сделай таблицу исключений", "tables", "compiled_tables_workflow_v1"),
    ]

    for prompt, category, schema in examples:
        draft = compile_agent_blueprint(prompt)

        assert draft["category"] == category
        assert draft["metadata"]["compiled_process"]["schema"] == schema
        assert draft["metadata"]["compiled_artifact_candidate"]["status"] == "validation_passed"
        assert draft["metadata"]["compiled_validation"]["valid"] is True
        assert draft["metadata"]["compiler_contract"]["runtime_llm_required"] is False


def test_agent_compiler_uses_gigachat_only_at_design_time(monkeypatch):
    from services import agent_blueprint_draft_builder

    def fake_llm_intent(description, *, business_id="", user_id="", planner_context=None):
        return {
            "status": "compiled_intent",
            "source": "gigachat",
            "intent": {
                "trigger": "schedule.daily",
                "source": "google_sheets",
                "destination": "localos_finance",
                "read_capability": "google_sheets.read_rows",
                "write_capability": "finance.transaction.create",
                "required_connectors": ["google_sheets", "localos_finance"],
                "approval_reasons": ["localos_finance_write", "ambiguous_data"],
                "limits": {"max_items_per_run": 100},
                "clarifying_questions": [],
                "confidence": 0.92,
            },
        }

    monkeypatch.setattr(agent_blueprint_draft_builder, "infer_agent_workflow_intent", fake_llm_intent)

    draft = agent_blueprint_draft_builder.compile_agent_blueprint(
        "Раз в день переноси оплаты из таблицы в финансы",
        use_ai=True,
        business_id="biz1",
        user_id="user1",
    )
    metadata = draft["metadata"]
    steps = draft["version_payload"]["steps"]

    assert metadata["compiler_source"] == "gigachat_intent_extractor"
    assert metadata["llm_intent"]["status"] == "compiled_intent"
    assert metadata["compiler_contract"]["llm_usage"] == "design_time_only"
    assert metadata["compiler_contract"]["runtime_llm_required"] is False
    assert metadata["compiled_process"]["runtime_truth"] == "agent_blueprint_versions.steps_json"
    assert [step["type"] for step in steps] == ["capability", "artifact", "approval", "capability", "artifact"]
    assert "gigachat" not in str(draft["version_payload"]).lower()


def test_agent_compiler_llm_intent_is_sanitized_to_allowed_capabilities():
    from services.agent_compiler_llm import infer_agent_workflow_intent

    def fake_generator(prompt, business_id, user_id):
        assert "LocalOS/OpenClaw policy envelope" in prompt
        return """
        {
          "trigger": "schedule.daily",
          "source": "google_sheets",
          "destination": "localos_finance",
          "read_capability": "google_sheets.read_rows",
          "write_capability": "dangerous.delete_everything",
          "required_connectors": ["google_sheets"],
          "approval_reasons": ["localos_finance_write"],
          "limits": {"max_items_per_run": 9999},
          "clarifying_questions": ["Какую вкладку читать?"],
          "confidence": 2
        }
        """

    result = infer_agent_workflow_intent(
        "Каждый вечер читай оплаты из Google Sheets и готовь финансы",
        business_id="biz1",
        user_id="user1",
        intent_generator=fake_generator,
    )

    assert result["status"] == "compiled_intent"
    assert result["intent"]["source"] == "google_sheets"
    assert result["intent"]["destination"] == "localos_finance"
    assert result["intent"]["read_capability"] == "google_sheets.read_rows"
    assert result["intent"]["write_capability"] == ""
    assert result["intent"]["limits"]["max_items_per_run"] == 500
    assert result["intent"]["confidence"] == 1.0


def test_agent_compiler_llm_prompt_includes_localos_openclaw_planner_context():
    from services.agent_compiler_llm import infer_agent_workflow_intent

    def fake_generator(prompt, business_id, user_id):
        assert business_id == "biz1"
        assert user_id == "user1"
        assert "localos_openclaw_planner_context_v1" in prompt
        assert '"tenant_boundary": "single_business"' in prompt
        assert '"missing_connections"' in prompt
        assert '"provider": "telegram"' in prompt
        assert "credential_extraction" in prompt
        assert "must_not_call_tools_directly" in prompt
        return """
        {
          "compiled_template_key": "google_sheets_to_telegram_post",
          "source": "google_sheets",
          "destination": "telegram",
          "read_capability": "google_sheets.read_rows",
          "write_capability": "communications.draft",
          "required_connectors": ["google_sheets", "telegram", "forbidden_provider"],
          "approval_reasons": ["external_publish"],
          "limits": {"max_items_per_run": 25},
          "clarifying_questions": ["Какую вкладку Google Sheets читать?"],
          "confidence": 0.7
        }
        """

    result = infer_agent_workflow_intent(
        "Подготовь пост в Telegram из заказа в Google Sheets",
        business_id="biz1",
        user_id="user1",
        planner_context={
            "schema": "localos_openclaw_planner_context_v1",
            "business_scope": {
                "business_id": "biz1",
                "user_id": "user1",
                "tenant_boundary": "single_business",
                "cross_business_access_allowed": False,
            },
            "allowed_capabilities": ["google_sheets.read_rows", "communications.draft"],
            "connection_state": {
                "missing_connections": [
                    {
                        "key": "telegram_delivery",
                        "provider": "telegram",
                        "provider_title": "Telegram",
                        "status": "missing",
                    }
                ]
            },
            "forbidden_action_classes": ["credential_extraction"],
            "output_contract": {"must_not_call_tools_directly": True},
        },
        intent_generator=fake_generator,
    )

    assert result["status"] == "compiled_intent"
    assert result["intent"]["compiled_template_key"] == "google_sheets_to_telegram_post"
    assert result["intent"]["required_connectors"] == ["google_sheets", "telegram"]
    assert result["intent"]["clarifying_questions"] == ["Какую вкладку Google Sheets читать?"]


def test_agent_compiler_llm_returns_planner_artifact_inside_policy_envelope():
    from services.agent_compiler_llm import infer_agent_workflow_intent

    def fake_generator(prompt, business_id, user_id):
        assert "workflow_draft" in prompt
        assert "approval_points" in prompt
        assert "unsupported_requests" in prompt
        return """
        {
          "compiled_template_key": "google_sheets_to_telegram_post",
          "source": "google_sheets",
          "destination": "telegram",
          "read_capability": "google_sheets.read_rows",
          "write_capability": "communications.draft",
          "required_connectors": ["google_sheets", "telegram", "unknown_provider"],
          "workflow_draft": {
            "trigger": "schedule.daily",
            "steps": [
              {"key": "read_orders", "capability": "google_sheets.read_rows", "secret": "drop"},
              {"key": "prepare_post", "type": "artifact"}
            ],
            "outputs": [{"key": "telegram_post_draft"}]
          },
          "approval_points": [
            {"key": "publish", "reason": "external publish", "raw_token": "drop"}
          ],
          "unsupported_requests": [
            {"request": "bypass approval", "reason": "LocalOS policy requires approval"}
          ],
          "approval_reasons": ["external_publish"],
          "limits": {"max_items_per_run": 15},
          "clarifying_questions": [],
          "confidence": 0.8
        }
        """

    result = infer_agent_workflow_intent(
        "Каждый день бери заказ из Google Sheets и готовь пост в Telegram",
        business_id="biz1",
        user_id="user1",
        intent_generator=fake_generator,
    )

    intent = result["intent"]
    assert result["status"] == "compiled_intent"
    assert intent["required_connectors"] == ["google_sheets", "telegram"]
    assert intent["workflow_draft"]["trigger"] == "schedule.daily"
    assert intent["workflow_draft"]["steps"][0]["key"] == "read_orders"
    assert "secret" not in intent["workflow_draft"]["steps"][0]
    assert intent["approval_points"][0]["reason"] == "external publish"
    assert "raw_token" not in intent["approval_points"][0]
    assert intent["unsupported_requests"][0]["reason"] == "LocalOS policy requires approval"


def test_agent_compiler_registry_drives_llm_template_selection():
    from services.agent_compiler_llm import infer_agent_workflow_intent
    from services.agent_compiler_registry import compiled_template_prompt_lines, get_compiled_agent_template

    def fake_generator(prompt, business_id, user_id):
        assert "google_sheets_to_localos_finance" in prompt
        assert "telegram_to_google_sheets" in prompt
        assert "google_sheets_to_telegram_post" in prompt
        return """
        {
          "compiled_template_key": "google_sheets_to_localos_finance",
          "source": "manual",
          "destination": "manual",
          "read_capability": "dangerous.read",
          "write_capability": "dangerous.write",
          "required_connectors": [],
          "approval_reasons": [],
          "limits": {"max_items_per_run": 100},
          "confidence": 0.8
        }
        """

    assert get_compiled_agent_template("google_sheets_to_localos_finance")["write_capability"] == "finance.transaction.create"
    assert any("communication:appointment_reminder" in line for line in compiled_template_prompt_lines())

    result = infer_agent_workflow_intent(
        "Импортируй оплаты из Google Sheets в финансы",
        business_id="biz1",
        user_id="user1",
        intent_generator=fake_generator,
    )

    assert result["status"] == "compiled_intent"
    assert result["intent"]["compiled_template_key"] == "google_sheets_to_localos_finance"
    assert result["intent"]["trigger"] == "schedule.daily"
    assert result["intent"]["source"] == "google_sheets"
    assert result["intent"]["destination"] == "localos_finance"
    assert result["intent"]["read_capability"] == "google_sheets.read_rows"
    assert result["intent"]["write_capability"] == "finance.transaction.create"
    assert result["intent"]["required_connectors"] == ["google_sheets", "localos_finance"]
    assert result["intent"]["approval_reasons"] == ["localos_finance_write", "ambiguous_data"]


def test_openclaw_capability_catalog_normalizes_actions_and_falls_back():
    from services.openclaw_capability_catalog import get_openclaw_capability_catalog

    fallback = get_openclaw_capability_catalog(fetcher=lambda: (_ for _ in ()).throw(RuntimeError("offline")))
    assert fallback["status"] == "fallback"
    assert any(
        action["localos_capability"] == "google_sheets.read_rows"
        and action["openclaw_action_ref"] == "openclaw.google_sheets.read_rows"
        for action in fallback["actions"]
    )

    catalog = get_openclaw_capability_catalog(
        fetcher=lambda: {
            "actions": [
                {
                    "name": "openclaw.custom.tool",
                    "capability": "custom.safe_action",
                    "provider": "custom",
                    "risk": "read",
                    "required_auth": ["custom_auth"],
                }
            ]
        }
    )
    assert catalog["source"] == "openclaw"
    assert catalog["actions"][0]["localos_capability"] == "custom.safe_action"
    assert catalog["actions"][0]["required_auth"] == ["custom_auth"]


def test_openclaw_capability_catalog_normalizes_current_capabilities_catalog_shape():
    from services.openclaw_capability_catalog import get_openclaw_capability_catalog

    catalog = get_openclaw_capability_catalog(
        fetcher=lambda: {
            "success": True,
            "capabilities": {
                "google_sheets.read_rows": {
                    "name": "google_sheets.read_rows",
                    "risk": "external_read",
                    "side_effects": "none; provider reads are resolved by the selected connector",
                    "approval_required": False,
                },
                "google_sheets.read": {
                    "name": "google_sheets.read",
                    "alias_for": "google_sheets.read_rows",
                    "risk": "external_read",
                    "side_effects": "none",
                    "approval_required": False,
                },
                "communications.send_offer": {
                    "name": "communications.send_offer",
                    "risk": "external_send_request",
                    "side_effects": "creates an offer send request only",
                    "approval_required": True,
                },
            },
        }
    )

    by_capability = {action["localos_capability"]: action for action in catalog["actions"]}
    assert catalog["source"] == "openclaw"
    assert catalog["discovery"]["provider_paths_preserved"] is True
    assert by_capability["google_sheets.read_rows"]["openclaw_action_ref"] == "openclaw.google_sheets.read_rows"
    assert by_capability["google_sheets.read_rows"]["required_auth"] == ["google_sheets"]
    assert "provider_candidates" in by_capability["google_sheets.read_rows"]
    assert any(item["provider"] == "native_localos" for item in by_capability["google_sheets.read_rows"]["provider_candidates"])
    assert by_capability["communications.send_offer"]["openclaw_action_ref"] == "openclaw.telegram.publish_message"
    assert by_capability["communications.send_offer"]["approval_class"] == "external_send_request"
    assert any(item["provider"] == "openclaw" for item in by_capability["communications.send_offer"]["provider_candidates"])


def test_openclaw_capability_catalog_can_use_sandbox_bridge_env(monkeypatch):
    from services import openclaw_capability_catalog

    captured = {}

    class FakeResponse:
        def raise_for_status(self):
            return None

        def json(self):
            return {
                "actions": [
                    {
                        "openclaw_action_ref": "openclaw.browser.supervised_publish",
                        "localos_capability": "social.post.publish_supervised_browser",
                        "service": "browser",
                        "status": "available",
                    }
                ]
            }

    def fake_get(url, headers=None, timeout=0):
        captured["url"] = url
        captured["headers"] = headers or {}
        captured["timeout"] = timeout
        return FakeResponse()

    monkeypatch.delenv("OPENCLAW_CAPABILITY_CATALOG_URL", raising=False)
    monkeypatch.delenv("OPENCLAW_BASE_URL", raising=False)
    monkeypatch.delenv("OPENCLAW_LOCALOS_TOKEN", raising=False)
    monkeypatch.delenv("OPENCLAW_TOKEN", raising=False)
    monkeypatch.setenv("OPENCLAW_SANDBOX_BRIDGE_URL", "http://openclaw.local/capabilities")
    monkeypatch.setenv("OPENCLAW_SANDBOX_BRIDGE_TOKEN", "bridge-token")
    monkeypatch.setattr(openclaw_capability_catalog.requests, "get", fake_get)

    catalog = openclaw_capability_catalog.get_openclaw_capability_catalog()

    assert captured["url"] == "http://openclaw.local/capabilities"
    assert captured["headers"]["Authorization"] == "Bearer bridge-token"
    assert catalog["source"] == "openclaw"
    assert catalog["actions"][0]["localos_capability"] == "social.post.publish_supervised_browser"


def test_openclaw_planner_loop_uses_catalog_without_tool_execution():
    from services.agent_openclaw_planner_loop import build_openclaw_planner_loop

    result = build_openclaw_planner_loop(
        {
            "schema": "localos_openclaw_planner_context_v1",
            "allowed_capabilities": ["google_sheets.read_rows", "communications.draft"],
            "required_bindings": [
                {
                    "key": "google_sheets_read",
                    "provider": "google_sheets",
                    "capability": "google_sheets.read_rows",
                    "required_config": ["spreadsheet_id", "sheet_name"],
                },
                {
                    "key": "telegram_delivery",
                    "provider": "telegram",
                    "capability": "communications.draft",
                    "required_config": ["bot_mode"],
                },
            ],
            "connection_state": {
                "missing_connections": [
                    {
                        "key": "google_sheets_read",
                        "provider": "google_sheets",
                        "provider_title": "Google Sheets",
                        "status": "missing",
                        "missing_config": ["spreadsheet_id", "sheet_name"],
                    }
                ],
            },
            "output_contract": {
                "format": "json_only",
                "compiled_workflow_owner": "localos",
            },
            "approval_required_action_classes": ["external_send"],
            "forbidden_action_classes": ["unauthorized_external_system_access"],
        },
        openclaw_catalog={
            "source": "openclaw",
            "actions": [
                {
                    "openclaw_action_ref": "openclaw.google_sheets.read_rows",
                    "localos_capability": "google_sheets.read_rows",
                    "service": "google_sheets",
                    "risk_class": "read",
                    "required_auth": ["google_sheets"],
                    "provider_candidates": [{"provider": "openclaw", "state": "available", "role": "planner_or_connector"}],
                },
                {
                    "openclaw_action_ref": "openclaw.telegram.create_draft",
                    "localos_capability": "communications.draft",
                    "service": "telegram",
                    "risk_class": "draft",
                    "required_auth": ["telegram"],
                    "provider_candidates": [{"provider": "openclaw", "state": "available", "role": "planner_or_connector"}],
                },
            ],
        },
    )

    assert result["schema"] == "localos_openclaw_planner_loop_v1"
    assert result["mode"] == "design_time_only"
    assert result["may_execute_tools"] is False
    assert result["must_compile_in_localos"] is True
    assert result["planner_contract"]["schema"] == "localos_openclaw_planner_contract_v1"
    assert result["planner_contract"]["tool_execution_allowed"] is False
    assert result["planner_contract"]["external_side_effects_allowed"] is False
    assert result["planner_contract"]["compiled_workflow_owner"] == "localos"
    assert result["planner_contract"]["required_response_schema"]["workflow_draft"] == "object"
    assert "execute_tools" in result["planner_contract"]["must_not"]
    assert result["status"] == "needs_clarification"
    assert result["clarifying_questions"][0]["key"] == "connect_google_sheets"
    assert any(item["key"] == "google_sheets_target" for item in result["clarifying_questions"])
    assert result["workflow_proposal"]["policy"] == "localos_envelope"
    assert {"capability": "google_sheets.read_rows", "provider_path": "openclaw:available"} in result["workflow_proposal"]["provider_paths"]
    assert result["workflow_proposal"]["openclaw_action_refs"] == [
        "openclaw.google_sheets.read_rows",
        "openclaw.telegram.create_draft",
    ]


def test_openclaw_planner_loop_requires_workflow_details_before_draft():
    from services.agent_openclaw_planner_loop import build_openclaw_planner_loop

    result = build_openclaw_planner_loop(
        {
            "schema": "localos_openclaw_planner_context_v1",
            "task": "Сделай агента: бери заказ из Google Sheets и делай пост в Telegram",
            "allowed_capabilities": ["google_sheets.read_rows", "communications.draft"],
            "required_bindings": [
                {
                    "key": "google_sheets_read",
                    "provider": "google_sheets",
                    "capability": "google_sheets.read_rows",
                    "required_config": ["spreadsheet_id", "sheet_name"],
                },
                {
                    "key": "telegram_delivery",
                    "provider": "telegram",
                    "capability": "communications.draft",
                    "required_config": ["bot_mode"],
                },
            ],
            "connection_state": {},
            "connection_answer_bindings": {},
            "output_contract": {
                "format": "json_only",
                "compiled_workflow_owner": "localos",
            },
        }
    )

    questions = {item["key"]: item for item in result["clarifying_questions"]}

    assert result["status"] == "needs_clarification"
    assert result["may_execute_tools"] is False
    assert questions["google_sheets_target"]["reason"] == "openclaw_workflow_detail_missing"
    assert questions["telegram_target"]["reason"] == "openclaw_workflow_detail_missing"
    assert questions["schedule_frequency"]["reason"] == "openclaw_workflow_detail_missing"
    assert questions["post_format"]["reason"] == "openclaw_workflow_detail_missing"
    assert "blocking_workflow_details_before_draft" in result["planner_contract"]["must_return"]


def test_openclaw_planner_loop_understands_review_schedule_and_localos_telegram_bot():
    from services.agent_openclaw_planner_loop import build_openclaw_planner_loop

    result = build_openclaw_planner_loop(
        {
            "schema": "localos_openclaw_planner_context_v1",
            "task": (
                "Создай агента, который каждю среду в 9 утра проверяет наличие новых отзывов - "
                "запускает парсер. Если они есть, то генерирует ответ. Оба - отзыв и ответ "
                "присылает мне в телеграм через бота"
            ),
            "allowed_capabilities": ["reviews.fetch", "communications.draft"],
            "required_bindings": [
                {
                    "key": "reviews_source",
                    "provider": "localos_reviews",
                    "capability": "reviews.fetch",
                    "required_config": [],
                },
                {
                    "key": "telegram_delivery",
                    "provider": "telegram",
                    "capability": "communications.draft",
                    "required_config": ["bot_mode"],
                },
            ],
            "connection_state": {},
            "connection_answer_bindings": {},
            "output_contract": {
                "format": "json_only",
                "compiled_workflow_owner": "localos",
            },
        }
    )

    questions = {item["key"]: item for item in result["clarifying_questions"]}

    assert "telegram_target" not in questions
    assert "schedule_frequency" not in questions
    assert "post_format" not in questions


def test_agent_builder_does_not_treat_scheduled_telegram_message_as_outreach():
    from services.agent_builder_session import build_agent_builder_state
    from services.agent_blueprint_draft_builder import compile_agent_blueprint, infer_blueprint_category

    prompt = 'Сощздай агента, который каждое утро в 9 утра шлёт мне сообщение "Привет" в телеграм'

    draft = compile_agent_blueprint(prompt)
    state = build_agent_builder_state([{"role": "user", "content": prompt}])
    questions_text = " ".join(str(item.get("question") or "") for item in state["missing_questions"]).lower()

    assert infer_blueprint_category(prompt) == "custom"
    assert draft["category"] == "custom"
    assert draft["version_payload"]["trigger"] == "schedule.daily"
    assert draft["version_payload"]["schedule"]["time"] == "09:00"
    assert state["category"] == "custom"
    assert "лид" not in questions_text
    assert "prospectingleads" not in questions_text


def test_agent_builder_understands_core_user_scenarios_without_cross_domain_questions():
    from services.agent_builder_session import build_agent_builder_state

    scenarios = [
        (
            "reviews_to_telegram",
            "Проверять новые отзывы каждый день, готовить черновик ответа и присылать отзыв + ответ владельцу в Telegram.",
            "custom",
            ["external_reviews", "telegram"],
            ["telegram_destination"],
            ["лид", "prospectingleads", "где искать клиентов"],
        ),
        (
            "daily_reminder",
            "Каждое утро в 9:00 отправлять владельцу короткое сообщение или чеклист дня в Telegram.",
            "custom",
            ["manual_context", "telegram"],
            ["telegram_destination"],
            ["лид", "prospectingleads", "где искать клиентов"],
        ),
        (
            "sheets_to_telegram",
            "Раз в день брать новую строку из Google Sheets и отправлять по ней краткое сообщение в Telegram.",
            "custom",
            ["google_sheets", "telegram"],
            ["google_sheets_target"],
            ["лид", "prospectingleads", "где искать клиентов"],
        ),
        (
            "orders_without_status",
            "Проверять таблицу заказов, находить заказы без статуса или ответственного и присылать список менеджеру.",
            "tables",
            ["uploaded_tables"],
            [],
            ["лид", "prospectingleads"],
        ),
        (
            "negative_reviews",
            "Отслеживать отзывы с оценкой 1-3, срочно уведомлять владельца и готовить аккуратный черновик ответа без обещаний скидок.",
            "reviews",
            ["external_reviews"],
            [],
            ["агент услуг", "prospectingleads", "где искать клиентов"],
        ),
        (
            "map_content_plan",
            "Раз в неделю предлагать 3 темы постов для карточек на картах на основе услуг, сезона и отзывов.",
            "custom",
            ["services", "external_reviews", "business_profile"],
            [],
            ["где искать клиентов", "какие лиды"],
        ),
        (
            "services_check",
            "Раз в неделю смотреть услуги в карточке бизнеса и находить пустые описания, плохие названия или отсутствующие цены.",
            "services",
            ["services"],
            [],
            ["где искать клиентов", "какие лиды"],
        ),
        (
            "finance_import",
            "Читать таблицу расходов, находить новые строки, нормализовать категории и готовить их к добавлению в финансы LocalOS после подтверждения.",
            "custom",
            ["google_sheets", "localos_finance"],
            ["google_sheets_target"],
            ["где искать клиентов", "какие лиды"],
        ),
        (
            "partner_search",
            "Найти потенциальных партнёров в городе, собрать shortlist, подготовить первое сообщение и ждать ручного подтверждения перед отправкой.",
            "partnerships",
            ["prospectingleads", "services"],
            [],
            ["где искать клиентов", "какие лиды"],
        ),
        (
            "booking_control",
            "Каждый день проверять ближайшие записи клиентов и готовить напоминания, но отправлять только после подтверждения человека.",
            "communications",
            ["appointments"],
            [],
            ["где искать клиентов", "prospectingleads"],
        ),
    ]

    for key, prompt, expected_category, expected_sources, expected_question_keys, forbidden_question_terms in scenarios:
        state = build_agent_builder_state([{"role": "user", "content": prompt}])
        preview = state["preview"]
        questions = state["missing_questions"]
        question_keys = {str(item.get("key") or "") for item in questions}
        questions_text = " ".join(str(item.get("question") or "") for item in questions).lower()
        sources = set(preview["data_sources"])

        assert state["category"] == expected_category, key
        for source in expected_sources:
            assert source in sources, key
        for question_key in expected_question_keys:
            assert question_key in question_keys, key
        for term in forbidden_question_terms:
            assert term not in questions_text, key


def test_agent_builder_understands_second_browser_scenario_pack_without_wrong_domains():
    from services.agent_builder_session import build_agent_builder_state

    scenarios = [
        (
            "overdue_invoices",
            "Каждый день проверяй неоплаченные счета, находи просроченные больше чем на 3 дня и присылай владельцу список в Telegram.",
            "custom",
            ["localos_finance", "telegram"],
            ["просроченных", "счет"],
            ["где искать клиентов", "какие лиды", "prospectingleads", "формат поста"],
        ),
        (
            "empty_customer_cards",
            "Раз в неделю находи клиентов без телефона, email или источника прихода и готовь список для менеджера.",
            "custom",
            ["clients"],
            ["пустыми полями", "телефон"],
            ["google", "таблиц", "письм", "лид", "prospectingleads"],
        ),
        (
            "expense_control",
            "Каждый вечер проверяй новые расходы в LocalOS, выделяй подозрительно крупные траты и проси владельца подтвердить категорию.",
            "custom",
            ["localos_finance"],
            ["подозрительных расходов", "категор"],
            ["где искать клиентов", "какие лиды", "prospectingleads"],
        ),
        (
            "bookings_no_prepayment",
            "Каждое утро проверяй записи на завтра, находи клиентов без предоплаты и готовь напоминание администратору.",
            "communications",
            ["appointments"],
            ["Черновики сообщений"],
            ["где искать клиентов", "prospectingleads"],
        ),
        (
            "weak_reviews_locations",
            "Раз в неделю сравнивай отзывы по всем точкам сети, находи филиалы с падением рейтинга и присылай короткий разбор.",
            "reviews",
            ["external_reviews", "locations"],
            ["филиалам", "рейтинг"],
            ["черновики ответов", "где искать клиентов", "prospectingleads"],
        ),
        (
            "content_from_reviews",
            "Каждую неделю бери новые положительные отзывы и предлагай 3 идеи постов на их основе для карточек на картах.",
            "custom",
            ["external_reviews", "services"],
            ["3 идеи", "положительных отзывов"],
            ["черновики ответов", "где искать клиентов", "prospectingleads"],
        ),
        (
            "duplicate_services",
            "Раз в неделю проверяй список услуг и находи дубли, похожие названия и услуги без категории.",
            "services",
            ["services"],
            ["Проверка услуг"],
            ["где искать клиентов", "prospectingleads"],
        ),
        (
            "old_clients_reactivation",
            "Каждый понедельник находи клиентов, которые не записывались больше 60 дней, и готовь мягкое сообщение для возврата. Не отправляй без подтверждения.",
            "communications",
            ["appointments"],
            ["клиентов"],
            ["telegram", "телеграм", "где искать клиентов", "prospectingleads"],
        ),
        (
            "partner_replies",
            "Каждый день проверяй ответы потенциальных партнёров, классифицируй их как интересно / отказ / нужен ручной ответ и показывай следующий шаг.",
            "partnerships",
            ["prospectingleads", "outreach_drafts"],
            ["ответов партнёров", "следующий шаг"],
            ["google", "таблиц", "агент отзывов", "черновики ответов"],
        ),
        (
            "daily_problem_digest",
            "Каждое утро собирай один короткий дайджест: новые негативные отзывы, отменённые записи, просроченные задачи и необычные расходы. Присылай в Telegram.",
            "custom",
            ["localos_digest", "telegram"],
            ["ежедневный дайджест", "проблем"],
            ["где искать клиентов", "prospectingleads", "формат поста"],
        ),
    ]

    for key, prompt, expected_category, expected_sources, output_terms, forbidden_terms in scenarios:
        state = build_agent_builder_state([{"role": "user", "content": prompt}], use_ai=True)
        preview = state["preview"]
        questions_text = " ".join(str(item.get("question") or "") for item in state["missing_questions"]).lower()
        surface_text = " ".join(
            [
                state["category"],
                preview["category_label"],
                ", ".join(preview["data_sources"]),
                preview["extraction_rules"],
                preview["output_format"],
                questions_text,
            ]
        ).lower()
        sources = set(preview["data_sources"])

        assert state["category"] == expected_category, key
        for source in expected_sources:
            assert source in sources, key
        for term in output_terms:
            assert term.lower() in surface_text, key
        for term in forbidden_terms:
            assert term.lower() not in questions_text, key


def test_agent_builder_understands_third_browser_scenario_pack_without_generic_outputs():
    from services.agent_builder_session import build_agent_builder_state

    scenarios = [
        (
            "photo_quality",
            "Раз в неделю проверяй фотографии в карточках филиалов, находи устаревшие, тёмные или нерелевантные фото и предлагай, что заменить.",
            "custom",
            ["business_cards", "photos", "locations"],
            ["фото", "замен"],
            ["лид", "prospectingleads", "подозрительных расходов", "черновики ответов"],
        ),
        (
            "competitor_prices",
            "Каждый понедельник сравнивай цены на ключевые услуги с конкурентами поблизости и присылай краткий список, где мы выше или ниже рынка.",
            "custom",
            ["services", "competitors"],
            ["конкурент", "выше или ниже"],
            ["проверка услуг", "черновики ответов", "где искать клиентов"],
        ),
        (
            "cancellation_risk",
            "Каждое утро находи записи клиентов, которые часто отменяют визиты, и готовь администратору список для ручного подтверждения.",
            "communications",
            ["appointments"],
            ["риском отмены", "администратор"],
            ["доставка", "статусы реакции", "prospectingleads"],
        ),
        (
            "new_services_control",
            "Когда в LocalOS появляется новая услуга, проверяй название, описание, цену и готовь улучшенную версию для карточек.",
            "services",
            ["services"],
            ["улучшенная версия", "новой услуги"],
            ["какой результат", "отзывы"],
        ),
        (
            "customer_questions_monitoring",
            "Каждый день собирай новые вопросы клиентов из Telegram/WhatsApp, группируй по темам и предлагай ответы для базы знаний.",
            "custom",
            ["telegram", "whatsapp", "customer_questions"],
            ["вопросов клиентов", "базы знаний"],
            ["какие темы вопросов", "prospectingleads", "финансы"],
        ),
        (
            "team_tasks_check",
            "Каждое утро находи просроченные задачи сотрудников и присылай владельцу короткий список: задача, ответственный, срок, следующий шаг.",
            "custom",
            ["localos_tasks", "team"],
            ["просроченных задач", "ответственный"],
            ["какие данные", "финансы localos", "prospectingleads"],
        ),
        (
            "no_discount_promos",
            "Раз в неделю предлагай 3 идеи продвижения без скидок на основе сезонности, услуг и отзывов клиентов.",
            "custom",
            ["services", "external_reviews", "seasonality"],
            ["без скидок", "3 идеи"],
            ["черновики ответов", "проверка услуг", "prospectingleads"],
        ),
        (
            "repeated_complaints",
            "Если в отзывах или сообщениях повторяется одна и та же проблема, собери примеры и предложи, что изменить в сервисе.",
            "custom",
            ["external_reviews", "customer_messages", "services"],
            ["повторяющиеся жалобы", "изменить в сервисе"],
            ["где человек должен проверить", "черновики ответов", "prospectingleads"],
        ),
        (
            "manager_report",
            "Каждую пятницу собирай отчёт по филиалам: отзывы, записи, выручка, расходы, проблемы и рекомендации на следующую неделю.",
            "custom",
            ["external_reviews", "appointments", "localos_finance", "locations"],
            ["отчёт по филиалам", "рекомендации"],
            ["подозрительных расходов", "проверка услуг", "prospectingleads"],
        ),
        (
            "holiday_readiness",
            "За две недели до праздников проверяй карточки, услуги, посты и расписание, находи пробелы и готовь чеклист подготовки.",
            "custom",
            ["business_cards", "services", "posts", "schedule"],
            ["чеклист подготовки", "праздникам"],
            ["проверка услуг", "черновики ответов", "prospectingleads"],
        ),
    ]

    for key, prompt, expected_category, expected_sources, output_terms, forbidden_terms in scenarios:
        state = build_agent_builder_state([{"role": "user", "content": prompt}], use_ai=True)
        preview = state["preview"]
        questions_text = " ".join(str(item.get("question") or "") for item in state["missing_questions"]).lower()
        surface_text = " ".join(
            [
                state["category"],
                preview["category_label"],
                ", ".join(preview["data_sources"]),
                preview["extraction_rules"],
                preview["output_format"],
                questions_text,
            ]
        ).lower()
        sources = set(preview["data_sources"])

        assert state["category"] == expected_category, key
        for source in expected_sources:
            assert source in sources, key
        for term in output_terms:
            assert term.lower() in surface_text, key
        for term in forbidden_terms:
            assert term.lower() not in surface_text, key


def test_agent_builder_understands_fourth_browser_scenario_pack_without_generic_outputs():
    from services.agent_builder_session import build_agent_builder_state

    scenarios = [
        (
            "inventory_control",
            "Каждый вечер проверяй остатки расходников и товаров, находи позиции ниже минимума и готовь список, что заказать.",
            "custom",
            ["inventory", "products", "supplies"],
            ["список для закупки", "сколько заказать"],
            ["подозрительных расходов", "финансы localos", "prospectingleads", "какой результат"],
        ),
        (
            "staff_schedule_check",
            "Раз в неделю проверяй расписание смен, находи пересечения, пустые окна и перегрузки по сотрудникам.",
            "custom",
            ["staff_schedule", "team"],
            ["расписании смен", "перегруз"],
            ["готовый результат по задаче", "prospectingleads", "какой результат"],
        ),
        (
            "cancellation_reasons",
            "Каждую неделю собирай отменённые записи, группируй причины и предлагай, что изменить в процессе записи.",
            "custom",
            ["appointments", "clients"],
            ["причин отмен", "процессе записи"],
            ["риск отмены", "черновики сообщений", "статусы реакции"],
        ),
        (
            "admin_response_control",
            "Каждый день проверяй чаты с клиентами и находи диалоги, где администратор долго не ответил или ответил неполно.",
            "custom",
            ["customer_chats", "team"],
            ["администратор", "долго не ответил"],
            ["готовый результат по задаче", "prospectingleads"],
        ),
        (
            "faq_from_chats",
            "Раз в неделю бери повторяющиеся вопросы из клиентских переписок и предлагай новые пункты для FAQ на сайте или в карточке.",
            "custom",
            ["customer_chats", "customer_questions", "business_cards"],
            ["пункты faq", "клиентских переписок"],
            ["какой результат", "telegram", "whatsapp", "prospectingleads"],
        ),
        (
            "new_employee_check",
            "Когда добавляется новый сотрудник, проверяй, заполнены ли фото, описание, услуги, график и привязка к филиалу.",
            "custom",
            ["team", "staff_profiles", "services", "schedule", "locations"],
            ["нового сотрудника", "график"],
            ["проблемных фото", "какой результат", "prospectingleads"],
        ),
        (
            "seasonal_services",
            "Раз в месяц проверяй, какие сезонные услуги пора добавить, скрыть или обновить в карточках и прайсе.",
            "services",
            ["services", "seasonality", "business_cards", "price_list"],
            ["сезонных услуг", "добавить, скрыть или обновить"],
            ["что агент должен понять", "проверка услуг", "prospectingleads"],
        ),
        (
            "revenue_anomalies",
            "Каждое утро сравнивай вчерашнюю выручку с обычным уровнем по дню недели и присылай владельцу резкие отклонения.",
            "custom",
            ["localos_finance", "revenue"],
            ["отклонений выручки", "обычный уровень"],
            ["готовый результат по задаче", "подозрительных расходов", "какой результат", "prospectingleads"],
        ),
        (
            "map_questions_answers",
            "Каждый день проверяй новые вопросы пользователей в Яндекс/Google-карточках и готовь ответы для ручного подтверждения.",
            "custom",
            ["business_cards", "map_questions"],
            ["ответов на вопросы", "яндекс/google-карточках"],
            ["готовый результат по задаче", "черновики ответов и причины ручной проверки", "prospectingleads"],
        ),
        (
            "location_description_quality",
            "Раз в неделю проверяй описания всех филиалов, находи устаревшую информацию, одинаковые тексты и слабые формулировки.",
            "custom",
            ["locations", "business_cards", "location_descriptions"],
            ["описаниях филиалов", "слабые формулировки"],
            ["готовый результат по задаче", "prospectingleads", "какой результат"],
        ),
    ]

    for key, prompt, expected_category, expected_sources, output_terms, forbidden_terms in scenarios:
        state = build_agent_builder_state([{"role": "user", "content": prompt}], use_ai=True)
        preview = state["preview"]
        questions_text = " ".join(str(item.get("question") or "") for item in state["missing_questions"]).lower()
        surface_text = " ".join(
            [
                state["category"],
                preview["category_label"],
                ", ".join(preview["data_sources"]),
                preview["extraction_rules"],
                preview["output_format"],
                questions_text,
            ]
        ).lower()
        sources = set(preview["data_sources"])

        assert state["category"] == expected_category, key
        for source in expected_sources:
            assert source in sources, key
        for term in output_terms:
            assert term.lower() in surface_text, key
        for term in forbidden_terms:
            assert term.lower() not in surface_text, key


def test_agent_builder_new_50_plus_legacy_40_scenario_corpus_is_complete():
    from tests.agent_builder_scenario_corpus import (
        AGENT_BUILDER_REGRESSION_PLAN,
        LEGACY_AGENT_BUILDER_COMBINATION_PACKS,
        LEGACY_BROWSER_SCENARIO_COUNT,
        MIXED_AGENT_BUILDER_COMBINATIONS,
        NEW_AGENT_BUILDER_SCENARIOS,
    )

    keys = [str(item["key"]) for item in NEW_AGENT_BUILDER_SCENARIOS]
    combination_keys = {key for _, items in MIXED_AGENT_BUILDER_COMBINATIONS for key in items}
    legacy_combination_count = sum(int(pack["scenario_count"]) for pack in LEGACY_AGENT_BUILDER_COMBINATION_PACKS)
    legacy_combination_names = {
        name
        for pack in LEGACY_AGENT_BUILDER_COMBINATION_PACKS
        for name in pack["combinations"]
    }

    assert LEGACY_BROWSER_SCENARIO_COUNT == 40
    assert len(NEW_AGENT_BUILDER_SCENARIOS) == 50
    assert len(set(keys)) == 50
    assert LEGACY_BROWSER_SCENARIO_COUNT + len(NEW_AGENT_BUILDER_SCENARIOS) == 90
    assert len(MIXED_AGENT_BUILDER_COMBINATIONS) >= 10
    assert combination_keys.issubset(set(keys))
    assert len(LEGACY_AGENT_BUILDER_COMBINATION_PACKS) == 4
    assert legacy_combination_count == 40
    assert len(legacy_combination_names) == 40
    assert AGENT_BUILDER_REGRESSION_PLAN["new_scenario_count"] == 50
    assert AGENT_BUILDER_REGRESSION_PLAN["legacy_scenario_count"] == 40
    assert AGENT_BUILDER_REGRESSION_PLAN["total_scenario_count"] == 90
    assert AGENT_BUILDER_REGRESSION_PLAN["new_combinations"] == MIXED_AGENT_BUILDER_COMBINATIONS
    assert AGENT_BUILDER_REGRESSION_PLAN["legacy_combinations"] == LEGACY_AGENT_BUILDER_COMBINATION_PACKS

    all_sources = {source for item in NEW_AGENT_BUILDER_SCENARIOS for source in item["expected_sources"]}
    for source in ["browser_use", "google_sheets", "telegram", "whatsapp", "external_reviews", "localos_finance"]:
        assert source in all_sources


def test_agent_builder_understands_new_50_mixed_scenarios_without_wrong_domains():
    from services.agent_builder_session import build_agent_builder_state
    from tests.agent_builder_scenario_corpus import NEW_AGENT_BUILDER_SCENARIOS

    for item in NEW_AGENT_BUILDER_SCENARIOS:
        key = str(item["key"])
        state = build_agent_builder_state([{"role": "user", "content": item["prompt"]}], use_ai=True)
        preview = state["preview"]
        questions_text = " ".join(str(question.get("question") or "") for question in state["missing_questions"]).lower()
        surface_text = " ".join(
            [
                state["category"],
                preview["category_label"],
                ", ".join(preview["data_sources"]),
                preview["extraction_rules"],
                preview["output_format"],
                questions_text,
            ]
        ).lower()
        sources = set(preview["data_sources"])

        assert state["category"] == item["expected_category"], key
        for source in item["expected_sources"]:
            assert source in sources, key
        for term in item["expected_terms"]:
            assert str(term).lower() in surface_text, key
        for term in item["forbidden_terms"]:
            assert str(term).lower() not in surface_text, key


def test_agent_builder_keeps_real_user_scenarios_on_the_obvious_next_step():
    from services.agent_builder_session import build_agent_builder_state

    scenarios = [
        (
            "sales_to_finance",
            "Каждый вечер проверяй Google-таблицу с продажами, находи новые строки и готовь их к добавлению во вкладку Финансы. Перед внесением показывай мне список на подтверждение.",
            ["что агент должен понять", "что нужно извлечь"],
            ["google_sheets_target"],
        ),
        (
            "telegram_content_reactions",
            "После публикации поста в Telegram проверяй реакции и комментарии через API, собирай выводы и предлагай, что изменить в следующем контент-плане.",
            ["кто будет принимать решение", "где человек должен проверить"],
            [],
        ),
        (
            "negative_review_event",
            "Если появляется отзыв с оценкой 1-3, сразу присылай мне уведомление в Telegram, кратко объясняй проблему клиента и предлагай аккуратный ответ без обещаний скидок.",
            ["когда запускать агента"],
            [],
        ),
        (
            "weekly_owner_report",
            "Каждую пятницу собирай краткий отчёт: новые отзывы, продажи, расходы, записи, проблемы в карточке и что нужно сделать на следующей неделе. Присылай в Telegram.",
            ["в какой telegram", "когда запускать агента", "какой формат поста"],
            [],
        ),
        (
            "answered_review_drafts",
            "Агент должен парсить отзывы каждую среду в 9 утра. Все отображать в аккаунте ЛокалОС. Если появляются новые, то генерировать ответ и оповещать меня в телеграмме + присылать отзыв и ответ\nОтдельные черновики человек проверяет в телегираме - по оповещению",
            ["нужны отдельные черновики ответов или общий план реакции"],
            [],
        ),
    ]

    for key, prompt, forbidden_fragments, expected_question_keys in scenarios:
        state = build_agent_builder_state([{"role": "user", "content": prompt}])
        questions = state["missing_questions"]
        question_keys = {str(item.get("key") or "") for item in questions}
        questions_text = " ".join(str(item.get("question") or "") for item in questions).lower()

        for fragment in forbidden_fragments:
            assert fragment not in questions_text, key
        for question_key in expected_question_keys:
            assert question_key in question_keys, key

    ai_state = build_agent_builder_state(
        [
            {
                "role": "user",
                "content": "После публикации поста в Telegram проверяй реакции и комментарии через API, собирай выводы и предлагай, что изменить в следующем контент-плане.",
            }
        ],
        use_ai=True,
    )
    ai_questions_text = " ".join(str(item.get("question") or "") for item in ai_state["missing_questions"]).lower()

    assert ai_state["category"] == "custom"
    assert "telegram" in set(ai_state["preview"]["data_sources"])
    assert ai_state["preview"]["feasibility"]["status"] != "forbidden"
    assert "какие данные агент должен использовать" not in ai_questions_text
    assert "в какой telegram" not in ai_questions_text


def test_agent_feasibility_resolver_reports_ready_missing_choice_and_forbidden():
    from services.agent_feasibility_resolver import resolve_agent_feasibility

    required_bindings = [
        {
            "key": "google_sheets_read",
            "provider": "google_sheets",
            "capability": "google_sheets.read_rows",
            "required_config": ["spreadsheet_id", "sheet_name"],
        },
        {
            "key": "telegram_delivery",
            "provider": "telegram",
            "capability": "communications.draft",
            "required_config": ["bot_mode"],
        },
    ]
    missing = resolve_agent_feasibility(
        description="Возьми заказ из Google Sheets и подготовь пост в Telegram",
        required_capabilities=["google_sheets.read_rows", "communications.draft"],
        required_bindings=required_bindings,
        connected_integrations=[
            {
                "id": "telegram-1",
                "provider": "telegram",
                "status": "active",
                "display_name": "Business bot",
                "config": {"bot_mode": "business_bot"},
            }
        ],
    )
    assert missing["status"] == "needs_connection"
    assert missing["ready"] is False
    assert [item["provider"] for item in missing["missing_connections"]] == ["google_sheets"]
    assert missing["missing_connections"][0]["route_state"] == "available"
    assert any(item["provider"] == "native_localos" for item in missing["missing_connections"][0]["provider_routes"])
    assert missing["ready_bindings"][0]["provider"] == "telegram"
    assert missing["ready_bindings"][0]["route_state"] == "connected"
    assert missing["capabilities"][0]["route_state"] == "available"
    assert any(item["provider"] == "openclaw" for item in missing["capabilities"][1]["provider_routes"])

    choice = resolve_agent_feasibility(
        required_capabilities=["google_sheets.read_rows"],
        required_bindings=[required_bindings[0]],
        connected_integrations=[
            {
                "id": "sheet-1",
                "provider": "google_sheets",
                "status": "active",
                "display_name": "Orders A",
                "config": {"spreadsheet_id": "a", "sheet_name": "Orders"},
            },
            {
                "id": "sheet-2",
                "provider": "google_sheets",
                "status": "active",
                "display_name": "Orders B",
                "config": {"spreadsheet_id": "b", "sheet_name": "Orders"},
            },
        ],
    )
    assert choice["status"] == "needs_choice"
    assert choice["connection_choices"][0]["connection_count"] == 2

    forbidden = resolve_agent_feasibility(
        description="Подключись к компьютерам Роскосмоса и забери данные",
        required_capabilities=["unknown.external_access"],
    )
    assert forbidden["status"] == "forbidden"
    assert forbidden["forbidden"][0]["term"] == "роскосмос"


def test_agent_feasibility_resolver_blocks_maton_until_api_key_connection_exists():
    from services.agent_feasibility_resolver import resolve_agent_feasibility

    result = resolve_agent_feasibility(
        description="Отправляй сообщения через Maton",
        required_capabilities=["communications.send_offer"],
        required_bindings=[
            {
                "key": "maton_delivery",
                "provider": "maton",
                "capability": "communications.send_offer",
                "required_config": ["channel"],
            }
        ],
        connected_integrations=[],
    )

    assert result["status"] == "needs_connection"
    assert result["missing_connections"][0]["provider"] == "maton"
    assert result["capabilities"][0]["status"] == "supported"
    assert any(action["service"] == "maton" for action in result["capabilities"][0]["openclaw_actions"])


def test_communication_agent_showcase_has_five_safe_mvp_blueprints():
    from services.agent_blueprint_draft_builder import build_communication_agent_showcase_blueprints

    drafts = build_communication_agent_showcase_blueprints()
    expected = {
        "appointment_reminder": ("appointment.reminder.before", "communications.send_reminder", "approved_batch_only"),
        "post_visit_followup": ("visit.completed.after", "communications.send_reminder", "approved_batch_only"),
        "inactive_client_winback": ("client.inactive.since", "communications.send_offer", "approved_batch_only"),
        "package_offer_after_service": ("service.completed.relevant", "communications.send_offer", "approved_batch_only"),
        "inbound_request_reply_draft": ("inbound.message.received", "communications.draft", "draft_only"),
    }

    assert len(drafts) == 5
    by_key = {draft["metadata"]["communication_template_key"]: draft for draft in drafts}
    assert set(by_key) == set(expected)

    for key, values in expected.items():
        trigger, capability, mode = values
        draft = by_key[key]
        payload = draft["version_payload"]
        steps = payload["steps"]

        assert draft["category"] == "communications"
        assert payload["trigger"] == trigger
        assert payload["send_capability"] == capability
        assert payload["mode"] == mode
        assert payload["audience_rules"]
        assert payload["consent_rules"]
        assert payload["message_template"]
        assert payload["persona"]
        assert payload["delivery_outcome_journal"]["external_dispatch_performed"] is False
        assert payload["limits"]["external_send_requires_approval"] is True
        assert payload["limits"]["autonomous_send_allowed"] is False
        assert payload["external_dispatch_performed"] is False
        assert draft["metadata"]["compiled_artifact_candidate"]["schema"] == "localos_compiled_artifact_candidate_v1"
        assert draft["metadata"]["compiled_artifact_candidate"]["status"] == "validation_passed"
        assert draft["metadata"]["compiled_validation"]["valid"] is True
        assert draft["metadata"]["compiled_process"]["schema"] == "compiled_communications_workflow_v1"
        assert "communications.draft" in payload["capability_allowlist"]
        if capability != "communications.draft":
            assert capability in payload["capability_allowlist"]
            send_step = [step for step in steps if step["key"] == "send_message"][0]
            assert send_step["type"] == "capability"
            assert send_step["requires_approval"] is True
            assert send_step["payload"]["external_dispatch_performed"] is False
        else:
            assert payload["capability_allowlist"] == ["appointments.read", "communications.draft"]
            send_step = [step for step in steps if step["key"] == "send_message"][0]
            assert send_step["type"] == "artifact"
            assert send_step["payload"]["delivery_state"] == "not_dispatched"


def test_communication_agent_compiler_selects_mvp_templates_from_text():
    from services.agent_blueprint_draft_builder import compile_agent_blueprint

    examples = [
        ("Сделай сообщение после визита", "post_visit_followup"),
        ("Вернуть клиента, который давно не был", "inactive_client_winback"),
        ("Пакетное предложение после релевантной услуги", "package_offer_after_service"),
        ("Черновик ответа на входящий запрос", "inbound_request_reply_draft"),
    ]

    for prompt, expected_key in examples:
        draft = compile_agent_blueprint(prompt)
        assert draft["category"] == "communications"
        assert draft["metadata"]["communication_template_key"] == expected_key


def test_agent_product_view_uses_aiagent_as_voice_persona():
    from services.agent_product_layer import (
        attach_persona_to_version,
        attach_product_agent_to_blueprint,
        parse_persona_row,
    )

    persona = parse_persona_row(
        {
            "id": "voice-1",
            "name": "Администратор Анна",
            "type": "communication",
            "description": "Голос администратора",
            "personality": "спокойная и внимательная",
            "identity": "администратор салона",
            "speech_style": "коротко и дружелюбно",
            "restrictions_json": "{\"no_promises\": true}",
            "variables_json": "{\"signature\": \"Анна\"}",
            "is_active": 1,
        }
    )
    personas = {"voice-1": persona}
    version = attach_persona_to_version(
        {
            "id": "version-1",
            "version_number": 2,
            "persona_agent_id": "voice-1",
        },
        personas,
    )
    blueprint = attach_product_agent_to_blueprint(
        {
            "id": "blueprint-1",
            "name": "Напоминания о записи",
            "category": "communications",
            "status": "draft",
            "metadata_json": "{\"compiler\": \"agent_compiler_v1\"}",
        },
        version,
        personas,
    )

    assert version["persona"]["source"] == "AIAgents"
    assert version["persona"]["role"] == "agent_voice"
    assert blueprint["product_agent"]["kind"] == "agent"
    assert blueprint["product_agent"]["source"] == "agent_blueprints"
    assert blueprint["product_agent"]["persona_agent_id"] == "voice-1"
    assert blueprint["product_agent"]["voice"]["name"] == "Администратор Анна"
    assert blueprint["product_agent"]["components"]["persona"]["role"] == "agent_voice"
    assert blueprint["product_agent"]["legacy"]["communication_agent_is_blueprint_category"] is True


def test_agent_builder_session_understands_document_task_and_asks_questions():
    from services.agent_builder_session import build_agent_builder_state

    state = build_agent_builder_state(
        [{"role": "user", "content": "Нужен агент, который проверяет договоры и ищет риски"}],
    )

    assert state["category"] == "documents"
    assert state["preview"]["category"] == "documents"
    assert "Понял задачу" in state["messages"][-1]["content"]
    assert state["preview"]["external_dispatch_performed"] is False
    assert state["missing_questions"]
    assert any("документ" in item["question"].lower() for item in state["missing_questions"])


def test_agent_builder_session_reduces_questions_after_clarification():
    from services.agent_builder_session import append_user_message, build_agent_builder_state

    messages = [{"role": "user", "content": "Сделай агента"}]
    initial = build_agent_builder_state(messages)
    clarified_messages = append_user_message(
        initial["messages"],
        "Он проверяет договоры из DOCX, извлекает суммы, сроки и риски, результат нужен как краткий отчёт, человек проверяет итог.",
    )
    clarified = build_agent_builder_state(clarified_messages)

    assert clarified["category"] == "documents"
    assert len(clarified["missing_questions"]) < len(initial["missing_questions"])
    assert clarified["preview"]["output_format"]


def test_agent_builder_session_does_not_repeat_review_draft_question_after_answer():
    from services.agent_builder_session import append_user_message, build_agent_builder_state

    messages = [
        {
            "role": "user",
            "content": (
                "Агент должен парсить отзывы каждую среду в 9 утра. Все отображать в аккаунте ЛокалОС. "
                "Если появляются новые, то генерировать ответ и оповещать меня в телеграмме + присылать отзыв и ответ"
            ),
        }
    ]
    initial = build_agent_builder_state(messages)
    clarified_messages = append_user_message(
        initial["messages"],
        "Отдельные чероновики человек проверяет в телегираме - по оповещению",
    )
    clarified = build_agent_builder_state(clarified_messages)

    repeated_questions = [
        item["question"]
        for item in clarified["missing_questions"]
        if "отдельные черновики ответов или общий план реакции" in item["question"].lower()
    ]
    assert repeated_questions == []
    assert all(item["key"] != "output" for item in clarified["missing_questions"])


def test_agent_builder_session_preview_includes_feasibility_for_required_connectors():
    from services.agent_builder_session import build_agent_builder_state

    state = build_agent_builder_state(
        [
            {
                "role": "user",
                "content": (
                    "Мне нужен агент, который из Google таблицы со списком заказов берёт один заказ "
                    "за предыдущий день и создаёт пост в Telegram. Человек проверяет перед публикацией. "
                    "Формат: короткий текст в стиле довольных пассажиров. "
                    "Таблица https://docs.google.com/spreadsheets/d/1s79gWCm7A8X1drwN6yAscetf0adpRkamHCJyHCKyIqY/edit?gid=0#gid=0, "
                    "вкладка Sheet1. Telegram канал @riderra_updates."
                ),
            }
        ],
        connected_integrations=[
            {
                "id": "telegram-1",
                "provider": "telegram",
                "status": "active",
                "display_name": "Business bot",
                "config": {"bot_mode": "business_bot"},
            }
        ],
    )

    preview = state["preview"]
    feasibility = preview["feasibility"]

    assert preview["category"] == "custom"
    assert preview["capability_allowlist"] == ["google_sheets.read_rows", "communications.draft"]
    assert [item["provider"] for item in preview["required_connectors"]] == ["google_sheets", "telegram"]
    assert preview["required_connectors"][0]["action"]["kind"] == "connect_after_draft"
    assert preview["required_connectors"][0]["action"]["after_draft"] == "open_agent_connections"
    assert preview["required_connectors"][1]["action"]["kind"] == "connected"
    assert feasibility["status"] == "needs_connection"
    assert [item["provider"] for item in feasibility["missing_connections"]] == ["google_sheets"]
    assert feasibility["ready_bindings"][0]["provider"] == "telegram"
    assert preview["setup_flow"]["schema"] == "localos_agent_builder_setup_flow_v1"
    assert preview["setup_flow"]["status"] == "needs_connection"
    assert preview["setup_flow"]["primary_action"] == "choose_route"
    assert preview["setup_flow"]["next_step"] == "create_draft_then_choose_route"
    assert preview["setup_flow"]["post_create_status"] == "needs_provider_route"
    assert preview["setup_flow"]["post_create_next_step"] == "choose_provider_route"
    assert preview["setup_flow"]["can_create_draft"] is True
    assert preview["setup_flow"]["can_activate"] is False
    assert preview["setup_flow"]["steps"][1]["key"] == "clarify"
    assert preview["setup_flow"]["steps"][1]["status"] == "done"
    assert preview["setup_flow"]["steps"][-2]["key"] == "preview"
    assert preview["setup_flow"]["steps"][-1]["key"] == "activate"
    assert not any(item["type"] == "clarification" for item in preview["setup_flow"]["activation_blockers"])
    assert any(item["type"] == "connection" and item["provider"] == "google_sheets" for item in preview["setup_flow"]["activation_blockers"])
    assert preview["connection_plan"]["schema"] == "localos_agent_connection_plan_v1"
    assert preview["connection_plan"]["status"] == "needs_action"
    assert preview["connection_plan"]["items"][0]["action"] == "choose_route"
    assert preview["connection_plan"]["items"][0]["recommended_route"]["provider"] == "openclaw"
    assert "OpenClaw" in preview["connection_plan"]["items"][0]["recommended_route_reason"]
    assert preview["connection_plan"]["items"][1]["action"] == "choose_route"
    assert preview["connection_summary"]["schema"] == "localos_agent_connection_summary_v1"
    assert preview["connection_summary"]["status"] == "needs_connection"
    assert preview["connection_summary"]["missing_count"] == 0
    assert preview["connection_summary"]["route_count"] == 2
    assert preview["connection_summary"]["ready_count"] == 0
    assert preview["connection_summary"]["next_action"] == "choose_provider_routes"
    assert preview["connection_summary"]["items"][0]["action"] == "choose_route"
    assert preview["connection_summary"]["items"][0]["setup_cta"]["mode"] == "choose_route"
    assert "маршрут" in preview["connection_summary"]["items"][0]["setup_cta"]["label"].lower()
    assert preview["connection_summary"]["items"][1]["action"] == "choose_route"
    assert preview["connection_readiness"]["schema"] == "localos_agent_connection_readiness_v1"
    assert preview["connection_readiness"]["next_action"] == "choose_provider_routes"
    assert preview["connection_readiness"]["post_create_workspace"] == "connections"
    assert preview["connection_readiness"]["required_count"] == 2
    assert preview["connection_readiness"]["missing_count"] == 0
    assert preview["connection_readiness"]["route_count"] == 2
    assert preview["connection_readiness"]["ready_count"] == 0
    assert preview["connection_readiness"]["missing_services"] == []
    assert preview["connection_readiness"]["route_services"][0]["provider"] == "google_sheets"
    assert preview["connection_readiness"]["route_services"][0]["recommended_route"]["provider"] == "openclaw"
    assert preview["connection_readiness"]["ready_services"] == []
    assert preview["connection_readiness"]["services"][0]["provider_route_cta"]
    assert preview["connection_resolver"]["schema"] == "localos_agent_connection_resolver_v1"
    assert preview["connection_resolver"]["next_action"] == "resolve_connections"
    assert preview["connection_resolver"]["required_count"] == 2
    assert preview["connection_resolver"]["unresolved_count"] == 2
    assert preview["connection_resolver"]["items"][0]["role_label"] == "Источник данных"
    assert preview["connection_resolver"]["items"][0]["service_label"] == "Google Sheets"
    assert preview["connection_resolver"]["items"][0]["recommended_provider"] == "openclaw"
    assert "OpenClaw boundary" in preview["connection_resolver"]["items"][0]["explanation"]
    assert preview["connection_resolver"]["items"][1]["role_label"] == "Куда подготовить результат"
    assert preview["connection_resolver"]["items"][1]["state"] == "choose_route"
    assert preview["connector_intelligence"]["schema"] == "localos_agent_connector_intelligence_v1"
    assert preview["connector_intelligence"]["status"] == "needs_connection"
    assert preview["connector_intelligence"]["can_compile_draft"] is True
    assert preview["connector_intelligence"]["can_preview_after_connections"] is False
    assert preview["connector_intelligence"]["bindings"][0]["action"] == "choose_route"
    assert preview["connector_intelligence"]["bindings"][0]["route_state"] == "available"
    assert preview["connector_intelligence"]["bindings"][0]["recommended_route"]["provider"] == "openclaw"
    assert any(item["provider"] == "native_localos" for item in preview["connector_intelligence"]["bindings"][0]["provider_routes"])
    assert preview["connector_intelligence"]["bindings"][0]["setup_cta"]["mode"] == "choose_route"
    assert preview["connector_intelligence"]["bindings"][1]["action"] == "choose_route"
    assert preview["connector_intelligence"]["bindings"][1]["route_state"] == "connected"
    assert preview["connector_intelligence"]["capabilities"][0]["route_state"] == "available"
    assert any(item["label"] == "OpenClaw" for item in preview["connector_intelligence"]["provider_paths"])
    assert preview["service_intelligence"]["schema"] == "localos_agent_service_intelligence_v1"
    assert preview["service_intelligence"]["status"] == "needs_connection"
    assert preview["service_intelligence"]["can_create_draft"] is True
    assert preview["service_intelligence"]["can_activate"] is False
    assert preview["service_intelligence"]["state_counts"]["route_choice"] == 2
    assert preview["service_intelligence"]["state_counts"].get("already_connected", 0) == 0
    service_states = {
        item["provider"]: item["state"]
        for item in preview["service_intelligence"]["items"]
        if item.get("kind") == "binding"
    }
    assert service_states["google_sheets"] == "route_choice"
    assert service_states["telegram"] == "route_choice"
    assert any(
        item["provider"] == "google_sheets" and item["next_action"] == "choose_provider_route"
        for item in preview["service_intelligence"]["items"]
    )
    assert preview["openclaw_planner_loop"]["schema"] == "localos_openclaw_planner_loop_v1"
    assert preview["openclaw_planner_loop"]["may_execute_tools"] is False
    assert "openclaw.google_sheets.read_rows" in preview["openclaw_planner_loop"]["workflow_proposal"]["openclaw_action_refs"]
    assert any(item["key"] == "google_sheets_target" for item in state["missing_questions"])
    assert any(item["reason"] == "connection_resolver" and item["provider"] == "google_sheets" for item in state["preview"]["connection_resolver_questions"])
    assert state["preview"]["connection_resolver_questions"][0]["key"] == "google_sheets_target"
    assert any("таблиц" in item["question"].lower() for item in state["missing_questions"])
    assert "Google Sheets" in state["messages"][-1]["content"]


def test_agent_builder_session_extracts_connection_answers_into_bindings():
    from services.agent_builder_session import build_agent_builder_state

    state = build_agent_builder_state(
        [
            {
                "role": "user",
                "content": (
                    "Каждый день бери заказ из Google Sheets за вчера и готовь пост в Telegram. "
                    "Результат нужен как черновик поста. Перед публикацией человек проверяет результат. "
                    "Таблица https://docs.google.com/spreadsheets/d/1s79gWCm7A8X1drwN6yAscetf0adpRkamHCJyHCKyIqY/edit?gid=0#gid=0, "
                    "вкладка Sheet1. Telegram канал @riderra_updates"
                ),
            }
        ],
    )

    answer_bindings = state["preview"]["connection_answer_bindings"]

    assert answer_bindings["google_sheets_read"]["spreadsheet_id"] == "1s79gWCm7A8X1drwN6yAscetf0adpRkamHCJyHCKyIqY"
    assert answer_bindings["google_sheets_read"]["sheet_name"] == "Sheet1"
    assert answer_bindings["google_sheets_read"]["gid"] == "0"
    assert answer_bindings["telegram_delivery"]["telegram_target"] == "@riderra_updates"
    assert answer_bindings["telegram_delivery"]["target_type"] == "chat_or_channel"


def test_agent_builder_answer_resources_allow_draft_but_not_preview_without_provider_route():
    from api import agent_builder_api
    from services.agent_blueprint_draft_builder import compile_agent_blueprint
    from services.agent_builder_session import build_agent_builder_state
    from services.agent_integration_preflight import build_agent_integration_preflight

    prompt = (
        "Каждый день бери заказ из Google Sheets за вчера и готовь пост в Telegram. "
        "Результат нужен как черновик поста. Перед публикацией человек проверяет результат. "
        "Таблица https://docs.google.com/spreadsheets/d/1s79gWCm7A8X1drwN6yAscetf0adpRkamHCJyHCKyIqY/edit?gid=0#gid=0, "
        "вкладка Sheet1. Telegram канал @riderra_updates"
    )
    state = build_agent_builder_state([{"role": "user", "content": prompt}], business_id="biz1", user_id="user1")
    preview = state["preview"]
    draft = compile_agent_blueprint(
        preview["understood_task"],
        preview["category"],
        business_id="biz1",
        user_id="user1",
        planner_context=preview["openclaw_planner_context"],
    )
    answer_bindings = preview["connection_answer_bindings"]
    metadata = agent_builder_api._apply_answer_connection_bindings(dict(draft["metadata"]), answer_bindings)
    version_payload = agent_builder_api._apply_answer_bindings_to_version_payload(dict(draft["version_payload"]), answer_bindings)

    class Cursor:
        def __init__(self):
            self.rows = []

        def execute(self, query, params=None):
            self.rows = []

        def fetchall(self):
            return self.rows

    preflight = build_agent_integration_preflight(
        Cursor(),
        business_id="biz1",
        metadata=metadata,
        input_payload={},
    )
    preflight_by_key = {
        item["key"]: item
        for item in preflight["items"]
        if isinstance(item, dict)
    }
    version_bindings = {
        item["key"]: item
        for item in version_payload["required_integration_bindings"]
        if isinstance(item, dict)
    }

    assert preview["setup_flow"]["can_create_draft"] is True
    assert preview["setup_flow"]["post_create_status"] == "needs_provider_route"
    assert state["missing_questions"]
    assert metadata["builder_answer_connection_bindings"]["google_sheets_read"]["spreadsheet_id"] == "1s79gWCm7A8X1drwN6yAscetf0adpRkamHCJyHCKyIqY"
    assert metadata["builder_answer_connection_bindings"]["telegram_delivery"]["telegram_target"] == "@riderra_updates"
    assert version_bindings["google_sheets_read"]["default_config"]["spreadsheet_id"] == "1s79gWCm7A8X1drwN6yAscetf0adpRkamHCJyHCKyIqY"
    assert version_bindings["google_sheets_read"]["answer_configured"] is True
    assert version_bindings["telegram_delivery"]["default_config"]["telegram_target"] == "@riderra_updates"
    assert preflight["ready"] is False
    assert preflight_by_key["google_sheets_read"]["resolution"] == "builder_answer_needs_provider_route"
    assert preflight_by_key["google_sheets_read"]["missing_config"] == []
    assert preflight_by_key["telegram_delivery"]["resolution"] == "builder_answer_needs_provider_route"


def test_agent_builder_setup_flow_points_ready_connectors_to_preview_run():
    from services.agent_builder_session import build_agent_builder_state

    state = build_agent_builder_state(
        [
            {
                "role": "user",
                "content": (
                    "Каждый день бери заказ из Google Sheets за вчера и готовь пост в Telegram. "
                    "Результат нужен как черновик поста. Перед публикацией человек проверяет результат."
                ),
            }
        ],
        connected_integrations=[
            {
                "id": "sheets-1",
                "provider": "google_sheets",
                "status": "active",
                "display_name": "Orders sheet",
                "config": {"spreadsheet_id": "sheet-1", "sheet_name": "Orders"},
            },
            {
                "id": "telegram-1",
                "provider": "telegram",
                "status": "active",
                "display_name": "Business bot",
                "config": {"bot_mode": "business_bot"},
            },
        ],
    )

    setup_flow = state["preview"]["setup_flow"]

    assert setup_flow["status"] == "ready_for_draft"
    assert setup_flow["next_step"] == "create_draft_then_choose_route"
    assert setup_flow["post_create_status"] == "needs_provider_route"
    assert setup_flow["post_create_next_step"] == "choose_provider_route"
    assert setup_flow["can_create_draft"] is True
    assert setup_flow["can_run_preview"] is False
    assert setup_flow["can_activate"] is False
    assert setup_flow["steps"][-2]["status"] == "blocked"
    assert "маршрут" in setup_flow["post_create_description"].lower()
    assert state["preview"]["connection_readiness"]["next_action"] == "choose_provider_routes"
    assert state["preview"]["connection_readiness"]["post_create_workspace"] == "connections"
    assert state["preview"]["connection_readiness"]["can_run_preview_after_create"] is False
    assert state["preview"]["connection_readiness"]["ready_count"] == 0
    assert [item["action"] for item in state["preview"]["connection_summary"]["items"]] == ["choose_route", "choose_route"]
    assert state["preview"]["connection_resolver"]["next_action"] == "resolve_connections"
    assert state["preview"]["connection_resolver"]["resolved_count"] == 0
    assert state["preview"]["connection_resolver"]["unresolved_count"] == 2


def test_compiled_agent_creation_contract_google_sheets_to_telegram():
    from api import agent_blueprints_api, agent_builder_api
    from services.agent_blueprint_draft_builder import compile_agent_blueprint
    from services.agent_integration_preflight import build_agent_integration_preflight
    from services.agent_builder_session import build_agent_builder_state

    prompt = (
        "Каждый день бери заказ из Google Sheets за вчера и готовь пост в Telegram. "
        "Результат нужен как черновик поста. Перед публикацией человек проверяет результат. "
        "Таблица https://docs.google.com/spreadsheets/d/1s79gWCm7A8X1drwN6yAscetf0adpRkamHCJyHCKyIqY/edit?gid=0#gid=0, "
        "вкладка Sheet1. Telegram канал @riderra_updates"
    )
    connected_integrations = [
        {
            "id": "sheets-1",
            "provider": "google_sheets",
            "status": "active",
            "display_name": "Orders sheet",
            "config": {"spreadsheet_id": "spreadsheet-1", "sheet_name": "Orders"},
        },
        {
            "id": "telegram-1",
            "provider": "telegram",
            "status": "active",
            "display_name": "Business bot",
            "config": {"bot_mode": "business_bot"},
        },
    ]
    state = build_agent_builder_state(
        [{"role": "user", "content": prompt}],
        connected_integrations=connected_integrations,
        business_id="biz1",
        user_id="user1",
    )

    preview = state["preview"]
    setup_flow = preview["setup_flow"]

    assert [item["reason"] for item in state["missing_questions"]] == [
        "connection_resolver",
        "connection_resolver",
    ]
    assert [item["provider"] for item in state["missing_questions"]] == ["google_sheets", "telegram"]
    assert setup_flow["status"] == "ready_for_draft"
    assert setup_flow["next_step"] == "create_draft_then_choose_route"
    assert setup_flow["post_create_next_step"] == "choose_provider_route"
    assert preview["connector_intelligence"]["status"] == "ready"
    assert [item["action"] for item in preview["connection_summary"]["items"]] == ["choose_route", "choose_route"]
    assert preview["openclaw_planner_loop"]["may_execute_tools"] is False
    assert "openclaw.google_sheets.read_rows" in preview["openclaw_planner_loop"]["workflow_proposal"]["openclaw_action_refs"]

    selected_bindings = agent_builder_api._selected_connection_bindings(
        {"selected_connection_bindings": {}},
        preview,
        connected_integrations,
    )
    draft = compile_agent_blueprint(
        preview["understood_task"],
        preview["category"],
        business_id="biz1",
        user_id="user1",
        planner_context=preview["openclaw_planner_context"],
    )
    metadata = dict(draft["metadata"])
    metadata["agent_builder_preview"] = preview
    metadata["required_connectors"] = preview["required_connectors"]
    metadata["builder_setup_flow"] = setup_flow
    metadata = agent_builder_api._apply_selected_connection_bindings(metadata, selected_bindings)

    class Cursor:
        def __init__(self):
            self.rows = []

        def execute(self, query, params=None):
            normalized_query = " ".join(query.split()).lower()
            self.rows = []
            if "from agent_integrations" in normalized_query:
                return None
            if "from agent_runs" in normalized_query:
                self.rows = [
                    {
                        "id": "run-1",
                        "status": "completed",
                        "input_json": {
                            "schema": "localos_agent_preview_input_v1",
                            "preview_mode": True,
                            "external_side_effects_allowed": False,
                            "source": "agent_preview",
                        },
                        "output_json": {},
                        "error_text": "",
                        "started_at": datetime(2026, 6, 11, tzinfo=timezone.utc),
                        "completed_at": datetime(2026, 6, 11, tzinfo=timezone.utc),
                        "updated_at": datetime(2026, 6, 11, tzinfo=timezone.utc),
                    }
                ]
                return None
            raise AssertionError(f"Unhandled SQL: {query}")

        def fetchall(self):
            return self.rows

    preflight = build_agent_integration_preflight(
        Cursor(),
        business_id="biz1",
        metadata=metadata,
        input_payload={},
    )
    post_create_handoff = agent_builder_api._build_post_create_handoff(preflight)
    blueprint = {
        "id": "bp1",
        "business_id": "biz1",
        "category": "custom",
        "description": preview["understood_task"],
        "metadata_json": metadata,
    }
    version = {
        "id": "version-1",
        "version_number": 1,
        "goal": draft["version_payload"]["goal"],
        "inputs_schema_json": draft["version_payload"]["inputs_schema"],
        "steps_json": draft["version_payload"]["steps"],
        "capability_allowlist_json": draft["version_payload"]["capability_allowlist"],
        "approval_policy_json": draft["version_payload"]["approval_policy"],
        "output_schema_json": draft["version_payload"]["output_schema"],
    }
    preview_input = agent_blueprints_api._build_agent_preview_run_input(
        blueprint,
        version,
        {"input": {}},
    )
    activation_gate = agent_blueprints_api._build_activation_gate_summary(
        Cursor(),
        blueprint,
        version,
        metadata,
    )

    assert selected_bindings["google_sheets_read"]["selection_source"] == "auto_single_connection"
    assert selected_bindings["telegram_delivery"]["selection_source"] == "auto_single_connection"
    assert metadata["agent_binding_integrations"]["google_sheets_read"]["integration_id"] == "sheets-1"
    assert metadata["agent_binding_integrations"]["telegram_delivery"]["integration_id"] == "telegram-1"
    assert preflight["ready"] is False
    assert {item["provider"]: item["resolution"] for item in preflight["items"]} == {
        "google_sheets": "provider_route_required",
        "telegram": "provider_route_required",
    }
    assert post_create_handoff["status"] == "needs_provider_route"
    assert post_create_handoff["next_step"] == "choose_provider_route"
    assert preview_input["schema"] == "localos_agent_preview_input_v1"
    assert preview_input["preview_mode"] is True
    assert preview_input["external_side_effects_allowed"] is False
    assert preview_input["google_sheets"]["read_only"] is True
    assert preview_input["telegram"]["draft_only"] is True
    assert activation_gate["can_activate"] is False
    assert activation_gate["next_step"] == "choose_provider_route"

    selected_provider_routes = agent_builder_api._selected_provider_routes(
        {"selected_provider_routes": {"google_sheets_read": "openclaw", "telegram_delivery": "openclaw"}},
        preview,
    )
    route_metadata = agent_builder_api._apply_selected_provider_routes(dict(metadata), selected_provider_routes)
    route_preflight = build_agent_integration_preflight(
        Cursor(),
        business_id="biz1",
        metadata=route_metadata,
        input_payload={},
    )

    assert selected_provider_routes["google_sheets_read"]["provider"] == "openclaw"
    assert selected_provider_routes["telegram_delivery"]["provider"] == "openclaw"
    assert route_metadata["agent_binding_provider_routes"]["google_sheets_read"]["route_provider"] == "openclaw"
    assert route_metadata["agent_binding_provider_routes"]["telegram_delivery"]["route_provider"] == "openclaw"
    assert route_preflight["ready"] is True
    route_items = {
        item["key"]: item
        for item in route_preflight["items"]
        if isinstance(item, dict)
    }
    assert route_items["google_sheets_read"]["resolution"] == "provider_route_openclaw_boundary"
    assert route_items["telegram_delivery"]["resolution"] == "provider_route_openclaw_boundary"


def test_agent_builder_create_blueprint_endpoint_returns_ready_preview_handoff(monkeypatch):
    from flask import Flask

    from api import agent_builder_api
    from services.agent_builder_session import build_agent_builder_state

    prompt = (
        "Каждый день бери заказ из Google Sheets за вчера и готовь пост в Telegram. "
        "Результат нужен как черновик поста. Перед публикацией человек проверяет результат. "
        "Таблица https://docs.google.com/spreadsheets/d/1s79gWCm7A8X1drwN6yAscetf0adpRkamHCJyHCKyIqY/edit?gid=0#gid=0, "
        "вкладка Sheet1. Telegram канал @riderra_updates"
    )
    connected_integrations = [
        {
            "id": "sheets-1",
            "business_id": "biz1",
            "provider": "google_sheets",
            "status": "active",
            "display_name": "Orders sheet",
            "config_json": {"spreadsheet_id": "spreadsheet-1", "sheet_name": "Orders"},
        },
        {
            "id": "telegram-1",
            "business_id": "biz1",
            "provider": "telegram",
            "status": "active",
            "display_name": "Business bot",
            "config_json": {"bot_mode": "business_bot"},
        },
    ]
    state = build_agent_builder_state(
        [{"role": "user", "content": prompt}],
        connected_integrations=[
            {
                "id": "sheets-1",
                "provider": "google_sheets",
                "status": "active",
                "display_name": "Orders sheet",
                "config": {"spreadsheet_id": "spreadsheet-1", "sheet_name": "Orders"},
            },
            {
                "id": "telegram-1",
                "provider": "telegram",
                "status": "active",
                "display_name": "Business bot",
                "config": {"bot_mode": "business_bot"},
            },
        ],
        business_id="biz1",
        user_id="user1",
    )

    class FakeConnection:
        def __init__(self):
            self.cursor_instance = FakeCursor()
            self.committed = False
            self.rolled_back = False

        def cursor(self):
            return self.cursor_instance

        def commit(self):
            self.committed = True

        def rollback(self):
            self.rolled_back = True

    class FakeDatabase:
        def __init__(self):
            self.conn = fake_connection
            self.closed = False

        def close(self):
            self.closed = True

    class FakeCursor:
        def __init__(self):
            self.sessions = {
                "session-1": {
                    "id": "session-1",
                    "business_id": "biz1",
                    "created_by_user_id": "user1",
                    "status": "draft",
                    "initial_prompt": prompt,
                    "category": state["category"],
                    "messages_json": state["messages"],
                    "preview_json": state["preview"],
                    "missing_questions_json": state["missing_questions"],
                    "blueprint_id": None,
                }
            }
            self.integrations = connected_integrations
            self.blueprints = {}
            self.versions = {}
            self.last_result = None
            self.last_results = []

        def execute(self, query, params=None):
            normalized_query = " ".join(query.split()).lower()
            params = params or ()
            self.last_result = None
            self.last_results = []
            if normalized_query.startswith("select * from agent_builder_sessions"):
                self.last_result = self.sessions.get(params[0])
                return None
            if normalized_query.startswith("select id, provider, status, display_name, config_json"):
                business_id = params[0]
                self.last_results = [item for item in self.integrations if item["business_id"] == business_id]
                return None
            if normalized_query.startswith("select id, source, display_name"):
                self.last_results = []
                return None
            if normalized_query.startswith("select telegram_bot_token"):
                self.last_result = {"telegram_bot_token": ""}
                return None
            if normalized_query.startswith("insert into agent_blueprints"):
                metadata = json.loads(params[7])
                self.blueprints[params[0]] = {
                    "id": params[0],
                    "business_id": params[1],
                    "name": params[2],
                    "category": params[3],
                    "description": params[4],
                    "status": params[5],
                    "created_by_user_id": params[6],
                    "metadata_json": metadata,
                }
                return None
            if normalized_query.startswith("select coalesce(max(version_number)"):
                self.last_result = {"next_version": 1}
                return None
            if normalized_query.startswith("insert into agent_blueprint_versions"):
                self.versions[params[0]] = {
                    "id": params[0],
                    "blueprint_id": params[1],
                    "version_number": params[2],
                    "goal": params[3],
                    "inputs_schema_json": json.loads(params[4]),
                    "steps_json": json.loads(params[5]),
                    "persona_agent_id": params[6],
                    "capability_allowlist_json": json.loads(params[7]),
                    "approval_policy_json": json.loads(params[8]),
                    "output_schema_json": json.loads(params[9]),
                    "created_by_user_id": params[10],
                }
                return None
            if normalized_query.startswith("select * from agent_blueprint_versions where id"):
                self.last_result = self.versions.get(params[0])
                return None
            if normalized_query.startswith("update agent_builder_sessions"):
                session = self.sessions.get(params[6])
                session["status"] = params[0]
                session["category"] = params[1]
                session["messages_json"] = json.loads(params[2])
                session["preview_json"] = json.loads(params[3])
                session["missing_questions_json"] = json.loads(params[4])
                if params[5]:
                    session["blueprint_id"] = params[5]
                return None
            if normalized_query.startswith("select b.*"):
                blueprint = self.blueprints.get(params[0])
                latest_version = next(
                    (version for version in self.versions.values() if version["blueprint_id"] == params[0]),
                    None,
                )
                self.last_result = {
                    **blueprint,
                    "latest_version_id": latest_version["id"],
                    "latest_version_number": latest_version["version_number"],
                    "latest_goal": latest_version["goal"],
                }
                return None
            if "from agent_integrations" in normalized_query:
                business_id = params[0]
                self.last_results = [item for item in self.integrations if item["business_id"] == business_id]
                return None
            raise AssertionError(f"Unhandled SQL: {query}")

        def fetchone(self):
            return self.last_result

        def fetchall(self):
            return self.last_results

    fake_connection = FakeConnection()
    monkeypatch.setattr(agent_builder_api, "DatabaseManager", FakeDatabase)
    monkeypatch.setattr(agent_builder_api, "_require_auth", lambda: ({"user_id": "user1"}, None))
    monkeypatch.setattr(agent_builder_api, "_require_business_access", lambda cursor, business_id, user_data: (True, None))
    monkeypatch.setattr(
        agent_builder_api,
        "charge_agent_creation_credits",
        lambda cursor, business_id, user_id, source_id, description: {
            "status": "charged",
            "action_key": "agent_creation",
            "actual_credits": 3,
        },
    )

    app = Flask(__name__)
    request_context = app.test_request_context(
        "/api/agent-builder/sessions/session-1/create-blueprint",
        method="POST",
        json={
            "selected_provider_routes": {
                "google_sheets_read": "openclaw",
                "telegram_delivery": "openclaw",
            },
            "accepted_provider_routes": True,
        },
    )
    request_context.push()
    try:
        response, status_code = agent_builder_api.create_blueprint_from_agent_builder_session("session-1")
    finally:
        request_context.pop()

    payload = response.get_json()
    blueprint_id = payload["blueprint"]["id"]
    metadata = fake_connection.cursor_instance.blueprints[blueprint_id]["metadata_json"]

    assert status_code == 201
    assert fake_connection.committed is True
    assert fake_connection.rolled_back is False
    assert payload["success"] is True
    assert payload["next_step"] == "run_preview"
    assert payload["post_create_handoff"]["status"] == "ready_for_preview"
    assert payload["post_create_handoff"]["workspace_mode"] == "run"
    assert payload["connection_preflight"]["ready"] is True
    assert {item["provider"]: item["resolution"] for item in payload["connection_preflight"]["items"]} == {
        "google_sheets": "provider_route_openclaw_boundary",
        "telegram": "provider_route_openclaw_boundary",
    }
    assert payload["session"]["status"] == "blueprint_created"
    assert payload["session"]["blueprint_id"] == blueprint_id
    assert metadata["builder_selected_connection_bindings"]["google_sheets_read"]["selection_source"] == "auto_single_connection"
    assert metadata["builder_provider_routes_accepted"] is True
    assert metadata["agent_binding_provider_routes"]["telegram_delivery"]["integration_id"] == "openclaw_boundary"
    assert metadata["agent_binding_provider_routes"]["telegram_delivery"]["execution_boundary"] == "localos_policy_envelope"
    assert metadata["builder_answer_connection_bindings"]["google_sheets_read"]["spreadsheet_id"] == "1s79gWCm7A8X1drwN6yAscetf0adpRkamHCJyHCKyIqY"
    assert metadata["builder_answer_connection_bindings"]["telegram_delivery"]["telegram_target"] == "@riderra_updates"
    assert metadata["agent_binding_integrations"]["google_sheets_read"]["answer_config"]["sheet_name"] == "Sheet1"
    assert metadata["agent_binding_integrations"]["telegram_delivery"]["answer_config"]["telegram_target"] == "@riderra_updates"
    assert metadata["custom_process"]["google_sheets"]["spreadsheet_id"] == "1s79gWCm7A8X1drwN6yAscetf0adpRkamHCJyHCKyIqY"
    assert metadata["custom_process"]["telegram"]["telegram_target"] == "@riderra_updates"
    assert metadata["openclaw_planner_loop"]["may_execute_tools"] is False
    assert payload["version"]["capability_allowlist_json"] == ["google_sheets.read_rows", "communications.draft"]


def test_agent_preview_run_and_activation_endpoints_enforce_safe_gate(monkeypatch):
    from flask import Flask

    from api import agent_blueprints_api
    from api.agent_blueprints_api import _build_agent_preview_run_input
    from services.agent_blueprint_draft_builder import compile_agent_blueprint

    draft = compile_agent_blueprint(
        "Каждый день бери заказ из Google Sheets за вчера и готовь пост в Telegram. "
        "Результат нужен как черновик поста. Перед публикацией человек проверяет результат."
    )
    metadata = dict(draft["metadata"])
    metadata["custom_process"] = {
        "google_sheets_read": {
            "integration_id": "sheets-1",
            "spreadsheet_id": "spreadsheet-1",
            "sheet_name": "Orders",
        },
        "telegram_delivery": {
            "integration_id": "telegram-1",
            "bot_mode": "business_bot",
        },
        "google_sheets": {
            "integration_id": "sheets-1",
            "spreadsheet_id": "spreadsheet-1",
            "sheet_name": "Orders",
        },
        "telegram": {
            "integration_id": "telegram-1",
            "bot_mode": "business_bot",
        },
    }
    metadata["agent_binding_integrations"] = {
        "google_sheets_read": {
            "integration_id": "sheets-1",
            "provider": "google_sheets",
        },
        "telegram_delivery": {
            "integration_id": "telegram-1",
            "provider": "telegram",
        },
    }
    metadata["agent_binding_provider_routes"] = {
        "google_sheets_read": {
            "route_provider": "openclaw",
            "provider": "openclaw",
            "status": "active",
            "integration_id": "openclaw_boundary",
        },
        "telegram_delivery": {
            "route_provider": "openclaw",
            "provider": "openclaw",
            "status": "active",
            "integration_id": "openclaw_boundary",
        },
    }
    cursor = FakeCursor()
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Sheets to Telegram",
        "category": "custom",
        "description": "Google Sheets to Telegram",
        "status": "draft",
        "metadata_json": metadata,
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "version_number": 1,
        "goal": draft["version_payload"]["goal"],
        "inputs_schema_json": draft["version_payload"]["inputs_schema"],
        "steps_json": draft["version_payload"]["steps"],
        "capability_allowlist_json": draft["version_payload"]["capability_allowlist"],
        "approval_policy_json": draft["version_payload"]["approval_policy"],
        "output_schema_json": draft["version_payload"]["output_schema"],
    }

    class FakeConnection:
        def __init__(self):
            self.cursor_instance = cursor
            self.commit_count = 0
            self.rollback_count = 0

        def cursor(self):
            return self.cursor_instance

        def commit(self):
            self.commit_count += 1

        def rollback(self):
            self.rollback_count += 1

    class FakeDatabase:
        def __init__(self):
            self.conn = fake_connection

        def close(self):
            return None

    fake_connection = FakeConnection()
    monkeypatch.setattr(agent_blueprints_api, "DatabaseManager", FakeDatabase)
    monkeypatch.setattr(agent_blueprints_api, "_require_auth", lambda: ({"user_id": "user1"}, None))
    monkeypatch.setattr(agent_blueprints_api, "_require_business_access", lambda cursor, business_id, user_data: (True, None))
    monkeypatch.setattr(agent_blueprints_api, "build_agent_blueprint_orchestrator", CountingOrchestrator)

    app = Flask(__name__)

    blocked_context = app.test_request_context(
        "/api/agent-blueprints/bp1/versions/ver1/activate",
        method="POST",
        json={},
    )
    blocked_context.push()
    try:
        blocked_response, blocked_status = agent_blueprints_api.activate_agent_blueprint_version("bp1", "ver1")
    finally:
        blocked_context.pop()
    blocked_payload = blocked_response.get_json()

    preflight_context = app.test_request_context(
        "/api/agent-blueprints/bp1/preflight",
        method="POST",
        json={"blueprint_version_id": "ver1", "input": {"preview_mode": True}},
    )
    preflight_context.push()
    try:
        preflight_response = agent_blueprints_api.preflight_agent_blueprint_run("bp1")
    finally:
        preflight_context.pop()
    preflight_payload = preflight_response.get_json()

    run_context = app.test_request_context(
        "/api/agent-blueprints/bp1/runs",
        method="POST",
        json={"blueprint_version_id": "ver1", "input": {"preview_mode": True}},
    )
    run_context.push()
    try:
        run_response, run_status = agent_blueprints_api.start_agent_blueprint_run("bp1")
    finally:
        run_context.pop()
    run_payload = run_response.get_json()
    run = run_payload["run"]

    activate_context = app.test_request_context(
        "/api/agent-blueprints/bp1/versions/ver1/activate",
        method="POST",
        json={"reason": "safe preview passed"},
    )
    activate_context.push()
    try:
        activate_response = agent_blueprints_api.activate_agent_blueprint_version("bp1", "ver1")
    finally:
        activate_context.pop()
    activate_payload = activate_response.get_json()

    preview_input = _build_agent_preview_run_input(
        cursor.tables["agent_blueprints"]["bp1"],
        cursor.tables["agent_blueprint_versions"]["ver1"],
        {"input": {"preview_mode": True}},
    )
    run_input = cursor.tables["agent_runs"][run["id"]]["input_json"]

    assert blocked_status == 400
    assert blocked_payload["code"] == "AGENT_ACTIVATION_GATE_BLOCKED"
    assert blocked_payload["activation_gate"]["next_step"] == "run_preview"
    assert preflight_payload["success"] is True
    assert preflight_payload["can_start"] is True
    assert preflight_payload["preview_run_gate"]["can_preview_run"] is True
    assert preflight_payload["preview_input"]["schema"] == "localos_agent_preview_input_v1"
    assert preflight_payload["preview_input"]["external_side_effects_allowed"] is False
    assert preflight_payload["preflight"]["ready"] is True
    assert run_status == 201
    assert run_payload["success"] is True
    assert run_input["schema"] == "localos_agent_preview_input_v1"
    assert run_input["preview_mode"] is True
    assert run_input["external_side_effects_allowed"] is False
    assert run["observability"]["preview_summary"]["safe_preview"] is True
    assert run["status"] in {"completed", "waiting_approval"}
    assert activate_payload["success"] is True
    assert activate_payload["active_version"]["id"] == "ver1"
    assert cursor.tables["agent_blueprints"]["bp1"]["status"] == "active"
    assert cursor.tables["agent_blueprints"]["bp1"]["metadata_json"]["active_version_id"] == "ver1"
    assert fake_connection.commit_count >= 2
    assert fake_connection.rollback_count == 0
    assert preview_input["external_side_effects_allowed"] is False


def test_browser_use_to_telegram_agent_preview_run_is_ready_after_connections(monkeypatch):
    from flask import Flask

    from api import agent_blueprints_api
    from services.agent_blueprint_draft_builder import compile_agent_blueprint

    draft = compile_agent_blueprint(
        "Через browser use открой сайт конкурента https://competitor.example, проверь изменения цен "
        "и подготовь короткое сообщение владельцу в Telegram.",
        use_ai=True,
    )
    metadata = dict(draft["metadata"])
    metadata["agent_integration_ids"] = ["browser-1", "telegram-1"]
    metadata["agent_binding_integrations"] = {
        "browser_use_read": {
            "integration_id": "browser-1",
            "provider": "browser_use",
        },
        "telegram_delivery": {
            "integration_id": "openclaw_boundary",
            "provider": "openclaw",
            "route_provider": "openclaw",
            "status": "active",
        },
    }
    metadata["agent_binding_provider_routes"] = {
        "telegram_delivery": {
            "route_provider": "openclaw",
            "provider": "openclaw",
            "status": "active",
            "integration_id": "openclaw_boundary",
            "execution_boundary": "localos_policy_envelope",
        },
    }
    metadata["custom_process"] = {
        "browser_use": {
            "integration_id": "browser-1",
            "target_urls": ["https://competitor.example"],
            "mode": "openclaw_browser_boundary",
        },
        "browser_use_read": {
            "integration_id": "browser-1",
            "target_urls": ["https://competitor.example"],
            "mode": "openclaw_browser_boundary",
        },
        "telegram": {
            "integration_id": "openclaw_boundary",
            "route_provider": "openclaw",
            "status": "active",
        },
    }
    cursor = FakeCursor()
    cursor.tables["agent_blueprints"]["bp-browser"] = {
        "id": "bp-browser",
        "business_id": "biz1",
        "name": "Мониторинг сайта конкурента",
        "category": "custom",
        "description": "Browser use проверяет сайт конкурента и готовит Telegram-отчёт.",
        "status": "draft",
        "metadata_json": metadata,
    }
    cursor.tables["agent_blueprint_versions"]["ver-browser"] = {
        "id": "ver-browser",
        "blueprint_id": "bp-browser",
        "version_number": 1,
        "goal": draft["version_payload"]["goal"],
        "inputs_schema_json": draft["version_payload"]["inputs_schema"],
        "steps_json": draft["version_payload"]["steps"],
        "capability_allowlist_json": draft["version_payload"]["capability_allowlist"],
        "approval_policy_json": draft["version_payload"]["approval_policy"],
        "output_schema_json": draft["version_payload"]["output_schema"],
    }
    cursor.tables["agent_integrations"]["browser-1"] = {
        "id": "browser-1",
        "business_id": "biz1",
        "provider": "browser_use",
        "status": "active",
        "display_name": "Browser use через OpenClaw",
        "auth_ref": None,
        "config_json": {
            "target_urls": ["https://competitor.example"],
            "mode": "openclaw_browser_boundary",
        },
        "limits_json": {"daily_page_check_cap": 12, "frequency_cap_minutes": 60},
        "connected_by_user_id": "user1",
        "created_at": "2026-06-21T10:00:00Z",
        "updated_at": "2026-06-21T10:00:00Z",
    }
    cursor.tables["agent_integrations"]["telegram-1"] = {
        "id": "telegram-1",
        "business_id": "biz1",
        "provider": "telegram",
        "status": "active",
        "display_name": "Бот владельца",
        "auth_ref": None,
        "config_json": {"bot_mode": "business_bot"},
        "limits_json": {"daily_message_cap": 30, "frequency_cap_minutes": 30},
        "connected_by_user_id": "user1",
        "created_at": "2026-06-21T10:00:00Z",
        "updated_at": "2026-06-21T10:00:00Z",
    }

    class FakeConnection:
        def __init__(self):
            self.cursor_instance = cursor
            self.commit_count = 0
            self.rollback_count = 0

        def cursor(self):
            return self.cursor_instance

        def commit(self):
            self.commit_count += 1

        def rollback(self):
            self.rollback_count += 1

    class FakeDatabase:
        def __init__(self):
            self.conn = fake_connection

        def close(self):
            return None

    fake_connection = FakeConnection()
    monkeypatch.setattr(agent_blueprints_api, "DatabaseManager", FakeDatabase)
    monkeypatch.setattr(agent_blueprints_api, "_require_auth", lambda: ({"user_id": "user1"}, None))
    monkeypatch.setattr(agent_blueprints_api, "_require_business_access", lambda cursor, business_id, user_data: (True, None))
    monkeypatch.setattr(agent_blueprints_api, "build_agent_blueprint_orchestrator", CountingOrchestrator)

    app = Flask(__name__)
    preflight_context = app.test_request_context(
        "/api/agent-blueprints/bp-browser/preflight",
        method="POST",
        json={"blueprint_version_id": "ver-browser", "input": {"preview_mode": True}},
    )
    preflight_context.push()
    try:
        preflight_response = agent_blueprints_api.preflight_agent_blueprint_run("bp-browser")
    finally:
        preflight_context.pop()
    preflight_payload = preflight_response.get_json()

    run_context = app.test_request_context(
        "/api/agent-blueprints/bp-browser/runs",
        method="POST",
        json={"blueprint_version_id": "ver-browser", "input": {"preview_mode": True}},
    )
    run_context.push()
    try:
        run_response, run_status = agent_blueprints_api.start_agent_blueprint_run("bp-browser")
    finally:
        run_context.pop()
    run_payload = run_response.get_json()
    run_input = cursor.tables["agent_runs"][run_payload["run"]["id"]]["input_json"]
    preflight_items = {
        item["key"]: item
        for item in preflight_payload["preflight"]["items"]
    }

    assert preflight_payload["success"] is True
    assert preflight_payload["can_start"] is True
    assert preflight_payload["preflight"]["ready"] is True
    assert preflight_items["browser_use_read"]["provider"] == "browser_use"
    assert preflight_items["browser_use_read"]["resolution"] == "blueprint_metadata"
    assert preflight_items["browser_use_read"]["execution_boundary"] == "connected_provider"
    assert preflight_items["telegram_delivery"]["resolution"] == "provider_route_openclaw_boundary"
    assert preflight_payload["preview_input"]["external_side_effects_allowed"] is False
    assert run_status == 201
    assert run_payload["success"] is True
    assert run_input["preview_mode"] is True
    assert run_input["external_side_effects_allowed"] is False
    assert run_payload["run"]["observability"]["preview_summary"]["safe_preview"] is True
    assert fake_connection.commit_count == 1
    assert fake_connection.rollback_count == 0


def test_agent_builder_setup_flow_blocks_draft_until_clarification_is_answered():
    from services.agent_builder_session import build_agent_builder_state

    state = build_agent_builder_state([{"role": "user", "content": "Сделай агента"}])
    setup_flow = state["preview"]["setup_flow"]

    assert setup_flow["status"] == "needs_clarification"
    assert setup_flow["primary_action"] == "answer_question"
    assert setup_flow["can_create_draft"] is False
    assert setup_flow["steps"][1]["key"] == "clarify"
    assert setup_flow["steps"][1]["status"] == "active"
    assert any(item["type"] == "clarification" for item in setup_flow["activation_blockers"])


def test_agent_builder_blocks_empty_google_sheets_to_telegram_workflow():
    from services.agent_builder_session import build_agent_builder_state

    state = build_agent_builder_state(
        [
            {
                "role": "user",
                "content": "Сделай агента: бери заказ из Google Sheets и делай пост в Telegram",
            }
        ]
    )
    setup_flow = state["preview"]["setup_flow"]
    questions = {item["key"]: item for item in state["missing_questions"]}

    assert setup_flow["status"] == "needs_clarification"
    assert setup_flow["primary_action"] == "answer_question"
    assert setup_flow["can_create_draft"] is False
    assert setup_flow["steps"][1]["status"] == "active"
    assert questions["google_sheets_target"]["reason"] == "openclaw_workflow_detail_missing"
    assert questions["telegram_target"]["reason"] == "openclaw_workflow_detail_missing"
    assert questions["schedule_frequency"]["reason"] == "openclaw_workflow_detail_missing"
    assert questions["post_format"]["reason"] == "openclaw_workflow_detail_missing"


def test_agent_builder_setup_flow_surfaces_compiler_clarifying_questions(monkeypatch):
    from services import agent_builder_session

    def fake_compile_agent_blueprint(description, category="", **kwargs):
        return {
            "name": "Google Sheets -> Telegram",
            "category": "custom",
            "description": description,
            "metadata": {
                "llm_intent": {
                    "status": "compiled_intent",
                    "source": "gigachat",
                    "intent": {
                        "clarifying_questions": [
                            "Какую вкладку Google Sheets читать?",
                            "В какой Telegram-канал готовить пост?",
                        ],
                    },
                },
                "required_integration_bindings": [
                    {
                        "key": "google_sheets_read",
                        "provider": "google_sheets",
                        "capability": "google_sheets.read_rows",
                        "required_config": ["spreadsheet_id", "sheet_name"],
                    },
                    {
                        "key": "telegram_delivery",
                        "provider": "telegram",
                        "capability": "communications.draft",
                        "required_config": ["bot_mode"],
                    },
                ],
            },
            "version_payload": {
                "steps": [],
                "capability_allowlist": ["google_sheets.read_rows", "communications.draft"],
                "required_integration_bindings": [
                    {
                        "key": "google_sheets_read",
                        "provider": "google_sheets",
                        "capability": "google_sheets.read_rows",
                        "required_config": ["spreadsheet_id", "sheet_name"],
                    },
                    {
                        "key": "telegram_delivery",
                        "provider": "telegram",
                        "capability": "communications.draft",
                        "required_config": ["bot_mode"],
                    },
                ],
            },
            "summary": {
                "sources": ["google_sheets", "telegram", "business_profile"],
                "capability_allowlist": ["google_sheets.read_rows", "communications.draft"],
            },
        }

    monkeypatch.setattr(agent_builder_session, "compile_agent_blueprint", fake_compile_agent_blueprint)

    state = agent_builder_session.build_agent_builder_state(
        [
            {
                "role": "user",
                "content": "Каждый день бери заказ из Google Sheets и готовь пост в Telegram.",
            }
        ],
        use_ai=True,
    )

    assert state["missing_questions"][0]["question"] == "Какую вкладку Google Sheets читать?"
    assert state["missing_questions"][0]["reason"] == "compiled_intent_clarification"
    assert state["preview"]["compiler_questions"][1]["question"] == "В какой Telegram-канал готовить пост?"
    assert state["preview"]["setup_flow"]["steps"][1]["questions"][0]["key"] == "compiler_question_1"
    assert state["preview"]["setup_flow"]["can_create_draft"] is False


def test_agent_builder_can_create_draft_when_sheet_id_is_deferred(monkeypatch):
    from services import agent_builder_session

    def fake_compile_agent_blueprint(description, category="", **kwargs):
        return {
            "name": "Orders status agent",
            "category": "custom",
            "description": description,
            "metadata": {
                "llm_intent": {
                    "status": "compiled_intent",
                    "source": "gigachat",
                    "intent": {
                        "workflow_draft": {
                            "trigger": "schedule.hourly",
                            "steps": [{"key": "read_orders", "capability": "google_sheets.read_rows"}],
                        },
                        "clarifying_questions": [
                            "Пожалуйста, уточните название Google Spreadsheet и Sheet, чтобы продолжить.",
                        ],
                        "approval_points": [{"key": "telegram_delivery", "reason": "external delivery requires approval"}],
                        "unsupported_requests": [],
                    },
                },
                "required_integration_bindings": [
                    {
                        "key": "google_sheets_read",
                        "provider": "google_sheets",
                        "capability": "google_sheets.read_rows",
                        "required_config": ["spreadsheet_id", "sheet_name"],
                    },
                    {
                        "key": "telegram_delivery",
                        "provider": "telegram",
                        "capability": "communications.draft",
                        "required_config": ["bot_mode"],
                    },
                ],
            },
            "summary": {
                "sources": ["google_sheets", "telegram", "business_profile"],
                "capability_allowlist": ["google_sheets.read_rows", "communications.draft"],
            },
        }

    monkeypatch.setattr(agent_builder_session, "compile_agent_blueprint", fake_compile_agent_blueprint)

    state = agent_builder_session.build_agent_builder_state(
        [
            {
                "role": "user",
                "content": (
                    "Каждый час бери новые строки из Google Sheets с заказами и отправляй краткий статус "
                    "владельцу в Telegram после проверки. Таблица называется Заказы, лист Новый поток. "
                    "ID сейчас нет. Создай черновик агента без ID, а подключение конкретной таблицы "
                    "оставь следующим шагом в доступах."
                ),
            }
        ],
        use_ai=True,
    )

    assert [item["reason"] for item in state["missing_questions"]] == ["connection_resolver", "connection_resolver"]
    assert state["preview"]["setup_flow"]["can_create_draft"] is True
    assert state["preview"]["setup_flow"]["next_step"] == "create_draft_then_choose_route"
    assert state["preview"]["setup_flow"]["can_run_preview"] is False


def test_agent_builder_setup_flow_blocks_compiler_unsupported_request(monkeypatch):
    from services import agent_builder_session

    def fake_compile_agent_blueprint(description, category="", **kwargs):
        return {
            "name": "Unsafe external workflow",
            "category": "custom",
            "description": description,
            "metadata": {
                "llm_intent": {
                    "status": "compiled_intent",
                    "source": "gigachat",
                    "intent": {
                        "workflow_draft": {
                            "trigger": "schedule.daily",
                            "steps": [{"key": "read_source", "capability": "google_sheets.read_rows"}],
                        },
                        "approval_points": [{"key": "external_action", "reason": "external action requires approval"}],
                        "unsupported_requests": [
                            {
                                "request": "send without approval",
                                "reason": "Нельзя отправлять внешние сообщения без approval gate.",
                            }
                        ],
                        "clarifying_questions": [],
                    },
                },
            },
            "version_payload": {
                "steps": [],
                "capability_allowlist": [],
                "required_integration_bindings": [],
            },
            "summary": {
                "sources": ["manual_context", "business_profile"],
                "capability_allowlist": [],
            },
        }

    monkeypatch.setattr(agent_builder_session, "compile_agent_blueprint", fake_compile_agent_blueprint)

    state = agent_builder_session.build_agent_builder_state(
        [
            {
                "role": "user",
                "content": (
                    "Из Google Sheets извлеки заказы, подготовь отчёт, "
                    "человек проверяет результат перед действием."
                ),
            }
        ],
        use_ai=True,
    )

    setup_flow = state["preview"]["setup_flow"]
    review = state["preview"]["compiler_policy_review"]

    assert review["schema"] == "localos_agent_compiler_policy_review_v1"
    assert review["status"] == "blocked"
    assert review["workflow_draft"]["trigger"] == "schedule.daily"
    assert setup_flow["status"] == "blocked"
    assert setup_flow["next_step"] == "cannot_create"
    assert setup_flow["can_create_draft"] is False
    assert setup_flow["steps"][3]["status"] == "blocked"
    assert setup_flow["activation_blockers"][0]["type"] == "compiler_unsupported"


def test_agent_builder_service_intelligence_marks_multiple_existing_routes():
    from services.agent_builder_session import build_agent_builder_state

    state = build_agent_builder_state(
        [
            {
                "role": "user",
                "content": (
                    "Каждый день бери заказ из Google Sheets и готовь пост в Telegram. "
                    "Человек проверяет перед публикацией."
                ),
            }
        ],
        connected_integrations=[
            {
                "id": "telegram-1",
                "provider": "telegram",
                "status": "active",
                "display_name": "Business bot",
                "config": {"bot_mode": "business_bot"},
            },
            {
                "id": "telegram-2",
                "provider": "telegram",
                "status": "active",
                "display_name": "Marketing channel",
                "config": {"bot_mode": "business_bot"},
            },
        ],
    )

    telegram_item = next(
        item
        for item in state["preview"]["service_intelligence"]["items"]
        if item.get("kind") == "binding" and item.get("provider") == "telegram"
    )

    assert state["preview"]["feasibility"]["status"] == "needs_choice"
    assert telegram_item["state"] == "route_choice"
    assert telegram_item["state_label"] == "Нужно выбрать маршрут"
    assert telegram_item["next_action"] == "choose_provider_route"
    assert telegram_item["connection_count"] == 2


def test_agent_builder_service_intelligence_marks_forbidden_request_impossible():
    from services.agent_builder_session import build_agent_builder_state

    state = build_agent_builder_state(
        [
            {
                "role": "user",
                "content": "Сделай агента, который подключается к компьютерам Роскосмоса и скачивает файлы.",
            }
        ],
    )
    intelligence = state["preview"]["service_intelligence"]

    assert intelligence["status"] == "forbidden"
    assert intelligence["can_create_draft"] is False
    assert intelligence["state_counts"]["impossible"] >= 1
    assert any(
        item["kind"] == "policy" and item["state"] == "impossible" and item["next_action"] == "explain_policy_boundary"
        for item in intelligence["items"]
    )


def test_agent_builder_api_uses_ai_compiler_by_default():
    from api.agent_builder_api import _use_ai_compiler

    assert _use_ai_compiler({}) is True
    assert _use_ai_compiler({"use_ai_compiler": False}) is False
    assert _use_ai_compiler({"use_ai_compiler": True}) is True


def test_agent_builder_api_requires_compiler_plan_confirmation_only_for_reviewable_plan():
    from api.agent_builder_api import _compiler_plan_requires_confirmation

    assert _compiler_plan_requires_confirmation({}) is False
    assert _compiler_plan_requires_confirmation({"compiler_policy_review": {"status": "ok"}}) is False
    assert _compiler_plan_requires_confirmation(
        {
            "compiler_policy_review": {
                "workflow_draft": {
                    "trigger": "schedule.daily",
                    "steps": [{"key": "read_orders"}],
                }
            }
        }
    ) is True
    assert _compiler_plan_requires_confirmation(
        {
            "compiler_policy_review": {
                "approval_points": [{"key": "publish", "reason": "external publish"}],
            }
        }
    ) is True
    assert _compiler_plan_requires_confirmation(
        {
            "compiler_unsupported_requests": [{"reason": "No approved provider path"}],
        }
    ) is True


def test_agent_builder_api_requires_selected_provider_routes_for_required_bindings():
    from api.agent_builder_api import _missing_required_provider_routes, _required_provider_route_bindings

    preview = {
        "connection_readiness": {
            "services": [
                {
                    "key": "google_sheets_read",
                    "recommended_route": {
                        "provider": "openclaw",
                        "state": "available",
                        "provider_action": {"available": True},
                    },
                },
                {
                    "key": "telegram_delivery",
                    "recommended_route": {
                        "provider": "composio",
                        "state": "planned",
                    },
                },
            ]
        }
    }

    required = _required_provider_route_bindings(preview)

    assert required == [{"key": "google_sheets_read", "available_routes": ["openclaw"]}]
    assert _missing_required_provider_routes(preview, {}) == required
    assert _missing_required_provider_routes(preview, {"google_sheets_read": {"provider": "openclaw"}}) == []


def test_agent_builder_provider_routes_create_action_handler_contracts():
    from api import agent_builder_api

    metadata = {
        "required_integration_bindings": [
            {"key": "google_sheets_read", "provider": "google_sheets"},
            {"key": "telegram_delivery", "provider": "telegram"},
        ],
        "custom_process": {},
    }
    selected = {
        "google_sheets_read": {
            "provider": "openclaw",
            "route_provider": "openclaw",
            "label": "OpenClaw",
            "connect_mode": "openclaw_policy_boundary",
        },
        "telegram_delivery": {
            "provider": "maton",
            "route_provider": "maton",
            "label": "Maton.ai",
            "connect_mode": "external_account_key",
            "external_account_id": "maton-account-1",
            "display_name": "Main Maton key",
        },
    }

    metadata = agent_builder_api._apply_selected_provider_routes(metadata, selected)

    assert metadata["agent_binding_provider_routes"]["google_sheets_read"]["route_provider"] == "openclaw"
    assert metadata["agent_binding_provider_routes"]["telegram_delivery"]["external_account_id"] == "maton-account-1"
    assert metadata["connector_action_handlers"]["google_sheets_read"]["handler"] == "openclaw_policy_boundary"
    assert metadata["connector_action_handlers"]["google_sheets_read"]["preflight_resolution"] == "provider_route_openclaw_boundary"
    assert metadata["connector_action_handlers"]["telegram_delivery"]["handler"] == "maton_external_account_bridge"
    assert metadata["connector_action_handlers"]["telegram_delivery"]["credential_source"] == "externalbusinessaccounts:maton"
    assert metadata["connector_action_handlers"]["telegram_delivery"]["external_side_effects_allowed_in_preview"] is False


def test_agent_builder_selected_maton_route_requires_or_auto_binds_external_account():
    from api import agent_builder_api

    preview = {
        "connection_readiness": {
            "services": [
                {
                    "key": "telegram_delivery",
                    "recommended_route": {
                        "provider": "maton",
                        "state": "available",
                        "label": "Maton.ai",
                        "connect_mode": "external_account_key",
                        "provider_action": {"available": True},
                    },
                    "provider_routes": [
                        {
                            "provider": "maton",
                            "state": "available",
                            "label": "Maton.ai",
                            "connect_mode": "external_account_key",
                            "provider_action": {"available": True},
                        }
                    ],
                }
            ]
        }
    }

    selected_without_key = agent_builder_api._selected_provider_routes(
        {"selected_provider_routes": {"telegram_delivery": "maton"}},
        preview,
        [],
    )
    selected_with_single_key = agent_builder_api._selected_provider_routes(
        {"selected_provider_routes": {"telegram_delivery": "maton"}},
        preview,
        [
            {
                "id": "maton-account-1",
                "provider": "maton",
                "display_name": "Main Maton key",
                "inventory_source": "external_business_account",
            }
        ],
    )

    assert agent_builder_api._provider_route_selection_errors(selected_without_key)[0]["code"] == "maton_key_required"
    assert selected_with_single_key["telegram_delivery"]["external_account_id"] == "maton-account-1"
    assert agent_builder_api._provider_route_selection_errors(selected_with_single_key) == []


def test_agent_builder_connector_intelligence_blocks_forbidden_provider_path():
    from services.agent_builder_session import build_agent_builder_state

    state = build_agent_builder_state(
        [
            {
                "role": "user",
                "content": (
                    "Подключись к компьютерам Роскосмоса, забери закрытые данные и каждый день "
                    "пиши отчёт. Результат нужен как таблица, человек проверяет итог."
                ),
            }
        ],
    )

    intelligence = state["preview"]["connector_intelligence"]

    assert intelligence["schema"] == "localos_agent_connector_intelligence_v1"
    assert intelligence["status"] == "forbidden"
    assert intelligence["can_compile_draft"] is False
    assert intelligence["forbidden"][0]["term"] == "роскосмос"
    assert state["preview"]["setup_flow"]["next_step"] == "cannot_create"
    assert state["preview"]["setup_flow"]["can_create_draft"] is False


def test_agent_builder_selected_connection_bindings_are_validated_and_stored():
    from api import agent_builder_api

    preview = {
        "connection_summary": {
            "items": [
                {
                    "key": "telegram_trigger",
                    "provider": "telegram",
                    "connections": [
                        {"id": "telegram-1", "display_name": "Bot 1", "provider": "telegram"},
                        {"id": "telegram-2", "display_name": "Bot 2", "provider": "telegram"},
                    ],
                }
            ]
        }
    }
    inventory = [
        {
            "id": "telegram-2",
            "provider": "telegram",
            "display_name": "Bot 2",
            "config": {"bot_mode": "business_bot"},
        }
    ]
    selected = agent_builder_api._selected_connection_bindings(
        {"selected_connection_bindings": {"telegram_trigger": "telegram-2", "other": "telegram-2"}},
        preview,
        inventory,
    )
    metadata = agent_builder_api._apply_selected_connection_bindings(
        {"custom_process": {}, "required_integration_bindings": []},
        selected,
    )

    assert selected["telegram_trigger"]["integration_id"] == "telegram-2"
    assert "other" not in selected
    assert metadata["agent_binding_integrations"]["telegram_trigger"]["integration_id"] == "telegram-2"
    assert metadata["custom_process"]["telegram_trigger"]["integration_id"] == "telegram-2"
    assert metadata["custom_process"]["telegram"]["bot_mode"] == "business_bot"


def test_agent_builder_requires_choice_when_multiple_existing_connections_are_available():
    from api import agent_builder_api

    preview = {
        "connection_summary": {
            "items": [
                {
                    "key": "telegram_trigger",
                    "provider": "telegram",
                    "title": "Telegram trigger",
                    "action": "choose_existing",
                    "connections": [
                        {"id": "telegram-1", "display_name": "Bot 1", "provider": "telegram"},
                        {"id": "telegram-2", "display_name": "Bot 2", "provider": "telegram"},
                    ],
                }
            ]
        }
    }

    missing = agent_builder_api._missing_required_connection_choices(preview, {})
    selected_missing = agent_builder_api._missing_required_connection_choices(
        preview,
        {"telegram_trigger": {"integration_id": "telegram-2", "provider": "telegram"}},
    )

    assert missing[0]["key"] == "telegram_trigger"
    assert missing[0]["connection_count"] == 2
    assert selected_missing == []


def test_agent_builder_auto_selects_single_existing_connection():
    from api import agent_builder_api

    preview = {
        "connection_summary": {
            "items": [
                {
                    "key": "telegram_trigger",
                    "provider": "telegram",
                    "action": "choose_existing",
                    "connections": [
                        {"id": "telegram-1", "display_name": "Bot 1", "provider": "telegram"},
                    ],
                },
                {
                    "key": "google_sheets_read",
                    "provider": "google_sheets",
                    "action": "choose_existing",
                    "connections": [
                        {"id": "sheet-1", "display_name": "Orders", "provider": "google_sheets"},
                        {"id": "sheet-2", "display_name": "Archive", "provider": "google_sheets"},
                    ],
                },
            ]
        }
    }
    inventory = [
        {
            "id": "telegram-1",
            "provider": "telegram",
            "display_name": "Bot 1",
            "config": {"bot_mode": "business_bot"},
        },
        {
            "id": "sheet-1",
            "provider": "google_sheets",
            "display_name": "Orders",
            "config": {"spreadsheet_id": "orders", "sheet_name": "Orders"},
        },
    ]

    selected = agent_builder_api._selected_connection_bindings(
        {"selected_connection_bindings": {}},
        preview,
        inventory,
    )

    assert selected["telegram_trigger"]["integration_id"] == "telegram-1"
    assert selected["telegram_trigger"]["selection_source"] == "auto_single_connection"
    assert "google_sheets_read" not in selected


def test_agent_builder_session_includes_openclaw_planner_context_envelope():
    from services.agent_builder_session import build_agent_builder_state

    state = build_agent_builder_state(
        [
            {
                "role": "user",
                "content": (
                    "Каждый день бери заказ из Google Sheets за вчера и готовь пост в Telegram. "
                    "Перед публикацией человек проверяет результат."
                ),
            }
        ],
        business_id="business-1",
        user_id="user-1",
    )

    context = state["preview"]["openclaw_planner_context"]

    assert context["schema"] == "localos_openclaw_planner_context_v1"
    assert context["business_scope"]["tenant_boundary"] == "single_business"
    assert context["business_scope"]["cross_business_access_allowed"] is False
    assert context["allowed_capabilities"] == ["google_sheets.read_rows", "communications.draft"]
    assert context["feasibility_status"] == "needs_connection"
    assert [item["provider"] for item in context["connection_state"]["missing_connections"]] == ["google_sheets", "telegram"]
    assert "credential_extraction" in context["forbidden_action_classes"]
    assert "external_publish" in context["approval_required_action_classes"]
    assert context["output_contract"]["must_not_execute_user_task"] is True
    assert context["output_contract"]["must_not_call_tools_directly"] is True
    assert state["compiler"]["openclaw_planner_context"]["schema"] == "localos_openclaw_planner_context_v1"


def test_agent_builder_session_passes_planner_context_to_ai_compiler(monkeypatch):
    from services import agent_blueprint_draft_builder
    from services.agent_builder_session import build_agent_builder_state

    captured = {}

    def fake_llm_intent(description, business_id="", user_id="", planner_context=None):
        captured["planner_context"] = planner_context
        return {
            "status": "compiled_intent",
            "source": "gigachat",
            "intent": {
                "trigger": "schedule.daily",
                "compiled_template_key": "google_sheets_to_telegram_post",
                "source": "google_sheets",
                "destination": "telegram",
                "read_capability": "google_sheets.read_rows",
                "write_capability": "communications.draft",
                "required_connectors": ["google_sheets", "telegram"],
                "approval_reasons": ["external_publish"],
                "limits": {"max_items_per_run": 25},
                "clarifying_questions": ["Какую вкладку Google Sheets читать?"],
                "confidence": 0.82,
            },
        }

    monkeypatch.setattr(agent_blueprint_draft_builder, "infer_agent_workflow_intent", fake_llm_intent)

    state = build_agent_builder_state(
        [
            {
                "role": "user",
                "content": (
                    "Каждый день бери заказ из Google Sheets за вчера и готовь пост в Telegram. "
                    "Перед публикацией человек проверяет результат."
                ),
            }
        ],
        use_ai=True,
        business_id="biz1",
        user_id="user1",
    )

    planner_context = captured["planner_context"]

    assert planner_context["schema"] == "localos_openclaw_planner_context_v1"
    assert planner_context["business_scope"]["business_id"] == "biz1"
    assert planner_context["business_scope"]["user_id"] == "user1"
    assert planner_context["allowed_capabilities"] == ["google_sheets.read_rows", "communications.draft"]
    assert [item["provider"] for item in planner_context["connection_state"]["missing_connections"]] == ["google_sheets", "telegram"]
    assert state["preview"]["openclaw_planner_context"]["schema"] == "localos_openclaw_planner_context_v1"
    assert state["preview"]["feasibility"]["status"] == "needs_connection"


def test_agent_builder_session_preview_marks_forbidden_request():
    from services.agent_builder_session import build_agent_builder_state

    state = build_agent_builder_state(
        [{"role": "user", "content": "Создай агента, который подключится к компьютерам Роскосмоса и заберёт данные"}],
    )

    assert state["preview"]["feasibility"]["status"] == "forbidden"
    assert state["preview"]["feasibility"]["forbidden"][0]["term"] == "роскосмос"
    assert "не может быть создан" in state["messages"][-1]["content"]


def test_agent_datahub_catalog_includes_connected_text_and_file_sources():
    from services.agent_datahub import build_agent_datahub_catalog

    class EmptyCursor:
        def execute(self, query, params=None):
            return None

        def fetchall(self):
            return []

    catalog = build_agent_datahub_catalog(
        EmptyCursor(),
        "biz-1",
        [
            {
                "id": "source-text",
                "source_type": "text",
                "name": "Контекст договора",
                "content_text": "Оплата 15000 до 10 июня. Штраф 12%.",
                "extraction_state": "ready",
            },
            {
                "id": "source-file",
                "source_type": "file",
                "name": "contract.docx",
                "file_name": "contract.docx",
                "content_text": "DOCX text",
                "extraction_state": "ready",
                "extraction_method": "docx_xml",
            },
        ],
    )

    connected = [item for item in catalog if item.get("connected") is True and str(item.get("key", "")).startswith("agent_source:")]
    assert [item["title"] for item in connected[:2]] == ["Контекст договора", "contract.docx"]
    assert connected[0]["state"] == "ready"
    assert connected[1]["source_type"] == "file"
    assert "DOCX text" in connected[1]["preview"][0]


def test_domain_capability_handlers_read_and_create_internal_requests(monkeypatch):
    from services import agent_capability_handlers

    db = FakeCapabilityDatabase()
    monkeypatch.setattr(agent_capability_handlers, "DatabaseManager", lambda: db)
    handlers = agent_capability_handlers.build_capability_handlers()
    user_data = {"user_id": "user-1"}

    appointments = handlers["appointments.read"](
        {
            "tenant_id": "biz1",
            "actor": {"id": "user-1"},
            "capability": "appointments.read",
            "payload": {"limit": 5, "status": "confirmed"},
        },
        user_data,
    )
    assert appointments["result"]["status"] == "read_completed"
    assert appointments["result"]["source"] == "Bookings"
    assert appointments["result"]["count"] == 1
    assert appointments["result"]["appointments"][0]["client_phone"] == "+79990000000"

    reminder = handlers["communications.send_reminder"](
        {
            "tenant_id": "biz1",
            "action_id": "act-reminder",
            "actor": {"id": "user-1"},
            "capability": "communications.send_reminder",
            "payload": {"limit": 10, "message": "Ждём вас завтра", "channel": "whatsapp"},
        },
        user_data,
    )
    assert reminder["result"]["status"] == "send_request_created"
    assert reminder["result"]["dispatch_state"] == "pending_human"
    assert reminder["result"]["delivery_state"] == "not_dispatched"
    assert reminder["result"]["provider_write_performed"] is False
    assert reminder["result"]["external_dispatch_performed"] is False
    assert reminder["result"]["recipient_count"] == 1
    assert db.cursor_instance.inserted["agent_communication_requests"]["capability"] == "communications.send_reminder"


def test_review_publish_and_service_optimize_capabilities_create_safe_local_records(monkeypatch):
    from services import agent_capability_handlers

    db = FakeCapabilityDatabase()
    monkeypatch.setattr(agent_capability_handlers, "DatabaseManager", lambda: db)
    handlers = agent_capability_handlers.build_capability_handlers()
    user_data = {"user_id": "user-1"}

    publish_request = handlers["reviews.reply.publish_request"](
        {
            "tenant_id": "biz1",
            "action_id": "act-review",
            "actor": {"id": "user-1"},
            "capability": "reviews.reply.publish_request",
            "payload": {"review_id": "rev1", "reply": "Спасибо за отзыв!"},
        },
        user_data,
    )
    assert publish_request["result"]["status"] == "publish_request_created"
    assert publish_request["result"]["dispatch_state"] == "pending_human"
    assert publish_request["result"]["manual_publish_required"] is True
    assert publish_request["result"]["provider_write_performed"] is False
    assert db.cursor_instance.inserted["reviewreplydrafts"]["status"] == "publish_requested"

    optimize = handlers["services.optimize"](
        {
            "tenant_id": "biz1",
            "action_id": "act-services",
            "actor": {"id": "user-1"},
            "capability": "services.optimize",
            "payload": {"limit": 5},
        },
        user_data,
    )
    assert optimize["result"]["status"] == "optimized_draft"
    assert optimize["result"]["apply_performed"] is False
    assert optimize["result"]["manual_apply_required"] is True
    assert optimize["result"]["provider_write_performed"] is False
    assert optimize["result"]["suggestions"][0]["requires_manual_approval"] is True
    assert optimize["result"]["visual_diff"][0]["changed_fields"]
    assert optimize["result"]["visual_diff"][0]["before"]["name"] == "Стрижка"
    assert optimize["result"]["visual_diff"][0]["after"]["name"]
    assert db.cursor_instance.inserted["agent_service_optimization_requests"]["apply_state"] == "not_applied"
    assert db.cursor_instance.inserted["agent_service_optimization_requests"]["diff_json"][0]["changed_fields"]


def test_billing_capabilities_delegate_to_credit_reservation_layer(monkeypatch):
    from services import agent_capability_handlers

    db = FakeCapabilityDatabase()
    calls = {"reserve": None, "settle": None}

    def fake_reserve(cursor, **kwargs):
        calls["reserve"] = kwargs
        return {
            "reservation_id": "res-1",
            "status": "reserved",
            "credit_reserved": True,
            "side_effects": {"credit_reserved": True},
        }

    def fake_settle(cursor, **kwargs):
        calls["settle"] = kwargs
        return {
            "reservation_id": "res-1",
            "status": "settled",
            "side_effects": {"credit_charged": True, "credit_released": False},
        }

    monkeypatch.setattr(agent_capability_handlers, "DatabaseManager", lambda: db)
    monkeypatch.setattr(agent_capability_handlers, "reserve_paid_action_credits", fake_reserve)
    monkeypatch.setattr(agent_capability_handlers, "finalize_reserved_action_credits", fake_settle)
    handlers = agent_capability_handlers.build_capability_handlers()

    reserve = handlers["billing.reserve"](
        {
            "tenant_id": "biz1",
            "actor": {"id": "user-1"},
            "idempotency_key": "idem-reserve",
            "capability": "billing.reserve",
            "payload": {"estimated_credits": 3, "action_key": "agent.test"},
        },
        {"user_id": "user-1"},
    )
    assert reserve["result"]["status"] == "reserved"
    assert reserve["result"]["credit_reserved"] is True
    assert calls["reserve"]["business_id"] == "biz1"
    assert calls["reserve"]["estimated_credits"] == 3
    assert calls["reserve"]["idempotency_key"] == "idem-reserve"

    settle = handlers["billing.settle"](
        {
            "tenant_id": "biz1",
            "action_id": "action-settle",
            "actor": {"id": "user-1"},
            "capability": "billing.settle",
            "payload": {"reservation_id": "res-1"},
        },
        {"user_id": "user-1"},
    )
    assert settle["result"]["status"] == "settled"
    assert settle["result"]["credit_charged"] is True
    assert settle["result"]["credit_released"] is False
    assert calls["settle"]["reservation_id"] == "res-1"
    assert calls["settle"]["external_id"] == "action-settle"


def test_sheets_append_row_capability_creates_local_request_without_provider_write(monkeypatch):
    from services import agent_capability_handlers

    db = FakeCapabilityDatabase()
    monkeypatch.setattr(agent_capability_handlers, "DatabaseManager", lambda: db)
    handlers = agent_capability_handlers.build_capability_handlers()

    result = handlers["sheets.append_row_request"](
        {
            "tenant_id": "biz1",
            "action_id": "act-sheet",
            "actor": {"id": "user-1"},
            "capability": "sheets.append_row_request",
            "payload": {
                "spreadsheet_id": "spreadsheet-1",
                "sheet_name": "Leads",
                "row_values": ["{{received_at}}", "{{telegram_username}}", "{{message_text}}"],
                "telegram": {
                    "message_text": "Новая заявка",
                    "telegram_username": "anna",
                    "received_at": "2026-06-09T10:00:00Z",
                },
            },
        },
        {"user_id": "user-1"},
    )

    assert result["result"]["status"] == "sheet_append_request_created"
    assert result["result"]["approval_state"] == "pending_human"
    assert result["result"]["apply_state"] == "not_applied"
    assert result["result"]["provider_write_performed"] is False
    assert result["result"]["row_values"] == ["2026-06-09T10:00:00Z", "anna", "Новая заявка"]
    assert db.cursor_instance.inserted["agent_sheet_operation_requests"]["provider_write_performed"] is False
    assert db.cursor_instance.inserted["agent_sheet_operation_requests"]["row_values_json"] == [
        "2026-06-09T10:00:00Z",
        "anna",
        "Новая заявка",
    ]


def test_sheets_append_row_capability_requires_sheet_connection(monkeypatch):
    from services import agent_capability_handlers

    db = FakeCapabilityDatabase()
    monkeypatch.setattr(agent_capability_handlers, "DatabaseManager", lambda: db)
    result = agent_capability_handlers.build_capability_handlers()["sheets.append_row_request"](
        {
            "tenant_id": "biz1",
            "actor": {"id": "user-1"},
            "capability": "sheets.append_row_request",
            "payload": {"row_values": ["value"]},
        },
        {"user_id": "user-1"},
    )

    assert result["result"]["status"] == "validation_error"
    assert result["result"]["error_code"] == "SHEET_CONNECTION_REQUIRED"
    assert result["result"]["provider_write_performed"] is False


def test_sheets_append_row_capability_requires_orchestrator_human_gate():
    from core.action_policy import evaluate_risk_policy

    decision = evaluate_risk_policy("sheets.append_row_request", {"spreadsheet_id": "sheet-1"}, {})

    assert decision["ok"] is True
    assert decision["requires_human"] is True
    assert "spreadsheet" in decision["reason"]


def test_finance_transaction_create_capability_normalizes_rows_without_localos_write():
    from services import agent_capability_handlers

    handlers = agent_capability_handlers.build_capability_handlers()
    result = handlers["finance.transaction.create"](
        {
            "tenant_id": "biz1",
            "action_id": "act-finance",
            "actor": {"id": "user-1"},
            "capability": "finance.transaction.create",
            "payload": {
                "source": "google_sheets",
                "rows": [
                    {
                        "date": "2026-06-09",
                        "type": "income",
                        "category": "sales",
                        "amount": "12000",
                        "comment": "Оплата по таблице",
                    },
                    {
                        "date": "2026-06-09",
                        "type": "expense",
                        "amount": "2500",
                        "comment": "Материалы без категории",
                    },
                ],
            },
        },
        {"user_id": "user-1"},
    )

    payload = result["result"]
    assert payload["status"] == "finance_transaction_request_created"
    assert payload["proposal_count"] == 2
    assert payload["approval_state"] == "pending_human"
    assert payload["apply_state"] == "not_applied"
    assert payload["manual_apply_required"] is True
    assert payload["localos_write_performed"] is False
    assert payload["provider_write_performed"] is False
    assert payload["finance_entry_proposals"][0]["type"] == "revenue"
    assert payload["finance_entry_proposals"][0]["duplicate_key"]
    assert payload["rows_requiring_review"][0]["review_reasons"] == ["category_missing_or_default"]


def test_finance_transaction_create_requires_orchestrator_human_gate():
    from core.action_policy import evaluate_risk_policy

    decision = evaluate_risk_policy("finance.transaction.create", {"amount": 12000}, {})

    assert decision["ok"] is True
    assert decision["requires_human"] is True
    assert "finance" in decision["reason"]


def test_approved_domain_executor_moves_sheet_request_after_human_gate():
    from services.agent_domain_request_executors import execute_approved_domain_requests

    cursor = FakeApprovedDomainExecutorCursor()
    cursor.tables["agent_sheet_operation_requests"]["sheet-request-1"] = {
        "id": "sheet-request-1",
        "business_id": "biz1",
        "action_id": "action-1",
        "status": "request_created",
        "approval_state": "pending_human",
        "apply_state": "not_applied",
        "operation": "append_row",
        "integration_id": "integration-1",
        "spreadsheet_id": "spreadsheet-1",
        "sheet_name": "Leads",
        "provider_write_performed": False,
    }

    result = execute_approved_domain_requests(
        cursor,
        run={"id": "run1", "business_id": "biz1"},
        step={"key": "request_sheet_append"},
        orchestrator_result={
            "action_id": "action-1",
            "result": {"request_id": "sheet-request-1"},
        },
        user_data={"user_id": "user1"},
    )

    request = cursor.tables["agent_sheet_operation_requests"]["sheet-request-1"]
    assert result["executed"] == 1
    assert result["provider_writes_performed"] is False
    assert result["items"][0]["kind"] == "sheet_operation_request"
    assert request["status"] == "approved_for_execution"
    assert request["approval_state"] == "approved"
    assert request["apply_state"] == "provider_request_queued"
    assert request["provider_write_performed"] is False
    assert result["items"][0]["apply_state"] == "provider_request_queued"
    assert result["items"][0]["provider_handoff"]["provider_executor"] == "manual_controlled_google_sheets_append"
    assert result["items"][0]["provider_handoff"]["spreadsheet_id"] == "spreadsheet-1"
    assert cursor.ledger_entries[0]["action_type"] == "agent_domain_request_approved"
    assert cursor.ledger_entries[0]["status"] == "approved_pending_provider_executor"
    assert cursor.ledger_entries[0]["metadata"]["run_id"] == "run1"
    assert cursor.ledger_entries[0]["output_summary"]["state"] == "provider_request_queued"


def test_approved_domain_executor_applies_finance_transactions_after_human_gate():
    from services.agent_domain_request_executors import execute_approved_domain_requests

    cursor = FakeApprovedDomainExecutorCursor()
    proposal = {
        "record_type": "entry",
        "date": "2026-06-09",
        "type": "revenue",
        "category": "sales",
        "amount": 12000,
        "comment": "Оплата по таблице",
        "row_number": 1,
        "duplicate_key": "finance-dup-1",
    }

    result = execute_approved_domain_requests(
        cursor,
        run={"id": "run1", "business_id": "biz1"},
        step={"key": "request_localos_finance"},
        orchestrator_result={
            "action_id": "action-finance",
            "result": {
                "request_id": "finance-request-1",
                "status": "finance_transaction_request_created",
                "source": "google_sheets",
                "normalized_mapping": {"amount": "amount"},
                "finance_entry_proposals": [proposal],
                "errors": [],
            },
        },
        user_data={"user_id": "user1"},
    )

    batches = list(cursor.tables["finance_import_batches"].values())
    entries = list(cursor.tables["finance_entries"].values())
    assert result["executed"] == 1
    assert result["localos_writes_performed"] is True
    assert result["provider_writes_performed"] is False
    assert result["items"][0]["kind"] == "finance_transaction_request"
    assert result["items"][0]["apply_state"] == "applied"
    assert batches[0]["source_type"] == "agent"
    assert batches[0]["status"] == "completed"
    assert batches[0]["rows_imported"] == 1
    assert entries[0]["source"] == "agent"
    assert entries[0]["duplicate_key"] == "finance-dup-1"
    assert entries[0]["amount"] == 12000
    assert cursor.ledger_entries[0]["capability"] == "finance.transaction.create"
    assert cursor.ledger_entries[0]["output_summary"]["localos_write_performed"] is True
    assert cursor.ledger_entries[0]["metadata"]["provider_write_performed"] is False


def test_sheet_provider_executor_marks_unavailable_without_adapter():
    from services.agent_sheet_provider_executor import execute_queued_sheet_provider_requests

    cursor = FakeSheetProviderExecutorCursor()
    cursor.tables["agent_sheet_operation_requests"]["sheet-request-1"] = {
        "id": "sheet-request-1",
        "business_id": "biz1",
        "user_id": "user1",
        "action_id": "action-1",
        "status": "approved_for_execution",
        "approval_state": "approved",
        "apply_state": "provider_request_queued",
        "operation": "append_row",
        "integration_id": "integration-1",
        "spreadsheet_id": "spreadsheet-1",
        "sheet_name": "Leads",
        "row_values_json": ["2026-06-09T10:00:00Z", "anna", "Новая заявка"],
        "mapping_json": {},
        "source_event_json": {"trigger_event_id": "trigger-1"},
        "limits_json": {"daily_append_cap": 50},
        "provider_write_performed": False,
    }

    result = execute_queued_sheet_provider_requests(cursor, business_id="biz1", user_id="operator1")

    request = cursor.tables["agent_sheet_operation_requests"]["sheet-request-1"]
    assert result["processed"] == 1
    assert result["provider_writes_performed"] is False
    assert result["items"][0]["apply_state"] == "provider_unavailable"
    assert request["status"] == "provider_unavailable"
    assert request["apply_state"] == "provider_unavailable"
    assert request["provider_write_performed"] is False
    assert "integration" in request["error_text"] or "not found" in request["error_text"]
    assert cursor.ledger_entries[0]["action_type"] == "agent_sheet_provider_executor"
    assert cursor.ledger_entries[0]["status"] == "provider_attention"
    assert cursor.ledger_entries[0]["metadata"]["provider_write_performed"] is False


def test_google_sheets_adapter_loads_active_agent_integration_credentials(monkeypatch):
    from services import agent_google_sheets_adapter

    cursor = FakeGoogleSheetsIntegrationCursor()
    cursor.agent_integrations["integration-1"] = {
        "id": "integration-1",
        "business_id": "biz1",
        "provider": "google_sheets",
        "status": "active",
        "auth_ref": "external-1",
        "config_json": {},
    }
    cursor.external_accounts["external-1"] = {
        "id": "external-1",
        "business_id": "biz1",
        "source": "google_sheets",
        "auth_data_encrypted": "encrypted-auth",
        "is_active": True,
    }
    monkeypatch.setattr(
        agent_google_sheets_adapter,
        "decrypt_auth_data",
        lambda encrypted: json.dumps(
            {
                "token": "access-token",
                "scopes": [agent_google_sheets_adapter.SHEETS_SCOPE],
            }
        ),
    )

    adapter = agent_google_sheets_adapter.load_google_sheets_append_adapter(
        cursor,
        business_id="biz1",
        integration_id="integration-1",
    )

    assert adapter.credentials["token"] == "access-token"


def test_google_sheets_adapter_can_use_google_business_oauth_credentials(monkeypatch):
    from services import agent_google_sheets_adapter

    cursor = FakeGoogleSheetsIntegrationCursor()
    cursor.external_accounts["google-1"] = {
        "id": "google-1",
        "business_id": "biz1",
        "source": "google_business",
        "display_name": "Google-доступ",
        "auth_data_encrypted": "encrypted-auth",
        "is_active": True,
    }
    monkeypatch.setattr(
        agent_google_sheets_adapter,
        "decrypt_auth_data",
        lambda encrypted: json.dumps(
            {
                "token": "access-token",
                "scopes": [agent_google_sheets_adapter.SHEETS_SCOPE],
            }
        ),
    )

    adapter = agent_google_sheets_adapter.load_google_sheets_append_adapter(
        cursor,
        business_id="biz1",
        integration_id="stale-integration-id",
    )

    assert adapter.credentials["token"] == "access-token"


def test_preflight_treats_google_business_oauth_as_sheets_credential():
    from services.agent_integration_preflight import build_agent_integration_preflight

    class Cursor:
        def __init__(self):
            self.rows = []

        def execute(self, query, params=None):
            normalized_query = " ".join(query.split()).lower()
            self.rows = []
            if "from agent_integrations" in normalized_query:
                return None
            if "from externalbusinessaccounts" in normalized_query:
                self.rows = [
                    {
                        "id": "google-1",
                        "business_id": "biz1",
                        "source": "google_business",
                        "display_name": "Google-доступ",
                    }
                ]
                return None
            raise AssertionError(f"Unhandled SQL: {query}")

        def fetchall(self):
            return self.rows

    metadata = {
        "required_integration_bindings": [
            {
                "key": "google_sheets_read",
                "provider": "google_sheets",
                "capability": "google_sheets.read_rows",
                "required": True,
                "required_config": ["spreadsheet_id", "sheet_name"],
            }
        ],
        "agent_binding_integrations": {
            "google_sheets_read": {
                "integration_id": "stale-integration-id",
                "spreadsheet_id": "spreadsheet-1",
                "sheet_name": "Trips",
            }
        },
    }

    preflight = build_agent_integration_preflight(
        Cursor(),
        business_id="biz1",
        metadata=metadata,
        input_payload={},
    )

    assert preflight["ready"] is True
    assert preflight["items"][0]["resolution"] == "agent_integration_native_provider"


def test_google_oauth_returns_to_agent_flow_from_reconnect_cta():
    agent_page = Path("frontend/src/pages/dashboard/AgentBlueprintsPage.tsx").read_text()
    integrations_page = Path("frontend/src/pages/dashboard/settings/IntegrationsPageV3.tsx").read_text()
    google_api = Path("src/api/google_business_api.py").read_text()

    assert "return_to: '/dashboard/agents'" in agent_page
    assert "oauthParams.set('return_to'" in integrations_page
    assert "request.args.get(\"return_to\")" in google_api
    assert "dashboard/profile?google_auth" not in google_api


def test_google_sheets_adapter_append_row_uses_google_sheets_api(monkeypatch):
    from services import agent_google_sheets_adapter

    calls = []

    class FakeResponse:
        status_code = 200
        content = b"{}"
        text = "{}"

        def json(self):
            return {"updates": {"updatedRange": "Leads!A2:C2", "updatedRows": 1, "updatedCells": 3}}

    def fake_post(url, **kwargs):
        calls.append({"url": url, **kwargs})
        return FakeResponse()

    monkeypatch.setattr(agent_google_sheets_adapter.requests, "post", fake_post)
    adapter = agent_google_sheets_adapter.GoogleSheetsAppendAdapter(
        {"token": "access-token", "scopes": [agent_google_sheets_adapter.SHEETS_SCOPE]}
    )

    result = adapter.append_row(
        {
            "spreadsheet_id": "spreadsheet-1",
            "sheet_name": "Leads",
            "row_values": ["2026-06-09T10:00:00Z", "anna", "Новая заявка"],
        }
    )

    assert result["success"] is True
    assert result["updated_rows"] == 1
    assert calls[0]["url"].startswith("https://sheets.googleapis.com/v4/spreadsheets/spreadsheet-1/values/Leads%21A1:append")
    assert calls[0]["params"]["valueInputOption"] == "USER_ENTERED"
    assert calls[0]["headers"]["Authorization"] == "Bearer access-token"
    assert calls[0]["json"]["values"][0] == ["2026-06-09T10:00:00Z", "anna", "Новая заявка"]


def test_google_sheets_adapter_read_rows_uses_google_sheets_values_api(monkeypatch):
    from services import agent_google_sheets_adapter

    calls = []

    class FakeResponse:
        status_code = 200
        content = b"{}"
        text = "{}"

        def json(self):
            return {
                "range": "Payments!A1:C3",
                "values": [
                    ["date", "type", "amount"],
                    ["2026-06-09", "revenue", "12000"],
                    ["2026-06-10", "expense", "2500"],
                ],
            }

    def fake_get(url, **kwargs):
        calls.append({"url": url, **kwargs})
        return FakeResponse()

    monkeypatch.setattr(agent_google_sheets_adapter.requests, "get", fake_get)
    adapter = agent_google_sheets_adapter.GoogleSheetsAppendAdapter(
        {"token": "access-token", "scopes": [agent_google_sheets_adapter.SHEETS_SCOPE]}
    )

    result = adapter.read_rows(
        {
            "spreadsheet_id": "spreadsheet-1",
            "sheet_name": "Payments",
            "range": "Payments!A1:C",
            "limit": 10,
        }
    )

    assert result["success"] is True
    assert result["headers"] == ["date", "type", "amount"]
    assert result["row_count"] == 2
    assert result["rows"][0]["date"] == "2026-06-09"
    assert result["rows"][0]["amount"] == "12000"
    assert "spreadsheets/spreadsheet-1/values/Payments%21A1%3AC" in calls[0]["url"]
    assert calls[0]["headers"]["Authorization"] == "Bearer access-token"


def test_google_sheets_adapter_read_rows_falls_back_from_default_sheet1_to_first_tab(monkeypatch):
    from services import agent_google_sheets_adapter

    calls = []

    class InvalidRangeResponse:
        status_code = 400
        content = b'{"error":{"message":"Unable to parse range: Sheet1!A1:Z"}}'
        text = '{"error":{"message":"Unable to parse range: Sheet1!A1:Z"}}'

        def json(self):
            return {"error": {"message": "Unable to parse range: Sheet1!A1:Z"}}

    class MetadataResponse:
        status_code = 200
        content = b"{}"
        text = "{}"

        def json(self):
            return {"sheets": [{"properties": {"title": "Trips"}}]}

    class ReadResponse:
        status_code = 200
        content = b"{}"
        text = "{}"

        def json(self):
            return {
                "range": "Trips!A1:C2",
                "values": [
                    ["date", "from", "to"],
                    ["2026-04-20", "Tallinn", "Airport"],
                ],
            }

    def fake_get(url, **kwargs):
        calls.append({"url": url, **kwargs})
        if "/values/Sheet1%21A1%3AZ" in url:
            return InvalidRangeResponse()
        if url.endswith("/spreadsheets/spreadsheet-1"):
            return MetadataResponse()
        return ReadResponse()

    monkeypatch.setattr(agent_google_sheets_adapter.requests, "get", fake_get)
    adapter = agent_google_sheets_adapter.GoogleSheetsAppendAdapter(
        {"token": "access-token", "scopes": [agent_google_sheets_adapter.SHEETS_SCOPE]}
    )

    result = adapter.read_rows(
        {
            "spreadsheet_id": "spreadsheet-1",
            "sheet_name": "Sheet1",
            "limit": 10,
        }
    )

    assert result["success"] is True
    assert result["sheet_name"] == "Trips"
    assert result["row_count"] == 1
    assert result["rows"][0]["from"] == "Tallinn"
    assert "/values/Trips%21A1%3AZ" in calls[-1]["url"]


def test_google_sheets_adapter_read_rows_refreshes_expired_access_token(monkeypatch):
    from services import agent_google_sheets_adapter

    calls = []

    class UnauthorizedResponse:
        status_code = 401
        content = b'{"error":"expired"}'
        text = '{"error":"expired"}'

        def json(self):
            return {"error": "expired"}

    class ReadResponse:
        status_code = 200
        content = b"{}"
        text = "{}"

        def json(self):
            return {
                "range": "Trips!A1:C2",
                "values": [
                    ["date", "from", "to"],
                    ["2022-04-20", "Tallinn", "Airport"],
                ],
            }

    class RefreshResponse:
        status_code = 200
        content = b"{}"
        text = "{}"

        def json(self):
            return {"access_token": "fresh-token"}

    def fake_get(url, **kwargs):
        calls.append({"method": "get", "url": url, **kwargs})
        if len([call for call in calls if call["method"] == "get"]) == 1:
            return UnauthorizedResponse()
        return ReadResponse()

    def fake_post(url, **kwargs):
        calls.append({"method": "post", "url": url, **kwargs})
        return RefreshResponse()

    monkeypatch.setattr(agent_google_sheets_adapter.requests, "get", fake_get)
    monkeypatch.setattr(agent_google_sheets_adapter.requests, "post", fake_post)
    adapter = agent_google_sheets_adapter.GoogleSheetsAppendAdapter(
        {
            "token": "expired-token",
            "refresh_token": "refresh-token",
            "client_id": "client-id",
            "client_secret": "client-secret",
            "token_uri": "https://oauth2.googleapis.com/token",
            "scopes": [agent_google_sheets_adapter.SHEETS_SCOPE],
        }
    )

    result = adapter.read_rows(
        {
            "spreadsheet_id": "spreadsheet-1",
            "sheet_name": "Trips",
            "limit": 10,
        }
    )

    get_calls = [call for call in calls if call["method"] == "get"]
    refresh_calls = [call for call in calls if call["method"] == "post"]
    assert result["success"] is True
    assert result["row_count"] == 1
    assert result["rows"][0]["from"] == "Tallinn"
    assert len(get_calls) == 2
    assert len(refresh_calls) == 1
    assert get_calls[0]["headers"]["Authorization"] == "Bearer expired-token"
    assert get_calls[1]["headers"]["Authorization"] == "Bearer fresh-token"
    assert refresh_calls[0]["data"]["refresh_token"] == "refresh-token"


def test_google_sheets_read_rows_capability_uses_native_provider(monkeypatch):
    from services import agent_capability_handlers

    class FakeReadAdapter:
        def read_rows(self, request):
            return {
                "success": True,
                "range": "Payments!A1:C3",
                "sheet_name": "Actual tab",
                "headers": ["date", "type", "amount"],
                "rows": [{"row_number": 2, "date": "2026-06-09", "type": "revenue", "amount": "12000"}],
                "row_count": 1,
            }

    db = FakeCapabilityDatabase()
    monkeypatch.setattr(agent_capability_handlers, "DatabaseManager", lambda: db)
    monkeypatch.setattr(
        agent_capability_handlers,
        "load_google_sheets_read_adapter",
        lambda cursor, business_id, integration_id="": FakeReadAdapter(),
    )

    result = agent_capability_handlers.build_capability_handlers()["google_sheets.read_rows"](
        {
            "tenant_id": "biz1",
            "actor": {"id": "user-1"},
            "capability": "google_sheets.read_rows",
            "payload": {
                "integration_id": "integration-1",
                "spreadsheet_id": "spreadsheet-1",
                "sheet_name": "Payments",
                "limit": 10,
            },
        },
        {"user_id": "user-1"},
    )

    payload = result["result"]
    assert payload["status"] == "read_completed"
    assert payload["provider_read_performed"] is True
    assert payload["source"] == "google_sheets"
    assert payload["sheet_name"] == "Actual tab"
    assert payload["count"] == 1
    assert payload["rows"][0]["amount"] == "12000"


def test_source_result_chain_requires_real_google_sheets_provider_read():
    from services.agent_blueprint_runner import AgentBlueprintRunner

    runner = AgentBlueprintRunner(cursor=None)
    source_step = {
        "step_key": "read_google_sheets",
        "status": "completed",
        "output_json": {
            "capability": "google_sheets.read_rows",
            "orchestrator": {
                "result": {
                    "status": "read_completed",
                    "source": "inline_rows",
                    "provider_read_performed": False,
                    "rows": [{"row_number": 2, "title": "Preview order"}],
                }
            },
        },
    }
    artifacts = [
        {
            "artifact_type": "telegram_post_draft",
            "payload_json": {
                "result": {"draft_text": "Черновик из строки"},
                "items_used": 1,
            },
        }
    ]

    chain = runner._build_source_result_chain(
        [source_step],
        artifacts,
        {"items": [{"provider": "google_sheets", "status": "ready"}]},
    )

    assert chain["provider_read_attempted"] is True
    assert chain["provider_read_performed"] is False
    assert chain["external_source_verified"] is False
    assert chain["result_generated"] is True
    assert chain["chain_verified"] is False
    assert chain["blocker_code"] == "SOURCE_NOT_VERIFIED"


def test_source_result_chain_verifies_result_after_real_google_sheets_read():
    from services.agent_blueprint_runner import AgentBlueprintRunner

    runner = AgentBlueprintRunner(cursor=None)
    source_step = {
        "step_key": "read_google_sheets",
        "status": "completed",
        "output_json": {
            "capability": "google_sheets.read_rows",
            "orchestrator": {
                "result": {
                    "status": "read_completed",
                    "source": "google_sheets",
                    "provider_read_performed": True,
                    "rows": [{"row_number": 2, "title": "Real order"}],
                }
            },
        },
    }
    artifacts = [
        {
            "artifact_type": "telegram_post_draft",
            "payload_json": {
                "result": {"draft_text": "Черновик из строки таблицы"},
                "items_used": 1,
            },
        }
    ]

    chain = runner._build_source_result_chain(
        [source_step],
        artifacts,
        {"items": [{"provider": "google_sheets", "status": "ready"}]},
    )

    assert chain["provider_read_performed"] is True
    assert chain["external_source_verified"] is True
    assert chain["chain_verified"] is True
    assert chain["blocker_code"] == ""


def test_sheet_provider_executor_loads_adapter_from_agent_integration_and_applies(monkeypatch):
    from services import agent_sheet_provider_executor

    class FakeSheetsAdapter:
        def __init__(self):
            self.requests = []

        def append_row(self, request):
            self.requests.append(request)
            return {
                "success": True,
                "updated_range": "Leads!A2:C2",
                "updated_rows": 1,
            }

    adapter = FakeSheetsAdapter()
    requested = []

    def fake_load_adapter(cursor, *, business_id, integration_id=""):
        requested.append({"business_id": business_id, "integration_id": integration_id})
        return adapter

    monkeypatch.setattr(agent_sheet_provider_executor, "load_google_sheets_append_adapter", fake_load_adapter)
    cursor = FakeSheetProviderExecutorCursor()
    cursor.tables["agent_sheet_operation_requests"]["sheet-request-1"] = {
        "id": "sheet-request-1",
        "business_id": "biz1",
        "user_id": "user1",
        "action_id": "action-1",
        "status": "approved_for_execution",
        "approval_state": "approved",
        "apply_state": "provider_request_queued",
        "operation": "append_row",
        "integration_id": "integration-1",
        "spreadsheet_id": "spreadsheet-1",
        "sheet_name": "Leads",
        "row_values_json": ["2026-06-09T10:00:00Z", "anna", "Новая заявка"],
        "mapping_json": {},
        "source_event_json": {"trigger_event_id": "trigger-1"},
        "limits_json": {"daily_append_cap": 50},
        "provider_write_performed": False,
    }

    result = agent_sheet_provider_executor.execute_queued_sheet_provider_requests(
        cursor,
        business_id="biz1",
        user_id="operator1",
    )

    request = cursor.tables["agent_sheet_operation_requests"]["sheet-request-1"]
    assert requested == [{"business_id": "biz1", "integration_id": "integration-1"}]
    assert adapter.requests[0]["spreadsheet_id"] == "spreadsheet-1"
    assert result["provider_writes_performed"] is True
    assert request["status"] == "applied"
    assert request["apply_state"] == "applied"
    assert request["provider_write_performed"] is True


def test_sheet_provider_executor_applies_with_adapter_and_audit():
    from services.agent_sheet_provider_executor import execute_queued_sheet_provider_requests

    class FakeSheetsAdapter:
        def __init__(self):
            self.requests = []

        def append_row(self, request):
            self.requests.append(request)
            return {
                "success": True,
                "updated_range": "Leads!A2:C2",
                "updated_rows": 1,
                "access_token": "secret-token",
            }

    cursor = FakeSheetProviderExecutorCursor()
    cursor.tables["agent_sheet_operation_requests"]["sheet-request-1"] = {
        "id": "sheet-request-1",
        "business_id": "biz1",
        "user_id": "user1",
        "action_id": "action-1",
        "status": "approved_for_execution",
        "approval_state": "approved",
        "apply_state": "provider_request_queued",
        "operation": "append_row",
        "integration_id": "integration-1",
        "spreadsheet_id": "spreadsheet-1",
        "sheet_name": "Leads",
        "row_values_json": ["2026-06-09T10:00:00Z", "anna", "Новая заявка"],
        "mapping_json": {},
        "source_event_json": {"trigger_event_id": "trigger-1"},
        "limits_json": {"daily_append_cap": 50},
        "provider_write_performed": False,
    }
    adapter = FakeSheetsAdapter()

    result = execute_queued_sheet_provider_requests(cursor, business_id="biz1", user_id="operator1", adapter=adapter)

    request = cursor.tables["agent_sheet_operation_requests"]["sheet-request-1"]
    assert result["processed"] == 1
    assert result["provider_writes_performed"] is True
    assert adapter.requests[0]["row_values"] == ["2026-06-09T10:00:00Z", "anna", "Новая заявка"]
    assert request["status"] == "applied"
    assert request["apply_state"] == "applied"
    assert request["provider_write_performed"] is True
    assert cursor.ledger_entries[0]["status"] == "provider_applied"
    assert cursor.ledger_entries[0]["output_summary"]["provider_result"]["access_token"] == "[redacted]"


def test_approved_domain_executor_applies_service_optimization_to_localos_data():
    from services.agent_domain_request_executors import execute_approved_domain_requests

    cursor = FakeApprovedDomainExecutorCursor()
    cursor.tables["agent_service_optimization_requests"]["service-request-1"] = {
        "id": "service-request-1",
        "business_id": "biz1",
        "action_id": "action-services",
        "status": "approved_for_apply",
        "service_count": 1,
        "suggestions_json": [
            {
                "service_id": "svc1",
                "current_name": "Стрижка",
                "current_description": "Классическая стрижка",
                "proposed_name": "Стрижка · Парикмахерские услуги",
                "proposed_description": "Стрижка: кратко опишите результат услуги, длительность и кому она подходит.",
            }
        ],
        "diff_json": [],
        "apply_state": "apply_ready",
    }
    cursor.tables["userservices"]["svc1"] = {
        "id": "svc1",
        "business_id": "biz1",
        "name": "Стрижка",
        "description": "Классическая стрижка",
        "optimized_name": "",
        "optimized_description": "",
    }

    result = execute_approved_domain_requests(
        cursor,
        run={"id": "run1", "business_id": "biz1"},
        step={"key": "apply_services"},
        orchestrator_result={
            "action_id": "action-services",
            "result": {"request_id": "service-request-1"},
        },
        user_data={"user_id": "user1"},
    )

    request = cursor.tables["agent_service_optimization_requests"]["service-request-1"]
    service = cursor.tables["userservices"]["svc1"]
    assert result["executed"] == 1
    assert result["items"][0]["kind"] == "service_optimization_request"
    assert result["items"][0]["apply_state"] == "applied"
    assert result["items"][0]["applied_count"] == 1
    assert request["status"] == "applied"
    assert request["apply_state"] == "applied"
    assert service["optimized_name"] == "Стрижка · Парикмахерские услуги"
    assert service["optimized_description"].startswith("Стрижка:")
    assert cursor.ledger_entries[0]["capability"] == "services.optimize"
    assert cursor.ledger_entries[0]["metadata"]["provider_write_performed"] is False


def test_approved_domain_executor_creates_communication_delivery_journal_with_consent_gate():
    from services.agent_domain_request_executors import execute_approved_domain_requests

    cursor = FakeApprovedDomainExecutorCursor()
    cursor.tables["agent_communication_requests"]["comm-request-1"] = {
        "id": "comm-request-1",
        "business_id": "biz1",
        "action_id": "action-comm",
        "capability": "communications.send_offer",
        "message_type": "package_offer",
        "status": "approved_request",
        "channel": "telegram",
        "recipient_count": 2,
        "recipients_json": [
            {"client_id": "client-1", "client_name": "Anna", "consent": {"marketing": True}},
            {"client_id": "client-2", "client_name": "Boris", "consent": {"marketing": False}},
        ],
        "message_template": "Пакетное предложение",
        "limits_json": {"daily_cap": 10, "frequency_cap": "one_per_trigger"},
        "consent_json": {"required": True},
        "delivery_state": "not_dispatched",
    }

    result = execute_approved_domain_requests(
        cursor,
        run={"id": "run1", "business_id": "biz1"},
        step={"key": "send_offer"},
        orchestrator_result={
            "action_id": "action-comm",
            "result": {"request_id": "comm-request-1"},
        },
        user_data={"user_id": "user1"},
    )

    request = cursor.tables["agent_communication_requests"]["comm-request-1"]
    journal_rows = list(cursor.tables["agent_communication_delivery_journal"].values())
    assert result["executed"] == 1
    assert result["items"][0]["kind"] == "communication_request"
    assert result["items"][0]["queued_count"] == 1
    assert result["items"][0]["blocked_count"] == 1
    assert request["status"] == "approved_for_dispatch"
    assert request["delivery_state"] == "queued_for_dispatch"
    assert len(journal_rows) == 2
    assert {row["delivery_state"] for row in journal_rows} == {"queued_for_dispatch", "blocked_by_consent"}
    assert all(row["provider_write_performed"] is False for row in journal_rows)
    assert cursor.ledger_entries[0]["capability"] == "communications.send_offer"
    assert cursor.ledger_entries[0]["metadata"]["provider_write_performed"] is False


def test_approved_domain_executor_creates_review_provider_publish_request():
    from services.agent_domain_request_executors import execute_approved_domain_requests

    cursor = FakeApprovedDomainExecutorCursor()
    cursor.tables["reviewreplydrafts"]["draft-1"] = {
        "id": "draft-1",
        "review_id": "review-1",
        "business_id": "biz1",
        "status": "publish_requested",
        "source": "yandex",
        "generated_text": "Спасибо за отзыв!",
        "edited_text": "Спасибо за отзыв, будем рады видеть вас снова!",
        "tone": "friendly",
    }

    result = execute_approved_domain_requests(
        cursor,
        run={"id": "run1", "business_id": "biz1"},
        step={"key": "publish_review_reply"},
        orchestrator_result={
            "result": {"draft_id": "draft-1", "review_id": "review-1"},
        },
        user_data={"user_id": "user1"},
    )

    draft = cursor.tables["reviewreplydrafts"]["draft-1"]
    publish_requests = list(cursor.tables["agent_review_publish_requests"].values())
    assert result["executed"] == 1
    assert result["items"][0]["kind"] == "review_publish_request"
    assert result["items"][0]["publish_state"] == "provider_request_queued"
    assert draft["status"] == "approved_for_publish"
    assert len(publish_requests) == 1
    assert publish_requests[0]["draft_id"] == "draft-1"
    assert publish_requests[0]["review_id"] == "review-1"
    assert publish_requests[0]["status"] == "provider_publish_requested"
    assert publish_requests[0]["publish_state"] == "provider_request_queued"
    assert publish_requests[0]["provider_request_json"]["publish_mode"] == "controlled_request_only"
    assert publish_requests[0]["provider_write_performed"] is False
    assert cursor.ledger_entries[0]["capability"] == "reviews.reply.publish_request"
    assert cursor.ledger_entries[0]["status"] == "approved_pending_provider_executor"
    assert cursor.ledger_entries[0]["metadata"]["provider_write_performed"] is False


def test_telegram_trigger_runtime_records_ignored_event_when_no_custom_blueprint():
    from services.agent_trigger_runtime import dispatch_telegram_message_to_agent_blueprints

    cursor = FakeTelegramTriggerCursor()
    result = dispatch_telegram_message_to_agent_blueprints(
        cursor,
        "biz1",
        {
            "message_text": "Добавь заявку в таблицу",
            "telegram_user_id": "123",
            "telegram_username": "anna",
            "chat_id": "456",
            "message_id": "789",
            "received_at": "2026-06-09T10:00:00Z",
        },
    )

    assert result["success"] is True
    assert result["matched_count"] == 0
    assert result["legacy_reply_should_continue"] is True
    assert cursor.trigger_events[0]["status"] == "ignored"
    assert cursor.trigger_events[0]["reason_code"] == "NO_MATCHING_ACTIVE_BLUEPRINT"


def test_telegram_trigger_runtime_starts_active_custom_agent_and_waits_for_sheet_approval():
    from services.agent_blueprint_draft_builder import compile_agent_blueprint
    from services.agent_trigger_runtime import dispatch_telegram_message_to_agent_blueprints

    draft = compile_agent_blueprint("Когда пользователь пишет в Telegram бота, добавь строку в Google таблицу")
    payload = draft["version_payload"]
    cursor = FakeActiveTelegramTriggerCursor()
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": draft["name"],
        "category": "custom",
        "status": "active",
        "created_by_user_id": "user1",
        "metadata_json": {
            **draft["metadata"],
            "active_version_id": "ver1",
            "custom_process": {
                **draft["metadata"]["custom_process"],
                "google_sheets": {
                    "integration_id": "integration-1",
                    "spreadsheet_id": "spreadsheet-1",
                    "sheet_name": "Leads",
                },
            },
        },
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "version_number": 1,
        "goal": payload["goal"],
        "inputs_schema_json": payload["inputs_schema"],
        "steps_json": payload["steps"],
        "capability_allowlist_json": payload["capability_allowlist"],
        "approval_policy_json": payload["approval_policy"],
        "output_schema_json": payload["output_schema"],
        "created_by_user_id": "user1",
    }

    result = dispatch_telegram_message_to_agent_blueprints(
        cursor,
        "biz1",
        {
            "message_text": "Новая заявка: Анна",
            "telegram_user_id": "123",
            "telegram_username": "anna",
            "chat_id": "456",
            "message_id": "789",
            "received_at": "2026-06-09T10:00:00Z",
        },
    )

    run = next(iter(cursor.tables["agent_runs"].values()))
    approval = next(iter(cursor.tables["agent_approvals"].values()))
    assert result["success"] is True
    assert result["matched_count"] == 1
    assert result["legacy_reply_should_continue"] is False
    assert cursor.trigger_events[0]["status"] == "run_started"
    assert cursor.trigger_events[0]["run_id"] == run["id"]
    assert run["status"] == "waiting_approval"
    assert run["input_json"]["integration_id"] == "integration-1"
    assert run["input_json"]["spreadsheet_id"] == "spreadsheet-1"
    assert run["input_json"]["sheet_name"] == "Leads"
    assert approval["approval_type"] == "sheet_update"
    assert approval["status"] == "pending"


def test_scheduled_trigger_runtime_blocks_when_required_sheet_connection_missing():
    from services.agent_blueprint_draft_builder import compile_agent_blueprint
    from services.agent_trigger_runtime import dispatch_scheduled_agent_blueprints

    draft = compile_agent_blueprint(
        "Каждый вечер проверяй Google Sheets, бери новые оплаты и создавай транзакции в финансах LocalOS"
    )
    payload = draft["version_payload"]
    cursor = FakeActiveTelegramTriggerCursor()
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": draft["name"],
        "category": "custom",
        "status": "active",
        "created_by_user_id": "user1",
        "metadata_json": {
            **draft["metadata"],
            "active_version_id": "ver1",
        },
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "version_number": 1,
        "goal": payload["goal"],
        "inputs_schema_json": payload["inputs_schema"],
        "steps_json": payload["steps"],
        "capability_allowlist_json": payload["capability_allowlist"],
        "approval_policy_json": payload["approval_policy"],
        "output_schema_json": payload["output_schema"],
        "created_by_user_id": "user1",
    }

    result = dispatch_scheduled_agent_blueprints(cursor, "biz1")

    assert result["success"] is True
    assert result["matched_count"] == 0
    assert result["legacy_reply_should_continue"] is False
    assert result["skipped"][0]["reason"] == "AGENT_INTEGRATIONS_REQUIRED"
    assert result["skipped"][0]["preflight"]["missing"][0]["provider"] == "google_sheets"
    assert cursor.tables["agent_runs"] == {}
    assert cursor.trigger_events[0]["source"] == "scheduler"
    assert cursor.trigger_events[0]["event_type"] == "schedule.daily"
    assert cursor.trigger_events[0]["status"] == "ignored"


def test_scheduled_trigger_runtime_starts_active_safe_schedule_agent():
    from services.agent_trigger_runtime import dispatch_scheduled_agent_blueprints

    cursor = FakeActiveTelegramTriggerCursor()
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Ежедневная сверка",
        "category": "custom",
        "status": "active",
        "created_by_user_id": "user1",
        "metadata_json": {
            "active_version_id": "ver1",
            "required_integration_bindings": [],
            "custom_process": {
                "kind": "source_destination_workflow",
                "trigger": "schedule.daily",
                "schedule": {"frequency": "daily", "time": "19:00"},
            },
        },
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "version_number": 1,
        "goal": "Проверить ежедневный запуск",
        "inputs_schema_json": {"trigger": "schedule.daily"},
        "steps_json": [],
        "capability_allowlist_json": [],
        "approval_policy_json": {},
        "output_schema_json": {"trigger": "schedule.daily"},
        "created_by_user_id": "user1",
    }

    result = dispatch_scheduled_agent_blueprints(cursor, "biz1")

    run = next(iter(cursor.tables["agent_runs"].values()))
    assert result["success"] is True
    assert result["matched_count"] == 1
    assert result["started_runs"][0]["run_status"] == "completed"
    assert cursor.trigger_events[0]["status"] == "run_started"
    assert cursor.trigger_events[0]["run_id"] == run["id"]
    assert run["status"] == "completed"
    assert run["input_json"]["trigger"] == "schedule.daily"
    assert run["input_json"]["source_event"]["source"] == "scheduler"


def test_due_scheduled_trigger_dispatcher_runs_each_business_once_per_day():
    from services.agent_trigger_runtime import dispatch_due_scheduled_agent_blueprints

    cursor = FakeActiveTelegramTriggerCursor()
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Ежедневная сверка",
        "category": "custom",
        "status": "active",
        "created_by_user_id": "user1",
        "metadata_json": {
            "active_version_id": "ver1",
            "required_integration_bindings": [],
            "custom_process": {
                "kind": "source_destination_workflow",
                "trigger": "schedule.daily",
                "schedule": {"frequency": "daily", "time": "19:00"},
            },
        },
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "version_number": 1,
        "goal": "Проверить ежедневный запуск",
        "inputs_schema_json": {"trigger": "schedule.daily"},
        "steps_json": [],
        "capability_allowlist_json": [],
        "approval_policy_json": {},
        "output_schema_json": {"trigger": "schedule.daily"},
        "created_by_user_id": "user1",
    }

    first = dispatch_due_scheduled_agent_blueprints(
        cursor,
        now=datetime(2026, 6, 10, 19, 5, tzinfo=timezone.utc),
    )
    second = dispatch_due_scheduled_agent_blueprints(
        cursor,
        now=datetime(2026, 6, 10, 19, 10, tzinfo=timezone.utc),
    )

    assert first["dispatched_count"] == 1
    assert first["dispatched"][0]["matched_count"] == 1
    assert second["dispatched_count"] == 0
    assert second["skipped"][0]["reason"] == "already_recorded_today"
    assert len(cursor.tables["agent_runs"]) == 1


def test_due_scheduled_trigger_dispatcher_waits_until_schedule_time():
    from services.agent_trigger_runtime import dispatch_due_scheduled_agent_blueprints

    cursor = FakeActiveTelegramTriggerCursor()
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Ежедневная сверка",
        "category": "custom",
        "status": "active",
        "created_by_user_id": "user1",
        "metadata_json": {
            "active_version_id": "ver1",
            "required_integration_bindings": [],
            "custom_process": {
                "trigger": "schedule.daily",
                "schedule": {"frequency": "daily", "time": "19:00"},
            },
        },
    }

    result = dispatch_due_scheduled_agent_blueprints(
        cursor,
        now=datetime(2026, 6, 10, 18, 59, tzinfo=timezone.utc),
    )

    assert result["checked_businesses"] == 0
    assert result["dispatched_count"] == 0
    assert cursor.tables["agent_runs"] == {}


def test_activate_version_marks_blueprint_active_for_trigger_runtime(monkeypatch):
    from api import agent_blueprints_api

    class ActivationCursor:
        def __init__(self):
            self.metadata = {}
            self.status = "draft"

        def execute(self, query, params=None):
            normalized_query = " ".join(query.split()).lower()
            params = params or ()
            if normalized_query.startswith("update agent_blueprints") and "set metadata_json" in normalized_query:
                self.metadata = json.loads(params[0])
                return None
            if normalized_query.startswith("update agent_blueprints") and "set status = 'active'" in normalized_query:
                self.status = "active"
                return None
            raise AssertionError(f"Unhandled activation SQL: {query}")

    cursor = ActivationCursor()
    blueprint = {"id": "bp1", "metadata_json": {"version_events": []}, "status": "draft"}
    version = {"id": "ver1", "version_number": 1}
    monkeypatch.setattr(agent_blueprints_api, "_load_blueprint", lambda _cursor, _blueprint_id: blueprint)

    event = agent_blueprints_api._remember_active_version(
        cursor,
        blueprint,
        version,
        {"user_id": "user1"},
        "activated",
        "ready for telegram trigger",
    )

    assert event["active_version_id"] == "ver1"
    assert cursor.metadata["active_version_id"] == "ver1"
    assert cursor.status == "active"


def test_agent_integration_binding_status_tracks_required_compiled_bindings():
    from api import agent_blueprints_api

    metadata = {
        "required_integration_bindings": [
            {"key": "telegram_trigger", "provider": "telegram", "direction": "trigger", "trigger": "telegram.message.received"},
            {
                "key": "google_sheets_append",
                "provider": "google_sheets",
                "direction": "external_write",
                "capability": "sheets.append_row_request",
                "required_config": ["spreadsheet_id", "sheet_name"],
            },
        ],
        "agent_binding_integrations": {
            "google_sheets_append": {
                "answer_config": {"spreadsheet_id": "spreadsheet-1", "sheet_name": "Leads"},
            }
        },
    }
    rows = [
        {
            "id": "telegram-1",
            "provider": "telegram",
            "status": "active",
            "config_json": {"bot_mode": "business_bot"},
        },
        {
            "id": "sheets-1",
            "provider": "google_sheets",
            "status": "active",
            "config_json": {"spreadsheet_id": "spreadsheet-1", "sheet_name": "Leads"},
        },
    ]

    status = agent_blueprints_api._agent_integration_binding_status(metadata, rows)

    assert [item["status"] for item in status] == ["connected", "connected"]
    assert status[1]["integration_id"] == "sheets-1"
    assert status[1]["missing_config"] == []
    assert status[1]["answer_config"]["sheet_name"] == "Leads"


def test_agent_integration_binding_status_treats_localos_finance_as_native_ready():
    from api import agent_blueprints_api

    metadata = {
        "required_integration_bindings": [
            {
                "key": "localos_finance",
                "provider": "localos_finance",
                "direction": "localos_write_request",
                "capability": "finance.transaction.create",
                "required_config": ["transaction_type"],
            },
        ]
    }

    status = agent_blueprints_api._agent_integration_binding_status(metadata, [])

    assert status[0]["status"] == "connected"
    assert status[0]["integration_id"] == "native_localos"
    assert status[0]["resolution"] == "native_localos"
    assert status[0]["missing_config"] == []


def test_direct_agent_draft_auto_selects_single_connection_and_requires_ambiguous_choice():
    from api import agent_blueprints_api

    preview = {
        "connection_summary": {
            "items": [
                {
                    "key": "telegram_delivery",
                    "provider": "telegram",
                    "title": "Telegram",
                    "action": "choose_existing",
                    "connections": [{"id": "tg1", "display_name": "Бот бизнеса"}],
                },
                {
                    "key": "google_sheets_read",
                    "provider": "google_sheets",
                    "title": "Google Sheets",
                    "action": "choose_existing",
                    "connections": [
                        {"id": "sheet1", "display_name": "Orders"},
                        {"id": "sheet2", "display_name": "Archive"},
                    ],
                },
            ]
        }
    }
    inventory = [
        {"id": "tg1", "provider": "telegram", "display_name": "Бот бизнеса", "config": {"bot_mode": "business_bot"}},
        {"id": "sheet1", "provider": "google_sheets", "display_name": "Orders", "config": {"sheet_name": "Sheet1"}},
        {"id": "sheet2", "provider": "google_sheets", "display_name": "Archive", "config": {"sheet_name": "Archive"}},
    ]

    selected = agent_blueprints_api._direct_selected_connection_bindings({}, preview, inventory)
    missing = agent_blueprints_api._direct_missing_required_connection_choices(preview, selected)

    assert selected["telegram_delivery"]["integration_id"] == "tg1"
    assert selected["telegram_delivery"]["selection_source"] == "auto_single_connection"
    assert "google_sheets_read" not in selected
    assert missing == [
        {
            "key": "google_sheets_read",
            "provider": "google_sheets",
            "title": "Google Sheets",
            "connection_count": 2,
        }
    ]

    selected = agent_blueprints_api._direct_selected_connection_bindings(
        {"selected_connection_bindings": {"google_sheets_read": "sheet2"}},
        preview,
        inventory,
    )
    metadata = agent_blueprints_api._apply_direct_selected_connection_bindings({}, selected)

    assert selected["google_sheets_read"]["integration_id"] == "sheet2"
    assert agent_blueprints_api._direct_missing_required_connection_choices(preview, selected) == []
    assert metadata["agent_binding_integrations"]["google_sheets_read"]["provider"] == "google_sheets"
    assert metadata["custom_process"]["google_sheets"]["sheet_name"] == "Archive"


def test_sync_blueprint_integration_metadata_records_selected_binding(monkeypatch):
    from api import agent_blueprints_api
    from services.agent_integration_preflight import build_agent_integration_preflight

    class Cursor:
        def __init__(self):
            self.saved_metadata = {}

        def execute(self, query, params=None):
            normalized_query = " ".join(query.split()).lower()
            if normalized_query.startswith("update agent_blueprints"):
                self.saved_metadata = json.loads(params[0])
                return None
            if "from agent_integrations" in normalized_query:
                return None
            raise AssertionError(f"Unhandled SQL: {query}")

        def fetchall(self):
            return []

    blueprint = {
        "id": "bp1",
        "business_id": "biz1",
        "metadata_json": {
            "required_integration_bindings": [
                {
                    "key": "google_sheets_read",
                    "provider": "google_sheets",
                    "capability": "google_sheets.read_rows",
                    "required_config": ["spreadsheet_id", "sheet_name"],
                }
            ],
            "custom_process": {},
        },
    }
    integration = {
        "id": "sheets-1",
        "provider": "google_sheets",
        "status": "active",
        "config_json": {"spreadsheet_id": "spreadsheet-1", "sheet_name": "Orders", "operation": "read_rows"},
    }
    monkeypatch.setattr(agent_blueprints_api, "_load_blueprint", lambda _cursor, _blueprint_id: blueprint)
    cursor = Cursor()

    metadata = agent_blueprints_api._sync_blueprint_integration_metadata(
        cursor,
        blueprint,
        integration,
        "google_sheets_read",
    )
    preflight = build_agent_integration_preflight(cursor, business_id="biz1", metadata=metadata, input_payload={})

    assert metadata["agent_binding_integrations"]["google_sheets_read"]["integration_id"] == "sheets-1"
    assert metadata["custom_process"]["google_sheets_read"]["spreadsheet_id"] == "spreadsheet-1"
    assert preflight["ready"] is False
    assert preflight["items"][0]["resolution"] == "provider_route_required"
    assert preflight["items"][0]["integration_id"] == "sheets-1"


def test_agent_preflight_reports_selected_binding_missing_config():
    from services.agent_integration_preflight import build_agent_integration_preflight

    class Cursor:
        def execute(self, query, params=None):
            return None

        def fetchall(self):
            return []

    metadata = {
        "required_integration_bindings": [
            {
                "key": "google_sheets_read",
                "provider": "google_sheets",
                "capability": "google_sheets.read_rows",
                "required_config": ["spreadsheet_id", "sheet_name"],
            }
        ],
        "agent_binding_integrations": {
            "google_sheets_read": {
                "integration_id": "sheets-1",
                "provider": "google_sheets",
            }
        },
    }

    preflight = build_agent_integration_preflight(Cursor(), business_id="biz1", metadata=metadata, input_payload={})

    assert preflight["ready"] is False
    assert preflight["status"] == "blocked"
    assert preflight["missing_count"] == 1
    assert preflight["items"][0]["status"] == "needs_config"
    assert preflight["items"][0]["resolution"] == "blueprint_metadata_missing_config"
    assert preflight["items"][0]["integration_id"] == "sheets-1"
    assert preflight["items"][0]["missing_config"] == ["spreadsheet_id", "sheet_name"]
    assert "spreadsheet_id" in preflight["items"][0]["summary"]


def test_agent_preflight_does_not_treat_answer_config_as_external_connection():
    from services.agent_integration_preflight import build_agent_integration_preflight

    class Cursor:
        def execute(self, query, params=None):
            return None

        def fetchall(self):
            return []

    metadata = {
        "required_integration_bindings": [
            {
                "key": "google_sheets_read",
                "provider": "google_sheets",
                "capability": "google_sheets.read_rows",
                "required_config": ["spreadsheet_id", "sheet_name"],
            }
        ],
        "agent_binding_integrations": {
            "google_sheets_read": {
                "answer_config": {
                    "spreadsheet_id": "spreadsheet-from-dialog",
                    "sheet_name": "Orders",
                }
            }
        },
        "custom_process": {
            "google_sheets_read": {
                "spreadsheet_id": "spreadsheet-from-dialog",
                "sheet_name": "Orders",
            },
            "google_sheets": {
                "spreadsheet_id": "spreadsheet-from-dialog",
                "sheet_name": "Orders",
            },
        },
    }

    preflight = build_agent_integration_preflight(Cursor(), business_id="biz1", metadata=metadata, input_payload={})

    assert preflight["ready"] is False
    assert preflight["status"] == "blocked"
    assert preflight["items"][0]["status"] == "needs_connection"
    assert preflight["items"][0]["resolution"] == "builder_answer_needs_provider_route"
    assert preflight["items"][0]["missing_config"] == []
    assert preflight["items"][0]["provider"] == "google_sheets"


def test_agent_preflight_merges_answer_config_into_selected_connection():
    from services.agent_integration_preflight import build_agent_integration_preflight

    class Cursor:
        def execute(self, query, params=None):
            return None

        def fetchall(self):
            return []

    metadata = {
        "required_integration_bindings": [
            {
                "key": "google_sheets_read",
                "provider": "google_sheets",
                "capability": "google_sheets.read_rows",
                "required_config": ["spreadsheet_id", "sheet_name"],
            }
        ],
        "agent_binding_integrations": {
            "google_sheets_read": {
                "integration_id": "sheets-1",
                "provider": "google_sheets",
                "answer_config": {
                    "spreadsheet_id": "spreadsheet-from-dialog",
                    "sheet_name": "Orders",
                },
            }
        },
    }

    preflight = build_agent_integration_preflight(Cursor(), business_id="biz1", metadata=metadata, input_payload={})

    assert preflight["ready"] is False
    assert preflight["items"][0]["status"] == "needs_connection"
    assert preflight["items"][0]["resolution"] == "provider_route_required"
    assert preflight["items"][0]["integration_id"] == "sheets-1"


def test_agent_preflight_allows_google_sheets_native_read_when_auth_ref_bound():
    from services.agent_integration_preflight import build_agent_integration_preflight

    class Cursor:
        def execute(self, query, params=None):
            return None

        def fetchall(self):
            return [
                {
                    "id": "sheets-1",
                    "provider": "google_sheets",
                    "status": "active",
                    "auth_ref": "google-account-1",
                    "config_json": {"spreadsheet_id": "spreadsheet-1", "sheet_name": "Orders"},
                }
            ]

    metadata = {
        "required_integration_bindings": [
            {
                "key": "google_sheets_read",
                "provider": "google_sheets",
                "capability": "google_sheets.read_rows",
                "required_config": ["spreadsheet_id", "sheet_name"],
            }
        ],
        "agent_binding_integrations": {
            "google_sheets_read": {
                "integration_id": "sheets-1",
                "provider": "google_sheets",
            }
        },
        "custom_process": {
            "google_sheets_read": {
                "integration_id": "sheets-1",
                "spreadsheet_id": "spreadsheet-1",
                "sheet_name": "Orders",
            },
            "google_sheets": {
                "integration_id": "sheets-1",
                "spreadsheet_id": "spreadsheet-1",
                "sheet_name": "Orders",
            },
        },
    }

    preflight = build_agent_integration_preflight(Cursor(), business_id="biz1", metadata=metadata, input_payload={})

    assert preflight["ready"] is True
    assert preflight["items"][0]["status"] == "ready"
    assert preflight["items"][0]["resolution"] == "agent_integration_native_provider"
    assert preflight["items"][0]["integration_id"] == "sheets-1"


def test_activation_connection_blocker_keeps_binding_route_context():
    from api import agent_blueprints_api

    preflight = {
        "items": [
            {
                "key": "google_sheets_read",
                "provider": "google_sheets",
                "capability": "google_sheets.read_rows",
                "status": "needs_connection",
                "resolution": "missing_integration",
                "required": True,
                "missing_config": ["spreadsheet_id", "sheet_name"],
            }
        ],
    }

    connection_plan = agent_blueprints_api._activation_connection_plan_from_preflight(
        preflight,
        available_integrations=[
            {
                "id": "sheet-existing",
                "provider": "google_sheets",
                "status": "active",
                "display_name": "Orders sheet",
            }
        ],
        provider_catalog=agent_blueprints_api._agent_integration_provider_catalog(),
    )
    plan_item = connection_plan["items"][0]
    blocker = agent_blueprints_api._activation_connection_blocker(preflight["items"][0], plan_item)
    human_blockers = agent_blueprints_api._activation_gate_human_blockers([blocker], preflight, {})

    assert blocker["binding_key"] == "google_sheets_read"
    assert blocker["missing_config"] == ["spreadsheet_id", "sheet_name"]
    assert blocker["route_state"] == "available"
    assert blocker["preferred_route"]["provider"] == "openclaw"
    assert plan_item["action"] == "choose_existing"
    assert plan_item["setup_cta"]["action"] == "choose_existing"
    assert "Orders sheet" == plan_item["existing_integrations"][0]["display_name"]
    assert "сохранённый доступ" in plan_item["why_blocked"]
    assert any(item["provider"] == "openclaw" for item in blocker["provider_routes"])
    assert human_blockers[0]["binding_key"] == "google_sheets_read"
    assert human_blockers[0]["preferred_route"]["provider"] == "openclaw"
    assert "Google Sheets" in human_blockers[0]["message"]


def test_agent_preflight_does_not_treat_compiled_defaults_as_external_connection():
    from services.agent_integration_preflight import build_agent_integration_preflight

    class Cursor:
        def execute(self, query, params=None):
            return None

        def fetchall(self):
            return []

    metadata = {
        "required_integration_bindings": [
            {
                "key": "telegram_delivery",
                "provider": "telegram",
                "capability": "communications.draft",
                "direction": "external_publish_request",
                "required_config": ["bot_mode"],
                "default_config": {"bot_mode": "business_bot"},
            }
        ]
    }

    preflight = build_agent_integration_preflight(Cursor(), business_id="biz1", metadata=metadata, input_payload={})

    assert preflight["ready"] is False
    assert preflight["status"] == "blocked"
    assert preflight["items"][0]["status"] == "needs_connection"
    assert preflight["items"][0]["resolution"] == "missing_integration"
    assert preflight["items"][0]["provider"] == "telegram"


def test_agent_preflight_accepts_openclaw_provider_route_for_binding(monkeypatch):
    from api import agent_blueprints_api
    from services.agent_integration_preflight import build_agent_integration_preflight

    class Cursor:
        def __init__(self):
            self.saved_metadata = {}

        def execute(self, query, params=None):
            normalized_query = " ".join(query.split()).lower()
            if normalized_query.startswith("update agent_blueprints"):
                self.saved_metadata = json.loads(params[0])
                return None
            if "from agent_integrations" in normalized_query:
                return None
            raise AssertionError(f"Unhandled SQL: {query}")

        def fetchall(self):
            return []

    blueprint = {
        "id": "bp1",
        "business_id": "biz1",
        "metadata_json": {
            "required_integration_bindings": [
                {
                    "key": "telegram_delivery",
                    "provider": "telegram",
                    "capability": "communications.send_offer",
                    "required_config": ["bot_mode"],
                }
            ],
        },
    }
    monkeypatch.setattr(agent_blueprints_api, "_load_blueprint", lambda _cursor, _blueprint_id: blueprint)
    cursor = Cursor()

    metadata = agent_blueprints_api._apply_agent_provider_route_metadata(
        cursor,
        blueprint,
        binding_key="telegram_delivery",
        route_provider="openclaw",
    )
    preflight = build_agent_integration_preflight(cursor, business_id="biz1", metadata=metadata, input_payload={})
    status = agent_blueprints_api._agent_integration_binding_status(metadata, [])

    assert preflight["ready"] is True
    assert preflight["items"][0]["resolution"] == "provider_route_openclaw_boundary"
    assert preflight["items"][0]["integration_id"] == "openclaw_boundary"
    assert preflight["items"][0]["execution_boundary"] == "openclaw_inside_localos_policy"
    assert preflight["items"][0]["autonomy_level"] == "supervised"
    assert preflight["items"][0]["approval_state"] == "approval_required"
    assert "OpenClaw" in preflight["items"][0]["policy_summary"]
    assert status[0]["status"] == "connected"
    assert status[0]["resolution"] == "provider_route_openclaw"


def test_agent_preflight_accepts_maton_provider_route_with_external_account(monkeypatch):
    from api import agent_blueprints_api
    from services.agent_integration_preflight import build_agent_integration_preflight

    class Cursor:
        def __init__(self):
            self.saved_metadata = {}

        def execute(self, query, params=None):
            normalized_query = " ".join(query.split()).lower()
            if normalized_query.startswith("update agent_blueprints"):
                self.saved_metadata = json.loads(params[0])
                return None
            if "from agent_integrations" in normalized_query:
                return None
            raise AssertionError(f"Unhandled SQL: {query}")

        def fetchall(self):
            return []

    blueprint = {
        "id": "bp1",
        "business_id": "biz1",
        "metadata_json": {
            "required_integration_bindings": [
                {
                    "key": "telegram_delivery",
                    "provider": "telegram",
                    "capability": "communications.send_offer",
                    "required_config": ["bot_mode"],
                }
            ],
        },
    }
    monkeypatch.setattr(agent_blueprints_api, "_load_blueprint", lambda _cursor, _blueprint_id: blueprint)
    cursor = Cursor()

    metadata = agent_blueprints_api._apply_agent_provider_route_metadata(
        cursor,
        blueprint,
        binding_key="telegram_delivery",
        route_provider="maton",
        external_account={"id": "maton-account-1", "display_name": "Main Maton key"},
    )
    preflight = build_agent_integration_preflight(cursor, business_id="biz1", metadata=metadata, input_payload={})
    status = agent_blueprints_api._agent_integration_binding_status(metadata, [])

    assert preflight["ready"] is True
    assert preflight["items"][0]["resolution"] == "provider_route_maton_external_account"
    assert preflight["items"][0]["integration_id"] == "maton-account-1"
    assert preflight["items"][0]["execution_boundary"] == "maton_bridge_inside_localos_policy"
    assert preflight["items"][0]["credential_state"] == "external_account_bound"
    assert "Maton" in preflight["items"][0]["policy_summary"]
    assert status[0]["status"] == "connected"
    assert status[0]["route_provider"] == "maton"


def test_agent_preflight_accepts_manual_provider_route_as_human_fallback(monkeypatch):
    from api import agent_blueprints_api
    from services.agent_integration_preflight import build_agent_integration_preflight

    class Cursor:
        def __init__(self):
            self.saved_metadata = {}

        def execute(self, query, params=None):
            normalized_query = " ".join(query.split()).lower()
            if normalized_query.startswith("update agent_blueprints"):
                self.saved_metadata = json.loads(params[0])
                return None
            if "from agent_integrations" in normalized_query:
                return None
            raise AssertionError(f"Unhandled SQL: {query}")

        def fetchall(self):
            return []

    blueprint = {
        "id": "bp1",
        "business_id": "biz1",
        "metadata_json": {
            "required_integration_bindings": [
                {
                    "key": "google_sheets_read",
                    "provider": "google_sheets",
                    "capability": "google_sheets.read_rows",
                    "required_config": ["spreadsheet_id", "sheet_name"],
                }
            ],
        },
    }
    monkeypatch.setattr(agent_blueprints_api, "_load_blueprint", lambda _cursor, _blueprint_id: blueprint)
    cursor = Cursor()

    metadata = agent_blueprints_api._apply_agent_provider_route_metadata(
        cursor,
        blueprint,
        binding_key="google_sheets_read",
        route_provider="manual",
    )
    preflight = build_agent_integration_preflight(cursor, business_id="biz1", metadata=metadata, input_payload={})
    status = agent_blueprints_api._agent_integration_binding_status(metadata, [])

    assert preflight["ready"] is True
    assert preflight["items"][0]["resolution"] == "provider_route_manual_fallback"
    assert preflight["items"][0]["execution_boundary"] == "human_operated_fallback"
    assert preflight["items"][0]["autonomy_level"] == "draft_only"
    assert preflight["items"][0]["approval_state"] == "human_action_required"
    assert "человек" in preflight["items"][0]["policy_summary"]
    assert status[0]["status"] == "connected"
    assert status[0]["route_provider"] == "manual"
    assert metadata["agent_binding_provider_routes"]["google_sheets_read"]["draft_only_until_human_action"] is True


def test_agent_preflight_reports_active_integration_missing_config():
    from services.agent_integration_preflight import build_agent_integration_preflight

    class Cursor:
        def __init__(self):
            self.results = []

        def execute(self, query, params=None):
            normalized_query = " ".join(query.split()).lower()
            if "from agent_integrations" in normalized_query:
                self.results = [
                    {
                        "id": "sheets-1",
                        "business_id": "biz1",
                        "provider": "google_sheets",
                        "status": "active",
                        "display_name": "Orders sheet",
                        "config_json": {"spreadsheet_id": "spreadsheet-1"},
                    }
                ]
                return None
            raise AssertionError(f"Unhandled SQL: {query}")

        def fetchall(self):
            return self.results

    metadata = {
        "required_integration_bindings": [
            {
                "key": "google_sheets_read",
                "provider": "google_sheets",
                "capability": "google_sheets.read_rows",
                "required_config": ["spreadsheet_id", "sheet_name"],
            }
        ],
    }

    preflight = build_agent_integration_preflight(Cursor(), business_id="biz1", metadata=metadata, input_payload={})

    assert preflight["ready"] is False
    assert preflight["items"][0]["status"] == "needs_config"
    assert preflight["items"][0]["resolution"] == "agent_integration_missing_config"
    assert preflight["items"][0]["integration_id"] == "sheets-1"
    assert preflight["items"][0]["missing_config"] == ["sheet_name"]


def test_agent_connection_plan_turns_bindings_into_user_next_actions():
    from api import agent_blueprints_api

    binding_status = [
        {
            "key": "google_sheets_read",
            "provider": "google_sheets",
            "capability": "google_sheets.read_rows",
            "status": "missing",
            "missing_config": ["spreadsheet_id", "sheet_name"],
        },
        {
            "key": "telegram_delivery",
            "provider": "telegram",
            "capability": "communications.draft",
            "status": "connected",
            "integration_id": "telegram-1",
            "resolution": "agent_integration",
        },
        {
            "key": "localos_finance",
            "provider": "localos_finance",
            "capability": "finance.transaction.create",
            "status": "connected",
            "integration_id": "native_localos",
            "resolution": "native_localos",
        },
    ]
    available = [
        {
            "id": "sheet-existing",
            "provider": "google_sheets",
            "status": "active",
            "display_name": "Orders sheet",
        }
    ]

    plan = agent_blueprints_api._agent_connection_plan(
        binding_status,
        [],
        available,
        agent_blueprints_api._agent_integration_provider_catalog(),
    )

    assert plan["schema"] == "localos_agent_connection_plan_v1"
    assert plan["status"] == "needs_action"
    assert plan["missing_count"] == 1
    assert plan["items"][0]["action"] == "choose_existing"
    assert plan["items"][0]["route_state"] == "available"
    assert "выберите" in plan["items"][0]["route_summary"].lower()
    assert plan["items"][0]["setup_cta"]["action"] == "choose_existing"
    assert "сохранённый доступ" in plan["items"][0]["why_blocked"]
    assert plan["items"][0]["existing_integrations"][0]["display_name"] == "Orders sheet"
    assert any(item["provider"] == "openclaw" for item in plan["items"][0]["provider_routes"])
    assert any(item["provider"] == "native_localos" for item in plan["items"][0]["provider_paths"])
    assert plan["items"][1]["action"] == "ready"
    assert plan["items"][1]["route_state"] == "connected"
    assert plan["items"][2]["action"] == "native_ready"


def test_agent_post_connect_handoff_routes_ready_connections_to_preview():
    from api import agent_blueprints_api

    ready = agent_blueprints_api._build_agent_post_connect_handoff(
        {
            "schema": "localos_agent_connection_plan_v1",
            "status": "ready",
            "missing_count": 0,
            "items": [],
        }
    )
    blocked = agent_blueprints_api._build_agent_post_connect_handoff(
        {
            "schema": "localos_agent_connection_plan_v1",
            "status": "needs_action",
            "missing_count": 2,
            "items": [
                {
                    "key": "google_sheets_read",
                    "provider": "google_sheets",
                    "action": "ready",
                },
                {
                    "key": "telegram_delivery",
                    "provider": "telegram",
                    "action": "connect_required",
                },
            ],
        }
    )

    assert ready["status"] == "ready_for_preview"
    assert ready["workspace_mode"] == "run"
    assert ready["next_step"] == "run_preview"
    assert "безопасный тест" in ready["description"]
    assert "workflow" not in ready["description"]
    assert "preflight" not in ready["description"]
    assert "approval gate" not in ready["description"]
    assert blocked["status"] == "needs_connections"
    assert blocked["workspace_mode"] == "connections"
    assert blocked["next_step"] == "connect_required_integrations"
    assert blocked["next_binding_key"] == "telegram_delivery"


def test_agent_builder_post_create_handoff_contains_connection_plan():
    from api import agent_builder_api

    handoff = agent_builder_api._build_post_create_handoff(
        {
            "ready": False,
            "items": [
                {
                    "key": "google_sheets_read",
                    "provider": "google_sheets",
                    "capability": "google_sheets.read_rows",
                    "status": "needs_connection",
                    "required": True,
                    "missing_config": ["spreadsheet_id", "sheet_name"],
                },
                {
                    "key": "telegram_delivery",
                    "provider": "telegram",
                    "capability": "communications.draft",
                    "status": "ready",
                    "resolution": "agent_integration",
                    "required": True,
                },
            ],
            "missing": [
                {
                    "key": "google_sheets_read",
                    "provider": "google_sheets",
                    "capability": "google_sheets.read_rows",
                    "status": "needs_connection",
                    "required": True,
                    "missing_config": ["spreadsheet_id", "sheet_name"],
                }
            ],
        }
    )

    assert handoff["schema"] == "localos_agent_post_create_handoff_v1"
    assert handoff["status"] == "needs_connections"
    assert handoff["next_binding_key"] == "google_sheets_read"
    assert handoff["next_binding"]["key"] == "google_sheets_read"
    assert handoff["next_binding"]["title"] == "Google Sheets"
    assert handoff["next_binding"]["route_state"] == "available"
    assert handoff["next_binding"]["route_summary"]
    assert handoff["next_binding"]["recommended_route"]["provider"] == "openclaw"
    assert "OpenClaw" in handoff["next_binding"]["recommended_route_reason"]
    assert handoff["next_route"]["provider"] == "openclaw"
    assert handoff["next_route"]["primary_cta"]
    assert handoff["connection_plan"]["schema"] == "localos_agent_connection_plan_v1"
    assert handoff["connection_plan"]["missing_count"] == 1
    assert handoff["connection_plan"]["items"][0]["action"] == "connect_required"
    assert handoff["connection_plan"]["items"][0]["provider_routes"]
    assert handoff["connection_plan"]["items"][1]["action"] == "ready"


def test_agent_builder_post_create_handoff_routes_ready_agent_to_preview_run():
    from api import agent_builder_api

    handoff = agent_builder_api._build_post_create_handoff(
        {
            "ready": True,
            "items": [
                {
                    "key": "google_sheets_read",
                    "provider": "google_sheets",
                    "capability": "google_sheets.read_rows",
                    "status": "ready",
                    "resolution": "agent_integration",
                    "required": True,
                },
                {
                    "key": "telegram_delivery",
                    "provider": "telegram",
                    "capability": "communications.draft",
                    "status": "ready",
                    "resolution": "agent_integration",
                    "required": True,
                },
            ],
            "missing": [],
        }
    )

    assert handoff["status"] == "ready_for_preview"
    assert handoff["next_step"] == "run_preview"
    assert handoff["workspace_mode"] == "run"
    assert handoff["next_binding"] == {}
    assert handoff["next_route"] == {}
    assert "preview run" in handoff["description"]


def test_activation_gate_summary_explains_missing_connector(monkeypatch):
    from api import agent_blueprints_api

    class Cursor:
        def execute(self, *args, **kwargs):
            return None

        def fetchall(self):
            return []

    monkeypatch.setattr(
        agent_blueprints_api,
        "validate_compiled_artifact_candidate",
        lambda payload, metadata: {"ready": True, "validation": {"status": "passed", "errors": []}},
    )
    monkeypatch.setattr(
        agent_blueprints_api,
        "build_version_payload_from_row",
        lambda row: {"steps": []},
    )

    metadata = {
        "required_integration_bindings": [
            {
                "key": "google_sheets_read",
                "provider": "google_sheets",
                "capability": "google_sheets.read_rows",
                "required_config": ["spreadsheet_id"],
            }
        ]
    }
    gate = agent_blueprints_api._build_activation_gate_summary(
        Cursor(),
        {"business_id": "biz1"},
        {"id": "version-1", "version_number": 1},
        metadata,
    )

    assert gate["schema"] == "localos_agent_activation_gate_v1"
    assert gate["can_activate"] is False
    assert gate["next_step"] == "connect_required_integrations"
    assert gate["primary_action_label"] == "Открыть подключения"
    assert gate["human_blockers"][0]["provider"] == "google_sheets"
    assert "Google Sheets" in gate["summary"]
    assert "spreadsheet_id" in gate["summary"]
    assert gate["connection_plan"]["schema"] == "localos_agent_connection_plan_v1"
    assert gate["connection_plan"]["items"][0]["action"] == "connect_required"
    assert gate["connection_plan"]["items"][0]["recommended_route"]["provider"] == "openclaw"
    assert gate["next_binding_key"] == "google_sheets_read"


def test_activation_gate_requires_safe_preview_run(monkeypatch):
    from api import agent_blueprints_api

    class Cursor:
        def __init__(self):
            self.rows = []

        def execute(self, query, params=None):
            normalized_query = " ".join(query.split()).lower()
            self.rows = []
            if "from agent_runs" in normalized_query:
                self.rows = []

        def fetchall(self):
            return self.rows

    monkeypatch.setattr(
        agent_blueprints_api,
        "validate_compiled_artifact_candidate",
        lambda payload, metadata: {"ready": True, "validation": {"status": "passed", "errors": []}},
    )
    monkeypatch.setattr(
        agent_blueprints_api,
        "build_version_payload_from_row",
        lambda row: {"steps": []},
    )

    gate = agent_blueprints_api._build_activation_gate_summary(
        Cursor(),
        {"id": "bp1", "business_id": "biz1"},
        {"id": "version-1", "version_number": 1},
        {"required_integration_bindings": []},
    )

    assert gate["can_activate"] is False
    assert gate["requires_preview_run"] is True
    assert gate["preview_run_status"]["ready"] is False
    assert gate["next_step"] == "run_preview"
    assert gate["primary_action_label"] == "Запустить preview"
    assert gate["human_blockers"][0]["type"] == "preview_run"


def test_activation_gate_accepts_completed_safe_preview_run(monkeypatch):
    from api import agent_blueprints_api

    class Cursor:
        def __init__(self):
            self.rows = []

        def execute(self, query, params=None):
            normalized_query = " ".join(query.split()).lower()
            self.rows = []
            if "from agent_runs" in normalized_query:
                self.rows = [
                    {
                        "id": "run-1",
                        "status": "completed",
                        "input_json": {
                            "schema": "localos_agent_preview_input_v1",
                            "preview_mode": True,
                            "external_side_effects_allowed": False,
                            "source": "agent_preview",
                        },
                        "output_json": {},
                        "error_text": "",
                        "started_at": datetime(2026, 6, 11, tzinfo=timezone.utc),
                        "completed_at": datetime(2026, 6, 11, tzinfo=timezone.utc),
                        "updated_at": datetime(2026, 6, 11, tzinfo=timezone.utc),
                    }
                ]

        def fetchall(self):
            return self.rows

    monkeypatch.setattr(
        agent_blueprints_api,
        "validate_compiled_artifact_candidate",
        lambda payload, metadata: {"ready": True, "validation": {"status": "passed", "errors": []}},
    )
    monkeypatch.setattr(
        agent_blueprints_api,
        "build_version_payload_from_row",
        lambda row: {"steps": []},
    )

    gate = agent_blueprints_api._build_activation_gate_summary(
        Cursor(),
        {"id": "bp1", "business_id": "biz1"},
        {"id": "version-1", "version_number": 1},
        {"required_integration_bindings": []},
    )

    assert gate["can_activate"] is True
    assert gate["preview_run_status"]["ready"] is True
    assert gate["preview_run_status"]["passed_run"]["id"] == "run-1"
    assert gate["next_step"] == "activate_version"


def test_activation_gate_blocks_autonomous_write_limits_even_after_preview(monkeypatch):
    from api import agent_blueprints_api

    class Cursor:
        def __init__(self):
            self.rows = []

        def execute(self, query, params=None):
            normalized_query = " ".join(query.split()).lower()
            self.rows = []
            if "from agent_runs" in normalized_query:
                self.rows = [
                    {
                        "id": "run-1",
                        "status": "completed",
                        "input_json": {
                            "schema": "localos_agent_preview_input_v1",
                            "preview_mode": True,
                            "external_side_effects_allowed": False,
                            "source": "agent_preview",
                        },
                        "output_json": {},
                        "error_text": "",
                        "started_at": datetime(2026, 6, 11, tzinfo=timezone.utc),
                        "completed_at": datetime(2026, 6, 11, tzinfo=timezone.utc),
                        "updated_at": datetime(2026, 6, 11, tzinfo=timezone.utc),
                    }
                ]

        def fetchall(self):
            return self.rows

    monkeypatch.setattr(
        agent_blueprints_api,
        "validate_compiled_artifact_candidate",
        lambda payload, metadata: {
            "ready": True,
            "validation": {"status": "passed", "errors": []},
            "candidate": {
                "dsl": {
                    "limits": {
                        "autonomous_external_write_allowed": True,
                    }
                }
            },
        },
    )

    gate = agent_blueprints_api._build_activation_gate_summary(
        Cursor(),
        {"id": "bp1", "business_id": "biz1"},
        {
            "id": "version-1",
            "version_number": 1,
            "steps_json": [
                {
                    "key": "write_sheet",
                    "type": "capability",
                    "capability": "sheets.append_row_request",
                    "requires_approval": True,
                    "required_approval_type": "sheet_update",
                },
                {
                    "key": "approve_sheet_update",
                    "type": "approval",
                    "approval_type": "sheet_update",
                },
            ],
            "approval_policy_json": {"sheet_update": "manual_approval_required"},
            "capability_allowlist_json": ["sheets.append_row_request"],
        },
        {"required_integration_bindings": []},
    )

    assert gate["can_activate"] is False
    assert gate["next_step"] == "fix_compiled_workflow"
    assert gate["approval_policy_status"]["ready"] is False
    assert gate["approval_policy_status"]["autonomous_writes_allowed"] is True
    assert gate["human_blockers"][0]["type"] == "approval_policy"


def test_google_sheets_integration_config_preserves_read_write_operation():
    from api import agent_blueprints_api

    read_config = agent_blueprints_api._sanitize_agent_integration_config(
        "google_sheets",
        {
            "config": {
                "spreadsheet_id": "spreadsheet-1",
                "sheet_name": "Payments",
                "operation": "read_rows",
            }
        },
    )
    invalid_config = agent_blueprints_api._sanitize_agent_integration_config(
        "google_sheets",
        {
            "config": {
                "spreadsheet_id": "spreadsheet-1",
                "sheet_name": "Payments",
                "operation": "delete_sheet",
            }
        },
    )

    assert read_config["operation"] == "read_rows"
    assert invalid_config["operation"] == "read_write"


def test_google_oauth_requests_sheets_scope_for_agent_runtime():
    from pathlib import Path

    from services.agent_google_sheets_adapter import SHEETS_SCOPE

    source = Path("src/google_business_auth.py").read_text(encoding="utf-8")
    assert SHEETS_SCOPE in source
    assert "autogenerate_code_verifier=False" in source


def test_google_sheets_integration_auto_binds_google_business_auth_ref():
    from api import agent_blueprints_api

    class Cursor:
        def __init__(self):
            self.last_result = None

        def execute(self, query, params=None):
            normalized_query = " ".join(query.split()).lower()
            params = params or ()
            if "from externalbusinessaccounts" in normalized_query:
                assert params[0] == "biz1"
                self.last_result = {"id": "google-account-1"}
                return None
            raise AssertionError(f"Unhandled SQL: {query}")

        def fetchone(self):
            return self.last_result

    assert (
        agent_blueprints_api._resolve_agent_integration_auth_ref(
            Cursor(),
            "biz1",
            "google_sheets",
            "",
        )
        == "google-account-1"
    )


def test_google_oauth_callback_syncs_google_sheets_runtime_auth_ref():
    from pathlib import Path

    source = Path("src/api/google_business_api.py").read_text(encoding="utf-8")
    assert "def _sync_google_sheets_agent_auth_refs" in source
    assert "UPDATE agent_integrations" in source
    assert "provider = 'google_sheets'" in source
    assert "COALESCE(auth_ref, '') = ''" in source
    assert "_sync_google_sheets_agent_auth_refs(cursor, business_id, account_id)" in source


def test_maton_integration_config_is_delivery_bridge_with_caps():
    from api import agent_blueprints_api

    config = agent_blueprints_api._sanitize_agent_integration_config(
        "maton",
        {"config": {"channel": "whatsapp_business"}},
    )
    limits = agent_blueprints_api._sanitize_agent_integration_limits(
        "maton",
        {"limits": {"daily_message_cap": 25, "frequency_cap_minutes": 15}},
    )

    assert config["channel"] == "whatsapp_business"
    assert config["mode"] == "approved_delivery_bridge"
    assert limits["daily_message_cap"] == 25
    assert limits["frequency_cap_minutes"] == 15


def test_whatsapp_integration_config_is_supported_channel_boundary():
    from api import agent_blueprints_api
    from services.agent_provider_registry import integration_provider_catalog

    config = agent_blueprints_api._sanitize_agent_integration_config(
        "whatsapp",
        {"config": {"channel_mode": "manual_whatsapp"}},
    )
    limits = agent_blueprints_api._sanitize_agent_integration_limits(
        "whatsapp",
        {"limits": {"daily_message_cap": 35, "frequency_cap_minutes": 20}},
    )
    catalog = {item["provider"]: item for item in integration_provider_catalog()}
    boundary = agent_blueprints_api._agent_integration_execution_boundary("whatsapp")

    assert config["channel_mode"] == "manual_whatsapp"
    assert config["trigger"] == "whatsapp.message.received"
    assert config["mode"] == "trigger_boundary"
    assert limits["daily_message_cap"] == 35
    assert limits["frequency_cap_minutes"] == 20
    assert catalog["whatsapp"]["status"] == "available"
    assert boundary["executor"] == "channel_router"
    assert boundary["external_write"] == "approved_delivery_only"


def test_browser_use_integration_config_and_generic_binding_use_openclaw_boundary():
    from api import agent_blueprints_api
    from services.agent_blueprint_draft_builder import compile_agent_blueprint
    from services.agent_provider_registry import integration_provider_catalog

    config = agent_blueprints_api._sanitize_agent_integration_config(
        "browser_use",
        {"config": {"target_urls": "https://example.com\nhttps://competitor.example"}},
    )
    limits = agent_blueprints_api._sanitize_agent_integration_limits(
        "browser_use",
        {"limits": {"daily_page_check_cap": 12, "frequency_cap_minutes": 45}},
    )
    catalog = {item["provider"]: item for item in integration_provider_catalog()}
    boundary = agent_blueprints_api._agent_integration_execution_boundary("browser_use")
    draft = compile_agent_blueprint(
        "Через browser use открой сайт конкурента https://example.com, проверь изменения цен и подготовь сообщение владельцу в Telegram.",
        use_ai=True,
    )
    metadata = draft["metadata"]
    required = metadata["required_integration_bindings"]
    browser_binding = next(item for item in required if item["provider"] == "browser_use")

    assert config["target_urls"] == ["https://example.com", "https://competitor.example"]
    assert config["mode"] == "openclaw_browser_boundary"
    assert limits["daily_page_check_cap"] == 12
    assert limits["frequency_cap_minutes"] == 45
    assert catalog["browser_use"]["status"] == "available"
    assert boundary["executor"] == "openclaw_browser_boundary"
    assert boundary["capabilities"] == ["browser_use.read_page"]
    assert browser_binding["key"] == "browser_use_read"
    assert browser_binding["capability"] == "browser_use.read_page"
    assert browser_binding["required_config"] == ["target_urls"]
    assert browser_binding["execution_boundary"] == "openclaw_browser_boundary"


def test_custom_process_mapping_updates_compiled_version_steps():
    from api import agent_blueprints_api
    from services.agent_blueprint_draft_builder import compile_agent_blueprint

    draft = compile_agent_blueprint("Когда пользователь пишет в Telegram бота, добавь строку в Google таблицу")
    version_payload = agent_blueprints_api._apply_custom_process_to_version_payload(
        draft["version_payload"],
        {
            "row_values": ["{{received_at}}", "{{message_text}}", "{{telegram_user_id}}"],
            "columns": ["received_at", "message_text", "telegram_user_id"],
            "daily_append_cap": 12,
            "google_sheets": {"sheet_name": "Requests"},
        },
    )

    request_step = [step for step in version_payload["steps"] if step.get("capability") == "sheets.append_row_request"][0]
    draft_step = [step for step in version_payload["steps"] if step["key"] == "prepare_sheet_row"][0]
    assert request_step["payload"]["sheet_name"] == "Requests"
    assert request_step["payload"]["daily_append_cap"] == 12
    assert request_step["payload"]["row_values"] == ["{{received_at}}", "{{message_text}}", "{{telegram_user_id}}"]
    assert draft_step["payload"]["columns"] == ["received_at", "message_text", "telegram_user_id"]
    assert version_payload["output_schema"]["sheet_name"] == "Requests"
    assert version_payload["limits"]["daily_append_cap"] == 12


def test_agent_blueprint_api_guards_version_blueprint_mismatch():
    api_source = Path("src/api/agent_blueprints_api.py").read_text(encoding="utf-8")
    workspace_source = Path("src/services/agent_blueprint_workspace.py").read_text(encoding="utf-8")
    legacy_migration_source = Path("src/services/agent_legacy_migration.py").read_text(encoding="utf-8")
    webhooks_source = Path("src/ai_agent_webhooks.py").read_text(encoding="utf-8")
    ai_agent_source = Path("src/ai_agent.py").read_text(encoding="utf-8")
    chats_api_source = Path("src/chats_api.py").read_text(encoding="utf-8")
    document_llm_source = Path("src/services/agent_document_llm.py").read_text(encoding="utf-8")
    email_llm_source = Path("src/services/agent_email_llm.py").read_text(encoding="utf-8")
    review_analysis_source = Path("src/services/agent_review_reply_analysis.py").read_text(encoding="utf-8")
    table_analysis_source = Path("src/services/agent_table_analysis.py").read_text(encoding="utf-8")
    capability_handlers_source = Path("src/services/agent_capability_handlers.py").read_text(encoding="utf-8")
    action_policy_source = Path("src/core/action_policy.py").read_text(encoding="utf-8")
    trigger_runtime_source = Path("src/services/agent_trigger_runtime.py").read_text(encoding="utf-8")
    worker_source = Path("src/worker.py").read_text(encoding="utf-8")
    compose_source = Path("docker-compose.yml").read_text(encoding="utf-8")
    telegram_webhook_source = Path("src/ai_agent_webhooks.py").read_text(encoding="utf-8")
    builder_api_source = Path("src/api/agent_builder_api.py").read_text(encoding="utf-8")
    agents_page_source = Path("frontend/src/pages/dashboard/AgentBlueprintsPage.tsx").read_text(encoding="utf-8")
    admin_page_source = Path("frontend/src/pages/dashboard/AdminPage.tsx").read_text(encoding="utf-8")

    assert "VERSION_BLUEPRINT_MISMATCH" in api_source
    assert "_load_blueprint_version_for_blueprint" in api_source
    assert "build_agent_blueprint_orchestrator" in api_source
    assert "run_status" in api_source
    assert "approval_queue" in api_source
    assert "last_run_status" in api_source
    assert "pending_approvals_count" in api_source
    assert "sources_count" in api_source
    assert "versions_count" in api_source
    assert "/api/agent-blueprints/draft" in api_source
    assert "build_agent_builder_state" in api_source
    assert "direct_draft_envelope" in api_source
    assert "_direct_selected_connection_bindings" in api_source
    assert "_direct_missing_required_connection_choices" in api_source
    assert "_apply_direct_selected_connection_bindings" in api_source
    assert "_selected_provider_routes" in api_source
    assert "_missing_required_provider_routes" in api_source
    assert "_required_provider_route_bindings" in api_source
    assert "_apply_selected_provider_routes" in api_source
    assert "_apply_answer_connection_bindings" in builder_api_source
    assert "builder_answer_connection_bindings" in builder_api_source
    assert "connection_answer_bindings" in Path("src/services/agent_builder_session.py").read_text(encoding="utf-8")
    assert "answer_config" in api_source
    assert "connectionResourceFacts" in agents_page_source
    assert "Ресурс из диалога" in agents_page_source
    assert "Поняли ресурс" in agents_page_source
    assert "BuilderExecutionBoundaryPanel" in agents_page_source
    assert "Execution boundary" in agents_page_source
    assert "OpenClaw action refs" in agents_page_source
    assert "openclaw_action_plan" in api_source
    assert "_preview_openclaw_action_plan" in api_source
    runner_source = Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "_preview_summary_openclaw_action_plan" in runner_source
    assert "_preview_summary_policy_envelope" in runner_source
    assert "\"approval_gate\": approval_gate" in runner_source
    assert "OpenClaw actions в safe preview" in agents_page_source
    assert "OpenClawPreviewActionPlanPanel" in agents_page_source
    assert "side effects выключены" in agents_page_source
    assert "BuilderRequiredConnectionsPanel" in agents_page_source
    assert "Доступы перед созданием агента" in agents_page_source
    assert "Ресурс из диалога" in agents_page_source
    assert "способ выбран" in agents_page_source
    assert "builderConnectionCardStatus" in agents_page_source
    assert "AGENT_CONNECTION_CHOICE_REQUIRED" in api_source
    assert "AGENT_PROVIDER_ROUTE_REQUIRED" in api_source
    assert "AGENT_PROVIDER_ROUTES_CONFIRMATION_REQUIRED" in api_source
    assert "\"post_create_handoff\": post_create_handoff" in api_source
    assert "metadata[\"agent_builder_preview\"] = preview" in api_source
    assert "metadata[\"openclaw_planner_context\"] = planner_context" in api_source
    assert "\"connection_summary\": preview.get(\"connection_summary\")" in api_source
    assert "_load_direct_builder_connection_inventory" in api_source
    assert "/api/agent-blueprints/legacy-migration/apply" in api_source
    assert "apply_legacy_ai_agent_migration" in api_source
    assert "learning_events" in api_source
    assert "version_events" in api_source
    assert "legacy_migration" in api_source
    assert "build_agent_blueprint_draft" in api_source
    assert "_insert_version(cursor, blueprint_id, version_payload, user_data)" in api_source
    assert "/api/agent-blueprints/<blueprint_id>/setup" in api_source
    assert "/api/agent-blueprints/<blueprint_id>/sources" in api_source
    assert "/api/agent-blueprints/<blueprint_id>/sources/catalog" in api_source
    assert "/api/agent-blueprints/<blueprint_id>/sources/upload" in api_source
    assert "build_agent_datahub_catalog" in api_source
    assert "build_agent_source_from_upload" in api_source
    assert "/api/agent-blueprints/<blueprint_id>/preflight" in api_source
    assert "preflight_agent_blueprint_run" in api_source
    assert "localos_agent_preview_run_gate_v1" in api_source
    assert '"external_side_effects_allowed": False' in api_source
    assert '"next_binding_key": _connection_plan_next_binding_key(connection_plan)' in api_source
    assert "_build_agent_preview_run_input" in api_source
    assert "localos_agent_preview_input_v1" in api_source
    assert "AGENT_INTEGRATIONS_REQUIRED" in api_source
    assert 'return jsonify(result), 400' in api_source
    assert "/api/agent-runs/<run_id>/feedback" in api_source
    assert "trigger_type" in api_source
    assert "auto_activate" in api_source
    assert "auto_activation_gate = _build_activation_gate_summary" in api_source
    assert "if auto_activation_gate.get(\"can_activate\")" in api_source
    assert "auto_activation_blocked" in api_source
    assert "auto_activation_applied" in api_source
    assert "SET status = 'active'" in api_source
    assert "build_learning_loop_summary" in api_source
    assert "/api/agent-runs/<run_id>/support-export" in api_source
    assert "build_run_support_export" in api_source
    assert "/api/agent-blueprints/<blueprint_id>/versions/<version_id>/diff" in api_source
    assert "/api/agent-blueprints/<blueprint_id>/versions/<version_id>/activate" in api_source
    assert "localos_agent_activation_gate_v1" in api_source
    assert "localos_agent_preview_run_status_v1" in api_source
    assert "_build_activation_gate_summary" in api_source
    assert "_activation_preview_run_status" in api_source
    assert "primary_action_label" in api_source
    assert "human_blockers" in api_source
    assert "requires_preview_run" in api_source
    assert "AGENT_ACTIVATION_GATE_BLOCKED" in api_source
    assert "_activation_connection_plan_from_preflight" in api_source
    assert "/api/agent-blueprints/<blueprint_id>/versions/<version_id>/rollback" in api_source
    assert "_resolve_active_version" in api_source
    assert "_remember_active_version" in api_source
    assert "_version_was_active_before" in api_source
    assert "AGENT_ROLLBACK_GATE_BLOCKED" in api_source
    assert "candidate_requires_preview" in api_source
    assert "Запустите безопасный preview run перед активацией." in api_source
    assert "build_agent_version_diff" in workspace_source
    assert "agent_learning_loop_v1" in workspace_source
    assert "versioned_review" in workspace_source
    assert "_review_journal" in workspace_source
    assert "journal" in workspace_source
    assert "analyze_document_sources_with_llm" in workspace_source
    assert "draft_email_with_llm" in workspace_source
    assert "draft_review_replies_with_llm" in workspace_source

    feedback_endpoint_start = api_source.index("def create_agent_run_feedback")
    feedback_endpoint = api_source[feedback_endpoint_start:api_source.index("@agent_blueprints_bp.route", feedback_endpoint_start + 1)]
    gate_position = feedback_endpoint.index("auto_activation_gate = _build_activation_gate_summary")
    gate_check_position = feedback_endpoint.index("if auto_activation_gate.get(\"can_activate\")")
    remember_position = feedback_endpoint.index("_remember_active_version")
    assert gate_position < gate_check_position < remember_position
    assert "build_learning_loop_summary(feedback, version, new_version, diff, auto_activation_applied)" in feedback_endpoint

    create_version_start = api_source.index("def create_agent_blueprint_version")
    create_version_endpoint = api_source[create_version_start:api_source.index("@agent_blueprints_bp.route", create_version_start + 1)]
    assert "candidate_version" in create_version_endpoint
    assert "rebuild_from_description" in create_version_endpoint
    assert "build_agent_blueprint_draft" in create_version_endpoint
    assert "_save_blueprint_metadata" in create_version_endpoint
    assert "_remember_active_version" not in create_version_endpoint

    setup_start = api_source.index("def setup_agent_blueprint")
    setup_endpoint = api_source[setup_start:api_source.index("@agent_blueprints_bp.route", setup_start + 1)]
    assert "candidate_version" in setup_endpoint
    assert "_build_activation_gate_summary" in setup_endpoint
    assert "_remember_active_version" not in setup_endpoint

    custom_process_start = api_source.index("def save_agent_blueprint_custom_process")
    custom_process_endpoint = api_source[custom_process_start:api_source.index("@agent_blueprints_bp.route", custom_process_start + 1)]
    assert "candidate_version" in custom_process_endpoint
    assert "_build_activation_gate_summary" in custom_process_endpoint
    assert "_remember_active_version" not in custom_process_endpoint

    direct_draft_start = api_source.index("def create_agent_blueprint_draft")
    direct_draft_endpoint = api_source[direct_draft_start:api_source.index("@agent_blueprints_bp.route", direct_draft_start + 1)]
    assert "connection_inventory = _load_direct_builder_connection_inventory" in direct_draft_endpoint
    assert "selected_bindings = _direct_selected_connection_bindings" in direct_draft_endpoint
    assert "missing_connection_choices = _direct_missing_required_connection_choices" in direct_draft_endpoint
    assert "selected_provider_routes = _selected_provider_routes" in direct_draft_endpoint
    assert "missing_provider_routes = _missing_required_provider_routes" in direct_draft_endpoint
    assert "accepted_provider_routes" in direct_draft_endpoint
    assert "metadata = _apply_direct_selected_connection_bindings" in direct_draft_endpoint
    assert "metadata = _apply_selected_provider_routes" in direct_draft_endpoint
    assert "metadata[\"builder_provider_routes_accepted\"]" in direct_draft_endpoint
    assert "connection_preflight = build_agent_integration_preflight" in direct_draft_endpoint
    assert "post_create_handoff = _build_agent_post_connect_handoff" in direct_draft_endpoint
    assert "analyze_table_with_llm" in workspace_source
    assert "analyze_text_with_gigachat" in document_llm_source
    assert "agent_email_draft" in email_llm_source
    assert "agent_review_replies" in review_analysis_source
    assert "agent_table_analysis" in table_analysis_source
    assert "external_dispatch_performed" in document_llm_source
    assert "external_dispatch_performed" in email_llm_source
    assert "external_dispatch_performed" in review_analysis_source
    assert "publish_state" in review_analysis_source
    assert "external_dispatch_performed" in table_analysis_source
    assert "Улучшение версии" in agents_page_source
    assert "Candidate-версия" in agents_page_source
    assert "Зафиксировать улучшение" in agents_page_source
    assert "auto_activate: false" in agents_page_source
    assert "Мои агенты" in agents_page_source
    assert "Состояние миграции" in agents_page_source
    assert "Ручные решения" in agents_page_source
    assert "Обучение" in agents_page_source
    assert "активной" in agents_page_source
    assert "explainApproval" in agents_page_source
    assert "Применить миграцию" in agents_page_source
    assert "Открыть Мои агенты" in admin_page_source
    assert "AIAgentsManagement" not in admin_page_source
    assert "AIAgentSettings" not in agents_page_source
    assert "AIAgentsManagement" not in agents_page_source
    assert "business_has_product_agent_runtime" in legacy_migration_source
    assert "business_agent_enabled_for_channel" in legacy_migration_source
    assert "apply_legacy_ai_agent_migration" in legacy_migration_source
    assert "legacy_ai_agent_migration_v1" in legacy_migration_source
    assert "business_agent_enabled_for_channel" in webhooks_source
    assert "WHERE ai_agent_enabled = 1" not in webhooks_source
    assert "legacy_workflow_context" in ai_agent_source
    assert "deprecated_not_runtime_truth" in ai_agent_source
    assert "Legacy AIAgents.workflow no longer drives runtime state transitions" in ai_agent_source
    assert "Ответь на сообщение клиента, учитывая workflow" not in ai_agent_source
    assert "state.get('init_state')" not in ai_agent_source
    assert "legacy_workflow_context" in chats_api_source
    assert "state.get('init_state')" not in chats_api_source
    assert "provenance" in document_llm_source
    assert "provenance" in email_llm_source
    assert "provenance" in review_analysis_source
    assert "provenance" in table_analysis_source
    assert "/api/agent-builder/sessions" in builder_api_source
    assert "build_agent_builder_state" in builder_api_source
    assert "create_blueprint_from_agent_builder_session" in builder_api_source
    assert builder_api_source.index("billing = charge_agent_creation_credits") < builder_api_source.index("draft = compile_agent_blueprint")
    assert "planner_context = preview.get(\"openclaw_planner_context\")" in builder_api_source
    assert "planner_context=planner_context" in builder_api_source
    assert "metadata[\"openclaw_planner_context\"] = planner_context" in builder_api_source
    assert "metadata[\"openclaw_planner_loop\"] = planner_loop" in builder_api_source
    assert "metadata[\"builder_setup_flow\"] = setup_flow" in builder_api_source
    assert "AGENT_SETUP_INCOMPLETE" in builder_api_source
    assert "\"missing_questions\": missing_questions" in builder_api_source
    assert "\"connector_intelligence\": preview.get(\"connector_intelligence\")" in builder_api_source
    assert "selected_connection_bindings" in builder_api_source
    assert "_selected_connection_bindings" in builder_api_source
    assert "_apply_selected_connection_bindings" in builder_api_source
    assert "_missing_required_connection_choices" in builder_api_source
    assert "AGENT_CONNECTION_CHOICE_REQUIRED" in builder_api_source
    assert "metadata[\"builder_selected_connection_bindings\"]" in builder_api_source
    assert "connection_preflight" in builder_api_source
    assert "post_create_handoff" in builder_api_source
    assert "localos_agent_post_create_handoff_v1" in builder_api_source
    assert "next_binding_key" in builder_api_source
    assert "_build_handoff_connection_plan" in builder_api_source
    assert "\"connection_plan\": connection_plan" in builder_api_source
    assert "next_step" in builder_api_source
    assert "use_ai_compiler: true" in agents_page_source
    assert "connect_required_integrations" in agents_page_source
    assert "recentPostCreateHandoff" in agents_page_source
    assert "recentPostCreateHandoff?.connection_plan || agentConnectionPlan" in agents_page_source
    assert "binding_key" in agents_page_source
    assert "selectedBuilderConnectionBindings" in agents_page_source
    assert "selected_connection_bindings: selectedBuilderConnectionBindings" in agents_page_source
    assert "onSelectConnectionBinding" in agents_page_source
    assert "autoSelectBuilderConnectionBindings" in agents_page_source
    assert "Выбрано автоматически" in agents_page_source
    assert "missingConnectionChoices" in agents_page_source
    assert "Сначала выберите подключение" in agents_page_source
    assert "Почему агента пока нельзя создать" in agents_page_source
    assert "createBlockers" in agents_page_source
    assert "LocalOS должен собрать проверяемый workflow" in agents_page_source
    assert "provider_routes" in agents_page_source
    assert "route_summary" in agents_page_source
    assert "providerRouteLabel" in agents_page_source
    assert "providerRouteTone" in agents_page_source
    assert "selectedConnectionBindingKey" in agents_page_source
    assert "onConfigureBinding" in agents_page_source
    assert "Настроить этот доступ" in agents_page_source
    assert "Сейчас настраивается" in agents_page_source
    assert "connectionPlan={agentConnectionPlan}" in agents_page_source
    assert "selectedPlanItem" in agents_page_source
    assert "binding_key: selectedBinding?.key || ''" in agents_page_source
    assert "Использовать для шага" in agents_page_source
    assert "Выбрано" in agents_page_source
    assert "Использовать" in agents_page_source
    assert "Preflight и preview run" in agents_page_source
    assert "preview_mode: true" in agents_page_source
    assert "external_side_effects_allowed: false" in agents_page_source
    assert "Тест без отправки" in agents_page_source
    assert "preflightResponse.data?.next_binding_key" in agents_page_source
    assert "setWorkspaceMode('connections')" in agents_page_source
    assert "activationGate" in agents_page_source
    assert "gate.summary" in agents_page_source
    assert "gate.connection_plan" in agents_page_source
    assert "activationGate?.next_binding_key" in agents_page_source
    assert "openConnectionsFromActivationGate" in agents_page_source
    assert "onDeleteAgent={onDeleteAgent}" in agents_page_source
    assert "Архивировать" in agents_page_source
    assert "primary_action_label" in agents_page_source
    assert "Активировать версию" in agents_page_source
    assert "Запустить preview" in agents_page_source
    assert "preview_run_status" in agents_page_source
    assert "BuilderPlannerLoopPanel" in agents_page_source
    assert "OpenClaw planner" in agents_page_source
    assert "ConnectorIntelligencePanel" in agents_page_source
    assert "Доступность сервисов" in agents_page_source
    assert "connector_intelligence" in agents_page_source
    assert "BuilderTechnicalDiagnostics" in agents_page_source
    assert "Техническая диагностика LocalOS/OpenClaw" in agents_page_source
    assert "Обычный следующий шаг показан выше" in agents_page_source
    assert "BuilderConnectionSummaryPanel" in agents_page_source
    assert "Подключения для агента" in agents_page_source
    assert "connection_summary" in agents_page_source
    assert "AgentProductCockpit" in agents_page_source
    assert "getBlueprintBuilderPreview" in agents_page_source
    assert "detailsBlueprint={blueprintDetails?.blueprint}" in agents_page_source
    assert "BuilderConnectionReadinessPanel" in agents_page_source
    assert "Что нужно агенту для работы" in agents_page_source
    assert "BuilderConnectionResolverPanel" in agents_page_source
    assert "Как LocalOS подключит сервисы" in agents_page_source
    assert "connection_resolver" in agents_page_source
    assert "resolverStateTone" in agents_page_source
    assert "connection_readiness" in agents_page_source
    assert "setup_cta" in agents_page_source
    assert "Настроить подключение" in agents_page_source
    assert "preview?.connection_plan" in agents_page_source
    assert "compact" in agents_page_source
    assert "_build_preview_connection_plan" in Path("src/services/agent_builder_session.py").read_text(encoding="utf-8")
    assert "_build_connection_readiness" in Path("src/services/agent_builder_session.py").read_text(encoding="utf-8")
    assert "localos_agent_connection_readiness_v1" in Path("src/services/agent_builder_session.py").read_text(encoding="utf-8")
    assert "_build_connection_resolver" in Path("src/services/agent_builder_session.py").read_text(encoding="utf-8")
    assert "localos_agent_connection_resolver_v1" in Path("src/services/agent_builder_session.py").read_text(encoding="utf-8")
    assert "_build_connection_summary" in Path("src/services/agent_builder_session.py").read_text(encoding="utf-8")
    assert "localos_agent_connection_summary_v1" in Path("src/services/agent_builder_session.py").read_text(encoding="utf-8")
    assert "_build_connector_intelligence" in Path("src/services/agent_builder_session.py").read_text(encoding="utf-8")
    assert "localos_agent_connector_intelligence_v1" in Path("src/services/agent_builder_session.py").read_text(encoding="utf-8")
    assert "onPreviewRun" in agents_page_source
    assert "onAttachExistingIntegration" in agents_page_source
    assert "bindingActionHint" in agents_page_source
    assert "AgentConnectionPlanPanel" in agents_page_source
    assert "План подключений" in agents_page_source
    assert "applyPostConnectHandoff" in agents_page_source
    assert "response.data?.post_connect_handoff" in agents_page_source
    assert "handoff.next_binding_key" in agents_page_source
    assert "setSelectedConnectionBindingKey(handoff.next_binding_key)" in agents_page_source
    assert "handoff.workspace_mode === 'run'" in agents_page_source
    assert "setSelectedConnectionBindingKey('')" in agents_page_source
    assert "next_binding" in builder_api_source
    assert "next_route" in builder_api_source
    assert "_preferred_handoff_route" in builder_api_source
    assert "Следующий доступ" in agents_page_source
    assert "recentPostCreateHandoff.next_binding" in agents_page_source
    assert "recentPostCreateHandoff.next_route" in agents_page_source
    assert "connection_plan" in api_source
    assert "localos_agent_connection_plan_v1" in api_source
    assert "agent_binding_integrations" in api_source
    assert "binding_key" in api_source
    assert "_build_agent_post_connect_handoff" in api_source
    assert "_connection_plan_next_binding_key" in api_source
    assert "_connection_plan_route_summary" in api_source
    assert "connector_provider_routes" in api_source
    assert "post_connect_handoff" in api_source
    assert "ready_for_preview" in api_source
    assert "connect_mode" in Path("src/services/agent_provider_registry.py").read_text(encoding="utf-8")
    assert "primary_cta" in Path("src/services/agent_provider_registry.py").read_text(encoding="utf-8")
    assert "provider_action" in Path("src/services/agent_provider_registry.py").read_text(encoding="utf-8")
    assert "use_openclaw_boundary" in Path("src/services/agent_provider_registry.py").read_text(encoding="utf-8")
    assert "select_external_account_key" in Path("src/services/agent_provider_registry.py").read_text(encoding="utf-8")
    assert "planned_oauth_connector" in Path("src/services/agent_provider_registry.py").read_text(encoding="utf-8")
    assert "ProviderActionPill" in agents_page_source
    assert "providerActionLabel" in agents_page_source
    assert "providerActionDescription" in agents_page_source
    assert "['openclaw', 'maton', 'manual'].includes(route.provider || '')" in agents_page_source
    assert "agentPolicyFacts(item)" in agents_page_source
    assert "item.policy_summary" in agents_page_source
    assert "RecommendedProviderRouteNote" in agents_page_source
    assert "recommended_route_reason" in agents_page_source
    assert "selected_provider_routes: acceptedBuilderProviderRoutes ? selectedBuilderProviderRoutes : {}" in agents_page_source
    assert "accepted_provider_routes: acceptedBuilderProviderRoutes" in agents_page_source
    assert "acceptedBuilderProviderRoutes" in agents_page_source
    assert "builderRequiredProviderRouteKeys" in agents_page_source
    assert "Подтвердить подключения" in agents_page_source
    assert "Подключения подтверждены" in agents_page_source
    assert "Что нужно агенту для работы" in agents_page_source
    assert "Использовать этот способ" in agents_page_source
    assert "Способ выбран" in agents_page_source
    assert "autoSelectBuilderProviderRoutes" in agents_page_source
    assert "builderConnectionStatusCopy" in agents_page_source
    assert "builderConnectionNextStepCopy" in agents_page_source
    assert "AGENT_PROVIDER_ROUTE_REQUIRED" in builder_api_source
    assert "AGENT_PROVIDER_ROUTES_CONFIRMATION_REQUIRED" in builder_api_source
    assert "builder_provider_routes_accepted" in builder_api_source
    assert "chooseProviderRoute" in agents_page_source
    assert "/provider-routes" in agents_page_source
    assert "onChooseProviderRoute" in agents_page_source
    assert "Выберите сохранённый Maton.ai key для этого шага." in agents_page_source
    assert "route?.primary_cta" in agents_page_source
    assert "saveMatonIntegration" in agents_page_source
    assert "Maton.ai bridge" in agents_page_source
    assert "matonAuthRef" in agents_page_source
    assert "onSaveMatonIntegration" in agents_page_source
    assert "AgentConnectionDecisionBanner" in agents_page_source
    assert "buildAgentConnectionDecision" in agents_page_source
    assert "BuilderCreationDecisionBanner" in agents_page_source
    assert "buildBuilderCreationDecision" in agents_page_source
    assert "BuilderCompilerPolicyReviewPanel" in agents_page_source
    assert "compiler_policy_review" in agents_page_source
    assert "compiler_workflow_draft" in agents_page_source
    assert "compiler_approval_points" in agents_page_source
    assert "compiler_unsupported_requests" in agents_page_source
    assert "accepted_compiler_plan: acceptedBuilderCompilerPlan" in agents_page_source
    assert "builderCompilerPlanRequiresConfirmation" not in agents_page_source
    assert "const canCreateDraft = setupFlowAllowsDraft" in agents_page_source
    assert "acceptedBuilderCompilerPlan" in agents_page_source
    assert "Принять план" in agents_page_source
    assert "План принят" in agents_page_source
    assert "План агента" in agents_page_source
    assert "compiled workflow candidate" in agents_page_source
    assert "Что нужно изменить в логике" in agents_page_source
    assert "Ответьте на уточнение" in agents_page_source
    assert "connection_resolver" in agents_page_source
    assert "подключение" in agents_page_source
    assert "Создать агента и открыть preview" in agents_page_source
    assert "У бизнеса уже есть несколько подходящих коннектов" in agents_page_source
    assert "Подключения готовы" in agents_page_source
    assert "Настройте один следующий доступ" in agents_page_source
    assert "Сохранить и перейти к тесту" in agents_page_source
    assert "LocalOS покажет следующий шаг" in agents_page_source
    assert "GenericRunProgress" in agents_page_source
    assert "Мои агенты" in agents_page_source
    assert "getAgentListStatus" in agents_page_source
    assert "AgentSummaryPill" in agents_page_source
    assert "Последний run" in agents_page_source
    assert "решений" in agents_page_source
    assert "Данные агента" in agents_page_source
    assert "Что будет делать агент" in agents_page_source
    assert "Голос и стиль" in agents_page_source
    assert "AgentVoiceStylePanel" in agents_page_source
    assert "AIAgents показываются как голоса" in agents_page_source
    assert "Путь {humanizeCategory(category).toLowerCase()}-агента" in agents_page_source
    assert "Технический журнал" in agents_page_source
    assert "AgentRunObservabilityPanel" in agents_page_source
    assert "PreviewRunSummaryPanel" in agents_page_source
    assert "CompiledPreviewSimulationPanel" in agents_page_source
    assert "previewSimulationTone" in agents_page_source
    assert "Симуляция compiled workflow" in agents_page_source
    assert "внешних действий не было" in agents_page_source
    assert "ActivationGateDecisionCard" in agents_page_source
    assert "buildActivationGateDecision" in agents_page_source
    assert "AgentActivationPathStrip" in agents_page_source
    assert "buildActivationPathSteps" in agents_page_source
    assert "workflow проверен" in agents_page_source
    assert "нужно подключить" in agents_page_source
    assert "activationBlockerText" in agents_page_source
    assert "Ждёт решения человека" in agents_page_source
    assert "AgentFourAnswerStrip" in agents_page_source
    assert "Что делает" in agents_page_source
    assert "Готов ли" in agents_page_source
    assert "Чего не хватает" in agents_page_source
    assert "Последний run" in agents_page_source
    assert "Единый billing ledger" in agents_page_source
    assert "Оценка до запуска" in agents_page_source
    assert "Факт после запуска" in agents_page_source
    assert "unified_billing_ledger" in agents_page_source
    assert "без решения человека агент не продолжит внешний шаг" in agents_page_source
    assert "Почему ждём" in agents_page_source
    assert "Preview: {gate.preview_run_status?.ready ? 'пройден' : 'нужен'}" in agents_page_source
    assert "Preflight: {gate.preflight?.ready ? 'готов' : 'проверить'}" in agents_page_source
    assert "Compiled: {gate.compiled_validation?.ready ? 'валиден' : 'проверить'}" in agents_page_source
    assert "Policy: {gate.approval_policy_status?.ready ? 'готова' : 'проверить'}" in agents_page_source
    assert "approvals и limits готовы" in agents_page_source
    assert "нужен human gate" in agents_page_source
    assert "Activation gate" in agents_page_source
    assert "Что показал preview run" in agents_page_source
    assert "activation_hint" in agents_page_source
    assert "Следующий шаг" in agents_page_source
    assert "next_step_label" in agents_page_source
    assert "previewNextStepActionLabel" in agents_page_source
    assert "onNextStepAction" in agents_page_source
    assert "canActivateFromPreview" in agents_page_source
    assert "Activation gate готов" in agents_page_source
    assert "onActivateVersion(activationVersionId)" in agents_page_source
    assert "Открыть подключения" in agents_page_source
    assert "Проверить активацию" in agents_page_source
    assert "Action ledger" in agents_page_source
    assert "localos_agent_preview_summary_v1" in Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "_build_preview_summary" in Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "_preview_next_step" in Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "domain_requests" in Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "agent_sheet_operation_requests" in Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "agent_communication_requests" in Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "provider_handoff" in Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "provider_handoff" in agents_page_source
    assert "Approvals" in agents_page_source
    assert "why_waiting" in agents_page_source
    assert "agent_review_publish_requests" in Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "publish_requests" in agents_page_source
    assert "Support export" in agents_page_source
    assert "agent_run_observability_v1" in Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "billing_ledger" in Path("src/services/agent_blueprint_runner.py").read_text(encoding="utf-8")
    assert "BillingActionItem" in agents_page_source
    assert "reserve ${item.billing_summary?.reserved_tokens" in agents_page_source
    assert "Использовано в последнем запуске" in agents_page_source
    assert "used_sources" in workspace_source
    assert "resultFieldLabels" in agents_page_source
    assert "Bookings" in capability_handlers_source
    assert "agent_communication_requests" in capability_handlers_source
    assert "reviewreplydrafts" in capability_handlers_source
    assert "agent_service_optimization_requests" in capability_handlers_source
    assert "agent_sheet_operation_requests" in capability_handlers_source
    assert "provider_write_performed=False" in capability_handlers_source
    assert "sheets.append_row_request" in action_policy_source
    assert "dispatch_telegram_message_to_agent_blueprints" in telegram_webhook_source
    assert "agent_trigger_events" in trigger_runtime_source
    assert "telegram.message.received" in trigger_runtime_source
    assert "dispatch_scheduled_agent_blueprints" in trigger_runtime_source
    assert "dispatch_due_scheduled_agent_blueprints" in trigger_runtime_source
    assert "schedule.daily" in trigger_runtime_source
    assert "scheduler" in trigger_runtime_source
    assert "AGENT_SCHEDULE_DISPATCH_ENABLED" in worker_source
    assert "_dispatch_agent_schedules_if_due" in worker_source
    assert "AGENT_SCHEDULE_DISPATCH_ENABLED: ${AGENT_SCHEDULE_DISPATCH_ENABLED:-false}" in compose_source
    assert "manual_publish_required=True" in capability_handlers_source
    assert "manual_apply_required=True" in capability_handlers_source
    assert "reserve_paid_action_credits" in capability_handlers_source
    assert "finalize_reserved_action_credits" in capability_handlers_source


def test_generic_document_runner_uses_sources_and_stops_for_final_approval():
    from services.agent_blueprint_draft_builder import build_agent_blueprint_draft
    from services.agent_blueprint_runner import AgentBlueprintRunner

    cursor = FakeCursor()
    draft = build_agent_blueprint_draft("Обработай договор и найди риски")
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Document agent",
        "category": "documents",
        "metadata_json": {
            **draft["metadata"],
            "agent_setup": {
                "workflow_description": "Проверить договор",
                "extraction_rules": "Найти сроки, оплату и ответственность",
                "processing_rules": "Не придумывать факты",
                "output_format": "Список рисков",
                "approval_boundaries": ["final_output", "external_delivery"],
            },
            "agent_sources": [
                {
                    "id": "src1",
                    "source_type": "text",
                    "name": "Договор",
                    "content_text": "Оплата 10000 рублей. Ответственность за просрочку: штраф 10%.",
                    "content_length": 68,
                    "extraction_state": "ready",
                }
            ],
        },
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "steps_json": draft["version_payload"]["steps"],
        "capability_allowlist_json": [],
    }

    result = AgentBlueprintRunner(cursor).start_run("ver1", {}, {"user_id": "user1"})

    assert result["success"] is True
    run = result["run"]
    assert run["status"] == "waiting_approval"
    assert [step["step_key"] for step in run["steps"]] == [
        "collect_inputs",
        "extract_context",
        "prepare_output",
        "approve_output",
    ]
    output = [item for item in run["artifacts"] if item["artifact_type"] == "agent_output_draft"][0]
    assert output["payload_json"]["external_dispatch_performed"] is False
    assert output["payload_json"]["result"]["title"] == "Разбор документа"
    assert output["payload_json"]["result"]["facts"]
    assert output["payload_json"]["result"]["fields"]["Оплата"]
    assert output["payload_json"]["result"]["fields"]["Ответственность"]
    assert output["payload_json"]["dispatch_state"] == "not_dispatched"
    assert run["approvals"][0]["approval_type"] == "final_output"


def test_generic_email_runner_prepares_draft_and_never_dispatches():
    from services.agent_blueprint_draft_builder import build_agent_blueprint_draft
    from services.agent_blueprint_runner import AgentBlueprintRunner

    cursor = FakeCursor()
    draft = build_agent_blueprint_draft("Подготовь письмо клиенту по контексту", "email")
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Email agent",
        "category": "email",
        "metadata_json": {
            **draft["metadata"],
            "agent_setup": {
                "workflow_description": "Подготовить письмо клиенту о новой услуге",
                "extraction_rules": "Взять услугу, выгоду и ограничение по тону",
                "processing_rules": "Писать дружелюбно, не обещать скидку без подтверждения",
                "output_format": "subject, body, checklist, missing_info",
                "approval_boundaries": ["final_output", "external_delivery"],
            },
            "agent_sources": [
                {
                    "id": "src1",
                    "source_type": "text",
                    "name": "Контекст письма",
                    "content_text": "Адресат: Анна. Услуга: уход для волос. Цель: пригласить на консультацию.",
                    "content_length": 77,
                    "extraction_state": "ready",
                }
            ],
        },
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "steps_json": draft["version_payload"]["steps"],
        "capability_allowlist_json": [],
    }

    result = AgentBlueprintRunner(cursor).start_run("ver1", {}, {"user_id": "user1"})

    assert result["success"] is True
    run = result["run"]
    assert run["status"] == "waiting_approval"
    assert not [step for step in run["steps"] if step["step_type"] == "capability"]
    output = [item for item in run["artifacts"] if item["artifact_type"] == "agent_output_draft"][0]
    payload = output["payload_json"]
    email_result = payload["result"]
    assert payload["category"] == "email"
    assert payload["external_dispatch_performed"] is False
    assert payload["dispatch_state"] == "not_dispatched"
    assert email_result["external_dispatch_performed"] is False
    assert email_result["delivery_state"] == "not_dispatched"
    assert email_result["subject"]
    assert email_result["body"]
    assert email_result["checklist"]
    assert email_result["provenance"] == ["Контекст письма"]
    assert run["approvals"][0]["approval_type"] == "final_output"


def test_generic_table_runner_prepares_report_and_never_dispatches():
    from services.agent_blueprint_draft_builder import build_agent_blueprint_draft
    from services.agent_blueprint_runner import AgentBlueprintRunner

    cursor = FakeCursor()
    draft = build_agent_blueprint_draft("Проверь CSV и найди ошибки", "tables")
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Table agent",
        "category": "tables",
        "metadata_json": {
            **draft["metadata"],
            "agent_setup": {
                "workflow_description": "Проверить таблицу клиентов",
                "extraction_rules": "Найти пустые email, дубли и строки к проверке",
                "processing_rules": "Не изменять таблицу, только показать проблемы",
                "output_format": "summary, exceptions, rows_to_review, recommendations",
                "approval_boundaries": ["final_output", "external_delivery"],
            },
            "agent_sources": [
                {
                    "id": "src1",
                    "source_type": "text",
                    "name": "clients.csv",
                    "content_text": "name,email,phone\nАнна,anna@example.com,+1\nАнна,anna@example.com,+1\nБорис,,+2\n",
                    "content_length": 83,
                    "extraction_state": "ready",
                }
            ],
        },
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "steps_json": draft["version_payload"]["steps"],
        "capability_allowlist_json": [],
    }

    result = AgentBlueprintRunner(cursor).start_run("ver1", {}, {"user_id": "user1"})

    assert result["success"] is True
    run = result["run"]
    assert run["status"] == "waiting_approval"
    assert not [step for step in run["steps"] if step["step_type"] == "capability"]
    output = [item for item in run["artifacts"] if item["artifact_type"] == "agent_output_draft"][0]
    payload = output["payload_json"]
    table_result = payload["result"]
    assert payload["category"] == "tables"
    assert payload["external_dispatch_performed"] is False
    assert payload["dispatch_state"] == "not_dispatched"
    assert table_result["external_dispatch_performed"] is False
    assert table_result["delivery_state"] == "not_dispatched"
    assert table_result["summary"]
    assert table_result["exceptions"]
    assert table_result["rows_to_review"]
    assert table_result["recommendations"]
    assert table_result["provenance"] == ["clients.csv"]
    assert run["approvals"][0]["approval_type"] == "final_output"


def test_generic_reviews_runner_prepares_reply_drafts_and_never_publishes():
    from services.agent_blueprint_draft_builder import build_agent_blueprint_draft
    from services.agent_blueprint_runner import AgentBlueprintRunner

    cursor = FakeCursor()
    draft = build_agent_blueprint_draft("Подготовь ответы на отзывы", "reviews")
    reviews_text = (
        "author_name,rating,text\n"
        "Анна,5,Очень понравился сервис\n"
        "Иван,2,Долго ждал и администратор был груб\n"
    )
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Reviews agent",
        "category": "reviews",
        "metadata_json": {
            **draft["metadata"],
            "agent_setup": {
                "workflow_description": "Подготовить ответы на отзывы",
                "extraction_rules": "Определить тон, проблему клиента и безопасный ответ",
                "processing_rules": "Не обещать скидку, компенсацию или публикацию без подтверждения",
                "output_format": "reply_drafts, manual_review_reasons, checklist",
                "approval_boundaries": ["final_output", "external_delivery"],
            },
            "agent_sources": [
                {
                    "id": "src1",
                    "source_type": "text",
                    "name": "Отзывы",
                    "content_text": reviews_text,
                    "content_length": len(reviews_text),
                    "extraction_state": "ready",
                }
            ],
        },
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "steps_json": draft["version_payload"]["steps"],
        "capability_allowlist_json": [],
    }

    result = AgentBlueprintRunner(cursor).start_run("ver1", {}, {"user_id": "user1"})

    assert result["success"] is True
    run = result["run"]
    assert run["status"] == "waiting_approval"
    assert not [step for step in run["steps"] if step["step_type"] == "capability"]
    output = [item for item in run["artifacts"] if item["artifact_type"] == "agent_output_draft"][0]
    payload = output["payload_json"]
    review_result = payload["result"]
    assert payload["category"] == "reviews"
    assert payload["external_dispatch_performed"] is False
    assert payload["dispatch_state"] == "not_dispatched"
    assert review_result["external_dispatch_performed"] is False
    assert review_result["publish_state"] == "not_published"
    assert review_result["delivery_state"] == "not_dispatched"
    assert review_result["reply_drafts"]
    assert review_result["manual_review_reasons"]
    assert review_result["checklist"]
    assert review_result["provenance"] == ["Отзывы"]
    assert run["approvals"][0]["approval_type"] == "final_output"


def test_message_result_needs_source_data_without_sheet_rows():
    from services.agent_blueprint_workspace import _render_output

    result = _render_output(
        "custom",
        {
            "workflow_description": "Открой таблицу поездок и подготовь сообщение владельцу",
            "processing_rules": "Не придумывать факты",
            "output_format": "Готовое сообщение для проверки",
        },
        [{"source_name": "business_profile", "summary": "ready; id: biz-1", "raw": {"id": "biz-1"}}],
        [],
        {},
    )

    assert result["status"] == "needs_source_data"
    assert "draft_text" not in result
    assert "не получил строку поездки" in result["summary"][0].lower()


def test_message_result_prompts_google_reconnect_when_sheet_auth_is_revoked():
    from services.agent_blueprint_workspace import _render_output

    result = _render_output(
        "custom",
        {
            "workflow_description": "Открой таблицу поездок и подготовь сообщение владельцу",
            "processing_rules": "Не придумывать факты",
            "output_format": "Готовое сообщение для проверки",
        },
        [
            {
                "source_name": "google_sheets_error",
                "summary": "Google token refresh failed with HTTP 400: invalid_grant",
                "raw": {
                    "provider_error": "GOOGLE_SHEETS_PROVIDER_NOT_READY",
                    "provider_error_message": "Google token refresh failed with HTTP 400: invalid_grant",
                    "next_action": "connect_or_repair_google_sheets_provider",
                },
            }
        ],
        [],
        {},
    )

    assert result["status"] == "needs_google_access"
    assert result["title"] == "Нужно переподключить Google-доступ"
    assert "Таблица выбрана" in result["summary"][0]
    assert "Переподключите Google-доступ" in result["next_questions"][0]
    assert "invalid_grant" in result["technical_reason"]


def test_message_result_for_disabled_google_sheets_api_guides_project_setup():
    from services.agent_blueprint_workspace import _render_output

    result = _render_output(
        "custom",
        {
            "workflow_description": "Открой таблицу поездок и подготовь сообщение владельцу",
            "processing_rules": "Не придумывать факты",
            "output_format": "Готовое сообщение для проверки",
        },
        [
            {
                "source_name": "google_sheets_error",
                "summary": "Google Sheets read failed with HTTP 403",
                "raw": {
                    "provider_error": "GOOGLE_SHEETS_PROVIDER_NOT_READY",
                    "provider_error_message": (
                        "Google Sheets read failed with HTTP 403: "
                        "Google Sheets API has not been used in project 304042072643 before or it is disabled. "
                        "Enable it by visiting https://console.developers.google.com/apis/api/sheets.googleapis.com/overview?project=304042072643"
                    ),
                },
            }
        ],
        [],
        {},
    )

    assert result["status"] == "needs_google_api_enabled"
    assert result["title"] == "Нужно включить Google Sheets API"
    assert "Google-доступ подключён" in result["summary"][0]
    assert "304042072643" in result["next_questions"][0]
    assert "Переподключите Google-доступ" not in " ".join(result["next_questions"])


def test_message_result_for_google_sheets_invalid_range_guides_sheet_tab_setup():
    from services.agent_blueprint_workspace import _render_output

    result = _render_output(
        "custom",
        {
            "workflow_description": "Открой таблицу поездок и подготовь сообщение владельцу",
            "processing_rules": "Не придумывать факты",
            "output_format": "Готовое сообщение для проверки",
        },
        [
            {
                "source_name": "google_sheets_error",
                "summary": "Google Sheets read failed with HTTP 400",
                "raw": {
                    "provider_error": "GOOGLE_SHEETS_PROVIDER_NOT_READY",
                    "provider_error_message": (
                        "Google Sheets read failed with HTTP 400: "
                        "Unable to parse range: Sheet1!A1:Z"
                    ),
                },
            }
        ],
        [],
        {},
    )

    assert result["status"] == "needs_sheet_tab"
    assert result["title"] == "Нужно выбрать лист таблицы"
    assert "лист таблицы" in result["summary"][0]
    assert "Переподключите Google-доступ" not in " ".join(result["next_questions"])


def test_message_result_uses_google_sheets_rows_for_concrete_draft(monkeypatch):
    import services.agent_blueprint_workspace as workspace

    monkeypatch.setattr(workspace, "analyze_text_with_gigachat", lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("llm offline")))

    result = workspace._render_output(
        "custom",
        {
            "workflow_description": "Выбери поездку на 20 апреля и подготовь сообщение владельцу",
            "processing_rules": "Не придумывать факты",
            "output_format": "Готовое сообщение для проверки",
        },
        [
            {
                "source_name": "google_sheets",
                "summary": "date: 2026-04-20; route: Airport -> Center; passenger: Anna",
                "raw": {
                    "date": "2026-04-20",
                    "route": "Airport -> Center",
                    "passenger": "Anna",
                },
            }
        ],
        [],
        {},
    )

    assert result["title"] == "Черновик сообщения"
    assert "Airport -> Center" in result["draft_text"]
    assert result["analysis_source"] == "deterministic_fallback"


def test_runner_propagates_google_sheets_run_rows_into_output_artifact(monkeypatch):
    import services.agent_blueprint_workspace as workspace
    from services.agent_blueprint_runner import AgentBlueprintRunner

    monkeypatch.setattr(workspace, "analyze_text_with_gigachat", lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("llm offline")))

    cursor = FakeCursor()
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Trips agent",
        "category": "custom",
        "metadata_json": {
            "agent_setup": {
                "workflow_description": "Выбери поездку на 20 апреля и подготовь сообщение владельцу",
                "processing_rules": "Не придумывать факты",
                "output_format": "Готовое сообщение для проверки",
            },
            "agent_sources": [
                {
                    "id": "business-profile",
                    "source_type": "internal",
                    "internal_source": "business_profile",
                    "name": "Профиль бизнеса",
                }
            ],
        },
    }
    cursor.tables["agent_runs"]["run1"] = {
        "id": "run1",
        "blueprint_id": "bp1",
        "blueprint_version_id": "ver1",
        "business_id": "biz1",
        "status": "running",
        "input_json": {},
        "output_json": {},
        "created_by_user_id": "user1",
    }
    cursor.tables["agent_run_steps"]["read-step"] = {
        "id": "read-step",
        "run_id": "run1",
        "step_index": 0,
        "step_key": "read_google_sheets",
        "step_type": "capability",
        "status": "completed",
        "input_json": {},
        "output_json": {
            "capability": "google_sheets.read_rows",
            "orchestrator": {
                "success": True,
                "result": {
                    "source": "google_sheets",
                    "provider_read_performed": True,
                    "rows": [
                        {
                            "date": "2026-04-20",
                            "route": "Tallinn Airport -> Old Town",
                            "passenger": "Anna",
                        }
                    ],
                },
            },
        },
    }

    payload = AgentBlueprintRunner(cursor)._build_artifact_payload(
        cursor.tables["agent_runs"]["run1"],
        {
            "key": "prepare_output",
            "artifact_type": "agent_output_draft",
            "payload": {
                "category": "custom",
                "rows_from_step": "read_google_sheets",
                "format": "Готовое сообщение для проверки",
            },
        },
    )

    assert payload["items_used"] >= 1
    assert payload["result"]["title"] == "Черновик сообщения"
    assert payload["result"].get("status") != "needs_source_data"
    assert "Tallinn Airport -> Old Town" in payload["result"]["draft_text"]
    assert "business_profile" not in payload["result"]["draft_text"]
    assert payload["external_dispatch_performed"] is False


def test_agents_page_normal_result_panel_does_not_dump_raw_artifact_payload():
    source = Path("frontend/src/pages/dashboard/AgentBlueprintsPage.tsx").read_text(encoding="utf-8")

    assert "stringifyBusinessValue(artifact.payload_json)" not in source
    assert "Результат не был сохранён. Запустите тест ещё раз" in source
    assert "const isBlocked = result.state === 'blocker';" in source
    assert "const canApprove = Boolean(pendingApproval && !isBlocked);" in source
    assert "needsScenarioRebuildForSourceResult" in source
    assert "Пересобрать сценарий" in source
    assert "Этот агент создан старой версией сценария" in source
    assert "needsGoogleSheetsSourceSetup" in source
    assert "Указать Google-таблицу" in source
    assert "Укажите Google-таблицу и лист со списком поездок" in source
    assert "needsGoogleAccessReconnect" in source
    assert "focus: 'google_sheets'" in source
    assert "return_to: '/dashboard/agents'" in source
    assert "Переподключить Google-доступ" in source
    assert "google_auth" in source
    assert "googleAccessJustConnected" in source
    assert "Google-доступ подключён. Теперь запустите тест ещё раз" in source
    assert "Этот результат был получен до переподключения" in source
    assert "blocked_result" in source
    assert "hasFreshGoogleSheetsAccessAfterResult" in source
    assert "Google-доступ обновлён. Запустите тест ещё раз" in source
    assert "Почему нельзя подтвердить результат" in source


def test_settings_integrations_first_layer_separates_google_sheets_from_google_business():
    external = Path("frontend/src/components/ExternalIntegrations.tsx").read_text(encoding="utf-8")
    hub_state = Path("frontend/src/pages/dashboard/settings/settingsHubState.ts").read_text(encoding="utf-8")
    hub_copy = Path("frontend/src/pages/dashboard/settings/settingsHubCopy.ts").read_text(encoding="utf-8")

    assert 'data-testid="settings-integrations-scenario"' in external
    assert "Google Таблицы" in external
    assert "Этот доступ нужен агентам для чтения Google Таблиц. Он не публикует ничего наружу." in external
    assert "Google Документы: позже" in external
    assert "Google-доступ" in external
    assert "Карточка" in external
    assert "Таблицы" in external
    assert "google_sheets" in hub_state
    assert "/dashboard/settings/integrations?focus=google_sheets" in hub_state
    assert "Agent access to table rows." in hub_copy
    assert "Доступ агентов к строкам таблиц." in hub_copy


def test_agent_source_ingestion_extracts_text_pdf_docx_xlsx_and_rejects_unsafe_files():
    import io
    import zipfile

    from openpyxl import Workbook

    from services.agent_source_ingestion import build_agent_source_from_upload

    def build_test_pdf_bytes(text):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("utf-8")
        objects = [
            b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
            b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n",
            (
                b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >> endobj\n"
            ),
            b"4 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n",
            b"5 0 obj << /Length "
            + str(len(stream)).encode("ascii")
            + b" >> stream\n"
            + stream
            + b"\nendstream endobj\n",
        ]
        output = bytearray(b"%PDF-1.4\n")
        offsets = [0]
        for item in objects:
            offsets.append(len(output))
            output.extend(item)
        xref_offset = len(output)
        output.extend(b"xref\n0 6\n0000000000 65535 f \n")
        for offset in offsets[1:]:
            output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
        output.extend(b"trailer << /Root 1 0 R /Size 6 >>\n")
        output.extend(f"startxref\n{xref_offset}\n%%EOF\n".encode("ascii"))
        return bytes(output)

    text_source, text_error = build_agent_source_from_upload(
        FakeUpload("contract.txt", "text/plain", "Оплата 10000. Ответственность: штраф.".encode("utf-8")),
        "Договор",
    )
    assert text_error == {}
    assert text_source["content_text"].startswith("Оплата 10000")
    assert text_source["extraction_state"] == "ready"

    pdf_bytes = build_test_pdf_bytes("Payment 15000. Penalty 12 percent.")
    pdf_source, pdf_error = build_agent_source_from_upload(
        FakeUpload("contract.pdf", "application/pdf", pdf_bytes),
    )
    assert pdf_error == {}
    assert "Payment 15000" in pdf_source["content_text"]
    assert pdf_source["extraction_method"] == "pypdf"

    docx_buffer = io.BytesIO()
    archive = zipfile.ZipFile(docx_buffer, "w")
    try:
        archive.writestr(
            "word/document.xml",
            (
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                "<w:body><w:p><w:r><w:t>Документ содержит срок и оплату.</w:t></w:r></w:p></w:body>"
                "</w:document>"
            ),
        )
    finally:
        archive.close()
    docx_source, docx_error = build_agent_source_from_upload(
        FakeUpload(
            "contract.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            docx_buffer.getvalue(),
        )
    )
    assert docx_error == {}
    assert "срок и оплату" in docx_source["content_text"]
    assert docx_source["extraction_method"] == "docx_xml"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Риски"
    sheet.append(["Поле", "Значение"])
    sheet.append(["Штраф", "10%"])
    xlsx_buffer = io.BytesIO()
    workbook.save(xlsx_buffer)
    xlsx_source, xlsx_error = build_agent_source_from_upload(
        FakeUpload(
            "risks.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            xlsx_buffer.getvalue(),
        )
    )
    assert xlsx_error == {}
    assert "Штраф" in xlsx_source["content_text"]
    assert xlsx_source["extraction_method"] == "openpyxl"

    unsafe_source, unsafe_error = build_agent_source_from_upload(
        FakeUpload("payload.exe", "application/octet-stream", b"bad"),
    )
    assert unsafe_source == {}
    assert unsafe_error["code"] == "UNSUPPORTED_FILE_TYPE"
    assert "поддерживается" in unsafe_error["message"].lower()

    empty_source, empty_error = build_agent_source_from_upload(
        FakeUpload("empty.txt", "text/plain", b""),
    )
    assert empty_source == {}
    assert empty_error["code"] == "EMPTY_FILE"
    assert "пустой" in empty_error["message"].lower()


def test_agent_datahub_catalog_returns_available_internal_sources():
    from services.agent_datahub import build_agent_datahub_catalog

    cursor = FakeDatahubCursor()
    catalog = build_agent_datahub_catalog(
        cursor,
        "biz1",
        [{"source_type": "internal", "internal_source": "services"}],
    )

    by_key = {item["key"]: item for item in catalog}
    assert by_key["business_profile"]["available_count"] == 1
    assert by_key["services"]["available_count"] == 2
    assert by_key["services"]["connected"] is True
    assert by_key["reviews"]["preview"]
    assert by_key["prospectingleads"]["state"] == "empty"


def test_agent_document_llm_analysis_uses_generator_rules_and_provenance():
    from services.agent_document_llm import analyze_document_sources_with_llm

    captured = {}

    def fake_generator(prompt, *, business_id="", user_id=""):
        captured["prompt"] = prompt
        captured["business_id"] = business_id
        captured["user_id"] = user_id
        return json.dumps(
            {
                "title": "LLM contract analysis",
                "summary": ["Оплата 10000 рублей"],
                "risks": ["Штраф 10% за просрочку"],
                "facts": ["Срок 30 дней", "Оплата 10000 рублей"],
                "fields": {"Оплата": "10000 рублей", "Срок": "30 дней"},
                "next_questions": ["Кто подписывает договор?"],
                "rules_applied": ["Не придумывать факты"],
            },
            ensure_ascii=False,
        )

    result = analyze_document_sources_with_llm(
        {
            "workflow_description": "Проверить договор",
            "extraction_rules": "Суммы, сроки, штрафы",
            "processing_rules": "Не придумывать факты",
            "output_format": "Краткий отчёт",
        },
        [
            {
                "source_name": "contract.txt",
                "summary": "Оплата 10000 рублей. Срок 30 дней. Штраф 10%.",
                "raw": {"text": "Оплата 10000 рублей. Срок 30 дней. Штраф 10%."},
            }
        ],
        business_id="biz1",
        user_id="user1",
        generator=fake_generator,
    )

    assert result["llm_analysis_used"] is True
    assert result["analysis_source"] == "gigachat"
    assert result["external_dispatch_performed"] is False
    assert result["fields"]["Оплата"] == "10000 рублей"
    assert result["provenance"] == ["contract.txt"]
    assert captured["business_id"] == "biz1"
    assert captured["user_id"] == "user1"
    assert "Не придумывать факты" in captured["prompt"]
    assert "contract.txt" in captured["prompt"]


def test_agent_document_llm_analysis_falls_back_without_external_dispatch():
    from services.agent_document_llm import analyze_document_sources_with_llm

    def failing_generator(prompt, *, business_id="", user_id=""):
        raise RuntimeError("provider unavailable")

    result = analyze_document_sources_with_llm(
        {"processing_rules": "Показывать риски", "output_format": "Отчёт"},
        [{"source_name": "contract.txt", "summary": "Оплата 10000. Ответственность: штраф.", "raw": {}}],
        generator=failing_generator,
    )

    assert result["llm_analysis_used"] is False
    assert result["analysis_source"] == "deterministic_fallback"
    assert result["external_dispatch_performed"] is False
    assert result["provenance"] == ["contract.txt"]
    assert result["risks"]


def test_agent_email_llm_draft_uses_generator_rules_and_provenance():
    from services.agent_email_llm import draft_email_with_llm

    captured = {}

    def fake_generator(prompt, *, business_id="", user_id=""):
        captured["prompt"] = prompt
        captured["business_id"] = business_id
        captured["user_id"] = user_id
        return json.dumps(
            {
                "title": "Email draft",
                "subject": "Анна, приглашаем на консультацию",
                "body": "Здравствуйте, Анна! Приглашаем на консультацию по уходу для волос.",
                "checklist": ["Проверить имя", "Проверить оффер"],
                "assumptions": ["Контекст взят из источника"],
                "missing_info": ["Дата консультации"],
                "rules_applied": ["Не обещать скидку"],
            },
            ensure_ascii=False,
        )

    result = draft_email_with_llm(
        {
            "workflow_description": "Подготовить письмо клиенту",
            "extraction_rules": "Адресат, услуга, цель",
            "processing_rules": "Не обещать скидку",
            "output_format": "subject/body/checklist",
        },
        [
            {
                "source_name": "Контекст",
                "summary": "Адресат: Анна. Услуга: уход для волос.",
                "raw": {"text": "Адресат: Анна. Услуга: уход для волос. Цель: консультация."},
            }
        ],
        business_id="biz1",
        user_id="user1",
        generator=fake_generator,
    )

    assert result["llm_analysis_used"] is True
    assert result["analysis_source"] == "gigachat"
    assert result["external_dispatch_performed"] is False
    assert result["delivery_state"] == "not_dispatched"
    assert result["subject"] == "Анна, приглашаем на консультацию"
    assert result["checklist"]
    assert result["provenance"] == ["Контекст"]
    assert captured["business_id"] == "biz1"
    assert captured["user_id"] == "user1"
    assert "Не обещать скидку" in captured["prompt"]


def test_agent_email_llm_falls_back_without_external_dispatch():
    from services.agent_email_llm import draft_email_with_llm

    def failing_generator(prompt, *, business_id="", user_id=""):
        raise RuntimeError("provider unavailable")

    result = draft_email_with_llm(
        {"workflow_description": "Подготовить письмо", "processing_rules": "Дружелюбно"},
        [{"source_name": "Контекст", "summary": "Адресат: Анна. Услуга: консультация.", "raw": {}}],
        generator=failing_generator,
    )

    assert result["llm_analysis_used"] is False
    assert result["analysis_source"] == "deterministic_fallback"
    assert result["external_dispatch_performed"] is False
    assert result["delivery_state"] == "not_dispatched"
    assert result["subject"]
    assert result["body"]


def test_agent_table_analysis_uses_generator_rules_and_provenance():
    from services.agent_table_analysis import analyze_table_with_llm

    captured = {}

    def fake_generator(prompt, *, business_id="", user_id=""):
        captured["prompt"] = prompt
        captured["business_id"] = business_id
        captured["user_id"] = user_id
        return json.dumps(
            {
                "title": "Table report",
                "summary": ["Проверено 3 строки"],
                "exceptions": ["Строка 3: пустой email"],
                "rows_to_review": [
                    {"row": 3, "reason": "пустой email", "source_name": "clients.csv", "values": {"name": "Борис"}}
                ],
                "recommendations": ["Заполнить email"],
                "rules_applied": ["Не изменять таблицу"],
            },
            ensure_ascii=False,
        )

    result = analyze_table_with_llm(
        {
            "workflow_description": "Проверить клиентов",
            "extraction_rules": "Пустые email и дубли",
            "processing_rules": "Не изменять таблицу",
            "output_format": "exceptions report",
        },
        [
            {"source_name": "clients.csv", "summary": "name: Борис; email: ; phone: +2", "raw": {"name": "Борис", "email": "", "phone": "+2"}}
        ],
        business_id="biz1",
        user_id="user1",
        generator=fake_generator,
    )

    assert result["llm_analysis_used"] is True
    assert result["analysis_source"] == "gigachat"
    assert result["external_dispatch_performed"] is False
    assert result["delivery_state"] == "not_dispatched"
    assert result["exceptions"] == ["Строка 3: пустой email"]
    assert result["rows_to_review"][0]["row"] == 3
    assert result["provenance"] == ["clients.csv"]
    assert captured["business_id"] == "biz1"
    assert captured["user_id"] == "user1"
    assert "Не изменять таблицу" in captured["prompt"]


def test_agent_table_analysis_falls_back_without_external_dispatch():
    from services.agent_table_analysis import analyze_table_with_llm

    def failing_generator(prompt, *, business_id="", user_id=""):
        raise RuntimeError("provider unavailable")

    result = analyze_table_with_llm(
        {"workflow_description": "Проверить таблицу", "processing_rules": "Только отчёт"},
        [
            {"source_name": "clients.csv", "summary": "name: Анна; email: anna@example.com", "raw": {"name": "Анна", "email": "anna@example.com"}},
            {"source_name": "clients.csv", "summary": "name: Борис; email: ", "raw": {"name": "Борис", "email": ""}},
        ],
        generator=failing_generator,
    )

    assert result["llm_analysis_used"] is False
    assert result["analysis_source"] == "deterministic_fallback"
    assert result["external_dispatch_performed"] is False
    assert result["delivery_state"] == "not_dispatched"
    assert result["exceptions"]
    assert result["rows_to_review"]


def test_agent_review_reply_analysis_uses_generator_rules_and_provenance():
    from services.agent_review_reply_analysis import draft_review_replies_with_llm

    captured = {}

    def fake_generator(prompt, *, business_id="", user_id=""):
        captured["prompt"] = prompt
        captured["business_id"] = business_id
        captured["user_id"] = user_id
        return json.dumps(
            {
                "title": "Review replies",
                "summary": ["Подготовлено 2 черновика"],
                "reply_drafts": [
                    {
                        "review_id": "rev1",
                        "author_name": "Иван",
                        "rating": "2",
                        "sentiment": "negative",
                        "reply": "Иван, спасибо за обратную связь. Мы разберём ситуацию с ожиданием.",
                        "manual_review_reason": "Негативный отзыв требует проверки менеджером.",
                    }
                ],
                "manual_review_reasons": ["Негативный отзыв требует проверки менеджером."],
                "checklist": ["Проверить тон", "Не обещать компенсацию"],
                "rules_applied": ["Не обещать скидку"],
            },
            ensure_ascii=False,
        )

    result = draft_review_replies_with_llm(
        {
            "workflow_description": "Подготовить ответы на отзывы",
            "extraction_rules": "Тональность и причина недовольства",
            "processing_rules": "Не обещать скидку",
            "output_format": "reply_drafts/checklist",
        },
        [
            {
                "source_name": "Отзывы",
                "summary": "Иван поставил 2: долго ждал",
                "raw": {"id": "rev1", "author_name": "Иван", "rating": 2, "text": "Долго ждал"},
            }
        ],
        business_id="biz1",
        user_id="user1",
        generator=fake_generator,
    )

    assert result["llm_analysis_used"] is True
    assert result["analysis_source"] == "gigachat"
    assert result["external_dispatch_performed"] is False
    assert result["publish_state"] == "not_published"
    assert result["delivery_state"] == "not_dispatched"
    assert result["reply_drafts"][0]["reply"]
    assert result["manual_review_reasons"]
    assert result["provenance"] == ["Отзывы"]
    assert captured["business_id"] == "biz1"
    assert captured["user_id"] == "user1"
    assert "Не обещать скидку" in captured["prompt"]


def test_agent_review_reply_analysis_falls_back_without_publish():
    from services.agent_review_reply_analysis import draft_review_replies_with_llm

    def failing_generator(prompt, *, business_id="", user_id=""):
        raise RuntimeError("provider unavailable")

    result = draft_review_replies_with_llm(
        {"workflow_description": "Ответить на отзывы", "processing_rules": "Проверить негатив вручную"},
        [
            {
                "source_name": "Отзывы",
                "summary": "Ольга поставила 2: плохо и долго",
                "raw": {"id": "rev2", "author_name": "Ольга", "rating": 2, "text": "Плохо и долго"},
            }
        ],
        generator=failing_generator,
    )

    assert result["llm_analysis_used"] is False
    assert result["analysis_source"] == "deterministic_fallback"
    assert result["external_dispatch_performed"] is False
    assert result["publish_state"] == "not_published"
    assert result["delivery_state"] == "not_dispatched"
    assert result["reply_drafts"]
    assert result["manual_review_reasons"]
    assert result["provenance"] == ["Отзывы"]


def test_agent_version_diff_shows_readable_changes():
    from services.agent_blueprint_workspace import build_agent_version_diff

    first_version = {
        "id": "ver1",
        "version_number": 1,
        "goal": "Проверить договор",
        "inputs_schema_json": {"agent_setup": {"processing_rules": "Показывать риски"}},
        "steps_json": [{"key": "prepare_output", "type": "artifact"}],
        "capability_allowlist_json": [],
        "approval_policy_json": {"required_for": ["final_output"]},
        "output_schema_json": {"format": "summary"},
    }
    second_version = {
        "id": "ver2",
        "version_number": 2,
        "goal": "Проверить договор и выделить санкции",
        "inputs_schema_json": {"agent_setup": {"processing_rules": "Показывать риски и санкции отдельно"}},
        "steps_json": [{"key": "prepare_output", "type": "artifact"}],
        "capability_allowlist_json": [],
        "approval_policy_json": {"required_for": ["final_output"]},
        "output_schema_json": {"format": "summary", "feedback_history": [{"feedback": "Выделяй санкции"}]},
    }

    diff = build_agent_version_diff(first_version, second_version)

    assert diff["change_type"] == "changed"
    assert "goal" in diff["changed_fields"]
    assert "inputs_schema" in diff["changed_fields"]
    assert "output_schema" in diff["changed_fields"]
    assert diff["summary"].startswith("Изменено:")


def test_agent_learning_loop_summary_requires_human_activation():
    from services.agent_blueprint_workspace import build_agent_version_diff, build_learning_loop_summary

    previous_version = {
        "id": "ver1",
        "version_number": 1,
        "goal": "Напоминать клиентам о записи",
        "inputs_schema_json": {},
        "steps_json": [{"key": "draft", "type": "artifact"}],
        "capability_allowlist_json": ["communications.draft"],
        "approval_policy_json": {"required_for": ["final_output"]},
        "output_schema_json": {"format": "drafts"},
    }
    candidate_version = {
        "id": "ver2",
        "version_number": 2,
        "goal": "Напоминать клиентам о записи",
        "inputs_schema_json": {},
        "steps_json": [{"key": "draft", "type": "artifact"}],
        "capability_allowlist_json": ["communications.draft"],
        "approval_policy_json": {"required_for": ["final_output"], "last_feedback_requires_review": True},
        "output_schema_json": {
            "format": "drafts",
            "feedback_history": [{"feedback": "Не обещать скидку без пакета"}],
        },
    }
    feedback = {
        "run_id": "run1",
        "trigger_type": "manual_edit",
        "feedback": "Не обещать скидку без пакета",
    }

    diff = build_agent_version_diff(previous_version, candidate_version)
    learning = build_learning_loop_summary(feedback, previous_version, candidate_version, diff)

    assert learning["schema"] == "agent_learning_loop_v1"
    assert learning["mode"] == "versioned_review"
    assert learning["trigger_label"] == "Ручная правка текста"
    assert learning["activation_state"] == "candidate"
    assert learning["human_gate_required"] is True
    assert learning["candidate_version_id"] == "ver2"
    assert "activate" in learning["available_actions"]
    assert "rollback" in learning["available_actions"]
    assert learning["diff"]["change_type"] == "changed"


def test_agent_run_review_journal_is_human_readable():
    from services.agent_blueprint_workspace import _review_journal

    journal = _review_journal(
        {"id": "run1", "input_json": {"source": "smoke"}},
        [
            {
                "artifact_type": "agent_output_draft",
                "payload_json": {
                    "status": "generated",
                    "analysis_source": "gigachat",
                    "llm_analysis_used": True,
                    "provenance": ["contract.txt"],
                    "external_dispatch_performed": False,
                    "result": {
                        "title": "Разбор документа",
                        "facts": ["Оплата 10000"],
                        "risks": ["Штраф 10%"],
                        "next_questions": ["Кто подписывает?"],
                    },
                },
            }
        ],
        [{"id": "approval1", "status": "pending", "title": "Подтвердить результат", "approval_type": "final_output", "payload_json": {}}],
        {
            "agent_setup": {
                "workflow_description": "Проверить договор",
                "processing_rules": "Не придумывать факты",
                "output_format": "summary/risks",
            },
            "agent_sources": [{"name": "contract.txt", "source_type": "file"}],
        },
    )

    kinds = [item["kind"] for item in journal]
    assert "input" in kinds
    assert "output" in kinds
    assert "approval" in kinds
    output_entry = [item for item in journal if item["kind"] == "output"][0]
    detail_labels = [item["label"] for item in output_entry["details"]]
    assert "Источник анализа" in detail_labels
    assert "Использованные источники" in detail_labels
    assert "Внешняя отправка" in detail_labels
    assert output_entry["payload"]["external_dispatch_performed"] is False


def test_agent_review_tracks_sources_used_by_latest_run():
    from services.agent_blueprint_workspace import _used_source_summaries

    used_sources = _used_source_summaries(
        {
            "agent_sources": [
                {
                    "id": "source-contract",
                    "name": "contract.txt",
                    "source_type": "file",
                    "file_name": "contract.txt",
                    "extraction_state": "ready",
                    "content_length": 120,
                },
                {
                    "id": "source-unused",
                    "name": "unused.txt",
                    "source_type": "file",
                    "extraction_state": "ready",
                },
            ],
        },
        [
            {
                "artifact_type": "agent_extracted_context",
                "payload_json": {
                    "items": [{"source_name": "contract.txt", "summary": "Оплата 10000"}],
                },
            },
            {
                "artifact_type": "agent_output_draft",
                "payload_json": {
                    "provenance": ["contract.txt"],
                    "result": {"summary": ["Оплата 10000"]},
                },
            },
        ],
    )

    assert len(used_sources) == 1
    assert used_sources[0]["name"] == "contract.txt"
    assert used_sources[0]["source_type"] == "file"
    assert used_sources[0]["content_length"] == 120


def test_outreach_run_review_journal_explains_pipeline_and_queue_boundary():
    from services.agent_blueprint_workspace import _review_journal

    journal = _review_journal(
        {"id": "run1", "input_json": {"source": "yandex_maps", "city": "Moscow", "limit": 5}},
        [
            {
                "artifact_type": "lead_source_plan",
                "payload_json": {
                    "status": "hydrated",
                    "source": "prospectingleads",
                    "count": 1,
                    "filters": {"source": "yandex_maps", "city": "Moscow", "intent": "client_outreach", "limit": 5},
                    "status_counts": {"new": 1},
                    "items": [{"id": "lead1", "name": "Fresh Lead", "status": "new"}],
                },
            },
            {
                "artifact_type": "lead_shortlist",
                "payload_json": {
                    "status": "hydrated",
                    "source": "prospectingleads",
                    "source_artifact": "lead_source_plan",
                    "count": 1,
                    "items": [{"id": "lead1", "name": "Fresh Lead", "selected_channel": "email"}],
                },
            },
            {
                "artifact_type": "message_drafts",
                "payload_json": {
                    "status": "hydrated",
                    "source": "outreachmessagedrafts",
                    "count": 1,
                    "items": [{"id": "draft1", "lead_name": "Fresh Lead", "channel": "email", "status": "generated"}],
                },
            },
            {
                "artifact_type": "outreach_outcomes",
                "payload_json": {
                    "status": "hydrated",
                    "source": "outreachsendqueue",
                    "count": 1,
                    "queued_count": 1,
                    "dispatch_state": "queued_not_dispatched",
                    "external_dispatch_performed": False,
                    "operator_note": "Queue rows are LocalOS handoff records. External dispatcher is a separate contour.",
                    "items": [{"id": "queue1", "delivery_status": "queued"}],
                },
            },
        ],
        [],
        {
            "agent_setup": {
                "workflow_description": "Найти клиентов и подготовить сообщения",
                "manual_control": "Подтверждать shortlist и черновики",
            },
        },
    )

    kinds = [item["kind"] for item in journal]
    assert "sourcing" in kinds
    assert "shortlist" in kinds
    assert "drafts" in kinds
    assert "queue" in kinds
    sourcing_labels = [item["label"] for item in [entry for entry in journal if entry["kind"] == "sourcing"][0]["details"]]
    shortlist_labels = [item["label"] for item in [entry for entry in journal if entry["kind"] == "shortlist"][0]["details"]]
    draft_labels = [item["label"] for item in [entry for entry in journal if entry["kind"] == "drafts"][0]["details"]]
    queue_labels = [item["label"] for item in [entry for entry in journal if entry["kind"] == "queue"][0]["details"]]
    assert "Источник данных" in sourcing_labels
    assert "Найдено лидов" in sourcing_labels
    assert "Лидов в shortlist" in shortlist_labels
    assert "Черновиков" in draft_labels
    assert "В очереди" in queue_labels
    assert "Dispatch" in queue_labels
    assert "Внешняя отправка" in queue_labels


def test_outreach_send_batch_handler_queues_approved_drafts_without_external_dispatch(monkeypatch):
    from services import outreach_send_capability

    connection = FakeOutreachConnection(
        [
            {
                "id": "draft1",
                "lead_id": "lead1",
                "channel": "email",
                "email": "owner@example.com",
            }
        ]
    )
    monkeypatch.setattr(outreach_send_capability, "get_db_connection", lambda: connection)

    result = outreach_send_capability.handle_outreach_send_batch(
        {
            "tenant_id": "biz1",
            "actor": {"user_id": "user1"},
            "payload": {"draft_ids": ["draft1"], "daily_limit": 99},
        },
        {"user_id": "user1"},
    )

    output = result["result"]
    assert output["status"] == "queued_for_dispatch"
    assert output["queue_count"] == 1
    assert output["draft_ids"] == ["draft1"]
    assert output["daily_limit"] == 10
    assert output["dispatch_state"] == "queued_not_dispatched"
    assert output["dispatcher_required"] is True
    assert output["external_dispatch_performed"] is False
    assert "External dispatcher did not run" in output["operator_note"]
    assert connection.committed is True
    assert any(item["kind"] == "batch" and item["status"] == "approved" for item in connection.inserted)
    assert any(item["kind"] == "queue" and item["draft_id"] == "draft1" for item in connection.inserted)


def test_runner_stops_on_first_approval_step():
    from services.agent_blueprint_runner import AgentBlueprintRunner, default_supervised_outreach_version_payload

    cursor = FakeCursor()
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Outreach",
        "category": "outreach",
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "steps_json": default_supervised_outreach_version_payload()["steps"],
        "capability_allowlist_json": ["outreach.send_batch"],
    }

    result = AgentBlueprintRunner(cursor).start_run("ver1", {"limit": 30}, {"user_id": "user1"})

    assert result["success"] is True
    run = result["run"]
    assert run["status"] == "waiting_approval"
    assert [step["step_key"] for step in run["steps"]] == ["source_leads", "shortlist", "approve_shortlist"]
    assert run["approvals"][0]["approval_type"] == "shortlist"
    assert run["approvals"][0]["status"] == "pending"


def test_runner_blocks_send_capability_without_required_drafts_approval():
    from services.agent_blueprint_runner import AgentBlueprintRunner

    cursor = FakeCursor()
    cursor.tables["agent_runs"]["run1"] = {
        "id": "run1",
        "blueprint_id": "bp1",
        "blueprint_version_id": "ver1",
        "business_id": "biz1",
        "status": "running",
        "input_json": {"draft_ids": ["draft1"]},
        "output_json": {},
        "created_by_user_id": "user1",
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "steps_json": [],
        "capability_allowlist_json": ["outreach.send_batch"],
    }
    cursor.tables["agent_approvals"]["approval1"] = {
        "id": "approval1",
        "run_id": "run1",
        "step_id": "step1",
        "status": "approved",
        "approval_type": "shortlist",
        "title": "Shortlist approved",
        "payload_json": {},
        "requested_by_user_id": "user1",
    }
    orchestrator = CountingOrchestrator()
    step = {
        "key": "send_limited_batch",
        "type": "capability",
        "capability": "outreach.send_batch",
        "requires_approval": True,
        "required_approval_type": "drafts",
        "payload": {"daily_limit": 10},
    }

    completed = AgentBlueprintRunner(cursor, orchestrator=orchestrator)._execute_capability_step(
        cursor.tables["agent_runs"]["run1"],
        cursor.tables["agent_blueprint_versions"]["ver1"],
        step,
        5,
        {"user_id": "user1"},
    )

    assert completed is False
    assert orchestrator.calls == 0
    assert cursor.tables["agent_runs"]["run1"]["status"] == "failed"
    blocked_steps = [item for item in cursor.tables["agent_run_steps"].values() if item["status"] == "blocked"]
    assert blocked_steps
    assert blocked_steps[0]["output_json"]["required_approval_type"] == "drafts"


def test_runner_creates_drafts_after_shortlist_approval_and_queues_after_drafts_approval():
    from services.agent_blueprint_runner import AgentBlueprintRunner, default_supervised_outreach_version_payload

    cursor = FakeCursor()
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Outreach",
        "category": "outreach",
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "steps_json": default_supervised_outreach_version_payload()["steps"],
        "capability_allowlist_json": ["outreach.send_batch"],
    }
    cursor.tables["prospectingleads"]["lead1"] = {
        "id": "lead1",
        "business_id": "biz1",
        "name": "Lead One",
        "category": "beauty",
        "city": "Moscow",
        "email": "owner@example.com",
        "status": "shortlist_approved",
        "selected_channel": "",
        "pipeline_status": "in_progress",
    }
    orchestrator = CountingOrchestrator()
    runner = AgentBlueprintRunner(cursor, orchestrator=orchestrator)

    result = runner.start_run("ver1", {"lead_ids": ["lead1"], "limit": 5}, {"user_id": "user1"})
    run = result["run"]
    source_artifact = [item for item in run["artifacts"] if item["artifact_type"] == "lead_source_plan"][-1]
    assert source_artifact["payload_json"]["source"] == "prospectingleads"
    assert source_artifact["payload_json"]["status"] == "hydrated"
    assert source_artifact["payload_json"]["count"] == 1
    assert source_artifact["payload_json"]["filters"]["lead_ids"] == ["lead1"]
    shortlist_approval = run["approvals"][0]
    assert shortlist_approval["payload_json"]["artifact_type"] == "lead_shortlist"
    assert shortlist_approval["payload_json"]["count"] == 1

    after_shortlist = runner.approve(run["id"], shortlist_approval["id"], {"user_id": "user1"})
    draft_approval = after_shortlist["run"]["approvals"][-1]
    assert draft_approval["approval_type"] == "drafts"
    assert draft_approval["payload_json"]["artifact_type"] == "message_drafts"
    assert draft_approval["payload_json"]["count"] == 1
    draft_id = draft_approval["payload_json"]["items"][0]["id"]
    assert cursor.tables["outreachmessagedrafts"][draft_id]["status"] == "generated"
    assert cursor.tables["prospectingleads"]["lead1"]["status"] == "channel_selected"

    after_drafts = runner.approve(after_shortlist["run"]["id"], draft_approval["id"], {"user_id": "user1"})

    assert after_drafts["run"]["status"] == "completed"
    assert cursor.tables["outreachmessagedrafts"][draft_id]["status"] == "approved"
    assert cursor.tables["prospectingleads"]["lead1"]["status"] == "draft_ready"
    assert orchestrator.calls == 1
    assert orchestrator.last_envelope["capability"] == "outreach.send_batch"
    assert orchestrator.last_envelope["payload"]["draft_ids"] == [draft_id]
    assert orchestrator.last_envelope["payload"]["daily_limit"] == 10


def test_runner_creates_maton_delivery_preview_draft_without_dispatch(monkeypatch):
    from services import agent_blueprint_runner
    from services.agent_blueprint_runner import AgentBlueprintRunner

    def _fail_dispatch(*args, **kwargs):
        raise AssertionError("Maton dispatch must not run during safe preview")

    monkeypatch.setattr(agent_blueprint_runner, "dispatch_with_routing", _fail_dispatch)

    cursor = FakeCursor()
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Maton delivery",
        "category": "communications",
        "metadata_json": {
            "connector_action_handlers": {
                "telegram_delivery": {
                    "handler": "maton_external_account_bridge",
                    "external_account_id": "maton-account-1",
                },
            },
            "agent_binding_provider_routes": {
                "telegram_delivery": {
                    "provider": "maton",
                    "external_account_id": "maton-account-1",
                },
            },
        },
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "steps_json": [],
        "capability_allowlist_json": ["communications.send_offer"],
    }
    cursor.tables["agent_runs"]["run1"] = {
        "id": "run1",
        "blueprint_id": "bp1",
        "blueprint_version_id": "ver1",
        "business_id": "biz1",
        "status": "running",
        "input_json": {"preview_mode": True, "external_side_effects_allowed": False},
        "output_json": {},
        "created_by_user_id": "user1",
    }
    cursor.tables["agent_approvals"]["approval1"] = {
        "id": "approval1",
        "run_id": "run1",
        "step_id": "approval-step",
        "status": "approved",
        "approval_type": "communications_send",
        "title": "Approve send",
        "payload_json": {},
        "requested_by_user_id": "user1",
    }
    orchestrator = CountingOrchestrator()
    step = {
        "key": "send_offer",
        "type": "capability",
        "capability": "communications.send_offer",
        "requires_approval": True,
        "required_approval_type": "communications_send",
        "payload": {
            "message": "Пакетное предложение для клиента",
            "channel": "telegram",
            "telegram_target": "@client",
        },
    }

    completed = AgentBlueprintRunner(cursor, orchestrator=orchestrator)._execute_capability_step(
        cursor.tables["agent_runs"]["run1"],
        cursor.tables["agent_blueprint_versions"]["ver1"],
        step,
        1,
        {"user_id": "user1"},
    )

    assert completed is True
    assert orchestrator.calls == 0
    artifact = next(item for item in cursor.tables["agent_artifacts"].values() if item["artifact_type"] == "maton_delivery_request")
    payload = artifact["payload_json"]
    assert payload["schema"] == "localos_maton_delivery_request_v1"
    assert payload["status"] == "preview_draft_created"
    assert payload["delivery_state"] == "preview_draft_only"
    assert payload["external_dispatch_performed"] is False
    assert payload["external_account_id"] == "maton-account-1"
    delivery_step = next(
        step
        for step in cursor.tables["agent_run_steps"].values()
        if step["step_key"] == "send_offer"
    )
    runtime_contract = delivery_step["output_json"]["production_action_contract"]
    assert runtime_contract["schema"] == "localos_production_action_contract_v1"
    assert runtime_contract["capability"] == "communications.send_offer"
    assert runtime_contract["preflight"]["status"] == "passed"
    assert runtime_contract["approval_policy"]["required"] is True
    assert runtime_contract["approval_policy"]["approval_type"] == "communications_send"
    assert runtime_contract["ledger"]["required"] is True
    assert runtime_contract["recovery"]["state"] == "ready"
    assert runtime_contract["side_effects"]["external_dispatch_performed"] is False


def test_runner_sends_maton_delivery_only_after_approval_and_explicit_dispatch(monkeypatch):
    from services import agent_blueprint_runner
    from services.agent_blueprint_runner import AgentBlueprintRunner

    captured = {}

    def _fake_load_context(cursor, business_id):
        captured["business_id"] = business_id
        return {"id": business_id, "name": "Riderra", "maton_api_key": "key", "maton_bridge_enabled": True}

    def _fake_dispatch(ctx, text, *, preferred_provider="telegram", force_channel_id=None):
        captured["text"] = text
        captured["preferred_provider"] = preferred_provider
        captured["force_channel_id"] = force_channel_id
        return {
            "success": True,
            "selected_channel_id": "maton_bridge",
            "selected_provider": "maton",
            "attempts": [{"channel_id": "maton_bridge", "provider": "maton", "success": True}],
        }

    monkeypatch.setattr(agent_blueprint_runner, "load_business_channel_context", _fake_load_context)
    monkeypatch.setattr(agent_blueprint_runner, "dispatch_with_routing", _fake_dispatch)

    cursor = FakeCursor()
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Maton delivery",
        "category": "communications",
        "metadata_json": {
            "connector_action_handlers": {
                "telegram_delivery": {
                    "handler": "maton_external_account_bridge",
                    "external_account_id": "maton-account-1",
                },
            },
            "agent_binding_provider_routes": {
                "telegram_delivery": {
                    "provider": "maton",
                    "external_account_id": "maton-account-1",
                },
            },
        },
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "steps_json": [],
        "capability_allowlist_json": ["communications.send_offer"],
    }
    cursor.tables["agent_runs"]["run1"] = {
        "id": "run1",
        "blueprint_id": "bp1",
        "blueprint_version_id": "ver1",
        "business_id": "biz1",
        "status": "running",
        "input_json": {},
        "output_json": {},
        "created_by_user_id": "user1",
    }
    cursor.tables["agent_approvals"]["approval1"] = {
        "id": "approval1",
        "run_id": "run1",
        "step_id": "approval-step",
        "status": "approved",
        "approval_type": "communications_send",
        "title": "Approve send",
        "payload_json": {},
        "requested_by_user_id": "user1",
    }
    step = {
        "key": "send_offer",
        "type": "capability",
        "capability": "communications.send_offer",
        "requires_approval": True,
        "required_approval_type": "communications_send",
        "payload": {
            "message": "Пакетное предложение для клиента",
            "dispatch_mode": "send_after_approval",
            "external_side_effects_allowed": True,
            "telegram_target": "@client",
        },
    }

    completed = AgentBlueprintRunner(cursor, orchestrator=CountingOrchestrator())._execute_capability_step(
        cursor.tables["agent_runs"]["run1"],
        cursor.tables["agent_blueprint_versions"]["ver1"],
        step,
        1,
        {"user_id": "user1"},
    )

    assert completed is True
    assert captured["business_id"] == "biz1"
    assert captured["text"] == "Пакетное предложение для клиента"
    assert captured["preferred_provider"] == "maton"
    assert captured["force_channel_id"] == "maton_bridge"
    artifact = next(item for item in cursor.tables["agent_artifacts"].values() if item["artifact_type"] == "maton_delivery_request")
    payload = artifact["payload_json"]
    assert payload["status"] == "sent"
    assert payload["delivery_state"] == "sent"
    assert payload["external_dispatch_performed"] is True
    assert payload["router_result"]["selected_channel_id"] == "maton_bridge"


def test_default_runner_is_wired_to_real_google_sheets_and_finance_handlers():
    from services.agent_blueprint_runner import AgentBlueprintRunner

    runner = AgentBlueprintRunner(FakeCursor())

    assert "google_sheets.read_rows" in runner.orchestrator.handlers
    assert "finance.transaction.create" in runner.orchestrator.handlers


def test_runner_build_capability_payload_hydrates_google_sheets_binding_from_metadata():
    from services.agent_blueprint_runner import AgentBlueprintRunner

    cursor = FakeCursor()
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "metadata_json": {
            "required_integration_bindings": [
                {
                    "key": "google_sheets_read",
                    "provider": "google_sheets",
                    "capability": "google_sheets.read_rows",
                    "required_config": ["spreadsheet_id", "sheet_name"],
                }
            ],
            "custom_process": {
                "google_sheets_read": {
                    "integration_id": "integration-1",
                    "spreadsheet_id": "spreadsheet-1",
                    "sheet_name": "Trips",
                }
            },
        },
    }
    run = {
        "id": "run1",
        "blueprint_id": "bp1",
        "input_json": {},
    }
    step = {
        "key": "read_google_sheets",
        "type": "capability",
        "capability": "google_sheets.read_rows",
        "payload": {
            "integration_binding": "google_sheets_read",
            "limit": 100,
        },
    }

    payload = AgentBlueprintRunner(cursor)._build_capability_payload(run, step)

    assert payload["integration_id"] == "integration-1"
    assert payload["spreadsheet_id"] == "spreadsheet-1"
    assert payload["sheet_name"] == "Trips"


def test_runner_passes_compiled_step_rows_to_next_capability_without_runtime_ai():
    from services.agent_blueprint_runner import AgentBlueprintRunner

    class FinanceCaptureOrchestrator:
        def __init__(self):
            self.calls = []

        def execute(self, envelope, user_data, *, allow_execute_when_approved=False):
            self.calls.append(envelope)
            return {
                "success": True,
                "status": "completed",
                "result": {
                    "status": "finance_transaction_request_created",
                    "proposal_count": len(envelope["payload"].get("rows") or []),
                    "rows": envelope["payload"].get("rows") or [],
                    "localos_write_performed": False,
                },
            }

    cursor = FakeCursor()
    cursor.tables["agent_runs"]["run1"] = {
        "id": "run1",
        "blueprint_id": "bp1",
        "blueprint_version_id": "ver1",
        "business_id": "biz1",
        "status": "running",
        "input_json": {
            "preview_mode": True,
            "external_side_effects_allowed": False,
            "goal": "Проверить таблицу и подготовить действие",
            "preview_context": {
                "understood_task": "Проверить Google Sheets и показать append request",
                "data_sources": ["Google Sheets", "Telegram"],
                "manual_control": "Подтвердить перед записью в таблицу.",
            },
        },
        "output_json": {},
        "created_by_user_id": "user1",
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "steps_json": [],
        "capability_allowlist_json": ["google_sheets.read_rows", "finance.transaction.create"],
    }
    cursor.tables["agent_run_steps"]["read-step"] = {
        "id": "read-step",
        "run_id": "run1",
        "step_index": 0,
        "step_key": "read_google_sheets",
        "step_type": "capability",
        "status": "completed",
        "input_json": {},
        "output_json": {
            "capability": "google_sheets.read_rows",
            "orchestrator": {
                "success": True,
                "result": {
                    "status": "read_completed",
                    "rows": [{"row_number": 2, "date": "2026-06-09", "type": "revenue", "amount": "12000"}],
                },
            },
        },
    }
    cursor.tables["agent_approvals"]["approval1"] = {
        "id": "approval1",
        "run_id": "run1",
        "step_id": "approval-step",
        "status": "approved",
        "approval_type": "finance_transaction_import",
        "title": "Approved",
        "payload_json": {},
        "requested_by_user_id": "user1",
    }
    orchestrator = FinanceCaptureOrchestrator()
    step = {
        "key": "request_localos_finance",
        "type": "capability",
        "capability": "finance.transaction.create",
        "requires_approval": True,
        "required_approval_type": "finance_transaction_import",
        "payload": {
            "input_mappings": [
                {
                    "target": "rows",
                    "from_step": "read_google_sheets",
                    "path": "orchestrator.result.rows",
                    "required": True,
                }
            ],
            "rows_from_step": "read_google_sheets",
            "localos_write_performed": False,
        },
    }

    completed = AgentBlueprintRunner(cursor, orchestrator=orchestrator)._execute_capability_step(
        cursor.tables["agent_runs"]["run1"],
        cursor.tables["agent_blueprint_versions"]["ver1"],
        step,
        3,
        {"user_id": "user1"},
    )

    assert completed is True
    assert len(orchestrator.calls) == 1
    assert orchestrator.calls[0]["capability"] == "finance.transaction.create"
    assert orchestrator.calls[0]["payload"]["rows"][0]["amount"] == "12000"
    assert "input_mappings" not in orchestrator.calls[0]["payload"]
    assert "rows_from_step" not in orchestrator.calls[0]["payload"]
    assert "gigachat" not in json.dumps(orchestrator.calls[0], ensure_ascii=False).lower()
    finance_step = next(
        step
        for step in cursor.tables["agent_run_steps"].values()
        if step["step_key"] == "request_localos_finance"
    )
    assert finance_step["output_json"]["approved_executor"]["localos_writes_performed"] is False
    runtime_contract = finance_step["output_json"]["production_action_contract"]
    assert runtime_contract["schema"] == "localos_production_action_contract_v1"
    assert runtime_contract["capability"] == "finance.transaction.create"
    assert runtime_contract["preflight"]["status"] == "passed"
    assert runtime_contract["approval_policy"]["required"] is True
    assert runtime_contract["approval_policy"]["approval_type"] == "finance_transaction_import"
    assert runtime_contract["ledger"]["required"] is True
    assert runtime_contract["limits"]["subscription_checked"] is True
    assert runtime_contract["recovery"]["idempotency_key"] == "agent-run:run1:request_localos_finance"
    assert runtime_contract["side_effects"]["localos_write_performed"] is False
    assert cursor.tables["finance_entries"] == {}


def test_runner_applies_finance_requests_only_after_explicit_apply():
    from services.agent_blueprint_runner import AgentBlueprintRunner

    class FinanceProposalOrchestrator:
        def execute(self, envelope, user_data, *, allow_execute_when_approved=False):
            return {
                "success": True,
                "status": "completed",
                "action_id": "action-finance-1",
                "result": {
                    "status": "finance_transaction_request_created",
                    "request_id": "finance-request-1",
                    "source": "google_sheets",
                    "normalized_mapping": {"amount": "amount"},
                    "proposal_count": 1,
                    "finance_entry_proposals": [
                        {
                            "record_type": "entry",
                            "date": "2026-06-09",
                            "type": "revenue",
                            "category": "sales",
                            "amount": 12000,
                            "comment": "Оплата по таблице",
                            "row_number": 1,
                            "duplicate_key": "finance-dup-apply-1",
                        }
                    ],
                    "rows_requiring_review": [],
                    "errors": [],
                    "approval_state": "pending_human",
                    "apply_state": "not_applied",
                    "localos_write_performed": False,
                },
            }

        def get_action_support_package(self, action_id, user_data, limit=100, full=False):
            return {
                "action_id": action_id,
                "action": {"status": "completed", "capability": "finance.transaction.create"},
                "timeline": {"count": 0, "events": []},
                "billing_summary": {},
                "delivery_stats": {},
            }

    cursor = FakeCursor()
    cursor.tables["agent_runs"]["run1"] = {
        "id": "run1",
        "blueprint_id": "bp1",
        "blueprint_version_id": "ver1",
        "business_id": "biz1",
        "status": "running",
        "input_json": {"preview_mode": False},
        "output_json": {},
        "created_by_user_id": "user1",
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "steps_json": [],
        "capability_allowlist_json": ["finance.transaction.create"],
    }
    cursor.tables["agent_approvals"]["approval1"] = {
        "id": "approval1",
        "run_id": "run1",
        "step_id": "approval-step",
        "status": "approved",
        "approval_type": "finance_transaction_import",
        "title": "Approved",
        "payload_json": {},
        "requested_by_user_id": "user1",
    }
    step = {
        "key": "request_localos_finance",
        "type": "capability",
        "capability": "finance.transaction.create",
        "requires_approval": True,
        "required_approval_type": "finance_transaction_import",
        "payload": {"rows": [{"amount": 12000}]},
    }
    runner = AgentBlueprintRunner(cursor, orchestrator=FinanceProposalOrchestrator())

    completed = runner._execute_capability_step(
        cursor.tables["agent_runs"]["run1"],
        cursor.tables["agent_blueprint_versions"]["ver1"],
        step,
        1,
        {"user_id": "user1"},
    )
    run_before_apply = runner.load_run("run1", {"user_id": "user1"})
    finance_request = run_before_apply["observability"]["domain_requests"]["items"][0]

    assert completed is True
    assert cursor.tables["finance_entries"] == {}
    assert finance_request["kind"] == "finance_transaction_request"
    assert finance_request["approval_state"] == "approved"
    assert finance_request["apply_state"] == "apply_ready"
    assert finance_request["can_apply"] is True

    result = runner.apply_finance_requests("run1", {"user_id": "user1"})

    assert result["success"] is True
    assert result["items"][0]["apply_state"] == "applied"
    assert len(cursor.tables["finance_entries"]) == 1
    assert next(iter(cursor.tables["finance_entries"].values()))["duplicate_key"] == "finance-dup-apply-1"
    assert cursor.ledger_entries[0]["capability"] == "finance.transaction.create"
    finance_step = next(
        step
        for step in cursor.tables["agent_run_steps"].values()
        if step["step_key"] == "request_localos_finance"
    )
    assert finance_step["output_json"]["approved_executor"]["localos_writes_performed"] is True


def test_runner_builds_finance_outcome_artifact_from_compiled_step_outputs():
    from services.agent_blueprint_runner import AgentBlueprintRunner

    cursor = FakeCursor()
    cursor.tables["agent_runs"]["run1"] = {
        "id": "run1",
        "blueprint_id": "bp1",
        "blueprint_version_id": "ver1",
        "business_id": "biz1",
        "status": "completed",
        "input_json": {
            "preview_mode": True,
            "external_side_effects_allowed": False,
            "goal": "Проверить таблицу и подготовить действие",
            "preview_context": {
                "understood_task": "Проверить Google Sheets и показать append request",
                "data_sources": ["Google Sheets", "Telegram"],
                "manual_control": "Подтвердить перед записью в таблицу.",
            },
        },
        "output_json": {},
        "created_by_user_id": "user1",
    }
    cursor.tables["agent_run_steps"]["read-step"] = {
        "id": "read-step",
        "run_id": "run1",
        "step_index": 0,
        "step_key": "read_google_sheets",
        "step_type": "capability",
        "status": "completed",
        "input_json": {},
        "output_json": {
            "capability": "google_sheets.read_rows",
            "orchestrator": {
                "success": True,
                "result": {
                    "status": "read_completed",
                    "count": 2,
                    "rows": [
                        {"row_number": 2, "date": "2026-06-09", "type": "revenue", "amount": "12000"},
                        {"row_number": 3, "date": "2026-06-10", "type": "expense", "amount": "2500"},
                    ],
                },
            },
        },
    }
    cursor.tables["agent_run_steps"]["finance-step"] = {
        "id": "finance-step",
        "run_id": "run1",
        "step_index": 3,
        "step_key": "request_localos_finance",
        "step_type": "capability",
        "status": "completed",
        "input_json": {},
        "output_json": {
            "capability": "finance.transaction.create",
            "orchestrator": {
                "success": True,
                "action_id": "action-finance-1",
                "result": {
                    "status": "finance_transaction_request_created",
                    "request_id": "finance-request-1",
                    "proposal_count": 2,
                    "review_count": 1,
                    "error_count": 0,
                    "finance_entry_proposals": [{"row_number": 2}, {"row_number": 3}],
                    "rows_requiring_review": [{"row_number": 3, "review_reasons": ["category_missing_or_default"]}],
                    "errors": [],
                    "approval_state": "pending_human",
                    "apply_state": "not_applied",
                    "localos_write_performed": False,
                },
            },
            "approved_executor": {
                "localos_writes_performed": 2,
                "items": [
                    {
                        "kind": "finance_transaction_request",
                        "batch_id": "batch-1",
                        "rows_imported": 2,
                        "rows_failed": 0,
                        "rows_skipped": 0,
                    }
                ],
            },
        },
    }
    runner = AgentBlueprintRunner(cursor, orchestrator=CountingOrchestrator())

    payload = runner._build_artifact_payload(
        cursor.tables["agent_runs"]["run1"],
        {
            "key": "record_localos_finance_outcome",
            "artifact_type": "localos_finance_outcome",
            "payload": {
                "source_step": "read_google_sheets",
                "request_step": "request_localos_finance",
                "journal": ["rows_read", "proposals", "rows_requiring_review", "errors"],
            },
        },
    )

    assert payload["rows_read"] == 2
    assert payload["proposal_count"] == 2
    assert payload["review_count"] == 1
    assert payload["error_count"] == 0
    assert payload["rows_imported"] == 2
    assert payload["apply_state"] == "applied"
    assert payload["localos_write_performed"] is True
    assert payload["recovery"]["idempotency"] == "rerun uses finance duplicate_key checks before inserting rows"
    assert "gigachat" not in json.dumps(payload, ensure_ascii=False).lower()


def test_finance_outcome_journal_is_human_readable():
    from services.agent_blueprint_workspace import _artifact_journal_entry

    entry = _artifact_journal_entry(
        "localos_finance_outcome",
        {
            "status": "applied",
            "rows_read": 2,
            "proposal_count": 2,
            "review_count": 1,
            "error_count": 0,
            "rows_imported": 2,
            "apply_state": "applied",
            "localos_write_performed": True,
        },
    )

    assert entry["kind"] == "finance_outcome"
    assert entry["title"] == "Итог записи в финансы"
    assert "Записано 2 финансовых строк" in entry["summary"]
    labels = [item["label"] for item in entry["details"]]
    assert "Прочитано строк" in labels
    assert "Требует проверки" in labels
    assert "Запись в LocalOS" in labels


def test_runner_blocks_custom_agent_start_when_required_external_binding_missing():
    from services.agent_blueprint_draft_builder import compile_agent_blueprint
    from services.agent_blueprint_runner import AgentBlueprintRunner

    draft = compile_agent_blueprint(
        "Каждый вечер проверяй Google Sheets, бери новые оплаты и создавай транзакции в финансах LocalOS"
    )
    cursor = FakeCursor()
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Finance import",
        "category": "custom",
        "metadata_json": draft["metadata"],
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "steps_json": draft["version_payload"]["steps"],
        "capability_allowlist_json": draft["version_payload"]["capability_allowlist"],
    }

    result = AgentBlueprintRunner(cursor, orchestrator=CountingOrchestrator()).start_run("ver1", {}, {"user_id": "user1"})

    assert result["success"] is False
    assert result["code"] == "AGENT_INTEGRATIONS_REQUIRED"
    assert result["preflight"]["status"] == "blocked"
    assert result["preflight"]["missing_count"] == 1
    assert result["preflight"]["missing"][0]["provider"] == "google_sheets"
    assert result["preflight"]["missing"][0]["missing_config"] == ["spreadsheet_id", "sheet_name"]
    assert cursor.tables["agent_runs"] == {}


def test_runner_creates_openclaw_preview_observations_from_route_contract():
    from services.agent_blueprint_runner import AgentBlueprintRunner

    cursor = FakeCursor()
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Sheets preview",
        "category": "custom",
        "metadata_json": {
            "required_integration_bindings": [
                {
                    "key": "google_sheets_read",
                    "provider": "google_sheets",
                    "capability": "google_sheets.read_rows",
                    "required_config": ["spreadsheet_id", "sheet_name"],
                }
            ],
            "agent_binding_provider_routes": {
                "google_sheets_read": {
                    "route_provider": "openclaw",
                    "status": "active",
                    "integration_id": "openclaw_boundary",
                }
            },
            "agent_binding_integrations": {
                "google_sheets_read": {
                    "route_provider": "openclaw",
                    "status": "active",
                    "integration_id": "openclaw_boundary",
                }
            },
        },
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "steps_json": [],
        "capability_allowlist_json": ["google_sheets.read_rows"],
    }
    orchestrator = CountingOrchestrator()
    result = AgentBlueprintRunner(cursor, orchestrator=orchestrator).start_run(
        "ver1",
        {
            "schema": "localos_agent_preview_input_v1",
            "preview_mode": True,
            "external_side_effects_allowed": False,
            "connector_action_handlers": [
                {
                    "binding_key": "google_sheets_read",
                    "route_provider": "openclaw",
                    "handler": "openclaw_policy_boundary",
                    "credential_source": "openclaw_m2m",
                    "preflight_resolution": "provider_route_openclaw_boundary",
                    "execution_boundary": "localos_policy_envelope",
                    "approval_required": True,
                    "audit_required": True,
                    "external_side_effects_allowed_in_preview": False,
                }
            ],
            "openclaw_preview_routes": [
                {
                    "binding_key": "google_sheets_read",
                    "capability": "google_sheets.read_rows",
                    "provider": "openclaw",
                    "provider_action_ref": "openclaw.google_sheets.read_rows",
                    "external_side_effects_allowed_in_preview": False,
                }
            ],
            "openclaw_action_plan": [
                {
                    "step_key": "google_sheets_read",
                    "capability": "google_sheets.read_rows",
                    "provider": "openclaw",
                    "provider_action_ref": "openclaw.google_sheets.read_rows",
                }
            ],
        },
        {"user_id": "user1"},
    )

    artifact = next(
        item
        for item in cursor.tables["agent_artifacts"].values()
        if item["artifact_type"] == "openclaw_preview_observations"
    )

    assert result["success"] is True
    assert result["run"]["status"] == "completed"
    assert orchestrator.calls == 0
    assert artifact["payload_json"]["schema"] == "localos_openclaw_preview_observations_v1"
    assert artifact["payload_json"]["external_actions_executed"] is False
    assert artifact["payload_json"]["handler_contracts"][0]["handler"] == "openclaw_policy_boundary"
    assert artifact["payload_json"]["observations"][0]["external_action_executed"] is False
    assert result["run"]["observability"]["preview_summary"]["openclaw_action_count"] == 1


def test_agent_integration_preflight_allows_inline_rows_and_native_finance():
    from services.agent_blueprint_draft_builder import compile_agent_blueprint
    from services.agent_integration_preflight import build_agent_integration_preflight

    draft = compile_agent_blueprint(
        "Каждый вечер проверяй Google Sheets, бери новые оплаты и создавай транзакции в финансах LocalOS"
    )
    cursor = FakeCursor()

    preflight = build_agent_integration_preflight(
        cursor,
        business_id="biz1",
        metadata=draft["metadata"],
        input_payload={"rows": [{"amount": "12000", "type": "revenue"}]},
    )

    assert preflight["status"] == "ready"
    assert preflight["missing_count"] == 0
    by_provider = {item["provider"]: item for item in preflight["items"]}
    assert by_provider["google_sheets"]["resolution"] == "input_payload"
    assert by_provider["localos_finance"]["resolution"] == "native_localos"


def test_runner_builds_shortlist_from_sourced_unprocessed_leads():
    from services.agent_blueprint_runner import AgentBlueprintRunner, default_supervised_outreach_version_payload

    cursor = FakeCursor()
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "name": "Outreach",
        "category": "outreach",
    }
    cursor.tables["agent_blueprint_versions"]["ver1"] = {
        "id": "ver1",
        "blueprint_id": "bp1",
        "steps_json": default_supervised_outreach_version_payload()["steps"],
        "capability_allowlist_json": ["outreach.send_batch"],
    }
    cursor.tables["prospectingleads"]["lead1"] = {
        "id": "lead1",
        "business_id": "biz1",
        "name": "Fresh Lead",
        "category": "beauty",
        "city": "Moscow",
        "email": "fresh@example.com",
        "source": "yandex_maps",
        "status": "new",
        "selected_channel": "",
        "pipeline_status": "unprocessed",
    }
    runner = AgentBlueprintRunner(cursor, orchestrator=CountingOrchestrator())

    result = runner.start_run(
        "ver1",
        {"source": "yandex_maps", "city": "Moscow", "intent": "client_outreach", "limit": 5},
        {"user_id": "user1"},
    )
    run = result["run"]
    shortlist_artifact = [item for item in run["artifacts"] if item["artifact_type"] == "lead_shortlist"][-1]
    shortlist_payload = shortlist_artifact["payload_json"]

    assert shortlist_payload["source"] == "prospectingleads"
    assert shortlist_payload["source_artifact"] == "lead_source_plan"
    assert shortlist_payload["count"] == 1
    assert shortlist_payload["items"][0]["id"] == "lead1"

    shortlist_approval = run["approvals"][0]
    after_shortlist = runner.approve(run["id"], shortlist_approval["id"], {"user_id": "user1"})

    assert cursor.tables["prospectingleads"]["lead1"]["status"] == "channel_selected"
    assert after_shortlist["run"]["approvals"][-1]["approval_type"] == "drafts"


def test_risk_policy_requires_human_for_dangerous_capabilities():
    from core.action_policy import evaluate_risk_policy

    for capability in ("outreach.send_batch", "content.publish", "billing.payment", "records.delete"):
        risk = evaluate_risk_policy(capability, {}, {})
        assert risk["requires_human"] is True
        assert risk["reason"] == "dangerous capability requires review"


def test_runner_load_run_includes_observability_envelope_for_openclaw_actions():
    from services.agent_blueprint_runner import AgentBlueprintRunner

    class FakeObservabilityOrchestrator:
        def execute(self, envelope, user_data, *, allow_execute_when_approved=False):
            return {"success": True, "status": "completed", "action_id": "action-1", "result": {}, "billing": {}}

        def get_action_support_package(self, action_id, user_data, **kwargs):
            return {
                "success": True,
                "action_id": action_id,
                "tenant_id": "biz1",
                "capability": "communications.send_reminder",
                "trace_id": "agent-run:run1:send",
                "status": "completed",
                "delivery_stats": {
                    "attempts_total": 1,
                    "attempts_success": 1,
                    "attempts_failed": 0,
                },
                "billing": {
                    "summary": {
                        "reserved_tokens": 2000,
                        "settled_tokens": 42,
                        "released_tokens": 1958,
                        "total_cost": 0.012,
                    },
                    "entries": [{"entry_type": "settle", "tokens_out": 42, "cost": 0.012}],
                },
                "timeline": {
                    "events": [
                        {"source": "action_transition", "event_type": "completed", "status": "completed", "details": {}},
                    ],
                },
            }

    cursor = FakeCursor()
    preview_openclaw_action_plan = [
        {
            "step_key": "read_sheet",
            "title": "Read Google Sheets rows",
            "capability": "sheets.read_rows",
            "provider": "openclaw",
            "provider_action_ref": "openclaw.google_sheets.read_rows",
            "provider_policy": "localos_envelope",
            "risk_class": "read",
            "approval_class": "none",
            "requires_approval": False,
        },
        {
            "step_key": "append_sheet",
            "title": "Prepare approved append request",
            "capability": "sheets.append_row_request",
            "provider": "openclaw",
            "provider_action_ref": "openclaw.google_sheets.append_row",
            "provider_policy": "localos_envelope",
            "risk_class": "external_write",
            "approval_class": "human_before_write",
            "requires_approval": True,
        },
    ]
    cursor.tables["agent_runs"]["run1"] = {
        "id": "run1",
        "blueprint_id": "bp1",
        "blueprint_version_id": "ver1",
        "business_id": "biz1",
        "status": "completed",
        "input_json": {
            "preview_mode": True,
            "external_side_effects_allowed": False,
            "goal": "Проверить таблицу и подготовить действие",
            "openclaw_action_plan": preview_openclaw_action_plan,
            "policy_envelope": {
                "execution_boundary": "openclaw_action_orchestrator",
                "external_side_effects_allowed_in_preview": False,
                "approval_owner": "LocalOS",
                "billing_owner": "LocalOS",
                "audit_owner": "LocalOS",
            },
            "preview_context": {
                "understood_task": "Проверить Google Sheets и показать append request",
                "data_sources": ["Google Sheets", "Telegram"],
                "manual_control": "Подтвердить перед записью в таблицу.",
            },
        },
        "output_json": {},
        "created_by_user_id": "user1",
    }
    cursor.tables["agent_run_steps"]["step1"] = {
        "id": "step1",
        "run_id": "run1",
        "step_index": 0,
        "step_key": "send",
        "step_type": "capability",
        "status": "completed",
        "input_json": {},
        "output_json": {
            "capability": "sheets.append_row_request",
            "orchestrator": {
                "action_id": "action-1",
                "result": {"request_id": "sheet-request-1"},
            },
        },
    }
    cursor.tables["agent_sheet_operation_requests"]["sheet-request-1"] = {
        "id": "sheet-request-1",
        "business_id": "biz1",
        "action_id": "action-1",
        "integration_id": "integration-1",
        "spreadsheet_id": "spreadsheet-1",
        "sheet_name": "Leads",
        "operation": "append_row",
        "status": "request_created",
        "approval_state": "pending_human",
        "apply_state": "not_applied",
        "row_values_json": ["Anna", "Telegram"],
        "mapping_json": {"name": "telegram_username"},
        "limits_json": {"daily_cap": 10},
        "provider_write_performed": False,
        "created_at": "2026-06-09T10:00:00Z",
    }

    run = AgentBlueprintRunner(cursor, FakeObservabilityOrchestrator()).load_run("run1", {"user_id": "user1"})
    observability = run["observability"]

    assert observability["schema"] == "agent_run_observability_v1"
    assert observability["preview_summary"]["schema"] == "localos_agent_preview_summary_v1"
    assert observability["preview_summary"]["is_preview"] is True
    assert observability["preview_summary"]["safe_preview"] is True
    assert observability["preview_summary"]["understood_task"] == "Проверить Google Sheets и показать append request"
    assert observability["preview_summary"]["external_actions_performed"] is False
    assert observability["preview_summary"]["openclaw_action_count"] == 2
    assert observability["preview_summary"]["openclaw_action_plan"][0]["provider_action_ref"] == "openclaw.google_sheets.read_rows"
    assert observability["preview_summary"]["openclaw_action_plan"][1]["requires_approval"] is True
    assert observability["preview_summary"]["policy_envelope"]["execution_boundary"] == "openclaw_action_orchestrator"
    assert observability["preview_summary"]["policy_envelope"]["approval_owner"] == "LocalOS"
    assert observability["preview_summary"]["approval_gate"]["waiting_actions_count"] == 1
    assert observability["preview_summary"]["approval_gate"]["external_actions_performed"] is False
    assert "approval" in observability["preview_summary"]["activation_hint"].lower()
    assert observability["preview_summary"]["next_step"] == "review_approvals"
    assert observability["preview_summary"]["next_step_label"] == "Проверить approval"
    assert observability["action_ids"] == ["action-1"]
    assert observability["action_ledger"]["items"][0]["billing_summary"]["settled_tokens"] == 42
    assert observability["billing_ledger"]["summary"]["reserved_tokens"] == 2000
    assert observability["billing_ledger"]["summary"]["settled_tokens"] == 42
    assert observability["billing_ledger"]["summary"]["released_tokens"] == 1958
    assert observability["billing_ledger"]["actions"][0]["action_id"] == "action-1"
    assert observability["billing_ledger"]["actions"][0]["capability"] == "communications.send_reminder"
    assert observability["billing_ledger"]["actions"][0]["entry_count"] == 1
    assert observability["billing_ledger"]["entries"][0]["entry_type"] == "settle"
    assert observability["billing_ledger"]["entries"][0]["tokens_out"] == 42
    assert observability["unified_billing_ledger"]["schema"] == "localos_agent_run_unified_billing_ledger_v1"
    assert observability["unified_billing_ledger"]["summary"]["actual_tokens"] == 42
    assert observability["unified_billing_ledger"]["summary"]["external_action_count"] == 1
    assert observability["unified_billing_ledger"]["items"][0]["key"] == "preview_run"
    assert observability["unified_billing_ledger"]["items"][0]["actual_tokens"] == 0
    assert observability["unified_billing_ledger"]["items"][1]["key"] == "external_action"
    assert observability["unified_billing_ledger"]["items"][1]["actual_tokens"] == 42
    assert observability["delivery_status"]["state"] == "delivered"
    assert observability["cost_tokens"]["total_cost"] == 0.012
    assert observability["domain_requests"]["count"] == 1
    assert observability["domain_requests"]["pending"] == 1
    domain_request = observability["domain_requests"]["items"][0]
    assert domain_request["kind"] == "sheet_operation_request"
    assert domain_request["approval_state"] == "pending_human"
    assert domain_request["apply_state"] == "not_applied"
    assert domain_request["provider_handoff"]["provider_executor"] == "manual_controlled_google_sheets_append"
    assert domain_request["provider_handoff"]["spreadsheet_id"] == "spreadsheet-1"
    assert domain_request["provider_write_performed"] is False
    assert "External spreadsheet write requires human approval" in domain_request["why_waiting"]
    assert observability["support_export"]["endpoint"] == "/api/agent-runs/run1/support-export"


def test_runner_observability_exposes_source_to_result_chain():
    from services.agent_blueprint_runner import AgentBlueprintRunner

    cursor = FakeCursor()
    cursor.tables["agent_blueprints"]["bp1"] = {
        "id": "bp1",
        "business_id": "biz1",
        "metadata_json": {
            "required_integration_bindings": [
                {
                    "key": "google_sheets_read",
                    "provider": "google_sheets",
                    "capability": "google_sheets.read_rows",
                    "required_config": ["spreadsheet_id", "sheet_name"],
                }
            ],
            "custom_process": {
                "google_sheets_read": {
                    "integration_id": "integration-1",
                    "spreadsheet_id": "spreadsheet-1",
                    "sheet_name": "Trips",
                }
            },
        },
    }
    cursor.tables["agent_runs"]["run1"] = {
        "id": "run1",
        "blueprint_id": "bp1",
        "blueprint_version_id": "ver1",
        "business_id": "biz1",
        "status": "completed",
        "input_json": {},
        "output_json": {},
        "created_by_user_id": "user1",
    }
    cursor.tables["agent_run_steps"]["read-step"] = {
        "id": "read-step",
        "run_id": "run1",
        "step_index": 0,
        "step_key": "read_google_sheets",
        "step_type": "capability",
        "status": "completed",
        "input_json": {},
        "output_json": {
            "capability": "google_sheets.read_rows",
            "orchestrator": {
                "success": True,
                "result": {
                    "status": "read_completed",
                    "provider_read_performed": True,
                    "rows": [{"date": "2026-04-20", "route": "Airport -> Center"}],
                },
            },
        },
    }
    cursor.tables["agent_artifacts"]["artifact1"] = {
        "id": "artifact1",
        "run_id": "run1",
        "step_id": "output-step",
        "artifact_type": "telegram_post_draft",
        "title": "Черновик поста",
        "payload_json": {
            "items_used": 1,
            "result": {
                "title": "Черновик сообщения",
                "draft_text": "Поездка на 20 апреля\n\nПодготовлен черновик сообщения:\nмаршрут: Airport -> Center.",
            },
        },
        "created_at": "2026-06-28 12:00:00",
    }
    cursor.tables["agent_integrations"]["integration-1"] = {
        "id": "integration-1",
        "business_id": "biz1",
        "provider": "google_sheets",
        "status": "active",
        "auth_ref": "ext-1",
        "config_json": {"spreadsheet_id": "spreadsheet-1", "sheet_name": "Trips"},
    }

    observability = AgentBlueprintRunner(cursor).load_run("run1", {"user_id": "user1"})["observability"]["source_result_chain"]

    assert observability["source_step_present"] is True
    assert observability["provider_connected"] is True
    assert observability["provider_read_attempted"] is True
    assert observability["rows_returned_count"] == 1
    assert observability["rows_used_for_output_count"] == 1
    assert observability["result_generated"] is True
    assert observability["blocker_code"] == ""


class FakeCursor:
    def __init__(self):
        self.tables = {
            "agent_blueprints": {},
            "agent_blueprint_versions": {},
            "agent_runs": {},
            "agent_run_steps": {},
            "agent_artifacts": {},
            "agent_approvals": {},
            "prospectingleads": {},
            "outreachmessagedrafts": {},
            "agent_sheet_operation_requests": {},
            "agent_communication_requests": {},
            "reviewreplydrafts": {},
            "agent_review_publish_requests": {},
            "agent_service_optimization_requests": {},
            "finance_import_batches": {},
            "finance_entries": {},
            "agent_action_ledger": {},
            "agent_integrations": {},
        }
        self.last_result = None
        self.last_results = []
        self.ledger_entries = []

    def execute(self, query, params=None):
        normalized_query = " ".join(query.split()).lower()
        params = params or ()
        if normalized_query.startswith("select to_regclass"):
            table_name = params[0]
            self.last_result = {"table_name": table_name if table_name in self.tables else None}
            return None
        if normalized_query.startswith("create table if not exists"):
            return None
        if normalized_query.startswith("create unique index if not exists"):
            return None
        if normalized_query.startswith("create index if not exists"):
            return None
        if "from agent_sheet_operation_requests" in normalized_query:
            business_id = params[0]
            request_ids = set(params[1])
            action_ids = set(params[2])
            self.last_results = [
                row
                for row in self.tables["agent_sheet_operation_requests"].values()
                if row.get("business_id") == business_id and (row.get("id") in request_ids or row.get("action_id") in action_ids)
            ]
            return None
        if "from agent_communication_requests" in normalized_query:
            business_id = params[0]
            request_ids = set(params[1])
            action_ids = set(params[2])
            self.last_results = [
                row
                for row in self.tables["agent_communication_requests"].values()
                if row.get("business_id") == business_id and (row.get("id") in request_ids or row.get("action_id") in action_ids)
            ]
            return None
        if "from reviewreplydrafts" in normalized_query:
            business_id = params[0]
            draft_ids = set(params[1])
            review_ids = set(params[2])
            self.last_results = [
                row
                for row in self.tables["reviewreplydrafts"].values()
                if row.get("business_id") == business_id and (row.get("id") in draft_ids or row.get("review_id") in review_ids)
            ]
            return None
        if "from agent_review_publish_requests" in normalized_query:
            business_id = params[0]
            draft_id = params[1]
            review_id = params[2]
            self.last_results = [
                row
                for row in self.tables["agent_review_publish_requests"].values()
                if row.get("business_id") == business_id and (row.get("draft_id") == draft_id or row.get("review_id") == review_id)
            ]
            return None
        if "from agent_service_optimization_requests" in normalized_query:
            business_id = params[0]
            request_ids = set(params[1])
            action_ids = set(params[2])
            self.last_results = [
                row
                for row in self.tables["agent_service_optimization_requests"].values()
                if row.get("business_id") == business_id and (row.get("id") in request_ids or row.get("action_id") in action_ids)
            ]
            return None
        if "from finance_import_batches" in normalized_query:
            business_id = params[0]
            batch_ids = set(params[1])
            file_hashes = set(params[2])
            self.last_results = [
                row
                for row in self.tables["finance_import_batches"].values()
                if row.get("business_id") == business_id and (row.get("id") in batch_ids or row.get("file_hash") in file_hashes)
            ]
            return None
        if "from finance_entries" in normalized_query and "duplicate_key" in normalized_query:
            business_id = params[0]
            duplicate_key = params[1]
            self.last_result = next(
                (
                    row
                    for row in self.tables["finance_entries"].values()
                    if row.get("business_id") == business_id and row.get("duplicate_key") == duplicate_key
                ),
                None,
            )
            return None
        if normalized_query.startswith("select * from agent_blueprint_versions where id"):
            self.last_result = self.tables["agent_blueprint_versions"].get(params[0])
            return None
        if normalized_query.startswith("select * from agent_blueprint_versions where blueprint_id"):
            blueprint_id = params[0]
            versions = [
                version
                for version in self.tables["agent_blueprint_versions"].values()
                if version.get("blueprint_id") == blueprint_id
            ]
            versions = sorted(versions, key=lambda item: item.get("version_number") or 0, reverse=True)
            self.last_result = versions[0] if versions else None
            return None
        if normalized_query.startswith("select * from agent_blueprints where id"):
            self.last_result = self.tables["agent_blueprints"].get(params[0])
            return None
        if "from agent_integrations" in normalized_query:
            business_id = params[0]
            integration_ids = []
            if "id = any" in normalized_query and len(params) > 1:
                integration_ids = list(params[1] or [])
            rows = [
                row
                for row in self.tables["agent_integrations"].values()
                if row.get("business_id") == business_id
                and (not integration_ids or row.get("id") in integration_ids)
            ]
            rows = sorted(rows, key=lambda item: str(item.get("updated_at") or item.get("created_at") or ""), reverse=True)
            self.last_results = rows[:100]
            return None
        if normalized_query.startswith("insert into agent_runs"):
            self.tables["agent_runs"][params[0]] = {
                "id": params[0],
                "blueprint_id": params[1],
                "blueprint_version_id": params[2],
                "business_id": params[3],
                "status": params[4],
                "input_json": json.loads(params[5]),
                "output_json": json.loads(params[6]),
                "created_by_user_id": params[7],
            }
            return None
        if normalized_query.startswith("select * from agent_runs where id"):
            self.last_result = self.tables["agent_runs"].get(params[0])
            return None
        if normalized_query.startswith("select id, status, input_json, output_json, error_text"):
            blueprint_id = params[0]
            version_id = params[1]
            self.last_results = [
                run
                for run in self.tables["agent_runs"].values()
                if run.get("blueprint_id") == blueprint_id and run.get("blueprint_version_id") == version_id
            ]
            return None
        if normalized_query.startswith("select step_index from agent_run_steps"):
            run_id = params[0]
            self.last_results = [
                {"step_index": step["step_index"]}
                for step in self.tables["agent_run_steps"].values()
                if step["run_id"] == run_id and step["status"] in {"completed", "rejected"}
            ]
            return None
        if normalized_query.startswith("insert into agent_run_steps"):
            self.tables["agent_run_steps"][params[0]] = {
                "id": params[0],
                "run_id": params[1],
                "step_index": params[2],
                "step_key": params[3],
                "step_type": params[4],
                "status": params[5],
                "input_json": json.loads(params[6]),
                "output_json": json.loads(params[7]),
            }
            return None
        if normalized_query.startswith("insert into agent_artifacts"):
            self.tables["agent_artifacts"][params[0]] = {
                "id": params[0],
                "run_id": params[1],
                "step_id": params[2],
                "artifact_type": params[3],
                "title": params[4],
                "payload_json": json.loads(params[5]),
            }
            return None
        if normalized_query.startswith("insert into agent_approvals"):
            self.tables["agent_approvals"][params[0]] = {
                "id": params[0],
                "run_id": params[1],
                "step_id": params[2],
                "status": "pending",
                "approval_type": params[3],
                "title": params[4],
                "payload_json": json.loads(params[5]),
                "requested_by_user_id": params[6],
            }
            return None
        if normalized_query.startswith("update agent_runs set status = 'waiting_approval'"):
            self.tables["agent_runs"][params[0]]["status"] = "waiting_approval"
            return None
        if normalized_query.startswith("update agent_runs set status = 'failed'"):
            self.tables["agent_runs"][params[1]]["status"] = "failed"
            self.tables["agent_runs"][params[1]]["error_text"] = params[0]
            return None
        if normalized_query.startswith("update agent_run_steps set status = 'failed'"):
            self.tables["agent_run_steps"][params[2]]["status"] = "failed"
            self.tables["agent_run_steps"][params[2]]["output_json"] = json.loads(params[0])
            self.tables["agent_run_steps"][params[2]]["error_text"] = params[1]
            return None
        if normalized_query.startswith("update agent_run_steps set status = 'blocked'"):
            self.tables["agent_run_steps"][params[2]]["status"] = "blocked"
            self.tables["agent_run_steps"][params[2]]["output_json"] = json.loads(params[0])
            self.tables["agent_run_steps"][params[2]]["error_text"] = params[1]
            return None
        if normalized_query.startswith("update agent_run_steps set output_json"):
            self.tables["agent_run_steps"][params[1]]["output_json"] = json.loads(params[0])
            return None
        if normalized_query.startswith("select step_key, output_json from agent_run_steps"):
            run_id = params[0]
            self.last_results = [
                {
                    "step_key": step["step_key"],
                    "output_json": step["output_json"],
                }
                for step in sorted(
                    self.tables["agent_run_steps"].values(),
                    key=lambda item: item["step_index"],
                )
                if step["run_id"] == run_id and step["status"] == "completed"
            ]
            return None
        if normalized_query.startswith("select * from agent_run_steps"):
            run_id = params[0]
            self.last_results = sorted(
                [step for step in self.tables["agent_run_steps"].values() if step["run_id"] == run_id],
                key=lambda item: item["step_index"],
            )
            return None
        if normalized_query.startswith("select output_json from agent_run_steps"):
            run_id = params[0]
            step_key = params[1]
            matches = [
                step
                for step in self.tables["agent_run_steps"].values()
                if step["run_id"] == run_id and step["step_key"] == step_key and step["status"] == "completed"
            ]
            matches = sorted(matches, key=lambda item: item["step_index"], reverse=True)
            self.last_result = {"output_json": matches[0]["output_json"]} if matches else None
            return None
        if normalized_query.startswith("select * from agent_artifacts"):
            run_id = params[0]
            self.last_results = [item for item in self.tables["agent_artifacts"].values() if item["run_id"] == run_id]
            return None
        if normalized_query.startswith("select * from agent_approvals where id"):
            approval_id = params[0]
            run_id = params[1]
            approval = self.tables["agent_approvals"].get(approval_id)
            self.last_result = approval if approval and approval["run_id"] == run_id else None
            return None
        if normalized_query.startswith("select * from agent_approvals"):
            run_id = params[0]
            self.last_results = [item for item in self.tables["agent_approvals"].values() if item["run_id"] == run_id]
            return None
        if normalized_query.startswith("select payload_json from agent_artifacts"):
            run_id = params[0]
            artifact_type = params[1]
            matches = [
                item
                for item in self.tables["agent_artifacts"].values()
                if item["run_id"] == run_id and item["artifact_type"] == artifact_type
            ]
            self.last_result = {"payload_json": matches[-1]["payload_json"]} if matches else None
            return None
        if normalized_query.startswith("insert into finance_import_batches"):
            self.tables["finance_import_batches"][params[0]] = {
                "id": params[0],
                "business_id": params[1],
                "source_type": "agent",
                "status": "processing",
                "file_name": params[2],
                "file_hash": params[3],
                "rows_total": params[4],
                "rows_imported": 0,
                "rows_skipped": 0,
                "rows_failed": 0,
                "mapping_json": json.loads(params[5]),
                "error_log": json.loads(params[6]),
            }
            return None
        if normalized_query.startswith("insert into finance_entries"):
            self.tables["finance_entries"][params[0]] = {
                "id": params[0],
                "business_id": params[1],
                "date": params[2],
                "type": params[3],
                "category": params[4],
                "amount": params[5],
                "source": "agent",
                "comment": params[6],
                "import_batch_id": params[7],
                "external_id": params[8],
                "duplicate_key": params[9],
            }
            return None
        if normalized_query.startswith("update finance_import_batches"):
            batch = self.tables["finance_import_batches"].get(params[5])
            if batch and batch.get("business_id") == params[6]:
                batch["status"] = params[0]
                batch["rows_imported"] = params[1]
                batch["rows_skipped"] = params[2]
                batch["rows_failed"] = params[3]
                batch["error_log"] = json.loads(params[4])
            return None
        if normalized_query.startswith("insert into agent_action_ledger"):
            metadata = json.loads(params[14])
            entry = {
                "id": params[0],
                "action_type": params[3],
                "capability": params[4],
                "risk_level": params[6],
                "output_summary": json.loads(params[8]),
                "status": params[10],
                "reason_code": params[11],
                "metadata": metadata,
            }
            self.ledger_entries.append(entry)
            self.tables["agent_action_ledger"][params[0]] = entry
            return None
        if normalized_query.startswith("select 1 from agent_approvals"):
            run_id = params[0]
            approval_type = params[1] if len(params) > 1 else None
            matches = [
                item
                for item in self.tables["agent_approvals"].values()
                if item["run_id"] == run_id
                and item["status"] == "approved"
                and (approval_type is None or item["approval_type"] == approval_type)
            ]
            self.last_result = {"?column?": 1} if matches else None
            return None
        if normalized_query.startswith("select id, name, city, category, source"):
            business_id = params[0]
            lead_ids = params[1] if len(params) > 2 and isinstance(params[1], list) else []
            rows = []
            for lead in self.tables["prospectingleads"].values():
                if lead.get("business_id") != business_id:
                    continue
                if lead_ids and lead.get("id") not in lead_ids:
                    continue
                rows.append(lead)
            self.last_results = rows
            return None
        if normalized_query.startswith("select id, name, city, email"):
            business_id = params[0]
            lead_ids = params[1] if len(params) > 2 else []
            rows = []
            for lead in self.tables["prospectingleads"].values():
                if lead.get("business_id") != business_id:
                    continue
                if lead_ids and lead.get("id") not in lead_ids:
                    continue
                rows.append(lead)
            self.last_results = rows
            return None
        if normalized_query.startswith("update prospectingleads set status = case"):
            business_id = params[3]
            lead_ids = params[4]
            for lead_id in lead_ids:
                lead = self.tables["prospectingleads"].get(lead_id)
                if lead and lead.get("business_id") == business_id and lead.get("status") not in {
                    "channel_selected",
                    "draft_ready",
                    "queued_for_send",
                    "sent",
                    "delivered",
                }:
                    lead["status"] = params[0]
                    lead["pipeline_status"] = params[1]
                    lead["last_manual_action_by"] = params[2]
            return None
        if normalized_query.startswith("select d.id"):
            business_id = params[0]
            draft_ids = params[1] if len(params) > 2 else []
            rows = []
            for draft in self.tables["outreachmessagedrafts"].values():
                lead = self.tables["prospectingleads"].get(draft.get("lead_id"))
                if not lead or lead.get("business_id") != business_id:
                    continue
                if draft_ids and draft.get("id") not in draft_ids:
                    continue
                rows.append(
                    {
                        **draft,
                        "lead_name": lead.get("name"),
                    }
                )
            self.last_results = rows
            return None
        if normalized_query.startswith("select id, name, category"):
            business_id = params[0]
            lead_ids = set(params[1])
            limit = params[2]
            rows = []
            for lead in self.tables["prospectingleads"].values():
                if lead.get("business_id") != business_id or lead.get("id") not in lead_ids:
                    continue
                has_draft = any(
                    draft.get("lead_id") == lead.get("id") and draft.get("status") in {"generated", "edited", "approved"}
                    for draft in self.tables["outreachmessagedrafts"].values()
                )
                if not has_draft:
                    rows.append(lead)
            self.last_results = rows[:limit]
            return None
        if normalized_query.startswith("insert into outreachmessagedrafts"):
            self.tables["outreachmessagedrafts"][params[0]] = {
                "id": params[0],
                "lead_id": params[1],
                "channel": params[2],
                "angle_type": params[3],
                "tone": params[4],
                "status": params[5],
                "generated_text": params[6],
                "edited_text": params[7],
                "learning_note_json": json.loads(params[8]),
                "created_by": params[9],
                "approved_text": None,
            }
            return None
        if normalized_query.startswith("update prospectingleads set status = %s, selected_channel"):
            lead = self.tables["prospectingleads"].get(params[3])
            if lead and lead.get("business_id") == params[4]:
                lead["status"] = params[0]
                lead["selected_channel"] = params[1]
                lead["pipeline_status"] = params[2]
            return None
        if normalized_query.startswith("update outreachmessagedrafts set status = %s"):
            draft_ids = params[2]
            business_id = params[3]
            for draft_id in draft_ids:
                draft = self.tables["outreachmessagedrafts"].get(draft_id)
                lead = self.tables["prospectingleads"].get((draft or {}).get("lead_id"))
                if draft and lead and lead.get("business_id") == business_id:
                    draft["status"] = params[0]
                    draft["approved_by"] = params[1]
                    draft["approved_text"] = draft.get("edited_text") or draft.get("generated_text")
                    draft["edited_text"] = draft.get("edited_text") or draft.get("generated_text")
            return None
        if normalized_query.startswith("update prospectingleads set status = %s, pipeline_status"):
            business_id = params[2]
            draft_ids = set(params[3])
            lead_ids = {
                draft.get("lead_id")
                for draft in self.tables["outreachmessagedrafts"].values()
                if draft.get("id") in draft_ids
            }
            for lead_id in lead_ids:
                lead = self.tables["prospectingleads"].get(lead_id)
                if lead and lead.get("business_id") == business_id:
                    lead["status"] = params[0]
                    lead["pipeline_status"] = params[1]
            return None
        if normalized_query.startswith("update agent_approvals set status = 'approved'"):
            approval = self.tables["agent_approvals"].get(params[2])
            if approval and approval["run_id"] == params[3]:
                approval["status"] = "approved"
                approval["decided_by_user_id"] = params[0]
                approval["decision_reason"] = params[1]
            return None
        if normalized_query.startswith("update agent_run_steps set status = 'completed'"):
            step = self.tables["agent_run_steps"].get(params[1])
            if step:
                step["status"] = "completed"
                step["output_json"] = json.loads(params[0])
            return None
        if normalized_query.startswith("update agent_runs set status = 'running'"):
            self.tables["agent_runs"][params[0]]["status"] = "running"
            return None
        if normalized_query.startswith("update agent_runs set status = 'completed'"):
            self.tables["agent_runs"][params[1]]["status"] = "completed"
            self.tables["agent_runs"][params[1]]["output_json"] = json.loads(params[0])
            return None
        if normalized_query.startswith("update agent_blueprints") and "set metadata_json" in normalized_query:
            blueprint = self.tables["agent_blueprints"].get(params[1])
            if blueprint:
                blueprint["metadata_json"] = json.loads(params[0])
            return None
        if normalized_query.startswith("update agent_blueprints") and "set status = 'active'" in normalized_query:
            blueprint = self.tables["agent_blueprints"].get(params[0])
            if blueprint and blueprint.get("status") != "archived":
                blueprint["status"] = "active"
            return None
        if normalized_query.startswith("select count(*) as count from agent_artifacts"):
            run_id = params[0]
            self.last_result = {"count": len([item for item in self.tables["agent_artifacts"].values() if item["run_id"] == run_id])}
            return None
        if normalized_query.startswith("select count(*) as count from agent_approvals"):
            run_id = params[0]
            self.last_result = {"count": len([item for item in self.tables["agent_approvals"].values() if item["run_id"] == run_id])}
            return None
        raise AssertionError(f"Unhandled SQL in fake cursor: {query}")

    def fetchone(self):
        return self.last_result

    def fetchall(self):
        return self.last_results


class FakeDatahubCursor:
    def __init__(self):
        self.last_results = []

    def execute(self, query, params=None):
        normalized_query = " ".join(query.split()).lower()
        if "from businesses" in normalized_query:
            self.last_results = [{"id": "biz1", "name": "Local Test", "business_type": "beauty", "city": "Moscow", "address": "Street"}]
            return None
        if "from userservices" in normalized_query:
            self.last_results = [
                {"id": "svc1", "name": "Haircut", "price": 1000, "description": "Cut"},
                {"id": "svc2", "name": "Color", "price": 3000, "description": "Color"},
            ]
            return None
        if "from externalbusinessreviews" in normalized_query:
            self.last_results = [{"id": "rev1", "author_name": "Anna", "rating": 5, "text": "Great"}]
            return None
        if "from prospectingleads" in normalized_query:
            self.last_results = []
            return None
        if "from outreachmessagedrafts" in normalized_query:
            self.last_results = [{"id": "draft1", "channel": "email", "status": "generated", "generated_text": "Hello"}]
            return None
        raise AssertionError(f"Unhandled Datahub SQL: {query}")

    def fetchall(self):
        return self.last_results


class FakeCapabilityDatabase:
    def __init__(self):
        self.cursor_instance = FakeCapabilityCursor()
        self.conn = self
        self.committed = False
        self.rolled_back = False
        self.closed = False

    def cursor(self):
        return self.cursor_instance

    def commit(self):
        self.committed = True

    def rollback(self):
        self.rolled_back = True

    def close(self):
        self.closed = True


class FakeCapabilityCursor:
    def __init__(self):
        self.last_result = None
        self.last_results = []
        self.description = []
        self.inserted = {}

    def execute(self, query, params=None):
        normalized_query = " ".join(query.split()).lower()
        params = params or ()
        if (
            normalized_query.startswith("create table")
            or normalized_query.startswith("create index")
            or normalized_query.startswith("create unique index")
            or normalized_query.startswith("alter table")
        ):
            return None
        if normalized_query.startswith("select to_regclass"):
            table_name = str(params[0])
            self.description = [("to_regclass",)]
            self.last_result = (table_name if table_name in {"bookings", "externalbusinessreviews", "userservices"} else None,)
            return None
        if "from information_schema.columns" in normalized_query:
            table_name = str(params[0])
            column_map = {
                "userservices": [
                    "id",
                    "business_id",
                    "name",
                    "description",
                    "category",
                    "price",
                    "is_active",
                    "updated_at",
                    "created_at",
                    "optimized_name",
                    "optimized_description",
                ],
            }
            self.description = [("column_name",)]
            self.last_results = [(column,) for column in column_map.get(table_name, [])]
            return None
        if "from bookings" in normalized_query:
            columns = [
                "id",
                "business_id",
                "client_phone",
                "client_name",
                "service_id",
                "service_name",
                "booking_date",
                "booking_time",
                "status",
                "notes",
                "created_at",
                "updated_at",
            ]
            self.description = [(column,) for column in columns]
            self.last_results = [
                (
                    "booking-1",
                    "biz1",
                    "+79990000000",
                    "Анна",
                    "svc1",
                    "Стрижка",
                    "2026-06-10",
                    "12:00",
                    "confirmed",
                    "",
                    None,
                    None,
                )
            ]
            return None
        if normalized_query.startswith("insert into agent_communication_requests"):
            self.inserted["agent_communication_requests"] = {
                "id": params[0],
                "action_id": params[1],
                "business_id": params[2],
                "user_id": params[3],
                "capability": params[4],
                "message_type": params[5],
                "channel": params[6],
                "recipient_count": params[7],
                "recipients_json": json.loads(params[8]),
                "message_template": params[9],
                "limits_json": json.loads(params[10]),
                "consent_json": json.loads(params[11]),
            }
            self.description = [("id",)]
            self.last_result = (params[0],)
            return None
        if "from externalbusinessreviews" in normalized_query:
            columns = [
                "id",
                "business_id",
                "source",
                "external_review_id",
                "rating",
                "author_name",
                "text",
                "response_text",
                "response_at",
            ]
            self.description = [(column,) for column in columns]
            self.last_result = ("rev1", "biz1", "yandex", "ext-rev1", 5, "Anna", "Great", None, None)
            return None
        if normalized_query.startswith("insert into reviewreplydrafts"):
            self.inserted["reviewreplydrafts"] = {
                "id": params[0],
                "business_id": params[1],
                "review_id": params[2],
                "user_id": params[3],
                "source": params[4],
                "rating": params[5],
                "author_name": params[6],
                "review_text": params[7],
                "generated_text": params[8],
                "status": "publish_requested",
            }
            self.description = [("id",)]
            self.last_result = (params[0],)
            return None
        if "from userservices" in normalized_query:
            columns = [
                "id",
                "business_id",
                "name",
                "description",
                "category",
                "price",
                "optimized_name",
                "optimized_description",
            ]
            self.description = [(column,) for column in columns]
            self.last_results = [
                ("svc1", "biz1", "Стрижка", "Классическая стрижка", "Парикмахерские услуги", 1500, "", "")
            ]
            return None
        if normalized_query.startswith("insert into agent_service_optimization_requests"):
            self.inserted["agent_service_optimization_requests"] = {
                "id": params[0],
                "action_id": params[1],
                "business_id": params[2],
                "user_id": params[3],
                "status": "draft_ready",
                "service_count": params[4],
                "suggestions_json": json.loads(params[5]),
                "diff_json": json.loads(params[6]),
                "apply_state": "not_applied",
            }
            self.description = [("id",)]
            self.last_result = (params[0],)
            return None
        if normalized_query.startswith("insert into agent_sheet_operation_requests"):
            self.inserted["agent_sheet_operation_requests"] = {
                "id": params[0],
                "action_id": params[1],
                "business_id": params[2],
                "user_id": params[3],
                "integration_id": params[4],
                "spreadsheet_id": params[5],
                "sheet_name": params[6],
                "operation": "append_row",
                "status": "request_created",
                "approval_state": "pending_human",
                "apply_state": "not_applied",
                "row_values_json": json.loads(params[7]),
                "mapping_json": json.loads(params[8]),
                "source_event_json": json.loads(params[9]),
                "limits_json": json.loads(params[10]),
                "provider_write_performed": False,
            }
            self.description = [("id",)]
            self.last_result = (params[0],)
            return None
        raise AssertionError(f"Unhandled capability SQL: {query}")

    def fetchone(self):
        return self.last_result

    def fetchall(self):
        return self.last_results


class FakeApprovedDomainExecutorCursor:
    def __init__(self):
        self.tables = {
            "agent_sheet_operation_requests": {},
            "agent_communication_requests": {},
            "reviewreplydrafts": {},
            "agent_review_publish_requests": {},
            "agent_service_optimization_requests": {},
            "agent_communication_delivery_journal": {},
            "agent_action_ledger": {},
            "userservices": {},
            "finance_import_batches": {},
            "finance_entries": {},
        }
        self.last_result = None
        self.last_results = []
        self.ledger_entries = []

    def execute(self, query, params=None):
        normalized_query = " ".join(query.split()).lower()
        params = params or ()
        if normalized_query.startswith("select to_regclass"):
            table_name = str(params[0])
            self.last_result = (table_name if table_name in self.tables else None,)
            return None
        if "from information_schema.columns" in normalized_query:
            table_name = str(params[0])
            columns = {
                "userservices": ["id", "business_id", "optimized_name", "optimized_description", "updated_at"],
            }.get(table_name, [])
            self.last_results = [(column,) for column in columns]
            return None
        if normalized_query.startswith("select id, action_id, status, approval_state"):
            business_id = params[0]
            request_ids = set(params[1])
            action_ids = set(params[2])
            self.last_results = [
                row
                for row in self.tables["agent_sheet_operation_requests"].values()
                if row.get("business_id") == business_id and (row.get("id") in request_ids or row.get("action_id") in action_ids)
            ]
            return None
        if "from agent_communication_requests" in normalized_query:
            business_id = params[0]
            request_ids = set(params[1])
            action_ids = set(params[2])
            self.last_results = [
                row
                for row in self.tables["agent_communication_requests"].values()
                if row.get("business_id") == business_id and (row.get("id") in request_ids or row.get("action_id") in action_ids)
            ]
            return None
        if "from reviewreplydrafts" in normalized_query:
            business_id = params[0]
            draft_ids = set(params[1])
            review_ids = set(params[2])
            self.last_results = [
                row
                for row in self.tables["reviewreplydrafts"].values()
                if row.get("business_id") == business_id and (row.get("id") in draft_ids or row.get("review_id") in review_ids)
            ]
            return None
        if "from agent_review_publish_requests" in normalized_query:
            business_id = params[0]
            draft_id = params[1]
            review_id = params[2]
            self.last_results = [
                row
                for row in self.tables["agent_review_publish_requests"].values()
                if row.get("business_id") == business_id and (row.get("draft_id") == draft_id or row.get("review_id") == review_id)
            ]
            return None
        if "from agent_service_optimization_requests" in normalized_query:
            business_id = params[0]
            request_ids = set(params[1])
            action_ids = set(params[2])
            self.last_results = [
                row
                for row in self.tables["agent_service_optimization_requests"].values()
                if row.get("business_id") == business_id and (row.get("id") in request_ids or row.get("action_id") in action_ids)
            ]
            return None
        if "from finance_entries" in normalized_query and "duplicate_key" in normalized_query:
            business_id = params[0]
            duplicate_key = params[1]
            self.last_result = next(
                (
                    row
                    for row in self.tables["finance_entries"].values()
                    if row.get("business_id") == business_id and row.get("duplicate_key") == duplicate_key
                ),
                None,
            )
            return None
        if normalized_query.startswith("update agent_sheet_operation_requests"):
            request = self.tables["agent_sheet_operation_requests"].get(params[0])
            if request and request.get("business_id") == params[1] and request.get("provider_write_performed") is False:
                request["status"] = "approved_for_execution"
                request["approval_state"] = "approved"
                request["apply_state"] = "provider_request_queued"
            return None
        if normalized_query.startswith("update userservices"):
            service = self.tables["userservices"].get(params[2])
            if service and service.get("business_id") == params[3]:
                if params[0]:
                    service["optimized_name"] = params[0]
                if params[1]:
                    service["optimized_description"] = params[1]
            return None
        if normalized_query.startswith("insert into agent_communication_delivery_journal"):
            row = {
                "id": params[0],
                "request_id": params[1],
                "action_id": params[2],
                "business_id": params[3],
                "run_id": params[4],
                "user_id": params[5],
                "recipient_key": params[6],
                "channel": params[7],
                "message_template": params[8],
                "status": params[9],
                "delivery_state": params[10],
                "consent_json": json.loads(params[11]),
                "limits_json": json.loads(params[12]),
                "router_handoff_json": json.loads(params[13]),
                "provider_write_performed": False,
            }
            self.tables["agent_communication_delivery_journal"][params[0]] = row
            return None
        if normalized_query.startswith("update agent_communication_requests"):
            request = self.tables["agent_communication_requests"].get(params[1])
            if request and request.get("business_id") == params[2] and request.get("delivery_state") != "dispatched":
                request["status"] = "approved_for_dispatch"
                request["delivery_state"] = params[0]
            return None
        if normalized_query.startswith("update reviewreplydrafts"):
            draft = self.tables["reviewreplydrafts"].get(params[0])
            if draft and draft.get("business_id") == params[1]:
                draft["status"] = "approved_for_publish"
            return None
        if normalized_query.startswith("insert into agent_review_publish_requests"):
            row = {
                "id": params[0],
                "draft_id": params[1],
                "review_id": params[2],
                "business_id": params[3],
                "run_id": params[4],
                "user_id": params[5],
                "source": params[6],
                "reply_text": params[7],
                "status": "provider_publish_requested",
                "publish_state": "provider_request_queued",
                "provider_request_json": json.loads(params[8]),
                "audit_json": json.loads(params[9]),
                "provider_write_performed": False,
            }
            self.tables["agent_review_publish_requests"][params[0]] = row
            return None
        if normalized_query.startswith("update agent_service_optimization_requests"):
            request = self.tables["agent_service_optimization_requests"].get(params[2])
            if request and request.get("business_id") == params[3] and request.get("apply_state") != "applied":
                request["status"] = params[0]
                request["apply_state"] = params[1]
            return None
        if normalized_query.startswith("insert into finance_import_batches"):
            self.tables["finance_import_batches"][params[0]] = {
                "id": params[0],
                "business_id": params[1],
                "source_type": "agent",
                "status": "processing",
                "file_name": params[2],
                "file_hash": params[3],
                "rows_total": params[4],
                "rows_imported": 0,
                "rows_skipped": 0,
                "rows_failed": 0,
                "mapping_json": json.loads(params[5]),
                "error_log": json.loads(params[6]),
            }
            return None
        if normalized_query.startswith("insert into finance_entries"):
            self.tables["finance_entries"][params[0]] = {
                "id": params[0],
                "business_id": params[1],
                "date": params[2],
                "type": params[3],
                "category": params[4],
                "amount": params[5],
                "source": "agent",
                "comment": params[6],
                "import_batch_id": params[7],
                "external_id": params[8],
                "duplicate_key": params[9],
            }
            return None
        if normalized_query.startswith("update finance_import_batches"):
            batch = self.tables["finance_import_batches"].get(params[5])
            if batch and batch.get("business_id") == params[6]:
                batch["status"] = params[0]
                batch["rows_imported"] = params[1]
                batch["rows_skipped"] = params[2]
                batch["rows_failed"] = params[3]
                batch["error_log"] = json.loads(params[4])
            return None
        if (
            normalized_query.startswith("create table")
            or normalized_query.startswith("create index")
            or normalized_query.startswith("create unique index")
        ):
            return None
        if normalized_query.startswith("insert into agent_action_ledger"):
            metadata = json.loads(params[14])
            entry = {
                "id": params[0],
                "action_type": params[3],
                "capability": params[4],
                "risk_level": params[6],
                "output_summary": json.loads(params[8]),
                "status": params[10],
                "reason_code": params[11],
                "metadata": metadata,
            }
            self.ledger_entries.append(entry)
            self.tables["agent_action_ledger"][params[0]] = entry
            return None
        raise AssertionError(f"Unhandled approved executor SQL: {query}")

    def fetchone(self):
        return self.last_result

    def fetchall(self):
        return self.last_results


class FakeSheetProviderExecutorCursor:
    def __init__(self):
        self.tables = {
            "agent_sheet_operation_requests": {},
            "agent_action_ledger": {},
        }
        self.last_result = None
        self.last_results = []
        self.ledger_entries = []

    def execute(self, query, params=None):
        normalized_query = " ".join(query.split()).lower()
        params = params or ()
        if normalized_query.startswith("select id, action_id, business_id"):
            business_id = params[0]
            limit = params[1]
            rows = [
                row
                for row in self.tables["agent_sheet_operation_requests"].values()
                if row.get("business_id") == business_id
                and row.get("status") == "approved_for_execution"
                and row.get("approval_state") == "approved"
                and row.get("apply_state") == "provider_request_queued"
                and row.get("provider_write_performed") is False
            ]
            self.last_results = rows[:limit]
            return None
        if "from agent_integrations" in normalized_query:
            self.last_result = None
            return None
        if normalized_query.startswith("update agent_sheet_operation_requests"):
            request = self.tables["agent_sheet_operation_requests"].get(params[4])
            if request and request.get("business_id") == params[5] and request.get("apply_state") == "provider_request_queued":
                request["status"] = params[0]
                request["apply_state"] = params[1]
                request["provider_write_performed"] = params[2]
                request["error_text"] = params[3] or None
            return None
        if (
            normalized_query.startswith("create table")
            or normalized_query.startswith("create index")
            or normalized_query.startswith("create unique index")
        ):
            return None
        if normalized_query.startswith("insert into agent_action_ledger"):
            entry = {
                "id": params[0],
                "action_type": params[3],
                "capability": params[4],
                "risk_level": params[6],
                "input_summary": json.loads(params[7]),
                "output_summary": json.loads(params[8]),
                "status": params[10],
                "reason_code": params[11],
                "metadata": json.loads(params[14]),
            }
            self.ledger_entries.append(entry)
            self.tables["agent_action_ledger"][params[0]] = entry
            return None
        raise AssertionError(f"Unhandled sheet provider SQL: {query}")

    def fetchall(self):
        return self.last_results

    def fetchone(self):
        return self.last_result


class FakeGoogleSheetsIntegrationCursor:
    def __init__(self):
        self.agent_integrations = {}
        self.external_accounts = {}
        self.last_result = None
        self.last_results = []

    def execute(self, query, params=None):
        normalized_query = " ".join(query.split()).lower()
        params = params or ()
        if "from agent_integrations" in normalized_query:
            business_id = params[1] if "where id = %s" in normalized_query else params[0]
            integration_id = params[0] if "where id = %s" in normalized_query else ""
            rows = [
                row
                for row in self.agent_integrations.values()
                if row.get("business_id") == business_id
                and row.get("provider") == "google_sheets"
                and row.get("status") == "active"
                and (not integration_id or row.get("id") == integration_id)
            ]
            self.last_result = rows[0] if rows else None
            return None
        if "from information_schema.columns" in normalized_query:
            self.last_results = [{"column_name": "auth_data_encrypted"}]
            return None
        if "from externalbusinessaccounts" in normalized_query:
            if "select id, business_id, source, display_name" in normalized_query:
                if "where id = %s" in normalized_query:
                    account = self.external_accounts.get(params[0])
                    self.last_result = (
                        account
                        if account and account.get("business_id") == params[1] and account.get("is_active") is True
                        else None
                    )
                    return None
                business_id = params[0]
                rows = [
                    account
                    for account in self.external_accounts.values()
                    if account.get("business_id") == business_id
                    and account.get("is_active") is True
                    and account.get("source") in {"google_sheets", "google_business"}
                ]
                self.last_result = rows[0] if rows else None
                self.last_results = rows
                return None
            auth_ref = params[0]
            business_id = params[1]
            account = self.external_accounts.get(auth_ref)
            self.last_result = (
                account
                if account and account.get("business_id") == business_id and account.get("is_active") is True
                else None
            )
            return None
        raise AssertionError(f"Unhandled Google Sheets integration SQL: {query}")

    def fetchone(self):
        return self.last_result

    def fetchall(self):
        return self.last_results


class FakeTelegramTriggerCursor:
    def __init__(self):
        self.last_result = None
        self.last_results = []
        self.trigger_events = []

    def execute(self, query, params=None):
        normalized_query = " ".join(query.split()).lower()
        params = params or ()
        if normalized_query.startswith("create table") or normalized_query.startswith("create index"):
            return None
        if normalized_query.startswith("insert into agent_trigger_events"):
            is_scheduler_event = len(params) == 4
            self.trigger_events.append(
                {
                    "id": params[0],
                    "business_id": params[1],
                    "source": "scheduler" if is_scheduler_event else "telegram",
                    "event_type": params[2] if is_scheduler_event else "telegram.message.received",
                    "status": "received",
                    "payload_json": json.loads(params[3] if is_scheduler_event else params[2]),
                    "reason_code": None,
                }
            )
            return None
        if "from agent_blueprints" in normalized_query:
            self.last_results = []
            return None
        if normalized_query.startswith("update agent_trigger_events set status = 'ignored'"):
            for item in self.trigger_events:
                if item["id"] == params[1]:
                    item["status"] = "ignored"
                    item["reason_code"] = params[0]
            return None
        raise AssertionError(f"Unhandled trigger SQL: {query}")

    def fetchone(self):
        return self.last_result

    def fetchall(self):
        return self.last_results


class FakeActiveTelegramTriggerCursor(FakeCursor):
    def __init__(self):
        super().__init__()
        self.trigger_events = []

    def execute(self, query, params=None):
        normalized_query = " ".join(query.split()).lower()
        params = params or ()
        if normalized_query.startswith("create table") or normalized_query.startswith("create index"):
            return None
        if normalized_query.startswith("insert into agent_trigger_events"):
            is_scheduler_event = len(params) == 4
            self.trigger_events.append(
                {
                    "id": params[0],
                    "business_id": params[1],
                    "source": "scheduler" if is_scheduler_event else "telegram",
                    "event_type": params[2] if is_scheduler_event else "telegram.message.received",
                    "status": "received",
                    "payload_json": json.loads(params[3] if is_scheduler_event else params[2]),
                    "reason_code": None,
                    "blueprint_id": None,
                    "run_id": None,
                }
            )
            return None
        if normalized_query.startswith("select id from agent_trigger_events"):
            business_id = params[0]
            event_type = params[1]
            self.last_result = next(
                (
                    item
                    for item in self.trigger_events
                    if item.get("business_id") == business_id
                    and item.get("source") == "scheduler"
                    and item.get("event_type") == event_type
                ),
                None,
            )
            return None
        if normalized_query.startswith("select id, business_id, metadata_json from agent_blueprints"):
            limit = int(params[0])
            self.last_results = [
                row
                for row in self.tables["agent_blueprints"].values()
                if row.get("status") == "active" and row.get("category") in {"custom", "tables"}
            ][:limit]
            return None
        if "from agent_blueprints" in normalized_query and "status = 'active'" in normalized_query:
            business_id = params[0]
            self.last_results = [
                row
                for row in self.tables["agent_blueprints"].values()
                if row.get("business_id") == business_id
                and row.get("status") == "active"
                and row.get("category") in {"custom", "tables"}
            ]
            return None
        if normalized_query.startswith("select * from agent_blueprint_versions where blueprint_id"):
            blueprint_id = params[0]
            rows = [
                row
                for row in self.tables["agent_blueprint_versions"].values()
                if row.get("blueprint_id") == blueprint_id
            ]
            rows.sort(key=lambda item: int(item.get("version_number") or 0), reverse=True)
            self.last_result = rows[0] if rows else None
            return None
        if normalized_query.startswith("update agent_trigger_events set blueprint_id"):
            for item in self.trigger_events:
                if item["id"] == params[2]:
                    item["blueprint_id"] = params[0]
                    item["run_id"] = params[1]
                    item["status"] = "run_started"
            return None
        if normalized_query.startswith("update agent_trigger_events set status = 'ignored'"):
            for item in self.trigger_events:
                if item["id"] == params[1]:
                    item["status"] = "ignored"
                    item["reason_code"] = params[0]
            return None
        return super().execute(query, params)


class CountingOrchestrator:
    def __init__(self):
        self.calls = 0
        self.last_envelope = None

    def execute(self, envelope, user_data, *, allow_execute_when_approved=False):
        self.calls += 1
        self.last_envelope = envelope
        return {
            "success": True,
            "status": "completed",
            "result": {
                "status": "queued_for_dispatch",
                "dispatch_state": "queued_not_dispatched",
                "external_dispatch_performed": False,
            },
        }


class FakeUpload:
    def __init__(self, filename, mimetype, data):
        self.filename = filename
        self.mimetype = mimetype
        self._data = data

    def read(self):
        return self._data


class FakeOutreachConnection:
    def __init__(self, draft_rows):
        self.cursor_instance = FakeOutreachCursor(draft_rows)
        self.inserted = self.cursor_instance.inserted
        self.committed = False
        self.rolled_back = False
        self.closed = False

    def cursor(self):
        return self.cursor_instance

    def commit(self):
        self.committed = True

    def rollback(self):
        self.rolled_back = True

    def close(self):
        self.closed = True


class FakeOutreachCursor:
    def __init__(self, draft_rows):
        self.draft_rows = draft_rows
        self.last_result = None
        self.last_results = []
        self.inserted = []

    def execute(self, query, params=None):
        normalized_query = " ".join(query.split()).lower()
        params = params or ()
        if normalized_query.startswith("select count(*) as cnt"):
            self.last_result = {"cnt": 0}
            return None
        if normalized_query.startswith("select d.id"):
            self.last_results = self.draft_rows
            return None
        if normalized_query.startswith("insert into outreachsendbatches"):
            self.inserted.append(
                {
                    "kind": "batch",
                    "id": params[0],
                    "daily_limit": params[1],
                    "status": params[2],
                    "created_by": params[3],
                    "approved_by": params[4],
                }
            )
            return None
        if normalized_query.startswith("insert into outreachsendqueue"):
            self.inserted.append(
                {
                    "kind": "queue",
                    "id": params[0],
                    "batch_id": params[1],
                    "lead_id": params[2],
                    "draft_id": params[3],
                    "channel": params[4],
                    "delivery_status": params[5],
                }
            )
            return None
        if normalized_query.startswith("update prospectingleads"):
            self.inserted.append(
                {
                    "kind": "lead_update",
                    "status": params[0],
                    "pipeline_status": params[1],
                    "lead_id": params[2],
                    "business_id": params[3],
                }
            )
            return None
        raise AssertionError(f"Unhandled SQL in fake outreach cursor: {query}")

    def fetchone(self):
        return self.last_result

    def fetchall(self):
        return self.last_results
